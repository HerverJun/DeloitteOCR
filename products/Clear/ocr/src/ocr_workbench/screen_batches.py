"""Durable screen-table batches, with bounded GPU work and explicit delivery state."""
from __future__ import annotations

from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import threading
import time
import zipfile

from ocr_workbench.store import Conflict, encoded, now, uid

VERSION = 'screen-batch-v1'
TERMINAL = {'completed', 'needs_review', 'failed', 'cancelled'}
ENGINES = {'glm', 'hunyuan', 'paddlevl'}


def fingerprint(value):
    return hashlib.sha256(encoded(value).encode('utf-8')).hexdigest()


def screen_seconds_spent(store, batch_id, version_id, attempt=0):
    timestamp = now()
    rows = store.rows('''SELECT COALESCE(SUM(s.spent_seconds + CASE WHEN t.started IS NULL THEN 0 ELSE MAX(0,(julianday(COALESCE(t.finished,?))-julianday(t.started))*86400) END),0) seconds
        FROM tasks t JOIN screen_tasks s ON s.task_id=t.id
        WHERE s.batch_id=? AND s.input_version_id=? AND s.attempt=?''',
        (timestamp, batch_id, version_id, attempt))
    spent = rows[0]['seconds']
    if attempt == 0:
        context = store.rows('''SELECT COALESCE(SUM(s.spent_seconds + CASE WHEN t.started IS NULL THEN 0 ELSE MAX(0,(julianday(COALESCE(t.finished,?))-julianday(t.started))*86400) END),0) seconds
            FROM tasks t JOIN screen_contexts s ON s.task_id=t.id
            WHERE s.batch_id=? AND s.input_version_id=?''', (timestamp, batch_id, version_id))
        spent += context[0]['seconds']
    return spent


def validate_policy(value, defaults=None, *, resolve_precision=True):
    value = value or {}
    if not isinstance(value, dict):
        raise ValueError('批次策略必须是对象')
    permitted = {'primary_engine', 'fallback_engines', 'max_upgrades', 'mode',
                 'group_size', 'budget_seconds', 'request_timeout_seconds', 'perspective', 'allow_text_fixes',
                 'visual_review', 'strategy_version', 'layout_evidence', 'batch_budget_seconds',
                 'upgrade_trigger', 'require_human_review', 'context_schedule', 'render_math_text',
                 'engine_route', 'technical_fallback_reserve_seconds'}
    if set(value) - permitted:
        raise ValueError('未知的屏拍批次策略参数')
    result = {**(defaults or {}), **value}
    result.setdefault('primary_engine', 'glm')
    if result['primary_engine'] == 'auto':
        result['primary_engine'] = (defaults or {}).get('primary_engine', 'glm')
    result.setdefault('fallback_engines', ['paddlevl', 'hunyuan'])
    result.setdefault('max_upgrades', 2)
    result.setdefault('mode', 'auto')
    result.setdefault('group_size', 32)
    result.setdefault('budget_seconds', 300)
    result.setdefault('request_timeout_seconds', 180)
    result.setdefault('batch_budget_seconds', 36000)
    result.setdefault('perspective', True)
    result.setdefault('allow_text_fixes', False)
    result.setdefault('visual_review', False)
    result.setdefault('layout_evidence', False)
    result.setdefault('strategy_version', VERSION + '-provisional')
    # Missing keys keep pre-existing persisted batches' interpretation. New
    # measured defaults explicitly opt in through config/screen-policy.json.
    result.setdefault('upgrade_trigger', 'quality')
    result.setdefault('require_human_review', False)
    result.setdefault('context_schedule', 'whole_batch')
    result.setdefault('render_math_text', False)
    result.setdefault('engine_route','fixed')
    result.setdefault('technical_fallback_reserve_seconds',0)
    if result['primary_engine'] not in ENGINES:
        raise ValueError('首轮模型必须具备表格结构输出能力')
    choices = result['fallback_engines']
    if (not isinstance(choices, list) or len(choices) > 3
            or len(set(choices)) != len(choices) or not set(choices) <= ENGINES):
        raise ValueError('升级模型清单无效')
    result['fallback_engines'] = [x for x in choices if x != result['primary_engine']]
    for key, lower, upper in [('max_upgrades', 0, 2), ('group_size', 1, 64),
                              ('budget_seconds', 10, 1800), ('request_timeout_seconds', 10, 900),
                              ('batch_budget_seconds', 10, 864000)]:
        if type(result[key]) is not int or not lower <= result[key] <= upper:
            raise ValueError('批次预算或分组大小无效')
    if result['mode'] not in {'auto', 'accuracy'}:
        raise ValueError('未知处理模式')
    if result['context_schedule'] not in {'whole_batch', 'grouped'}:
        raise ValueError('未知表区证据分组方式')
    if not isinstance(result['upgrade_trigger'], str) or result['upgrade_trigger'] not in {'failure_only', 'quality'}:
        raise ValueError('升级触发条件无效')
    for key in ('perspective', 'allow_text_fixes', 'visual_review', 'layout_evidence', 'require_human_review', 'render_math_text'):
        if type(result[key]) is not bool:
            raise ValueError('处理开关必须为布尔值')
    if result['visual_review']:
        raise ValueError('局部视觉审校仍在评测中，请通过复核工作流调用；批次自动审校尚未启用')
    if not isinstance(result['strategy_version'], str) or not 1 <= len(result['strategy_version']) <= 120:
        raise ValueError('策略版本无效')
    # Preserve the historical accuracy + auto API interpretation. The current
    # checkbox submits the explicit business-language route; explicit choices
    # and frozen requests must retain their original meaning.
    if (resolve_precision and 'upgrade_trigger' in (defaults or {})
            and value.get('mode') == 'accuracy'
            and value.get('primary_engine', 'auto') == 'auto'
            and 'upgrade_trigger' not in value):
        alternatives = [engine for engine in [result['primary_engine'], *result['fallback_engines']]
                        if engine != 'glm']
        result.update(primary_engine='glm', fallback_engines=alternatives,
            max_upgrades=min(1, len(alternatives)), upgrade_trigger='quality',
            require_human_review=True, strategy_version='screen-glm-quality-human-v1-20261007')
    if not isinstance(result['engine_route'], str) or result['engine_route'] not in {'fixed','business_language_v1'}:
        raise ValueError('未知的逐图模型路线')
    reserve=result['technical_fallback_reserve_seconds']
    if type(reserve) is not int or not 0<=reserve<=900:
        raise ValueError('技术兜底预留须为0–900整数秒')
    if result['engine_route']=='business_language_v1':
        if (result['primary_engine']!='glm' or result['fallback_engines']!=['hunyuan']
                or result['max_upgrades']!=1 or result['mode']!='accuracy'
                or result['upgrade_trigger']!='failure_only' or not result['require_human_review']
                or not result['layout_evidence'] or not 10<=reserve<=result['budget_seconds']-10):
            raise ValueError('逐图精度路线须明确GLM/H、一次技术兜底、定位、人工复核和原额度内预留')
    elif reserve:
        raise ValueError('固定路线不接受逐图精度兜底预留')
    return result


class ScreenBatches:
    def __init__(self, store, queue, bundle, *, prepare=None, assess=None, select=None, hold_wal=True):
        self.store, self.queue, self.bundle = store, queue, Path(bundle)
        self.prepare_override, self.assess_override, self.select_override = prepare, assess, select
        self.guard = threading.RLock()
        # Serialize coordinators, without making controls wait for CPU work.
        self.step_guard = threading.Lock()
        self.stopping, self.wake = threading.Event(), threading.Event()
        self.thread = None
        self.worker_error = None
        self.wal_keeper = None
        config = self.bundle / 'config/engines.json'
        source = Path(__file__).resolve().parent
        self.model_identity = fingerprint({'configuration': hashlib.sha256(config.read_bytes()).hexdigest() if config.is_file() else None,
            'worker_source': [(p.name, hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(source.glob('*.py'))]})
        self.root = store.root / 'screen-batches'
        if self.root.absolute().drive.upper() != 'F:' or self.root.resolve().drive.upper() != 'F:':
            raise ValueError('屏拍批次数据必须保存于 F 盘')
        self.root.mkdir(exist_ok=True)
        from ocr_workbench.screen_schema import migrate_screen_schema
        with store.transaction() as db:
            migrate_screen_schema(db)
        if hold_wal:
            self._hold_wal()

    def _hold_wal(self):
        # Keep the WAL open without holding a transaction. FULL commits remain
        # durable; repeated short connections no longer checkpoint every write.
        if self.wal_keeper is None:
            self.wal_keeper = sqlite3.connect(self.store.root / 'workbench.sqlite3', check_same_thread=False)
            self.wal_keeper.execute('PRAGMA journal_mode=WAL').fetchall()
            self.wal_keeper.execute('SELECT count(*) FROM projects').fetchall()

    def defaults(self):
        path = Path(__file__).resolve().parents[2] / 'config/screen-policy.json'
        return json.loads(path.read_text('utf-8')) if path.is_file() else {}

    def start(self):
        if self.thread and self.thread.is_alive():
            return
        self.stopping.clear()
        self._hold_wal()
        self.thread = threading.Thread(target=self._run, name='screen-batch-coordinator', daemon=True)
        self.thread.start()

    def stop(self):
        self.stopping.set()
        self.wake.set()
        if self.thread:
            self.thread.join(timeout=15)
        with self.store.lock:
            if self.wal_keeper is not None:
                self.wal_keeper.close()
                self.wal_keeper = None

    def _run(self):
        while not self.stopping.is_set():
            try:
                worked = self.step()
                self.worker_error = None
            except Exception as error:
                import traceback
                self.worker_error = str(error)
                with (self.root / 'coordinator-errors.log').open('a', encoding='utf-8') as out:
                    out.write(traceback.format_exc() + '\n')
                worked = False
            if not worked:
                self.wake.wait(.35)
                self.wake.clear()

    def _batch(self, project_id, batch_id):
        records = self.store.rows('SELECT * FROM screen_batches WHERE id=?', (batch_id,))
        if not records:
            raise KeyError('批次不存在')
        row = records[0]
        if row['project_id'] != project_id:
            raise KeyError('批次不属于当前项目')
        row['policy'] = json.loads(row['policy'])
        return row

    def list(self, project_id):
        self.store.one('projects', project_id)
        summaries = []
        for row in self.store.rows('SELECT id FROM screen_batches WHERE project_id=? ORDER BY created DESC LIMIT 100', (project_id,)):
            summary = self.get(project_id, row['id'])
            summaries.append({key: value for key, value in summary.items() if key not in {'items', 'task_ids'}})
        return {'batches': summaries}

    def submit(self, project_id, body, *, parent_batch_id=None):
        if not isinstance(body, dict) or set(body) - {'version_ids', 'request_id', 'policy', 'export_options'}:
            raise ValueError('无效的屏拍批次请求')
        versions, request_id = body.get('version_ids'), body.get('request_id')
        if (not isinstance(versions, list) or not 1 <= len(versions) <= 1000
                or any(not isinstance(x, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,81}', x) for x in versions)
                or len(set(versions)) != len(versions)):
            raise ValueError('请选择 1–1000 个不同的图片版本')
        if not isinstance(request_id, str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,120}', request_id):
            raise ValueError('批次缺少有效请求编号')
        if body.get('export_options') is not None:
            raise ValueError('请在下载时明确选择导出选项；创建批次不接受导出参数')
        self.store.one('projects', project_id)
        with self.guard, self.store.file_lock, self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            previous = db.execute('SELECT * FROM screen_batches WHERE project_id=? AND request_id=?', (project_id, request_id)).fetchone()
            # Recover the original policy after defaults change for future batches.
            saved_policy = json.loads(previous['policy']) if previous else None
            policy = validate_policy(body.get('policy'), saved_policy if previous else self.defaults(),
                                     resolve_precision=not bool(previous))
            if previous:
                for key in ('request_timeout_seconds', 'upgrade_trigger', 'require_human_review',
                            'engine_route','technical_fallback_reserve_seconds'):
                    if key not in saved_policy and key not in (body.get('policy') or {}):
                        policy.pop(key)
            signature = fingerprint({'version_ids': versions, 'policy': policy, 'export_options': None, 'parent_batch_id': parent_batch_id})
            if previous:
                if previous['payload_hash'] != signature:
                    raise Conflict('请求编号已绑定不同的批次范围或策略')
                batch_id = previous['id']
            else:
                if (policy.get('engine_route')=='business_language_v1'
                        and policy['batch_budget_seconds']<len(versions)*policy['budget_seconds']):
                    raise ValueError('逐图精度路线的整批原额度须覆盖每图完整额度，不能挤占其它图片的兜底预留')
                batch_id = uid()
                db.execute('INSERT INTO screen_batches(id,project_id,request_id,payload_hash,policy,status,created,updated,parent_batch_id,deadline) VALUES(?,?,?,?,?,?,?,?,?,?)',
                           (batch_id, project_id, request_id, signature, encoded(policy), 'running', now(), now(), parent_batch_id, time.time() + policy['batch_budget_seconds']))
                for position, key in enumerate(versions):
                    version = db.execute('SELECT v.*,i.project_id,i.name,i.active_version FROM versions v JOIN images i ON i.id=v.image_id WHERE v.id=?', (key,)).fetchone()
                    if not version or version['project_id'] != project_id or version['active_version'] != key:
                        raise Conflict('图片范围或当前版本已变化，请重新选择')
                    db.execute('''INSERT INTO screen_items(batch_id,input_version_id,image_id,position,source_hash,source_name,status,stage,updated)
                                  VALUES(?,?,?,?,?,?,?,?,?)''',
                               (batch_id, key, version['image_id'], position, version['sha256'], version['name'], 'queued', 'preparing', now()))
        self.wake.set()
        return self.get(project_id, batch_id)

    def lookup(self, project_id, request_id):
        self.store.one('projects', project_id)
        rows = self.store.rows('SELECT id FROM screen_batches WHERE project_id=? AND request_id=?', (project_id, request_id))
        return {'submitted': bool(rows), 'batch_id': rows[0]['id'] if rows else None}

    def _items(self, batch_id):
        return self.store.rows('SELECT * FROM screen_items WHERE batch_id=? ORDER BY position', (batch_id,))

    def _update(self, batch_id, version_id, **changes):
        permitted = {'status', 'stage', 'preparation', 'candidates', 'issues', 'pass_index',
                     'result_id', 'result_revision', 'review_state', 'error', 'started', 'attempt', 'evidence'}
        if set(changes) - permitted:
            raise ValueError('无效内部批次状态')
        changes['updated'] = now()
        with self.store.transaction() as db:
            db.execute('UPDATE screen_items SET ' + ','.join(key + '=?' for key in changes)
                       + ' WHERE batch_id=? AND input_version_id=?', (*changes.values(), batch_id, version_id))

    def _prepare_item(self, batch, item):
        from ocr_workbench.screen_imaging import prepare_screen_image
        prepare = self.prepare_override or prepare_screen_image
        with self.guard, self.store.file_lock:
            current = self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?',
                                      (batch['id'], item['input_version_id']))[0]
            if (current['status'] != 'queued' or current['attempt'] != item['attempt']
                    or self._batch(batch['project_id'], batch['id'])['status'] != 'running'):
                return None
            # An existing precision request can retain a historical candidate
            # for its frozen input. A change during this computation must not
            # be mistaken for that original request's unchanged input state.
            active_version_at_start = self.store.one('images', item['image_id'])['active_version']
        source = self.store.one('versions', item['input_version_id'])
        actual = hashlib.sha256(self.store.file(source['path']).read_bytes()).hexdigest()
        if actual != item['source_hash']:
            raise Conflict('输入图像内容已变化')
        folder = self.root / batch['id'] / item['input_version_id'] / 'prepared'
        folder.mkdir(parents=True, exist_ok=True)
        evidence = json.loads(item['evidence']) if item.get('evidence') else None
        tick = time.perf_counter()
        info = prepare(self.store.file(source['path']), folder, evidence=evidence, perspective=batch['policy']['perspective'])
        info['preparation_seconds'] = time.perf_counter() - tick
        info.setdefault('warnings', []).extend((evidence or {}).get('warnings', []))
        if any(w.get('code') == 'context_failed' for w in info['warnings']):
            info['quality']['flags'].append({'code': 'context_failed', 'severity': 'high', 'message': '布局或文字证据不可用，请核对业务范围及界面排除'})
        for prepared in info['prepared_images']:
            path = Path(prepared['path']).resolve(strict=True)
            if not path.is_relative_to(folder.resolve()):
                raise ValueError('预处理图像路径越界')
            prepared['path'] = str(path)
        with self.guard, self.store.file_lock:
            current = self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?',
                                      (batch['id'], item['input_version_id']))[0]
            if current['attempt'] != item['attempt'] or current['status'] not in {'queued', 'paused', 'cancelled'}:
                return None
            image = self.store.one('images', item['image_id'])
            if (image['active_version'] != active_version_at_start
                    or hashlib.sha256(self.store.file(source['path']).read_bytes()).hexdigest() != item['source_hash']):
                raise Conflict('预处理期间原图版本或内容已变化')
            # Retain successful preparation for a paused item, so continuing
            # does not recompute it. Never undo a pause/cancel with late work.
            changes = {'preparation': encoded(info), 'started': current['started'] or now()}
            if current['status'] == 'queued':
                changes['status'] = 'prepared'
            self._update(batch['id'], item['input_version_id'], **changes)
            if current['status'] != 'queued':
                return None
            return self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?',
                                   (batch['id'], item['input_version_id']))[0]

    def _contexts(self, batch, items):
        """Complete the resident layout/text stage before switching to OCR."""
        from ocr_workbench.screen_evidence import normalize_screen_context
        worked = False
        with self.guard, self.store.file_lock:
            for item in items:
                if item['status'] != 'queued' or item.get('evidence'):
                    continue
                old=self.store.rows('SELECT * FROM screen_contexts WHERE batch_id=? AND input_version_id=?',(batch['id'],item['input_version_id']))
                if old:
                    continue
                key = uid()
                sources=self.store.rows('''SELECT c.task_id FROM screen_contexts c
                    JOIN tasks t ON t.id=c.task_id JOIN screen_items i
                    ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                    WHERE c.batch_id=? AND i.source_hash=? AND t.kind='screen_context'
                    AND t.status IN ('queued','running','succeeded') ORDER BY i.position LIMIT 1''',
                    (batch['id'], item['source_hash']))
                source=sources[0] if sources else None
                allowance=0
                if not source:
                    budget=self._execution_budget(batch,item)
                    from ocr_workbench.screen_engine_route import reserve_seconds
                    allowance=int(max(0,budget['remaining_seconds']-reserve_seconds(batch['policy'])))
                    if allowance<1:
                        if reserve_seconds(batch['policy']) and budget['remaining_seconds']>=1:
                            self._update(batch['id'],item['input_version_id'],evidence=encoded({
                                'layout':[],'text':[],'provenance':{'image_version':item['input_version_id'],
                                    'contains_reference_input':False},
                                'warnings':[{'code':'context_budget_reserved','message':'来源剩余额度保留给识别；未调用定位模型。'}]}))
                            worked=True
                            continue
                        self._budget_stop(batch,item,budget)
                        continue
                pending_raw=encoded({'cache_pending_source_task_id':source['task_id'],
                    'cache_key':fingerprint([item['source_hash'],self.model_identity,'context']),
                    'budget_reservation_seconds':0}) if source else encoded({'budget_reservation_seconds':allowance})
                with self.store.transaction() as db:
                    db.execute('''INSERT INTO tasks(id,project_id,image_id,version_id,input_version_id,engine,batch,status,phase,kind,created)
                              VALUES(?,?,?,?,?,?,?,?,?,?,?)''', (key, batch['project_id'], item['image_id'], item['input_version_id'],
                    item['input_version_id'], 'geometry', batch['id'] + '-stage-0-context', 'queued',
                    '等待相同图片定位结果' if source else '定位业务表区',
                    'screen_context_cached' if source else 'screen_context', now()))
                    db.execute('INSERT INTO screen_contexts(task_id,batch_id,input_version_id,raw) VALUES(?,?,?,?)',
                        (key,batch['id'],item['input_version_id'],pending_raw))
                worked = True
        worked = self._resolve_context_reuse(batch, items) or worked
        rows = self.store.rows('''SELECT c.*,t.status,t.error FROM screen_contexts c JOIN tasks t ON t.id=c.task_id WHERE c.batch_id=?''', (batch['id'],))
        by_version = {item['input_version_id']: item for item in items}
        pending=any(item['status']=='queued' and not item.get('evidence') and
            not self.store.rows('SELECT 1 FROM screen_contexts WHERE batch_id=? AND input_version_id=?',
                (batch['id'],item['input_version_id'])) for item in self._items(batch['id'])
            if item['input_version_id'] in by_version)
        for row in rows:
            item = by_version.get(row['input_version_id'])
            if item is None:
                continue
            if item['status'] != 'queued' or item.get('evidence'):
                continue
            if row['status'] == 'succeeded' and row['raw']:
                context = json.loads(row['raw'])
                evidence = normalize_screen_context(context)
                if context.get('cache_hit'):
                    evidence['provenance']['cache_hit'] = context['cache_hit']
                self._update(batch['id'], row['input_version_id'], evidence=encoded(evidence))
                worked = True
            elif row['status'] in {'failed', 'cancelled'}:
                self._update(batch['id'], row['input_version_id'], evidence=encoded({'layout': [], 'text': [],
                    'warnings': [{'code': 'context_failed', 'message': row['error'] or row['status']}]}))
                worked = True
            elif row['status'] in {'interrupted', 'paused'}:
                self._update(batch['id'], row['input_version_id'], status='paused', stage='interrupted')
                worked = True
            else:
                pending = True
        self.queue.wake.set()
        return worked, pending

    def _resolve_context_reuse(self, batch, items):
        """Reuse only exact original bytes; every target keeps its own binding."""
        worked = False
        by_version = {item['input_version_id']: item for item in items}
        waiters = self.store.rows('''SELECT c.*,t.status FROM screen_contexts c JOIN tasks t ON t.id=c.task_id
            WHERE c.batch_id=? AND t.kind='screen_context_cached' AND t.status='queued' ORDER BY t.created,t.id''', (batch['id'],))
        for row in waiters:
            if row['input_version_id'] not in by_version:
                continue
            item = by_version[row['input_version_id']]
            pending = json.loads(row['raw'] or '{}')
            source_id = pending.get('cache_pending_source_task_id')
            sources = self.store.rows('''SELECT c.raw,c.artifact,c.input_version_id,t.status,t.id,i.source_hash
                FROM screen_contexts c JOIN tasks t ON t.id=c.task_id JOIN screen_items i
                ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                WHERE c.task_id=? AND c.batch_id=?''', (source_id, batch['id']))
            source = sources[0] if sources else None
            context = json.loads(source['raw']) if source and source['status']=='succeeded' and source['raw'] else None
            if source and source['status'] in {'queued','running'}:
                continue
            if (source and source['source_hash']==item['source_hash'] and context
                    and context.get('status')=='success'):
                version = self.store.one('versions', item['input_version_id'])
                if hashlib.sha256(self.store.file(version['path']).read_bytes()).hexdigest()!=item['source_hash']:
                    raise Conflict('重复图片的输入内容已变化')
                raw = deepcopy(context)
                raw.update(image_version=item['input_version_id'], inference_seconds=0, load_seconds=0,
                    screen_context_input={'image_version':item['input_version_id'],
                        'input_path':str(self.store.file(version['path'])), 'source_sha256':item['source_hash'],
                        'coordinate_space':'exif_normalized_original'},
                    cache_hit={'source_task_id':source['id'], 'source_image_version':source['input_version_id'],
                        'source_artifact':source['artifact'], 'source_raw_sha256':fingerprint(context),
                        'key':pending['cache_key'], 'reuse_scope':'same_batch_exact_input'})
                from ocr_workbench.atomic_files import write_json
                artifact = self.root/batch['id']/item['input_version_id']/'context-reuse.json'
                artifact.parent.mkdir(parents=True,exist_ok=True)
                write_json(artifact,raw,durable=True)
                with self.store.transaction() as db:
                    if db.execute("UPDATE tasks SET status='succeeded',phase='复用相同图片定位结果',finished=? WHERE id=? AND status='queued'",
                                  (now(),row['task_id'])).rowcount:
                        db.execute('UPDATE screen_contexts SET raw=?,artifact=? WHERE task_id=?',
                            (encoded(raw),str(artifact.relative_to(self.store.root)),row['task_id']))
                        worked = True
            else:
                # A failed/cancelled/individually paused leader must not strand
                # an unrelated active input or manufacture successful evidence.
                with self.store.transaction() as db:
                    replacement = db.execute('''SELECT c.task_id FROM screen_contexts c JOIN tasks t ON t.id=c.task_id
                        JOIN screen_items i ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                        WHERE c.batch_id=? AND i.source_hash=? AND t.kind='screen_context'
                        AND t.status IN ('queued','running') AND t.id!=? ORDER BY i.position LIMIT 1''',
                        (batch['id'],item['source_hash'],row['task_id'])).fetchone()
                    if replacement:
                        pending['cache_pending_source_task_id']=replacement['task_id']
                        db.execute('UPDATE screen_contexts SET raw=? WHERE task_id=?',(encoded(pending),row['task_id']))
                    else:
                        budget=self._execution_budget(batch,item)
                        from ocr_workbench.screen_engine_route import reserve_seconds
                        allowance=int(max(0,budget['remaining_seconds']-reserve_seconds(batch['policy'])))
                        if allowance<1:
                            if budget['potential_remaining_seconds']<1 or reserve_seconds(batch['policy']) and budget['remaining_seconds']>=1:
                                db.execute("UPDATE tasks SET status='failed',phase='定位额度不可用',error='剩余额度保留给有界识别或来源累计预算已用尽',finished=? WHERE id=? AND status='queued'",(now(),row['task_id']))
                                worked=True
                            continue
                        db.execute("UPDATE tasks SET kind='screen_context',phase='定位业务表区' WHERE id=? AND status='queued'",(row['task_id'],))
                        db.execute('UPDATE screen_contexts SET raw=?,artifact=NULL WHERE task_id=?',
                            (encoded({'budget_reservation_seconds':allowance}),row['task_id']))
                    worked = True
        return worked

    def _resolve_ocr_reuse(self, batch):
        """Resolve exact prepared-pixel dependencies without sharing source maps."""
        worked = False
        waiters = self.store.rows('''SELECT s.*,t.status,t.engine FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? AND t.kind='screen_cached' AND t.status IN ('queued','running') ORDER BY t.created,t.id''',(batch['id'],))
        for row in waiters:
            metadata = json.loads(row['metadata'])
            source_id = metadata.get('cache_pending_source_task_id')
            sources = self.store.rows('''SELECT t.id,t.status,t.result_id,t.input_version_id,s.metadata FROM tasks t
                JOIN screen_tasks s ON s.task_id=t.id WHERE t.id=?''',(source_id,))
            source = sources[0] if sources else None
            if source and source['status'] in {'queued','running'}:
                continue
            if (source and source['status']=='succeeded' and source['result_id']
                    and json.loads(source['metadata']).get('cache_key')==metadata.get('cache_key')):
                if (metadata.get('prepared_pixel_sha256') and hashlib.sha256(
                        self.store.file(row['prepared_path']).read_bytes()).hexdigest()!=metadata['prepared_pixel_sha256']):
                    with self.store.transaction() as db:
                        db.execute("UPDATE tasks SET status='failed',phase='输入已变化',error='重复复用的预处理像素已变化',finished=? WHERE id=? AND status IN ('queued','running')",
                                   (now(),row['task_id']))
                    worked = True
                    continue
                raw = deepcopy(self.store.result(source['result_id'])['original'])
                version = self.store.one('versions',row['input_version_id'])
                raw.update(project_image_version=row['input_version_id'], screen_input=metadata,
                    version_operations=version['operations'], elapsed_seconds=0,
                    cache_hit={'source_task_id':source['id'],'source_result_id':source['result_id'],
                        'source_image_version':source['input_version_id'],'key':metadata['cache_key']})
                with self.store.transaction() as db:
                    changed = db.execute("UPDATE tasks SET status='running',phase='复用相同图片结果' WHERE id=? AND status IN ('queued','running')",
                                         (row['task_id'],)).rowcount
                if changed:
                    worked = self.store.complete(row['task_id'],raw) or worked
            else:
                with self.store.transaction() as db:
                    replacement = db.execute('''SELECT s.task_id FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
                        WHERE s.batch_id=? AND json_extract(s.metadata,'$.cache_key')=?
                        AND t.kind!='screen_cached' AND t.status IN ('queued','running') AND t.id!=?
                        ORDER BY t.created,t.id LIMIT 1''',(batch['id'],metadata.get('cache_key'),row['task_id'])).fetchone()
                    if replacement:
                        metadata['cache_pending_source_task_id']=replacement['task_id']
                    else:
                        item=next(item for item in self._items(batch['id']) if item['input_version_id']==row['input_version_id'])
                        budget=self._execution_budget(batch,item)
                        reserve=metadata.get('technical_fallback_reserved_seconds',0)
                        allowance=int(max(0,budget['remaining_seconds']-reserve))
                        if allowance<1:
                            if budget['potential_remaining_seconds']<1 or reserve and budget['remaining_seconds']>=1:
                                phase = '首轮额度不可用' if reserve else '来源链预算已用尽'
                                error = '首轮可用额度已用尽，剩余额度保留给有界技术兜底' if reserve else '相同原图累计处理预算已用尽'
                                db.execute("UPDATE tasks SET status='failed',phase=?,error=?,finished=? WHERE id=? AND status IN ('queued','running')",(phase,error,now(),row['task_id']))
                                worked=True
                            continue
                        metadata.pop('cache_pending_source_task_id',None)
                        metadata.update(timeout_seconds=allowance,budget_reservation_seconds=allowance,
                            request_timeout_seconds=min(allowance,batch['policy'].get('request_timeout_seconds',900)))
                        db.execute("UPDATE tasks SET kind='ocr',status='queued',phase='屏拍表格识别' WHERE id=? AND status IN ('queued','running')",(row['task_id'],))
                    db.execute('UPDATE screen_tasks SET metadata=? WHERE task_id=?',(encoded(metadata),row['task_id']))
                    worked = True
        if worked:
            self.queue.wake.set()
        return worked

    def _execution_budget(self, batch, item, *, task_id=None):
        from ocr_workbench.screen_resolution import execution_budget
        return execution_budget(self,batch['project_id'],batch['id'],item['input_version_id'],task_id=task_id)

    def _budget_stop(self,batch,item,budget,*,waiting_status=None):
        waiting=budget['potential_remaining_seconds']>=1
        changes={'stage':'waiting_budget' if waiting else 'source_budget_reached',
            'error':'来源链剩余额度已由在途任务预留，等待其结束后再处理。' if waiting else
                    '此原图或来源链累计处理预算已用尽，保留已有结果及复核状态，不再发起模型任务。'}
        if not waiting and budget.get('wall_deadline_remaining_seconds')==0:
            changes.update(stage='batch_budget_reached',error='整批墙钟时间预算已到期，保留已有结果及复核状态，不再发起模型任务。')
        if waiting and waiting_status:changes['status']=waiting_status
        if not waiting:changes['status']='needs_review' if item['result_id'] else 'failed'
        self._update(batch['id'],item['input_version_id'],**changes)
        return not waiting

    def _enqueue(self, batch, items, engine, pass_index):
        from ocr_workbench.screen_imaging import select_screen_first_pass
        from ocr_workbench.screen_engine_route import reserve_seconds
        with self.guard, self.store.file_lock:
            for item in items:
                choices = self._engine_choices(batch, item)
                if not 0 <= pass_index < len(choices) or choices[pass_index] != engine:
                    raise ValueError('派发模型与已冻结的逐图路线不一致')
                info = json.loads(item['preparation'])
                chosen = select_screen_first_pass(info)
                if not chosen:
                    raise ValueError('预处理没有提供有效识别图像')
                planned=[]
                for image in chosen:
                    prior=self.store.rows('SELECT task_id FROM screen_tasks WHERE batch_id=? AND input_version_id=? AND pass_index=? AND prepared_id=? AND attempt=?',
                        (batch['id'],item['input_version_id'],pass_index,image['id'],item['attempt']))
                    if prior:continue
                    pixel_hash = hashlib.sha256(Path(image['path']).read_bytes()).hexdigest()
                    key=fingerprint([pixel_hash,engine,self.model_identity,batch['policy']['strategy_version']])
                    hits=self.store.rows('SELECT r.task_id FROM screen_cache c JOIN results r ON r.id=c.result_id WHERE c.cache_key=?',(key,))
                    if not hits:
                        hits=self.store.rows('''SELECT s.task_id FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
                            WHERE s.batch_id=? AND json_extract(s.metadata,'$.cache_key')=?
                            AND t.kind!='screen_cached' AND t.status IN ('queued','running','succeeded')
                            ORDER BY t.created,t.id LIMIT 1''',(batch['id'],key))
                    planned.append((image,pixel_hash,key,hits[0]['task_id'] if hits else None))
                computing=sum(hit is None for _,_,_,hit in planned)
                allowance=0
                reserve=reserve_seconds(batch['policy']) if pass_index==0 and len(choices)>1 else 0
                if computing:
                    budget=self._execution_budget(batch,item)
                    allowance=int(max(0,budget['remaining_seconds']-reserve)//computing)
                    if allowance<1:
                        if reserve and budget['remaining_seconds']>=1:
                            info['engine_route_primary_skip'] = {'reason': 'primary_envelope_exhausted',
                                'original_remaining_seconds': budget['remaining_seconds'],
                                'fallback_reserve_seconds': reserve,
                                'original_budget': deepcopy(budget), 'model_invoked': False}
                            self._update(batch['id'],item['input_version_id'],status='awaiting_upgrade',
                                stage='primary_budget_reserved',pass_index=0,
                                preparation=encoded(info),
                                error='独立上下文已耗用首轮可用额度；首轮识别未调用，剩余额度留给有界技术兜底。')
                            continue
                        self._budget_stop(batch,item,budget,waiting_status='prepared')
                        continue
                # Commit each input's frozen allocation before sizing the next
                # input. Independent connections must see sibling reservations.
                with self.store.transaction() as db:
                    db.execute('BEGIN IMMEDIATE')
                    for image,pixel_hash,key,hit in planned:
                        task_id=self.store._enqueue(db,batch['project_id'],[item['input_version_id']],[engine],
                            request_id='screen-'+fingerprint([batch['id'],item['input_version_id'],pass_index,image['id'],item['attempt']])[:64])[0]
                        rank = ('0-' if engine == 'hunyuan' else '1-') if batch['policy'].get('engine_route') == 'business_language_v1' else ''
                        db.execute('UPDATE tasks SET batch=?,phase=? WHERE id=?',
                            (batch['id']+'-stage-'+str(pass_index+1)+'-'+rank+engine,'屏拍表格识别',task_id))
                        frozen=0 if hit else allowance
                        metadata = {**image, 'input_version_id': item['input_version_id'], 'batch_id': batch['id'],
                            'prepared_pixel_sha256':pixel_hash,'strategy_version':batch['policy']['strategy_version'],
                            'model_identity':self.model_identity,
                            'timeout_seconds':frozen,'budget_reservation_seconds':frozen,
                            'request_timeout_seconds':min(frozen,batch['policy'].get('request_timeout_seconds',900)),
                            'attempt':item['attempt'],'cache_key':key}
                        if info.get('engine_route_snapshot'):
                            metadata['engine_route_snapshot']=deepcopy(info['engine_route_snapshot'])
                            metadata['technical_fallback_reserved_seconds']=reserve
                        if hit:metadata['cache_pending_source_task_id']=hit
                        path=str(Path(image['path']).relative_to(self.store.root))
                        db.execute('INSERT INTO screen_tasks(task_id,batch_id,input_version_id,pass_index,prepared_id,prepared_path,metadata,attempt) VALUES(?,?,?,?,?,?,?,?)',
                            (task_id,batch['id'],item['input_version_id'],pass_index,image['id'],path,encoded(metadata),item['attempt']))
                        if hit:db.execute("UPDATE tasks SET kind='screen_cached',phase='等待相同图片结果' WHERE id=?",(task_id,))
                    db.execute('UPDATE screen_items SET status=?,stage=?,pass_index=?,updated=? WHERE batch_id=? AND input_version_id=?',
                        ('running','first_pass' if pass_index==0 else 'upgrading',pass_index,now(),batch['id'],item['input_version_id']))
        self._resolve_ocr_reuse(batch)
        self.queue.wake.set()

    def _engine_choices(self,batch,item):
        if batch['policy'].get('engine_route')!='business_language_v1':
            return [batch['policy']['primary_engine'],*batch['policy']['fallback_engines']]
        from ocr_workbench.screen_engine_route import frozen_engine_route
        info=json.loads(item['preparation'])
        evidence=json.loads(item['evidence']) if item.get('evidence') else {}
        route=frozen_engine_route(info,evidence,item['source_hash'],item['input_version_id'])
        if info.get('engine_route_snapshot') is None:
            info['engine_route_snapshot']=route
            item['preparation']=encoded(info)
            self._update(batch['id'],item['input_version_id'],preparation=item['preparation'])
        return route['engines']

    def _route_binding_stop(self, batch, item, error):
        tasks = self.store.rows('''SELECT t.id FROM tasks t JOIN screen_tasks s ON s.task_id=t.id
            WHERE s.batch_id=? AND s.input_version_id=? AND t.status NOT IN ('succeeded','failed','cancelled')''',
            (batch['id'], item['input_version_id']))
        if tasks:
            self.queue.action(batch['project_id'], 'cancel', [task['id'] for task in tasks])
        self._update(batch['id'], item['input_version_id'],
            status='needs_review' if item['result_id'] else 'failed',
            stage='route_binding_failed', error=str(error)[:500])

    def _candidate(self, batch, item, rows):
        from ocr_workbench.screen_quality import map_screen_result_to_source
        tables, texts, blocks, sources = [], [], [], []
        warnings, finishes = [], []
        truncated = False
        preparation = json.loads(item['preparation'])
        evidence = json.loads(item['evidence']) if item.get('evidence') else {}
        elapsed = 0.0
        for row in rows:
            result = self.store.result(row['result_id'])
            raw = deepcopy(result['original'])
            metadata = json.loads(row['metadata'])
            if metadata.get('cache_key') and not raw.get('cache_hit'):
                with self.store.transaction() as db:
                    db.execute('INSERT OR IGNORE INTO screen_cache VALUES(?,?,?)', (metadata['cache_key'], result['id'], now()))
            raw['tables'] = deepcopy(result['edited'].get('tables', []))
            raw = map_screen_result_to_source(raw, metadata)
            tables.extend(raw.get('tables', []))
            texts.append(result['edited'].get('text', ''))
            blocks.extend(raw.get('blocks', []))
            elapsed += float(raw.get('elapsed_seconds', 0) or 0)
            warnings.extend(raw.get('warnings', []))
            finishes.append(raw.get('finish_reason'))
            truncated = truncated or bool(raw.get('truncated')) or raw.get('finish_reason') in {'length', 'max_tokens', 'token_limit'}
            sources.append({'result_id': result['id'], 'revision': result['revision'], 'engine': row['engine'],
                            'text_sha256':hashlib.sha256(result['edited'].get('text','').encode('utf-8')).hexdigest(),
                            'table_caption':raw['tables'][0].get('caption','') if len(raw.get('tables',[]))==1 else None,
                            'task_id': row['task_id'], 'prepared_id': row['prepared_id'],
                            'transform': metadata.get('prepared_to_source'), 'prepared_metadata': metadata,
                            'warnings': raw.get('warnings', []), 'finish_reason': raw.get('finish_reason'),
                            'cache_hit': deepcopy(raw.get('cache_hit'))})
        candidate = {'id': fingerprint([x['result_id'] for x in sources]), 'engine': rows[0]['engine'],
                'raw': {'status': 'success', 'engine': rows[0]['engine'], 'text': '\n\n'.join(texts), 'tables': tables,
                        'blocks': blocks, 'warnings': warnings, 'truncated': truncated,
                        'finish_reason': 'length' if truncated else next((v for v in finishes if v), None),
                        'elapsed_seconds': elapsed, 'project_image_version': item['input_version_id'],
                        'origin': 'screen_batch', 'screen': {'version': VERSION, 'batch_id': batch['id'],
                        'input_sha256': item['source_hash'], 'strategy_version': batch['policy']['strategy_version'],
                        'sources': sources, 'prepared': preparation,
                        'context_provenance': deepcopy(evidence.get('provenance'))}}}
        original_pixels = [image for image in preparation.get('prepared_images',[]) if image.get('kind')=='original'
                           and image.get('source_to_prepared') == [1,0,0,0,1,0,0,0,1]
                           and image.get('prepared_to_source') == [1,0,0,0,1,0,0,0,1]]
        if len(original_pixels)==1:
            source_path = Path(original_pixels[0]['path'])
            candidate['raw']['screen']['context_pixel_source'] = {'path':str(source_path),
                'sha256':hashlib.sha256(source_path.read_bytes()).hexdigest(),
                'image_version':item['input_version_id'],'input_sha256':item['source_hash']}
        from ocr_workbench.screen_evidence import apply_screen_business_metadata
        candidate['raw'] = apply_screen_business_metadata(candidate['raw'], preparation, evidence)
        if batch['policy'].get('render_math_text', False) and all(source['revision'] == 0 for source in sources):
            from ocr_workbench.screen_text_rendering import render_screen_table_text
            candidate['raw'] = render_screen_table_text(candidate['raw'])
        return candidate

    def _final_result(self, batch, item, raw):
        content = fingerprint(raw)
        with self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            previous = db.execute('SELECT task_id FROM screen_final_tasks WHERE batch_id=? AND input_version_id=? AND content_hash=?',
                                  (batch['id'], item['input_version_id'], content)).fetchone()
            if previous:
                task_id = previous['task_id']
                task = db.execute('SELECT * FROM tasks WHERE id=?', (task_id,)).fetchone()
                if task['result_id']:
                    return task['result_id']
                db.execute("UPDATE tasks SET status='running',error=NULL WHERE id=?", (task_id,))
            else:
                task_id = uid()
                db.execute('''INSERT INTO tasks(id,project_id,image_id,version_id,input_version_id,engine,batch,status,phase,kind,created)
                              VALUES(?,?,?,?,?,?,?,?,?,?,?)''',
                           (task_id, batch['project_id'], item['image_id'], item['input_version_id'], item['input_version_id'],
                            raw.get('engine', batch['policy']['primary_engine']), batch['id'], 'running', '保存最终结果', 'screen_final', now()))
                db.execute('INSERT INTO screen_final_tasks VALUES(?,?,?,?)', (batch['id'], item['input_version_id'], content, task_id))
        if not self.store.complete(task_id, raw):
            raise Conflict('最终结果保存被中断，请恢复批次')
        result_id = self.store.one('tasks', task_id)['result_id']
        # Preserve manual selections/edits. Replace only untouched intermediate results of this batch.
        with self.store.transaction() as db:
            selected = db.execute('SELECT r.id,r.revision,r.task_id FROM selections s JOIN results r ON r.id=s.result_id WHERE s.image_id=?', (item['image_id'],)).fetchone()
            owned = selected and (db.execute('SELECT 1 FROM screen_tasks WHERE task_id=? AND batch_id=?', (selected['task_id'], batch['id'])).fetchone()
                                  or db.execute('SELECT 1 FROM screen_final_tasks WHERE task_id=? AND batch_id=?', (selected['task_id'], batch['id'])).fetchone())
            if owned and selected['revision'] == 0 and item['review_state'] != 'accepted':
                db.execute('UPDATE selections SET result_id=? WHERE image_id=?', (result_id, item['image_id']))
        return result_id

    def _collect(self, batch, item):
        from ocr_workbench.screen_quality import select_screen_candidates
        rows = self.store.rows('''SELECT s.*,t.status,t.engine,t.result_id,t.error FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? AND s.input_version_id=? AND s.pass_index=? AND s.attempt=? ORDER BY s.prepared_id''',
                               (batch['id'], item['input_version_id'], item['pass_index'], item['attempt']))
        if not rows or any(row['status'] not in {'succeeded', 'failed', 'cancelled', 'interrupted', 'paused'} for row in rows):
            return False
        if any(row['status'] in {'paused', 'interrupted'} for row in rows):
            self._update(batch['id'], item['input_version_id'], status='paused', stage='interrupted')
            return True
        candidates = json.loads(item['candidates'])
        spent = screen_seconds_spent(self.store, batch['id'], item['input_version_id'], item['attempt'])
        spent += json.loads(item['preparation']).get('preparation_seconds', 0) if item['preparation'] and item['attempt'] == 0 else 0
        successful = [row for row in rows if row['status'] == 'succeeded' and row['result_id']]
        if len(successful) == len(rows):
            candidate = self._candidate(batch, item, successful)
            if not any(c['id'] == candidate['id'] for c in candidates):
                candidates.append(candidate)
        if not candidates:
            if item['pass_index']+1 < len(self._engine_choices(batch,item)) and item['pass_index'] < batch['policy']['max_upgrades']:
                budget=self._execution_budget(batch,item)
                if budget['remaining_seconds']>=1:
                    self._update(batch['id'],item['input_version_id'],status='awaiting_upgrade',stage='technical_retry',
                        error='；'.join(str(r.get('error') or r['status']) for r in rows)[:500])
                elif budget['potential_remaining_seconds']>=1:
                    self._update(batch['id'],item['input_version_id'],status='awaiting_upgrade',stage='waiting_budget',
                        error='来源链剩余额度已由在途任务预留，等待其结束后再处理。')
                else:self._budget_stop(batch,item,budget)
            else:
                self._update(batch['id'], item['input_version_id'], status='failed', stage='failed',
                             error='；'.join(str(r.get('error') or r['status']) for r in rows)[:500])
            return True
        select = self.select_override or select_screen_candidates
        evidence = json.loads(item['evidence']) if item.get('evidence') else None
        selected = select(candidates, json.loads(item['preparation']), text_evidence=(evidence or {}).get('text'), allow_text_fixes=batch['policy']['allow_text_fixes'])
        raw = selected['selected_result']
        raw.setdefault('screen', {})['selection'] = {key: selected.get(key) for key in
            ('selected_id', 'conflicts', 'modifications', 'automatic_delivery', 'review_required', 'upgrade_kind')}
        table_checks_passed = bool(selected.get('automatic_delivery', False))
        metadata_items = raw.get('screen',{}).get('business_metadata',[])
        metadata_issues = [{'code':'business_metadata_review_required','category':'business_metadata',
            'severity':'medium','scope':'metadata','item_index':index,'kind':part.get('kind'),
            'selected_text':part.get('text'),'alternative_text':part.get('alternative_text'),
            'selected_source':part.get('source'),'alternative_source':part.get('alternative_source'),
            'primary_source':deepcopy(part.get('primary_source')),
            'auxiliary_source':deepcopy(part.get('transcription_evidence',{}).get('auxiliary_source')),
            'human_confirmed':False,
            'reason':part.get('review_reason') or part.get('classification_basis') or 'ambiguous_business_context',
            'message':('业务说明待核对；当前候选（'+str(part.get('source') or '来源未确定')+'）：'+str(part.get('text') or '')
                + ('；另一候选（'+str(part.get('alternative_source') or '来源未确定')+'）：'+str(part['alternative_text'])
                   if part.get('alternative_text') is not None else '')+'。初始候选选择不代表人工确认。')}
            for index,part in enumerate(metadata_items) if part.get('review_required')]
        metadata_pending = bool(metadata_issues) and item['review_state'] != 'accepted'
        checks_passed = table_checks_passed and not metadata_pending
        human_required = (batch['policy'].get('require_human_review', False) or metadata_pending) and item['review_state'] != 'accepted'
        raw['screen']['selection'].update(automatic_checks_passed=checks_passed,
            table_checks_passed=table_checks_passed,business_metadata_review_required=metadata_pending,
            human_review_required=human_required)
        if metadata_pending:
            raw['screen']['selection'].update(automatic_delivery=False,review_required=True)
        chosen_assessment = selected.get('selected_assessment') or next((a for a in selected.get('assessments', []) if a['id'] == selected.get('selected_id')), {})
        issues = list(chosen_assessment.get('issues', []))
        if selected.get('conflicts'):
            issues.extend(selected['conflicts'])
        issues.extend(metadata_issues)
        result_id = self._final_result(batch, item, raw)
        legacy_accuracy_pass = ('upgrade_trigger' not in batch['policy']
            and batch['policy']['mode'] == 'accuracy' and item['pass_index'] == 0)
        needs_upgrade = (batch['policy'].get('upgrade_trigger', 'quality') == 'quality'
            and (not table_checks_passed or legacy_accuracy_pass))
        can_upgrade = (item['pass_index'] < batch['policy']['max_upgrades']
                       and item['pass_index']+1 < len(self._engine_choices(batch,item))
                       and spent < batch['policy']['budget_seconds'])
        if needs_upgrade and can_upgrade:
            state, stage = 'awaiting_upgrade', selected.get('upgrade_kind') or 'accuracy_check'
        else:
            state = 'completed' if checks_passed and not human_required else 'needs_review'
            stage = 'completed' if state == 'completed' else 'quality_review' if not checks_passed else 'human_confirmation'
            if not checks_passed:
                issues.append({'code': 'quality_review_required', 'category': 'quality',
                    'message': '自动质量检查存在问题，请对照原图复核；该状态不代表技术失败'})
            if human_required:
                issues.append({'code': 'human_confirmation_required', 'category': 'human_confirmation',
                    'message': '候选 Excel 尚未经人工确认；自动检查通过不等于已复核通过'})
            if needs_upgrade and not can_upgrade:
                issues.append({'code': 'upgrade_budget_reached', 'message': '已达到本图升级预算，保留最佳结果待核对'})
        self._update(batch['id'], item['input_version_id'], candidates=encoded(candidates), issues=encoded(issues),
                     result_id=result_id, result_revision=self.store.result(result_id)['revision'], status=state, stage=stage,
                     error=None)
        return True

    def step(self):
        with self.step_guard:
            batches = self.store.rows("SELECT id,project_id FROM screen_batches WHERE status='running' ORDER BY created")
            worked = False
            for record in batches:
                worked = self._step_batch(record['project_id'], record['id']) or worked
            return worked

    def _expire_batch(self, batch, items):
        if not batch.get('deadline') or time.time() < batch['deadline']:
            return False
        self.action(batch['project_id'], batch['id'], {'action': 'cancel'})
        for item in items:
            if item['status'] not in TERMINAL:
                self._update(batch['id'], item['input_version_id'], status='needs_review' if item['result_id'] else 'failed',
                             stage='batch_budget_reached', error='已达到整批时间预算，成果已保留')
        self._reconcile(batch)
        return True

    def _step_batch(self, project_id, batch_id):
        worked = False
        with self.guard, self.store.file_lock:
            batch = self._batch(project_id, batch_id)
            if batch['status'] != 'running':
                return False
            items = self._items(batch_id)
            if self._expire_batch(batch, items):
                return True
            if batch['policy'].get('engine_route') == 'business_language_v1':
                for item in items:
                    if item['status'] in TERMINAL or not item.get('preparation'):
                        continue
                    try:
                        self._engine_choices(batch, item)
                    except (ValueError, OSError, TypeError, KeyError) as error:
                        self._route_binding_stop(batch, item, error)
                        worked = True
                items = self._items(batch_id)
            worked = self._resolve_ocr_reuse(batch) or worked
            grouped_context = batch['policy'].get('context_schedule') == 'grouped' and batch['policy']['layout_evidence']
            resident_group = grouped_context and any(item['status'] in {'prepared', 'running'} for item in items)
            if batch['policy']['layout_evidence']:
                context_items = ([] if resident_group else
                    [item for item in items if item['status'] == 'queued'][:batch['policy']['group_size']]) if grouped_context else items
                changed, pending = self._contexts(batch, context_items) if context_items else (False, False)
                worked = worked or changed
                if pending:
                    return worked
                items = self._items(batch_id)
            todo = [] if resident_group else [item for item in items if item['status'] == 'queued'][:batch['policy']['group_size']]
            eligible = []
            for item in todo:
                budget = self._execution_budget(batch, item)
                if budget['remaining_seconds'] >= 1:
                    eligible.append(item)
                else:
                    worked = self._budget_stop(batch, item, budget) or worked
        # CPU work may take minutes for a group. Controls and status readers
        # operate on persisted state while these futures run; each late result
        # checks its attempt, input and pause/cancel state before committing.
        if eligible:
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures = [(item, pool.submit(self._prepare_item, batch, item)) for item in eligible]
                for item, future in futures:
                    try:
                        future.result()
                    except Exception as error:
                        with self.guard, self.store.file_lock:
                            current = next(row for row in self._items(batch_id) if row['input_version_id'] == item['input_version_id'])
                            if current['status'] == 'queued' and current['attempt'] == item['attempt']:
                                self._update(batch_id, item['input_version_id'], status='failed',
                                             stage='preparation_failed', error=str(error)[:500])
            worked = True
        with self.guard, self.store.file_lock:
            batch = self._batch(project_id, batch_id)
            if batch['status'] != 'running':
                return worked
            current = self._items(batch_id)
            if self._expire_batch(batch, current):
                return True
            # Prepared items also survive a crash between preparation/enqueue.
            prepared = [item for item in current if item['status'] == 'prepared']
            if prepared:
                first_groups={}
                for item in prepared:
                    try:
                        first_groups.setdefault(self._engine_choices(batch,item)[0],[]).append(item)
                    except (ValueError, OSError, TypeError, KeyError) as error:
                        self._route_binding_stop(batch, item, error)
                for engine,group in sorted(first_groups.items(),key=lambda pair:pair[0]!='hunyuan'):
                    self._enqueue(batch,group,engine,0)
                worked = True
            for item in self._items(batch_id):
                if item['status'] == 'running':
                    worked = self._collect(batch, item) or worked
            current = self._items(batch_id)
            # Finish the resident engine group before scheduling another model.
            active = any(i['status'] in {'queued', 'prepared', 'running'} for i in current)
            if not active:
                grouped = {}
                for item in [i for i in current if i['status'] == 'awaiting_upgrade']:
                    next_pass = item['pass_index'] + 1
                    choices = self._engine_choices(batch,item)
                    if next_pass >= len(choices):
                        self._update(batch_id, item['input_version_id'], status='needs_review' if item['result_id'] else 'failed', stage='review')
                        continue
                    grouped.setdefault((choices[next_pass], next_pass), []).append(item)
                for (engine, index), group in grouped.items():
                    self._enqueue(batch, group, engine, index)
                    worked = True
            self._reconcile(batch)
        return worked

    def _reconcile(self, batch):
        states = Counter(item['status'] for item in self._items(batch['id']))
        if states and all(state in TERMINAL for state in states):
            status = ('succeeded' if states['completed'] == sum(states.values()) else
                      'cancelled' if states['cancelled'] == sum(states.values()) else
                      'failed' if states['failed'] == sum(states.values()) else 'partial')
            with self.store.transaction() as db:
                db.execute('UPDATE screen_batches SET status=?,updated=? WHERE id=?', (status, now(), batch['id']))

    def get(self, project_id, batch_id, *, offset=0, limit=30, status=None):
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 100:
            raise ValueError('批次分页范围无效')
        batch = self._batch(project_id, batch_id)
        items = self._items(batch_id)
        stale = False
        active_versions = {row['id']: row['active_version'] for row in self.store.rows('SELECT id,active_version FROM images WHERE project_id=?', (project_id,))}
        for item in items:
            if active_versions.get(item['image_id']) != item['input_version_id'] and item['status'] in {'completed', 'needs_review'}:
                stale = True
                item.update(status='needs_review', stage='input_changed',
                    issues=encoded([*json.loads(item['issues']), {'code': 'input_version_changed', 'severity': 'high', 'message': '原图版本已变化，旧结果不能交付'}]))
        counts = Counter(item['status'] for item in items)
        coverage = {key: counts[key] for key in ('completed', 'needs_review', 'failed', 'cancelled')}
        coverage.update(total=len(items), pending=len(items) - sum(coverage.values()))
        if status is not None and status not in TERMINAL | {'pending', 'queued', 'running', 'paused', 'awaiting_upgrade', 'prepared'}:
            raise ValueError('未知图片状态筛选')
        filtered = [item for item in items if not status or item['status'] == status or status == 'pending' and item['status'] not in TERMINAL]
        page = [self._display(item) for item in filtered[offset:offset + limit]]
        from ocr_workbench.screen_eta import estimate_screen_eta
        timed_tasks = self.store.rows('''SELECT t.id,t.kind,t.engine,t.status,t.started,t.finished,t.created,s.pass_index,s.spent_seconds
            FROM screen_tasks s JOIN tasks t ON t.id=s.task_id WHERE s.batch_id=?
            UNION ALL SELECT t.id,t.kind,t.engine,t.status,t.started,t.finished,t.created,-1,c.spent_seconds
            FROM screen_contexts c JOIN tasks t ON t.id=c.task_id WHERE c.batch_id=?''',(batch_id,batch_id))
        eta = estimate_screen_eta(timed_tasks,batch_status=batch['status'],pending_images=coverage['pending'])
        delivery = self.delivery_snapshot(project_id, batch_id)
        return {'batch_id': batch_id, 'project_id': project_id, 'status': 'partial' if stale and batch['status'] in {'succeeded', 'partial'} else batch['status'],
                'delivery_snapshot_id': delivery['snapshot_id'], 'delivery_fresh': delivery['fresh'],
                'stale_item_count': len(delivery['stale_items']),
                'parent_batch_id': batch.get('parent_batch_id'),
                'counts': coverage, 'policy': batch['policy'], 'items': page, 'total_items': len(filtered), 'eta':eta,
                'next_offset': offset + limit if offset + limit < len(filtered) else None,
                'created': batch['created'], 'updated': batch['updated'], 'worker_error': self.worker_error,
                'task_ids': [r['task_id'] for r in self.store.rows('SELECT task_id FROM screen_tasks WHERE batch_id=?', (batch_id,))][:100]}

    @staticmethod
    def _display(item, *, detail=False):
        value = {key: item[key] for key in ('input_version_id', 'image_id', 'source_name', 'status', 'stage',
                 'pass_index', 'attempt', 'result_id', 'result_revision', 'review_state', 'error', 'updated')}
        value['issues'] = json.loads(item['issues'])
        value['candidates'] = [{'id': c['id'], 'engine': c['engine'], 'table_count': len(c['raw'].get('tables', []))}
                               for c in json.loads(item['candidates'])]
        value['human_confirmed'] = item['review_state'] == 'accepted'
        if detail:
            value['preparation'] = json.loads(item['preparation']) if item['preparation'] else None
        return value

    def detail(self, project_id, batch_id, version_id):
        batch = self._batch(project_id, batch_id)
        rows = self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?', (batch_id, version_id))
        if not rows:
            raise KeyError('批次没有该图片')
        value = self._display(rows[0], detail=True)
        value['result'] = self.store.result(rows[0]['result_id']) if rows[0]['result_id'] else None
        value['candidate_results'] = json.loads(rows[0]['candidates'])
        if value['result']:
            with self.store.transaction() as db:
                context = self._review_context(db, rows[0])
            value['review_binding_token'] = fingerprint(context)
            value['business_metadata'] = self._current_business_metadata(value['result'])
            value['business_metadata_source'] = ('edited' if 'business_metadata' in
                (value['result']['edited'].get('screen') or {}) else 'original')
            value['current_result_revision'] = value['result']['revision']
            value['human_confirmed'] = (rows[0]['review_state'] == 'accepted'
                and context['human_review_status'] == 'confirmed'
                and context['review_result_id'] == value['result']['id']
                and context['review_revision'] == value['result']['revision']
                and context['review_version_id'] == version_id and context['active_version'] == version_id
                and context['selected_result_id'] in {None, value['result']['id']})
            value['review_stale'] = (value['result']['revision'] != rows[0]['result_revision']
                or rows[0]['review_state'] == 'accepted' and not value['human_confirmed']
                or context['active_version'] != version_id
                or context['selected_result_id'] not in {None, value['result']['id']})
            if not value['human_confirmed']:
                for part in value['business_metadata']:
                    if part.get('human_confirmed'):
                        part.update(human_confirmed=False, review_required=True, review_reason='human_review_binding_changed')
            from ocr_workbench.screen_diagnostics import build_screen_diagnostics
            from ocr_workbench.screen_resolution import precision_capability
            result = value['result']
            raw = {**deepcopy(result['original']), **deepcopy(result['edited'])}
            raw['screen'] = {**deepcopy(result['original'].get('screen') or {}),
                             **deepcopy(result['edited'].get('screen') or {})}
            evidence = json.loads(rows[0]['evidence']) if rows[0].get('evidence') else {}
            budget = self._execution_budget(batch, rows[0])
            remaining = budget['remaining_seconds']
            value['diagnostics'] = build_screen_diagnostics(raw, value['preparation'], evidence,
                remaining_seconds=remaining, capabilities={'structure_compare': len(value['candidate_results']) > 1,
                    'precision_subbatch': precision_capability(self)['available'], 'local_reread': False})
            value['diagnostics']['result_id'] = result['id']
            value['diagnostics']['result_revision'] = result['revision']
            value['diagnostics']['input_version_id'] = version_id
            value['diagnostics']['execution_budget'] = {
                'basis': budget['basis'], 'root_batch_id': budget['root_batch_id'],
                'remaining_seconds': remaining, 'actual_seconds': budget['actual_seconds'],
                'inflight_reserved_seconds': budget['inflight_reserved_seconds'],
                'local_input_actual_seconds': budget['local_input_actual_seconds'],
                'source_actual_seconds': budget['per_source_actual_seconds'].get(rows[0]['source_hash'], 0),
                'wall_deadline_remaining_seconds': budget['wall_deadline_remaining_seconds']}
        value['image_url'] = '/api/versions/' + version_id + '/image'
        from ocr_workbench.screen_cell_review import review_summary
        value['cell_review'] = review_summary(self, batch_id, version_id)
        return value

    def resolve(self, project_id, batch_id, body, *, allow_new=True):
        from ocr_workbench.screen_resolution import resolve_screen_issues
        return resolve_screen_issues(self, project_id, batch_id, body, allow_new=allow_new)

    def compare(self, project_id, batch_id, version_id, body):
        from ocr_workbench.screen_diagnostics import compare_screen_candidates
        with self.guard, self.store.file_lock:
            value = self.detail(project_id, batch_id, version_id)
            result = value.get('result')
            if (not result or body.get('result_id') != result['id']
                    or type(body.get('expected_revision')) is not int
                    or body['expected_revision'] != result['revision']):
                raise Conflict('当前结果版本已变化，请重新读取候选')
            candidate = next((item for item in value['candidate_results'] if item['id'] == body.get('candidate_id')), None)
            if candidate is None:
                raise ValueError('候选不属于当前图片或批次')
            row = next(item for item in self._items(batch_id) if item['input_version_id'] == version_id)
            evidence = json.loads(row['evidence']) if row.get('evidence') else {}
            current = {**deepcopy(result['original']), **deepcopy(result['edited'])}
            current['screen'] = {**deepcopy(result['original'].get('screen') or {}),
                                 **deepcopy(result['edited'].get('screen') or {})}
            report = compare_screen_candidates(current,
                candidate['raw'], value['preparation'], evidence, current_id=result['id'], alternative_id=candidate['id'],
                remaining_seconds=value['diagnostics']['budget']['remaining_seconds'],
                capabilities={'structure_compare': True, 'precision_subbatch': None, 'local_reread': False})
            return {**report, 'input_version_id': version_id, 'result_id': result['id'],
                    'result_revision': result['revision']}

    def action(self, project_id, batch_id, body, *, allow_new=True):
        if not isinstance(body, dict):
            raise ValueError('批次操作无效')
        action = body.get('action')
        if action not in {'pause', 'resume', 'cancel', 'retry_failed', 'review', 'upgrade'}:
            raise ValueError('未知批次操作')
        with self.guard, self.store.file_lock:
            batch = self._batch(project_id, batch_id)
            request_id = body.get('request_id')
            if request_id is not None:
                if not isinstance(request_id, str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,120}', request_id):
                    raise ValueError('操作请求编号无效')
                old = self.store.rows('SELECT * FROM screen_actions WHERE batch_id=? AND request_id=?', (batch_id, request_id))
                if old:
                    if old[0]['payload_hash'] != fingerprint(body):
                        raise Conflict('操作编号已用于不同的参数')
                    return json.loads(old[0]['receipt'])
            items = self._items(batch_id)
            selected = body.get('version_ids')
            if selected is not None:
                if not isinstance(selected, list) or not selected or len(set(selected)) != len(selected) or not set(selected) <= {i['input_version_id'] for i in items}:
                    raise ValueError('操作范围不属于该批次')
                items = [i for i in items if i['input_version_id'] in selected]
            if action == 'upgrade':
                if selected is None or not request_id:
                    raise ValueError('提高精度必须指定图片和操作请求编号')
                if len(items) > 50:
                    raise ValueError('一次精度处理最多选择 50 张图片')
                resolution_id = 'screen-' + fingerprint([batch_id, request_id, body])[:64]
                resolution_rows = self.store.rows('SELECT receipt FROM screen_actions WHERE batch_id=? AND request_id=?',
                                                 (batch_id, resolution_id))
                if not allow_new:
                    if not resolution_rows:
                        return None
                    frozen = json.loads(resolution_rows[0]['receipt'])
                    targets = [{key: binding[key] for key in ('input_version_id', 'result_id', 'result_revision', 'issue_ids')}
                               for binding in frozen['bindings']]
                    return self.resolve(project_id, batch_id, {'request_id': resolution_id, 'action': 'precision_subbatch',
                        'budget_seconds': batch['policy']['batch_budget_seconds'], 'targets': targets}, allow_new=False)
                if any(i['status'] not in TERMINAL for i in items):
                    raise Conflict('请等待指定图片的当前尝试结束')
                targets = []
                for item in items:
                    detail = self.detail(project_id, batch_id, item['input_version_id'])
                    result = detail.get('result')
                    issues = detail.get('diagnostics', {}).get('items', [])
                    if not result or not issues:
                        receipt = {'resolution_version': 'screen-resolution-v1', 'status': 'blocked', 'action': 'upgrade',
                            'request_id': request_id, 'parent_batch_id': batch_id, 'batch_id': None,
                            'reason': 'no_current_result_or_diagnosed_issue', 'input_version_id': item['input_version_id'],
                            'values_modified': False, 'selection_applied': False, 'human_review': 'not_performed'}
                        break
                    targets.append({'input_version_id': item['input_version_id'], 'result_id': result['id'],
                        'result_revision': result['revision'], 'issue_ids': [issue['id'] for issue in issues[:50]]})
                else:
                    receipt = self.resolve(project_id, batch_id, {'request_id': resolution_id, 'action': 'precision_subbatch',
                        'budget_seconds': batch['policy']['batch_budget_seconds'], 'targets': targets})
                    if receipt.get('attempt'):
                        receipt = {**receipt, 'policy': receipt['attempt']['policy']}
                with self.store.transaction() as db:
                    db.execute('INSERT INTO screen_actions VALUES(?,?,?,?)', (batch_id, request_id, fingerprint(body), encoded(receipt)))
                return receipt
            if not allow_new:
                return None
            if action == 'review':
                if body.get('metadata_edits') and (selected is None or len(items)!=1):
                    raise ValueError('编辑业务说明必须明确指定单张图片')
                if len(items) != 1:
                    raise ValueError('单次复核必须指定一张图片')
                self._review(batch, items[0], body)
            elif action == 'retry_failed':
                retried=False
                for item in items:
                    if item['status'] != 'failed':
                        continue
                    budget=self._execution_budget(batch,item)
                    if budget['remaining_seconds']<1:
                        self._budget_stop(batch,item,budget)
                        continue
                    # A technical retry gets a new pass identity, preserving every failed task.
                    self._update(batch_id, item['input_version_id'], status='prepared' if item['preparation'] else 'queued',
                                 stage='retry', pass_index=0, attempt=item['attempt'] + 1, error=None)
                    retried=True
                if retried:
                    with self.store.transaction() as db:
                        db.execute("UPDATE screen_batches SET status='running',updated=? WHERE id=?", (now(), batch_id))
            else:
                active = [i for i in items if i['status'] not in TERMINAL]
                # Read the batch once and commit item controls together. Per-item
                # connections/FULL commits made cancelling 240 queued inputs
                # exceed an HTTP timeout before the durable receipt was saved.
                task_rows = self.store.rows('SELECT input_version_id,task_id,attempt FROM screen_tasks WHERE batch_id=?', (batch_id,))
                context_rows = self.store.rows('SELECT input_version_id,task_id FROM screen_contexts WHERE batch_id=?', (batch_id,))
                active_versions = {i['input_version_id'] for i in active}
                task_ids = [r['task_id'] for r in [*task_rows, *context_rows] if r['input_version_id'] in active_versions]
                if task_ids:
                    self.queue.action(project_id, action, task_ids)
                existing_attempts = {(r['input_version_id'], r['attempt']) for r in task_rows}
                updated = now()
                with self.store.transaction() as db:
                    db.executemany('UPDATE screen_items SET status=?,stage=?,updated=? WHERE batch_id=? AND input_version_id=?', [
                        ('cancelled' if action == 'cancel' else 'paused' if action == 'pause' else
                         'running' if (item['input_version_id'], item['attempt']) in existing_attempts else
                         'prepared' if item['preparation'] else 'queued',
                         'cancelled' if action == 'cancel' else 'paused' if action == 'pause' else 'resuming',
                         updated, batch_id, item['input_version_id']) for item in active])
                    if selected is None and action == 'pause':
                        db.execute('UPDATE screen_batches SET paused_at=COALESCE(paused_at,?) WHERE id=?', (time.time(), batch_id))
                    if selected is None and action == 'resume' and batch.get('paused_at'):
                        db.execute('UPDATE screen_batches SET deadline=deadline+?,paused_at=NULL WHERE id=?', (time.time()-batch['paused_at'], batch_id))
                    db.execute('UPDATE screen_batches SET status=?,updated=? WHERE id=?',
                               ('paused' if action == 'pause' and selected is None else 'running', now(), batch_id))
            self._reconcile(batch)
            receipt = self.get(project_id, batch_id)
            if request_id:
                with self.store.transaction() as db:
                    db.execute('INSERT INTO screen_actions VALUES(?,?,?,?)', (batch_id, request_id, fingerprint(body), encoded(receipt)))
            self.wake.set()
            return receipt

    def _review(self, batch, item, body):
        if item['status'] not in {'completed', 'needs_review'}:
            raise Conflict('图片仍在处理或未产生结果，请等待当前尝试结束')
        state = body.get('state')
        if state not in {'accepted', 'retain', 'deferred'}:
            raise ValueError('复核状态无效')
        result_id = body.get('result_id', item['result_id'])
        if not result_id or result_id != item['result_id']:
            raise Conflict('最终候选已变化，请重新加载；局部中间结果不能替代整图结果')
        result = self.store.result(result_id)
        if type(body.get('expected_revision')) is not int or result['revision'] != body['expected_revision']:
            raise Conflict('复核内容已变化，请重新加载')
        image = self.store.one('images', item['image_id'])
        if image['active_version'] != item['input_version_id']:
            raise Conflict('原图版本已变化，不能复核旧批次')
        with self.store.transaction() as db:
            review_context = self._review_context(db, item)
        if (review_context['selected_result_id'] not in {None, result_id}
                or body.get('expected_review_token') is not None
                and body['expected_review_token'] != fingerprint(review_context)):
            raise Conflict('采用结果或人工复核状态已变化，请重新加载')
        candidate_id = body.get('candidate_id')
        metadata_edits = body.get('metadata_edits',[])
        if not isinstance(metadata_edits,list) or len(metadata_edits)>100:
            raise ValueError('业务说明编辑范围无效')
        if candidate_id and metadata_edits:
            raise Conflict('切换整图候选后请先重新加载当前业务说明，再编辑')
        if candidate_id:
            candidate = next((c for c in json.loads(item['candidates']) if c['id'] == candidate_id), None)
            if not candidate:
                raise ValueError('整图候选不存在')
            result_id = self._final_result(batch, item, deepcopy(candidate['raw']))
            result = self.store.result(result_id)
            # _final_result may select the new untouched candidate itself.
            with self.store.transaction() as db:
                switched_context = self._review_context(db, item)
            if (switched_context['selected_result_id'] not in {review_context['selected_result_id'], result_id}
                    or {key: value for key, value in switched_context.items() if key != 'selected_result_id'}
                    != {key: value for key, value in review_context.items() if key != 'selected_result_id'}):
                raise Conflict('切换候选期间人工复核或原图已变化')
            review_context = switched_context
        edits = body.get('edits') or []
        if not isinstance(edits, list) or len(edits) > 100:
            raise ValueError('复核编辑范围无效')
        metadata = self._current_business_metadata(result)
        if edits or metadata_edits or metadata and state == 'accepted':
            value = deepcopy(result['edited'])
            for edit in edits:
                if (not isinstance(edit, dict) or any(type(edit.get(k)) is not int or edit[k] < 0 for k in ('table_index', 'row', 'column'))
                        or edit['table_index'] >= len(value.get('tables', []))):
                    raise ValueError('单元格编辑范围无效')
                table = value['tables'][edit['table_index']]
                cells = [c for c in table['cells'] if (c['row'], c['column']) == (edit['row'], edit['column'])]
                if len(cells) != 1 or cells[0]['text'] != edit['before'] or not isinstance(edit['after'], str):
                    raise Conflict('待修改单元格已变化')
                cells[0]['text'] = edit['after']
            metadata_log = deepcopy((value.get('screen') or {}).get('metadata_edit_log',[]))
            original_metadata = result['original'].get('screen',{}).get('business_metadata',[])
            seen = set()
            for edit in metadata_edits:
                if (not isinstance(edit,dict) or set(edit)-{'item_index','before','after','choice'}
                        or type(edit.get('item_index')) is not int or edit['item_index']<0
                        or edit['item_index']>=len(metadata) or edit['item_index'] in seen
                        or not isinstance(edit.get('before'),str) or not isinstance(edit.get('after'),str)
                        or len(edit['after'])>10000 or any(ord(c)<32 and c not in '\n\r\t' for c in edit['after'])):
                    raise ValueError('业务说明编辑内容或位置无效')
                index=edit['item_index'];seen.add(index);part=metadata[index]
                if part.get('text')!=edit['before']:
                    raise Conflict('待修改业务说明已变化，请重新加载')
                choice=edit.get('choice','manual')
                source_part=original_metadata[index] if index<len(original_metadata) else None
                if not source_part:
                    raise Conflict('业务说明原始证据不存在，不能按旧位置修改')
                evidence=source_part.get('transcription_evidence') or {}
                if (choice=='primary' and not evidence.get('primary_source')
                        or choice=='auxiliary' and 'auxiliary_text' not in evidence):
                    raise Conflict('所选识别候选缺少原始来源证据')
                source_text=(evidence.get('auxiliary_text') if choice=='auxiliary' else
                    evidence.get('primary_text',source_part.get('text')) if choice=='primary' else edit['after'])
                if choice not in {'manual','primary','auxiliary'} or source_text!=edit['after']:
                    raise Conflict('所选识别候选与原始证据不一致')
                metadata_log.append({'item_index':index,'before':edit['before'],'after':edit['after'],
                    'choice':choice,'base_revision':result['revision'],'image_version':item['input_version_id'],
                    'original_selected_text':source_part.get('text'),
                    'original_evidence_refs':deepcopy(source_part.get('evidence_refs',[])),
                    'original_transcription_evidence':deepcopy(evidence),
                    'actor':'human','at':now(),'confirmed_in_this_review':state=='accepted'})
                part.update(text=edit['after'],source='human_edit' if choice=='manual' else 'human_selected_'+choice,
                    selected_text_basis='human_metadata_'+choice,review_required=state!='accepted',
                    review_reason='human_edit_pending_confirmation' if state!='accepted' else None,
                    human_confirmed=state=='accepted',automatic_metadata_selection=False)
                part.update(alternative_text=(evidence.get('primary_text',source_part['text']) if choice!='primary'
                    else evidence.get('auxiliary_text',source_part['text'])),
                    alternative_source=(source_part.get('source') if choice!='primary' else
                        (evidence.get('auxiliary_source') or {}).get('source')))
            if metadata_edits or metadata and state=='accepted':
                if state=='accepted':
                    for part in metadata:
                        part.update(human_confirmed=True,review_required=False,review_reason=None)
                value.setdefault('screen',{}).update(business_metadata=metadata,metadata_edit_log=metadata_log,
                    metadata_review={'state':state,'actor':'human','input_version_id':item['input_version_id'],
                                     'base_revision':result['revision'],'at':now()})
            screen_metadata = value.pop('screen', None)
            result = self.store.save(result_id, value, result['revision'], screen_metadata=screen_metadata)
        issues = json.loads(item['issues'])
        if metadata_edits:
            current_metadata = self._current_business_metadata(result)
            for issue in issues:
                if issue.get('code')=='business_metadata_review_required':
                    index=issue.get('item_index')
                    if type(index) is int and 0<=index<len(current_metadata):
                        part=current_metadata[index]
                        issue.update(selected_text=part['text'],selected_source=part.get('source'),
                            alternative_text=part.get('alternative_text'),alternative_source=part.get('alternative_source'),
                            human_confirmed=part.get('human_confirmed',False),reason=part.get('review_reason'),
                            message='业务说明已人工编辑，当前文字：'+part['text']+'；尚需确认当前版本。')
        if state == 'accepted':
            issues = [issue for issue in issues if issue.get('code') not in {
                'human_confirmation_required', 'quality_review_required'}]
            issues = [issue for issue in issues if issue.get('code') != 'business_metadata_review_required']
        with self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            current = db.execute('SELECT revision FROM results WHERE id=?', (result_id,)).fetchone()
            if (current['revision'] != result['revision']
                    or self._review_context(db, item) != review_context):
                raise Conflict('复核内容、采用结果或人工复核状态已变化')
            db.execute('INSERT INTO selections VALUES(?,?) ON CONFLICT(image_id) DO UPDATE SET result_id=excluded.result_id', (item['image_id'], result_id))
            target = self.store._review_target(db, item['image_id'])
            if (target['result_id'] != result_id or target['revision'] != result['revision']
                    or target['active_version'] != item['input_version_id'] or target['version_id'] != item['input_version_id']):
                raise Conflict('当前候选来源版本已变化')
            stamp = now()
            db.execute('INSERT INTO reviews VALUES(?,?,?,?,?,?) ON CONFLICT(image_id) DO UPDATE SET result_id=excluded.result_id,version_id=excluded.version_id,revision=excluded.revision,status=excluded.status,updated=excluded.updated',
                (item['image_id'], result_id, item['input_version_id'], result['revision'], 'confirmed' if state == 'accepted' else 'question', stamp))
            db.execute('UPDATE screen_items SET result_id=?,result_revision=?,issues=?,review_state=?,status=?,stage=?,updated=? WHERE batch_id=? AND input_version_id=?',
                (result_id, result['revision'], encoded(issues), state, 'completed' if state == 'accepted' else 'needs_review',
                 'completed' if state == 'accepted' else 'review', stamp, batch['id'], item['input_version_id']))

    @staticmethod
    def _review_context(db, item):
        row = db.execute('''SELECT im.active_version,s.result_id selected_result_id,
            rv.result_id review_result_id,rv.version_id review_version_id,rv.revision review_revision,
            rv.status human_review_status,rv.updated review_epoch FROM images im
            LEFT JOIN selections s ON s.image_id=im.id LEFT JOIN reviews rv ON rv.image_id=im.id WHERE im.id=?''',
            (item['image_id'],)).fetchone()
        return dict(row)

    def file(self, project_id, batch_id, version_id, file_id):
        detail = self.detail(project_id, batch_id, version_id)
        for image in (detail['preparation'] or {}).get('prepared_images', []):
            if image['id'] == file_id:
                path = Path(image['path']).resolve(strict=True)
                if not path.is_relative_to((self.root / batch_id / version_id).resolve()):
                    raise ValueError('图像路径越界')
                return path
        raise KeyError('预处理图像不存在')

    def delivery_snapshot(self, project_id, batch_id):
        from ocr_workbench.screen_delivery import delivery_snapshot
        return delivery_snapshot(self, project_id, batch_id)

    def export(self, project_id, batch_id, body=None):
        from ocr_workbench.tables import export_xlsx
        body = body or {}
        if set(body) - {'include_pending_review', 'numeric_columns'} or type(body.get('include_pending_review', False)) is not bool:
            raise ValueError('导出参数无效')
        batch = self._batch(project_id, batch_id)
        numeric = body.get('numeric_columns')
        if numeric is not None and (not isinstance(numeric, dict)
                or not set(numeric) <= {i['input_version_id'] for i in self._items(batch_id)}
                or any(not isinstance(v, dict) for v in numeric.values())):
            raise ValueError('数值列须按本批次图片版本逐项指定')
        include_review = body.get('include_pending_review', False)
        options = {'include_pending_review': include_review,
                   'numeric_columns': body.get('numeric_columns')}
        folder = self.root / batch_id / 'exports' / uid()
        folder.mkdir(parents=True)
        archive_path = folder / 'screen-tables.zip'
        manifest = {'format': VERSION + '-export', 'batch_id': batch_id, 'project_id': project_id,
                    'policy': batch['policy'], 'created': now(), 'items': [], 'success_files': 0}
        with self.guard, self.store.file_lock, zipfile.ZipFile(archive_path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            delivery = self.delivery_snapshot(project_id, batch_id)
            if not delivery['fresh']:
                raise Conflict('图片、采用结果或复核状态已变化，请刷新当前结果后重新导出')
            manifest['snapshot_id'] = delivery['snapshot_id']
            manifest['export_options'] = options
            manifest['version_vector'] = {**{key: delivery[key] for key in
                ('snapshot_id', 'batch_id', 'project_id', 'policy_sha256', 'items')}, 'export_options': options}
            snapshots = []
            for item in self._items(batch_id):
                entry = self._display(item)
                entry['file_name'] = None
                if item['result_id'] and (item['status'] == 'completed' or include_review and item['status'] == 'needs_review'):
                    result = self.store.result(item['result_id'])
                    if result['revision'] != item['result_revision']:
                        raise Conflict('最终结果已编辑，请刷新后重新导出')
                    if self.store.one('images', item['image_id'])['active_version'] != item['input_version_id']:
                        raise Conflict('图片版本已变化，请重新处理或复核')
                    if result['edited'].get('tables'):
                        stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', Path(item['source_name']).stem).strip(' .') or 'image'
                        name = f'{item["position"] + 1:04d}-{stem[:100]}.xlsx'
                        target = folder / name
                        export_xlsx(result['edited']['tables'], target,
                                    source_rows=self._export_source_rows(batch, item, result, name))
                        metadata = self._current_business_metadata(result)
                        if metadata:
                            self._export_metadata(target, metadata)
                        numeric = (body.get('numeric_columns') or {}).get(item['input_version_id'])
                        if numeric:
                            if item['review_state'] != 'accepted':
                                raise Conflict('数值列转换须先确认当前图片结果')
                            self._numeric_columns(target, result['edited']['tables'], numeric)
                        archive.write(target, name)
                        snapshots.append((item['image_id'], item['input_version_id'], result['id'], result['revision']))
                        entry['file_name'] = name
                        manifest['success_files'] += 1
                    else:
                        entry['export_error'] = '没有结构化表格，不生成空白 Excel'
                manifest['items'].append(entry)
            if not manifest['success_files']:
                raise ValueError('尚无可交付表格，可先复核或明确导出待复核结果')
            for image_id, version_id, result_id, revision in snapshots:
                if (self.store.one('images', image_id)['active_version'] != version_id
                        or self.store.result(result_id)['revision'] != revision):
                    raise Conflict('导出期间结果或图片已变化，请重新导出')
            if self.delivery_snapshot(project_id, batch_id)['snapshot_id'] != delivery['snapshot_id']:
                raise Conflict('导出期间采用结果、复核状态或批次范围已变化，请重新导出')
            archive.writestr('manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2))
        (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        return archive_path, manifest

    @staticmethod
    def _export_source_rows(batch, item, result, workbook):
        original = result['original']
        engine_info = original.get('engine_info') or {}
        engine_package = original.get('engine_package') or {}
        metadata = {'workbook': workbook, 'image_name': item['source_name'],
            'image_id': item['image_id'], 'image_version': item['input_version_id'],
            'input_sha256': item['source_hash'], 'engine': original.get('engine', ''),
            'engine_name': engine_info.get('name', ''), 'engine_package': engine_package.get('id', 'builtin'),
            'model_revisions': encoded(original.get('model_revisions', {})),
            'result_id': result['id'], 'revision': result['revision'],
            'origin': original.get('origin', 'screen_batch'),
            'policy_version': batch['policy']['strategy_version'], 'policy_sha256': fingerprint(batch['policy']),
            'review_summary': encoded({'review_state': item['review_state'], 'status': item['status']})}
        return [{**metadata, 'sheet': f'Table {index+1}', 'table_index': index+1,
                 'structure_review': encoded(table.get('structure_review', {})),
                 'header_cells': encoded([[c['row'], c['column'], c['row_span'], c['column_span']]
                                          for c in table['cells'] if c.get('is_header')])}
                for index, table in enumerate(result['edited']['tables'])]

    @staticmethod
    def _current_business_metadata(result):
        screen = result.get('edited',{}).get('screen') or {}
        metadata = screen.get('business_metadata') if 'business_metadata' in screen else (
            result.get('original',{}).get('screen') or {}).get('business_metadata',[])
        if not isinstance(metadata,list) or any(not isinstance(item,dict) or not isinstance(item.get('text'),str) for item in metadata):
            raise Conflict('当前业务说明结构无效，请重新加载')
        metadata = deepcopy(metadata)
        review = screen.get('metadata_review') or {}
        if review.get('state') == 'accepted' and result.get('revision') != review.get('base_revision', -2) + 1:
            for part in metadata:
                part.update(human_confirmed=False, review_required=True, review_reason='metadata_confirmation_revision_changed')
        return metadata

    @staticmethod
    def _export_metadata(path, metadata):
        from openpyxl import load_workbook
        book = load_workbook(path)
        sheet = book.create_sheet('业务说明')
        sheet.append(['类别', '识别原文', '来源', '原图坐标', '核对状态', '另一识别候选（未人工确认）', '处理依据', '另一候选来源'])
        for part in metadata:
            status = ('人工已确认' if part.get('human_confirmed') else
                      '需要人工核对' if part.get('review_required') else '识别候选，未逐项人工确认')
            sheet.append([part['kind'], part['text'], part.get('source', ''), encoded(part.get('box')), status,
                          part.get('alternative_text',''),part.get('review_reason',part.get('selected_text_basis','')),
                          part.get('alternative_source','')])
            sheet.cell(sheet.max_row, 2).data_type = 's'
            sheet.cell(sheet.max_row, 6).data_type = 's'
        sheet.column_dimensions['B'].width = 70
        book.save(path)

    @staticmethod
    def _numeric_columns(path, tables, specification):
        from openpyxl import load_workbook
        from ocr_workbench.financial_checks import parse_amount, IDENTIFIERS
        from ocr_workbench.screen_quality import _field_kind
        if not isinstance(specification, dict) or len(specification) > 100:
            raise ValueError('请明确指定已确认的数值列')
        book = load_workbook(path)
        for table_index, columns in specification.items():
            if not str(table_index).isdecimal() or not 0 <= int(table_index) < len(tables):
                raise ValueError('数值列的表格范围无效')
            table = tables[int(table_index)]
            if not isinstance(columns, list) or any(type(c) is not int or not 0 <= c < table['columns'] for c in columns):
                raise ValueError('数值列无效')
            sheet = book.worksheets[int(table_index)]
            offset = 1 if table.get('caption') else 0
            for cell in table['cells']:
                header = ' '.join(c.get('text', '') for c in table['cells'] if c['column'] == cell['column'] and (c.get('is_header') or c['row'] == 0))
                if cell['column'] not in columns or cell.get('is_header') or _field_kind(table, cell) == 'id' or IDENTIFIERS.search(header):
                    continue
                parsed = parse_amount(cell['text'])
                if parsed['kind'] not in {'amount', 'number', 'percent'}:
                    continue
                amount = parsed.get('value')
                if amount is None or sum(ch.isdigit() for ch in cell['text']) > 15:
                    continue
                target = sheet.cell(cell['row'] + 1 + offset, cell['column'] + 1)
                target.value = float(amount)
                target.number_format = '0' + ('.' + '0' * parsed.get('places', 0) if parsed.get('places', 0) else '') + ('%' if parsed.get('percent') else '')
        book.save(path)
