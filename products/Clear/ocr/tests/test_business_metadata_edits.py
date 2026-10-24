"""Current metadata edit/review/ZIP flow and version races use actual SQLite."""
from copy import deepcopy
import io
from unittest.mock import patch
import unittest
import zipfile
from openpyxl import load_workbook
from ocr_workbench.store import Conflict
import test_business_metadata_delivery as fixture

class MetadataEditTests(unittest.TestCase):
    setUp=fixture.MetadataDeliveryTests.setUp
    tearDown=fixture.MetadataDeliveryTests.tearDown
    execute=fixture.MetadataDeliveryTests.execute

    def prepare_conflict(self):
        auxiliary='表注：本季度调整金额为-20.00万元。'
        primary='表注：本季度调整金额为+20.00万元。'
        state,item,result,tasks=self.execute(auxiliary,primary)
        index=next(n for n,p in enumerate(result['original']['screen']['business_metadata']) if p['kind']=='note' and p['table']==0)
        return state,item,result,index,auxiliary,primary

    def body(self,state,item,result,**extra):
        return {'action':'review','state':state,'version_ids':[item['input_version_id']],
            'result_id':result['id'],'expected_revision':result['revision'],**extra}

    def test_auxiliary_selection_and_table_edit_confirm_current_zip_together(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        original=deepcopy(result['original']);before_snapshot=self.manager.delivery_snapshot(self.project,state['batch_id'])
        body=self.body('accepted',item,result,request_id='metadata-edit-confirm-001',
            metadata_edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}],
            edits=[{'table_index':0,'row':1,'column':1,'before':'123.45','after':'124.00'}])
        receipt=self.manager.action(self.project,state['batch_id'],body)
        revised=self.store.result(result['id'])
        self.assertEqual(revised['revision'],result['revision']+1)
        self.assertEqual(revised['original'],original)
        current=revised['edited']['screen']['business_metadata'][index]
        self.assertEqual(current['text'],aux);self.assertEqual(current['alternative_text'],primary)
        self.assertTrue(current['human_confirmed']);self.assertFalse(current['review_required'])
        self.assertEqual(current['source'],'human_selected_auxiliary')
        self.assertEqual(revised['edited']['screen']['metadata_edit_log'][0]['original_selected_text'],primary)
        self.assertEqual(revised['edited']['tables'][0]['cells'][3]['text'],'124.00')
        self.assertFalse(any(i['code']=='business_metadata_review_required' for i in receipt['items'][0]['issues']))
        detail=self.manager.detail(self.project,state['batch_id'],item['input_version_id'])
        self.assertEqual(detail['business_metadata_source'],'edited');self.assertEqual(detail['business_metadata'][index]['text'],aux)
        self.assertEqual(detail['diagnostics']['result_revision'],revised['revision'])
        self.assertNotEqual(self.manager.delivery_snapshot(self.project,state['batch_id'])['snapshot_id'],before_snapshot['snapshot_id'])
        archive,manifest=self.manager.export(self.project,state['batch_id'])
        self.assertEqual(manifest['success_files'],1)
        with zipfile.ZipFile(archive) as bundle:
            name=next(n for n in bundle.namelist() if n.endswith('.xlsx'))
            book=load_workbook(io.BytesIO(bundle.read(name)))
            self.assertEqual(book['业务说明'].cell(index+2,2).value,aux)
            self.assertEqual(book['业务说明'].cell(index+2,5).value,'人工已确认')
            self.assertEqual(book['业务说明'].cell(index+2,6).value,primary)
            self.assertEqual(book['Table 1']['B3'].value,'124.00')
        again=self.manager.action(self.project,state['batch_id'],body)
        self.assertEqual(again,receipt);self.assertEqual(self.store.result(result['id'])['revision'],revised['revision'])

    def test_manual_edit_remains_pending_then_confirmation_exports_current_text(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        text='表注：审批完成后调整金额为-20.00万元。'
        self.manager.action(self.project,state['batch_id'],self.body('retain',item,result,
            metadata_edits=[{'item_index':index,'before':primary,'after':text,'choice':'manual'}]))
        revised=self.store.result(result['id'])
        part=revised['edited']['screen']['business_metadata'][index]
        self.assertEqual(part['text'],text);self.assertFalse(part['human_confirmed']);self.assertTrue(part['review_required'])
        detail=self.manager.detail(self.project,state['batch_id'],item['input_version_id'])
        self.assertEqual(detail['status'],'needs_review')
        issue=next(i for i in detail['issues'] if i.get('item_index')==index and i['code']=='business_metadata_review_required')
        self.assertEqual(issue['selected_text'],text)
        with self.assertRaises(ValueError):self.manager.export(self.project,state['batch_id'])
        self.manager.action(self.project,state['batch_id'],self.body('accepted',detail,revised))
        final=self.store.result(result['id'])
        self.assertEqual(final['revision'],revised['revision']+1)
        self.assertEqual(final['original'],result['original'])
        self.assertEqual(self.manager.detail(self.project,state['batch_id'],item['input_version_id'])['business_metadata'][index]['text'],text)
        self.assertTrue(self.manager.delivery_snapshot(self.project,state['batch_id'])['fresh'])
        archive,_=self.manager.export(self.project,state['batch_id'])
        with zipfile.ZipFile(archive) as bundle:
            book=load_workbook(io.BytesIO(bundle.read(next(n for n in bundle.namelist() if n.endswith('.xlsx')))))
            self.assertEqual(book['业务说明'].cell(index+2,2).value,text)
            self.assertEqual(book['业务说明'].cell(index+2,5).value,'人工已确认')

    def test_invalid_metadata_batch_does_not_commit_any_table_or_metadata_edit(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        common={'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}
        for edits in (
            [{**common,'before':'stale value'}],
            [{**common,'item_index':True}],
            [{**common,'item_index':999}],
            [{**common,'after':primary+' forged','choice':'auxiliary'}],
            [{**common,'after':'forged','choice':'primary'}],
            [common,common],
            [{**common,'after':'bad\x00text','choice':'manual'}],
            False,
        ):
            with self.subTest(edits=edits):
                with self.assertRaises((ValueError,Conflict)):
                    self.manager.action(self.project,state['batch_id'],self.body('accepted',item,result,
                        edits=[{'table_index':0,'row':1,'column':1,'before':'123.45','after':'124.00'}],metadata_edits=edits))
                unchanged=self.store.result(result['id'])
                self.assertEqual(unchanged['revision'],result['revision']);self.assertEqual(unchanged['edited'],result['edited'])

    def test_stale_revision_candidate_switch_and_unscoped_metadata_edits_rejected(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}]
        body=self.body('accepted',item,result,metadata_edits=edits)
        for changed in (
            {**body,'expected_revision':result['revision']-1},
            {**body,'candidate_id':item['candidates'][0]['id']},
            {k:v for k,v in body.items() if k!='version_ids'},
            {**body,'result_id':'another-final-result'},
        ):
            with self.subTest(body=changed),self.assertRaises((ValueError,Conflict)):
                self.manager.action(self.project,state['batch_id'],changed)
        self.assertEqual(self.store.result(result['id'])['revision'],result['revision'])

    def test_revision_race_refuses_human_confirmation_and_export(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        real_save=self.store.save
        def racing_save(key,value,revision,**kwargs):
            saved=real_save(key,value,revision,**kwargs)
            another=deepcopy(saved['edited']);another['text']+=' concurrent'
            real_save(key,another,saved['revision'])
            return saved
        with patch.object(self.store,'save',side_effect=racing_save),self.assertRaises(Conflict):
            self.manager.action(self.project,state['batch_id'],self.body('accepted',item,result,
                metadata_edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}]))
        self.assertNotEqual(self.manager.get(self.project,state['batch_id'])['items'][0]['review_state'],'accepted')
        with self.assertRaises(Conflict):self.manager.export(self.project,state['batch_id'],{'include_pending_review':True})

    def test_ordinary_table_edits_preserve_metadata_and_cannot_forge_confirmation(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        self.manager.action(self.project,state['batch_id'],self.body('retain',item,result,
            metadata_edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}]))
        revised=self.store.result(result['id']);current=deepcopy(revised['edited']['screen'])
        forged=deepcopy(revised['edited']);forged['screen']['business_metadata'][index]['human_confirmed']=True
        with self.assertRaises(Conflict):self.store.save(result['id'],forged,revised['revision'])
        erased=deepcopy(revised['edited']);erased['screen']['business_metadata']=[]
        with self.assertRaises(Conflict):self.store.save(result['id'],erased,revised['revision'])
        plain={'text':revised['edited']['text']+'\nordinary edit','tables':revised['edited']['tables']}
        saved=self.store.save(result['id'],plain,revised['revision'])
        self.assertEqual(saved['edited']['screen'],current)
        self.assertEqual(saved['original'],result['original'])
        echoed=deepcopy(saved['edited']);echoed['text']+='\necho unchanged metadata'
        saved=self.store.save(result['id'],echoed,saved['revision'])
        self.assertEqual(saved['edited']['screen'],current)
        from ocr_workbench.editing import validate_edit,present_result,export_text,export_markdown
        validate_edit(saved['edited']);present_result(saved)
        self.assertIn('ordinary edit',export_text(saved['edited']))
        self.assertIn('ordinary edit',export_markdown(saved['edited']))
        with self.assertRaises(Conflict):self.manager.export(self.project,state['batch_id'],{'include_pending_review':True})

    def test_metadata_history_and_revision_snapshots_survive_reopening(self):
        from ocr_workbench.store import Store
        state,item,result,index,aux,primary=self.prepare_conflict()
        self.manager.action(self.project,state['batch_id'],self.body('retain',item,result,
            metadata_edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}]))
        revised=self.store.result(result['id'])
        reopened=Store(self.store.root)
        self.assertEqual(reopened.result(result['id'])['edited'],revised['edited'])
        undone=reopened.history(result['id'],-1,revised['revision'])
        self.assertEqual(self.manager._current_business_metadata(undone)[index]['text'],primary)
        redone=reopened.history(result['id'],1,undone['revision'])
        self.assertEqual(redone['edited']['screen'],revised['edited']['screen'])
        self.assertEqual(redone['original'],result['original'])
        old=reopened.result_snapshot(result['id'],revised['revision'])
        self.assertEqual(old['edited']['screen'],revised['edited']['screen'])
        with self.assertRaises(Conflict):self.manager.export(self.project,state['batch_id'],{'include_pending_review':True})

    def test_review_epoch_and_selection_races_do_not_confirm_new_state(self):
        for race in ('review_epoch', 'selected_result'):
            with self.subTest(race=race):
                state,item,result,index,aux,primary=self.prepare_conflict()
                detail=self.manager.detail(self.project,state['batch_id'],item['input_version_id'])
                real_save=self.store.save
                def racing_save(key,value,revision,**kwargs):
                    saved=real_save(key,value,revision,**kwargs)
                    if race=='review_epoch':
                        self.store.set_review(item['image_id'],'question',key,saved['revision'],item['input_version_id'])
                    else:
                        intermediate=self.store.rows('SELECT t.result_id FROM screen_tasks s JOIN tasks t ON t.id=s.task_id WHERE s.batch_id=?',(state['batch_id'],))[0]['result_id']
                        with self.store.transaction() as db:
                            db.execute('UPDATE selections SET result_id=? WHERE image_id=?',(intermediate,item['image_id']))
                    return saved
                with patch.object(self.store,'save',side_effect=racing_save),self.assertRaises(Conflict):
                    self.manager.action(self.project,state['batch_id'],self.body('accepted',item,result,
                        expected_review_token=detail['review_binding_token'],
                        metadata_edits=[{'item_index':index,'before':primary,'after':aux,'choice':'auxiliary'}]))
                self.assertNotEqual(self.manager.get(self.project,state['batch_id'])['items'][0]['review_state'],'accepted')
                with self.assertRaises(Conflict):self.manager.export(self.project,state['batch_id'],{'include_pending_review':True})
                # A new test input is needed: fixture submit id is scoped to a
                # fixed original, so use the clean setup again for the next race.
                if race=='review_epoch':
                    self.tearDown();self.setUp()

    def test_metadata_confirmation_cannot_carry_to_a_later_table_revision(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        self.manager.action(self.project,state['batch_id'],self.body('accepted',item,result))
        accepted=self.store.result(result['id'])
        self.manager.action(self.project,state['batch_id'],self.body('retain',item,accepted,
            edits=[{'table_index':0,'row':1,'column':1,'before':'123.45','after':'124.00'}]))
        detail=self.manager.detail(self.project,state['batch_id'],item['input_version_id'])
        self.assertFalse(detail['business_metadata'][index]['human_confirmed'])
        self.assertTrue(detail['business_metadata'][index]['review_required'])
        archive,_=self.manager.export(self.project,state['batch_id'],{'include_pending_review':True})
        with zipfile.ZipFile(archive) as bundle:
            book=load_workbook(io.BytesIO(bundle.read(next(n for n in bundle.namelist() if n.endswith('.xlsx')))))
            self.assertEqual(book['业务说明'].cell(index+2,5).value,'需要人工核对')

    def test_same_revision_human_review_change_is_visible_in_metadata_detail(self):
        state,item,result,index,aux,primary=self.prepare_conflict()
        self.manager.action(self.project,state['batch_id'],self.body('accepted',item,result))
        current=self.store.result(result['id'])
        self.store.set_review(item['image_id'],'question',result['id'],current['revision'],item['input_version_id'])
        detail=self.manager.detail(self.project,state['batch_id'],item['input_version_id'])
        self.assertTrue(detail['review_stale']);self.assertFalse(detail['human_confirmed'])
        self.assertFalse(detail['business_metadata'][index]['human_confirmed'])
        self.assertTrue(detail['business_metadata'][index]['review_required'])

if __name__=='__main__':unittest.main()
