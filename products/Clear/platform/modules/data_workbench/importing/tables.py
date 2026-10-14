"""Import selected CSV/XLSX tables without inferring away source literals.

The caller must supply an already authorized, verified local file and its
registered SHA256. This library has no route that accepts a browser path.
JSON export carries explicit cell kinds, so zero, empty text, absent values,
and formula results without a cache remain distinct.
"""

from __future__ import annotations

import hashlib
import io
import json
import posixpath
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as StdElementTree

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook


_HASH = re.compile(r"[0-9a-f]{64}\Z")
_CELL = re.compile(r"[A-Z]{1,3}[1-9][0-9]{0,6}\Z")
_MAX_BYTES = 64 * 1024 * 1024
_MAX_ROWS = 100_000
_MAX_COLS = 256
_MAX_CELLS = 1_000_000
_ENCODINGS = frozenset({"utf-8", "utf-8-sig", "gb18030", "utf-16"})
_DELIMITERS = frozenset({",", "\t", ";", "|"})


class TableImportError(ValueError):
    """A fixed failure suitable for an import boundary."""


@dataclass(frozen=True)
class Cell:
    kind: str
    raw: str | None
    value: str | bool | None
    formula: str | None = None
    warning: str | None = None


@dataclass(frozen=True)
class Column:
    column_id: str
    header: str | None


@dataclass(frozen=True)
class Row:
    row_id: str
    source_row: int
    cells: tuple[Cell, ...]


@dataclass(frozen=True)
class ImportedTable:
    source_sha256: str
    revision: str
    format: str
    selector: dict[str, str | int | None]
    columns: tuple[Column, ...]
    rows: tuple[Row, ...]
    warnings: tuple[str, ...]

    def to_dict(self) -> dict:
        return asdict(self)

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _verified(path: str | Path, source_sha256: str, suffix: str) -> bytes:
    if not isinstance(source_sha256, str) or not _HASH.fullmatch(source_sha256):
        raise TableImportError("Invalid source identity")
    source = Path(path)
    try:
        if source.is_symlink() or source.suffix.lower() != suffix or not source.is_file():
            raise TableImportError("Source is unavailable")
        if source.stat().st_size > _MAX_BYTES:
            raise TableImportError("Source exceeds import limit")
        payload = source.read_bytes()
        if len(payload) > _MAX_BYTES:
            raise TableImportError("Source exceeds import limit")
        current = hashlib.sha256(payload).hexdigest()
    except OSError:
        raise TableImportError("Source is unavailable") from None
    if current != source_sha256:
        raise TableImportError("Source content changed")
    return payload


def _identity(source_sha256: str, format: str, selector: dict) -> str:
    encoded = json.dumps({"source": source_sha256, "format": format, "selector": selector},
                         sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _row_id(revision: str, row_number: int) -> str:
    return "r" + hashlib.sha256(f"{revision}:{row_number}".encode("ascii")).hexdigest()[:32]


def _csv_records(text: str, delimiter: str):
    """RFC 4180 style parser that retains whether an empty field was quoted."""
    if text.startswith("\ufeff"):
        text = text[1:]
    fields: list[tuple[str, bool]] = []
    chars: list[str] = []
    state, quoted, index, row_number = "start", False, 0, 1
    while index < len(text):
        char = text[index]
        newline = char in "\r\n"
        width = 2 if char == "\r" and index + 1 < len(text) and text[index + 1] == "\n" else 1
        if state == "quoted":
            if char == '"':
                if index + 1 < len(text) and text[index + 1] == '"':
                    chars.append('"')
                    index += 2
                    continue
                state = "after_quote"
            else:
                chars.append(text[index:index + width] if newline else char)
        elif char == delimiter:
            fields.append(("".join(chars), quoted))
            chars, quoted, state = [], False, "start"
        elif newline:
            fields.append(("".join(chars), quoted))
            yield row_number, tuple(fields)
            row_number += 1
            fields, chars, quoted, state = [], [], False, "start"
        elif state == "after_quote":
            raise TableImportError("Malformed quoted field")
        elif state == "start" and char == '"':
            quoted, state = True, "quoted"
        elif char == '"':
            raise TableImportError("Malformed quoted field")
        else:
            chars.append(char)
            state = "unquoted"
        if len(chars) > 1_000_000:
            raise TableImportError("Cell exceeds import limit")
        index += width
    if state == "quoted":
        raise TableImportError("Unclosed quoted field")
    if fields or chars or state == "after_quote" or text and text[-1] == delimiter:
        fields.append(("".join(chars), quoted))
        yield row_number, tuple(fields)


def _make_columns(count: int, headers: tuple[Cell, ...] | None) -> tuple[Column, ...]:
    if count < 1 or count > _MAX_COLS:
        raise TableImportError("Invalid column count")
    return tuple(Column(f"c{index + 1}", headers[index].raw if headers else None)
                 for index in range(count))


def import_csv(path: str | Path, *, source_sha256: str, encoding: str = "utf-8",
               delimiter: str = ",", header_row: int | None = 1) -> ImportedTable:
    if (encoding not in _ENCODINGS or delimiter not in _DELIMITERS
            or (header_row is not None and (type(header_row) is not int or not 1 <= header_row <= 1000))):
        raise TableImportError("Invalid CSV selection")
    payload = _verified(path, source_sha256, ".csv")
    try:
        text = payload.decode(encoding, errors="strict")
    except UnicodeError:
        raise TableImportError("CSV encoding differs") from None
    selector = {"encoding": encoding, "delimiter": delimiter, "header_row": header_row}
    revision = _identity(source_sha256, "csv", selector)
    rows: list[Row] = []
    columns: tuple[Column, ...] | None = None
    seen = 0
    for row_number, fields in _csv_records(text, delimiter):
        if row_number > _MAX_ROWS + (header_row or 0) or len(fields) > _MAX_COLS:
            raise TableImportError("Table exceeds import limit")
        cells = tuple(Cell("text", value, value) if quoted or value else Cell("null", None, None)
                      for value, quoted in fields)
        if header_row is not None and row_number < header_row:
            continue
        if header_row == row_number:
            columns = _make_columns(len(cells), cells)
            continue
        if columns is None:
            columns = _make_columns(len(cells), None)
        if len(cells) != len(columns):
            raise TableImportError("CSV row width differs")
        seen += len(cells)
        if seen > _MAX_CELLS:
            raise TableImportError("Table exceeds import limit")
        rows.append(Row(_row_id(revision, row_number), row_number, cells))
    if columns is None:
        raise TableImportError("Selected CSV header or data is missing")
    return ImportedTable(source_sha256, revision, "csv", selector, columns, tuple(rows), ())


def _sheet_xml(archive: zipfile.ZipFile, sheet_name: str) -> dict[str, tuple[str | None, str | None, str | None]]:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "p": "http://schemas.openxmlformats.org/package/2006/relationships"}
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    selected = [item for item in workbook.findall("m:sheets/m:sheet", ns)
                if item.get("name") == sheet_name]
    if len(selected) != 1:
        raise TableImportError("Worksheet is unavailable")
    relationship_id = selected[0].get("{" + ns["r"] + "}id")
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = [item.get("Target") for item in relationships.findall("p:Relationship", ns)
               if item.get("Id") == relationship_id]
    if len(targets) != 1 or not targets[0]:
        raise TableImportError("Worksheet is unavailable")
    target = targets[0]
    path = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                              else posixpath.join("xl", target))
    if not path.startswith("xl/worksheets/") or path not in archive.namelist():
        raise TableImportError("Worksheet is unavailable")
    sheet = ElementTree.fromstring(archive.read(path))
    result = {}
    for cell in sheet.findall(".//m:sheetData/m:row/m:c", ns):
        coordinate = cell.get("r", "")
        if not _CELL.fullmatch(coordinate) or coordinate in result:
            raise TableImportError("Invalid worksheet cell")
        value = cell.find("m:v", ns)
        formula = cell.find("m:f", ns)
        result[coordinate] = (value.text if value is not None else None,
                              formula.text if formula is not None else None,
                              cell.get("t"))
    return result


def _numeric(raw: str, *, formula: str | None = None) -> Cell:
    try:
        numeric = Decimal(raw)
        if not numeric.is_finite():
            raise InvalidOperation()
    except (InvalidOperation, TypeError):
        raise TableImportError("Invalid numeric cell") from None
    digits = len(numeric.as_tuple().digits)
    warning = "excel_numeric_precision_unverifiable" if digits > 15 else None
    return Cell("decimal", raw, str(numeric), formula, warning)


def _excel_cell(cell, xml: tuple[str | None, str | None, str | None] | None) -> Cell:
    raw_value, _, xml_type = xml if xml is not None else (None, None, None)
    value = cell.value
    if cell.data_type == "f":
        formula = str(value)
        if raw_value is None or raw_value == "":
            return Cell("unknown_formula", formula, None, formula)
        if xml_type == "b":
            return Cell("boolean", raw_value, raw_value == "1", formula)
        if xml_type in {"str", "inlineStr"}:
            return Cell("text", raw_value, raw_value, formula)
        if xml_type == "e":
            return Cell("error", raw_value, raw_value, formula)
        return _numeric(raw_value, formula=formula)
    if value is None:
        if xml_type in {"inlineStr", "s", "str"} and xml is not None:
            return Cell("text", "", "")
        return Cell("null", None, None)
    if isinstance(value, bool):
        return Cell("boolean", raw_value or ("1" if value else "0"), value)
    if isinstance(value, (datetime, date)):
        return Cell("date", raw_value or value.isoformat(), value.isoformat())
    if isinstance(value, (int, float, Decimal)):
        return _numeric(raw_value or str(value))
    if cell.data_type == "e":
        return Cell("error", str(value), str(value))
    return Cell("text", str(value), str(value))


def import_xlsx(path: str | Path, *, source_sha256: str, sheet_name: str,
                header_row: int | None = 1) -> ImportedTable:
    if (not isinstance(sheet_name, str) or not sheet_name or len(sheet_name) > 31
            or header_row is not None and (type(header_row) is not int or not 1 <= header_row <= 1000)):
        raise TableImportError("Invalid worksheet selection")
    payload = _verified(path, source_sha256, ".xlsx")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            parts = archive.infolist()
            if len(parts) > 10_000 or sum(part.file_size for part in parts) > 256 * 1024 * 1024:
                raise TableImportError("Workbook exceeds import limit")
            xml = _sheet_xml(archive, sheet_name)
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        try:
            if sheet_name not in book.sheetnames:
                raise TableImportError("Worksheet is unavailable")
            sheet = book[sheet_name]
            if (sheet.max_row is None or sheet.max_column is None
                    or sheet.max_row > _MAX_ROWS + (header_row or 0)
                    or sheet.max_column > _MAX_COLS):
                raise TableImportError("Table exceeds import limit")
            selector = {"sheet_name": sheet_name, "header_row": header_row}
            revision = _identity(source_sha256, "xlsx", selector)
            columns: tuple[Column, ...] | None = None
            rows = []
            warnings = set()
            seen = 0
            for row_number, cells_in_row in enumerate(sheet.iter_rows(), start=1):
                if not cells_in_row:
                    continue
                cells = tuple(_excel_cell(cell, xml.get(getattr(cell, "coordinate", "")))
                              for cell in cells_in_row)
                if header_row is not None and row_number < header_row:
                    continue
                if header_row == row_number:
                    columns = _make_columns(len(cells), cells)
                    continue
                if columns is None:
                    columns = _make_columns(len(cells), None)
                if len(cells) != len(columns):
                    raise TableImportError("Worksheet row width differs")
                seen += len(cells)
                if seen > _MAX_CELLS:
                    raise TableImportError("Table exceeds import limit")
                for cell in cells:
                    if cell.warning:
                        warnings.add(cell.warning)
                rows.append(Row(_row_id(revision, row_number), row_number, cells))
            if columns is None:
                raise TableImportError("Selected worksheet header or data is missing")
            return ImportedTable(source_sha256, revision, "xlsx", selector,
                                 columns, tuple(rows), tuple(sorted(warnings)))
        finally:
            book.close()
    except (OSError, zipfile.BadZipFile, KeyError, StdElementTree.ParseError,
            DefusedXmlException):
        raise TableImportError("Workbook is unavailable") from None
