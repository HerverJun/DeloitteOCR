"""Atomic publication, version-pinned readback and native reference failures."""

import hashlib
import io
import json
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from packages.contracts.models import JobRef, SourceLocator
from workbench_platform.artifacts import ArtifactStore, ArtifactUnavailable
import workbench_platform.artifacts.store as artifact_store_module
from test_native_batches import disposition_bridge, disposition_request


EXAMPLES = json.loads((Path(__file__).parents[2] / "packages/contracts/examples/positive.json").read_text(
    encoding="utf-8"))
JOB = JobRef.model_validate(EXAMPLES["JobRef"])
SOURCE = SourceLocator.model_validate(EXAMPLES["SourceLocator"])


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@pytest.fixture
def store(tmp_path):
    blobs = tmp_path / "blobs"
    blobs.mkdir()
    return ArtifactStore(tmp_path / "artifacts.sqlite", blobs)


def publish(store, data=b"published bytes"):
    return store.publish_managed("workspaceA", io.BytesIO(data), expected_hash=sha(data),
                                 expected_size=len(data), job_ref=JOB, source_refs=[SOURCE],
                                 media_type="application/octet-stream",
                                 application_version="1.0.0", rule_or_model_version="rules-1")


def reference(store, data=b"native bytes", resolver=None):
    resolver = resolver or (lambda app, instance, artifact: io.BytesIO(data))
    expiry = (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat().replace("+00:00", "Z")
    return store.register_reference("workspaceA", "nativeArtifactA", resolver,
                                    expected_hash=sha(data), expected_size=len(data),
                                    expires_at=expiry, job_ref=JOB, source_refs=[SOURCE],
                                    media_type="application/json",
                                    application_version="1.0.0", rule_or_model_version="model-1")


def test_managed_publication_and_download_readback_are_exact(store):
    data = b"published bytes" * 100
    artifact = publish(store, data)
    assert artifact.availability == "available" and artifact.blob_hash == sha(data)
    assert artifact.input_manifest_hash == JOB.input_manifest_hash
    assert store.provenance("workspaceA", artifact.artifact_id) == {
        "storage_mode": "managed", "origin_app_id": JOB.app_id,
        "origin_instance_id": JOB.instance_id, "native_artifact_id": None,
        "expires_at": None}
    with store.stage_download("workspaceA", artifact.artifact_id, expected_revision=1,
                              expected_hash=artifact.blob_hash) as path:
        assert path.read_bytes() == data
    assert store.get("workspaceA", artifact.artifact_id) == artifact
    with pytest.raises(ArtifactUnavailable) as wrong:
        with store.stage_download("otherWorkspace", artifact.artifact_id,
                                  expected_revision=1, expected_hash=artifact.blob_hash):
            pytest.fail("Another workspace must not read this artifact")
    assert wrong.value.code == "NOT_FOUND"


def test_artifact_and_review_mapping_registration_roll_back_together(store):
    def fail_mapping(db, artifact):
        raise OSError('mapping transaction failed')
    with pytest.raises(OSError, match='mapping transaction failed'):
        store.publish_managed('workspaceA', io.BytesIO(b'whole'), expected_hash=sha(b'whole'), expected_size=5,
            job_ref=JOB, source_refs=[SOURCE], media_type='text/plain', application_version='1',
            rule_or_model_version='review', on_insert=fail_mapping)
    with sqlite3.connect(store.db_path) as db:
        assert db.execute('SELECT count(*) FROM published_artifact').fetchone()[0] == 0


@pytest.mark.parametrize('decision', ('classification_confirm_A', 'classification_adopt_B', 'ocr_confirm_original', 'ocr_adopt_revision_2'))
def test_review_export_uses_actual_native_adopted_snapshot_not_latest(disposition_bridge, decision):
    import csv
    from contextlib import contextmanager
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from packages.contracts.automation import ProcessingBatch, ReviewBundle
    from workbench_platform.artifacts.center import ArtifactCenter
    fixture = disposition_bridge('ocr' if decision.startswith('ocr') else 'classification')
    service, native = fixture.service, fixture.native
    actor = disposition_request(fixture, 'manual')['actor']
    if decision == 'classification_adopt_B':
        raw = fixture.issues[0]['raw_output']
        native.dispose_results(fixture.job_id, [raw['native_item_id']], 'confirmed', 'B', '', {raw['native_item_id']: 0},
            operation_id='manualB', request_id='manualB', actor=actor, instance_id='instanceA', intent='manual')
        fixed = service.dispositions(fixture.batch_id, disposition_request(fixture, 'adoptB', 'adopt',
            adoption={'native_revision': 1, 'human_label_id': 'B'}))
        expected = 'B'
        native.dispose_results(fixture.job_id, [raw['native_item_id']], 'confirmed', 'A', '', {raw['native_item_id']: 2},
            operation_id='laterA', request_id='laterA', actor=actor, instance_id='instanceA', intent='manual')
    elif decision == 'ocr_adopt_revision_2':
        result_id = fixture.state['result_id']
        native.save(result_id, {'text': 'revision 1', 'tables': []}, 0)
        native.save(result_id, {'text': 'revision 2', 'tables': []}, 1)
        fixed = service.dispositions(fixture.batch_id, disposition_request(fixture, 'adopt2', 'adopt',
            adoption={'native_revision': 2, 'native_result_id': result_id,
                      'native_version_id': fixture.issues[0]['raw_output']['native_version_id']}))
        expected = 'revision 2'
        native.save(result_id, {'text': 'revision 3 later edit', 'tables': []}, 2)
    elif decision == 'ocr_confirm_original':
        fixed = service.dispositions(fixture.batch_id, disposition_request(fixture, 'confirmOriginalOCR'))
        expected = 'original A'
    else:
        fixed = service.dispositions(fixture.batch_id, disposition_request(fixture, 'confirmA'))
        expected = 'A'
    # The native fixture already created a real source PNG. Its platform
    # registry boundary verifies that exact pinned file without registering a
    # second copy or changing the native/platform identities under review.
    @contextmanager
    def pinned(workspace_id, asset_id, *, revision, blob_hash):
        assert (workspace_id, asset_id, revision) == ('workspaceA', 'assetA', 1)
        source = (native.file(native.one('images', fixture.state['image_id'])['original_path'])
                  if decision.startswith('ocr') else fixture.source_path)
        assert sha(source.read_bytes()) == blob_hash == fixture.source_hash
        yield source
    center = ArtifactCenter(service.store.path.parent / 'review-artifacts', SimpleNamespace(stage_verified=pinned), None, None)
    center.batch_store = service.store
    service.app.state.artifact_center = center
    calls = list(fixture.state['calls'])
    payload = {'request_id': 'nativeReview', 'expected_revision': fixed['revision'], 'actor': actor, 'kind': 'final'}
    response = service.review_export(fixture.batch_id, payload)
    ProcessingBatch.model_validate(response)
    bundle = response['artifacts'][-1]
    ReviewBundle.model_validate(bundle)
    assert bundle['status'] == 'current' and fixture.state['calls'] == calls
    def content(name):
        artifact = center.file('workspaceA', bundle['bundle_id'], name)
        with center.store.stage_download('workspaceA', artifact.artifact_id, expected_revision=artifact.revision,
                                         expected_hash=artifact.blob_hash) as path:
            return path.read_bytes()
    rows = list(csv.DictReader(io.StringIO(content('results.csv').decode('utf-8-sig'))))
    assert len(rows) == 1 and rows[0]['adopted_value'] == expected and rows[0]['category'] == 'review'
    original = json.loads(rows[0]['original_json'])
    assert original == fixture.issues[0]['raw_output']
    workbook = load_workbook(io.BytesIO(content('results.xlsx')), read_only=True, data_only=False)
    try:
        assert workbook.active.cell(2, 7).value == expected
        assert all(cell.data_type != 'f' for row in workbook.active.iter_rows() for cell in row)
    finally:
        workbook.close()
    review = json.loads(content('review.json'))
    assert review['history'][0]['actor'] == actor and review['items'][0]['raw_output'] == original
    if decision == 'ocr_adopt_revision_2':
        assert review['items'][0]['adoption']['snapshot']['edited']['text'] == 'revision 2'
        assert review['items'][0]['adoption']['native_revision'] == 2
        assert '原生复核资格独立' in content('review.md').decode()
    assert service.review_export(fixture.batch_id, payload) == response


def test_interrupted_or_mismatched_write_never_publishes_metadata(store):
    class BrokenStream:
        calls = 0

        def read(self, _size):
            self.calls += 1
            if self.calls == 1:
                return b"partial"
            raise OSError("injected interruption")

    with pytest.raises(ArtifactUnavailable) as failure:
        store.publish_managed("workspaceA", BrokenStream(), expected_hash=sha(b"whole"),
                              expected_size=5, job_ref=JOB, source_refs=[SOURCE],
                              media_type="application/octet-stream",
                              application_version="1", rule_or_model_version="rule")
    assert failure.value.code == "SOURCE_FAILED"
    with pytest.raises(ArtifactUnavailable) as mismatch:
        store.publish_managed("workspaceA", io.BytesIO(b"other"), expected_hash=sha(b"whole"),
                              expected_size=5, job_ref=JOB, source_refs=[SOURCE],
                              media_type="application/octet-stream",
                              application_version="1", rule_or_model_version="rule")
    assert mismatch.value.code == "INPUT_CHANGED"
    with sqlite3.connect(store.db_path) as db:
        assert db.execute("SELECT COUNT(*) FROM published_artifact").fetchone()[0] == 0
    assert not list(store.blob_root.glob("**/*.tmp"))


def test_interruption_after_blob_publication_still_has_no_downloadable_artifact(store, monkeypatch):
    def interrupted(*_args, **_kwargs):
        raise OSError("injected metadata write interruption")

    monkeypatch.setattr(store, "_insert", interrupted)
    with pytest.raises(OSError):
        publish(store, b"fully written but uncommitted")
    with sqlite3.connect(store.db_path) as db:
        assert db.execute("SELECT COUNT(*) FROM published_artifact").fetchone()[0] == 0


def test_missing_or_corrupt_managed_blob_fails_closed(store):
    missing = publish(store, b"missing")
    path = store.blob_root / missing.blob_hash[:2] / missing.blob_hash
    path.unlink()
    with pytest.raises(ArtifactUnavailable) as error:
        with store.stage_download("workspaceA", missing.artifact_id,
                                  expected_revision=1, expected_hash=missing.blob_hash):
            pytest.fail("Missing bytes must not be yielded")
    assert error.value.code == "MISSING"
    assert store.get("workspaceA", missing.artifact_id).availability == "missing"
    with pytest.raises(ArtifactUnavailable) as stale:
        with store.stage_download("workspaceA", missing.artifact_id,
                                  expected_revision=1, expected_hash=missing.blob_hash):
            pytest.fail("Old revision must not be yielded")
    assert stale.value.code == "REVISION_CONFLICT"

    corrupt = publish(store, b"corrupt target")
    path = store.blob_root / corrupt.blob_hash[:2] / corrupt.blob_hash
    path.write_bytes(b"changed")
    with pytest.raises(ArtifactUnavailable) as error:
        with store.stage_download("workspaceA", corrupt.artifact_id,
                                  expected_revision=1, expected_hash=corrupt.blob_hash):
            pytest.fail("Corrupt bytes must not be yielded")
    assert error.value.code == "CORRUPT"
    assert store.get("workspaceA", corrupt.artifact_id).availability == "corrupt"


def test_expiration_is_explicit_and_revision_bound(store):
    artifact = publish(store)
    expired = store.expire("workspaceA", artifact.artifact_id, expected_revision=1)
    assert expired.availability == "expired" and expired.revision == 2
    with pytest.raises(ArtifactUnavailable) as stale:
        with store.stage_download("workspaceA", artifact.artifact_id,
                                  expected_revision=1, expected_hash=artifact.blob_hash):
            pytest.fail("Expired result must not be yielded")
    assert stale.value.code == "REVISION_CONFLICT"
    with pytest.raises(ArtifactUnavailable) as error:
        with store.stage_download("workspaceA", artifact.artifact_id,
                                  expected_revision=2, expected_hash=artifact.blob_hash):
            pytest.fail("Expired result must not be yielded")
    assert error.value.code == "EXPIRED"


def test_native_reference_auto_expires_before_any_download(store, monkeypatch):
    artifact = reference(store)

    class AfterRetention(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime.now(tz) + timedelta(hours=2)

    monkeypatch.setattr(artifact_store_module, "datetime", AfterRetention)
    with pytest.raises(ArtifactUnavailable) as error:
        with store.stage_download("workspaceA", artifact.artifact_id,
                                  expected_revision=1, expected_hash=artifact.blob_hash,
                                  open_native=lambda *_: io.BytesIO(b"native bytes")):
            pytest.fail("Expired native content must not be downloaded")
    assert error.value.code == "EXPIRED"
    assert store.get("workspaceA", artifact.artifact_id).availability == "expired"


def test_native_reference_keeps_identity_and_rechecks_same_bytes(store):
    calls = []
    payload = {"data": b"native bytes"}

    def resolver(app, instance, native_id):
        calls.append((app, instance, native_id))
        return io.BytesIO(payload["data"])

    artifact = reference(store, payload["data"], resolver)
    assert store.provenance("workspaceA", artifact.artifact_id)["native_artifact_id"] == "nativeArtifactA"
    with store.stage_download("workspaceA", artifact.artifact_id, expected_revision=1,
                              expected_hash=artifact.blob_hash, open_native=resolver) as path:
        assert path.read_bytes() == b"native bytes"
    assert calls == [(JOB.app_id, JOB.instance_id, "nativeArtifactA")] * 2
    payload["data"] = b"changed native bytes"
    with pytest.raises(ArtifactUnavailable) as error:
        with store.stage_download("workspaceA", artifact.artifact_id, expected_revision=1,
                                  expected_hash=artifact.blob_hash, open_native=resolver):
            pytest.fail("Changed native bytes must not be yielded")
    assert error.value.code == "CORRUPT"
    assert store.get("workspaceA", artifact.artifact_id).availability == "corrupt"


def test_download_consumer_error_does_not_mark_valid_native_source_missing(store):
    artifact = reference(store)
    with pytest.raises(ValueError, match="consumer failure"):
        with store.stage_download("workspaceA", artifact.artifact_id,
                                  expected_revision=1, expected_hash=artifact.blob_hash,
                                  open_native=lambda *_: io.BytesIO(b"native bytes")):
            raise ValueError("consumer failure")
    assert store.get("workspaceA", artifact.artifact_id).availability == "available"


def test_unavailable_native_resolver_never_returns_an_old_path(store):
    artifact = reference(store)
    with pytest.raises(ArtifactUnavailable) as absent:
        with store.stage_download("workspaceA", artifact.artifact_id,
                                  expected_revision=1, expected_hash=artifact.blob_hash):
            pytest.fail("No resolver must not yield")
    assert absent.value.code == "SOURCE_UNAVAILABLE"

    def gone(_app, _instance, _id):
        raise FileNotFoundError("native source removed")

    with pytest.raises(ArtifactUnavailable) as missing:
        with store.stage_download("workspaceA", artifact.artifact_id, expected_revision=1,
                                  expected_hash=artifact.blob_hash, open_native=gone):
            pytest.fail("Missing native source must not yield")
    assert missing.value.code == "MISSING"
    assert store.get("workspaceA", artifact.artifact_id).availability == "missing"
