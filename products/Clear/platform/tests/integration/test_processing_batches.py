"""Real worker automation: fixed inputs/results, replay, recovery and dispositions."""
import json
from pathlib import Path
import sqlite3
import threading
import time

from fastapi.testclient import TestClient
import pytest
from packages.contracts.automation import ProcessingBatch
from workbench_platform.app import app
from workbench_platform.batches.store import encoded, fingerprint
from workbench_platform.data_api.service import DataWorkbenchService, now

LOCAL = {'host': '127.0.0.1:8766', 'origin': 'http://127.0.0.1:8766'}
PREFIX = '/api/platform/v1'
DATA_PYTHON = Path('F:/OCR/output/clear-workbench-preview-20261002/program/runtimes/data/python.exe')


@pytest.fixture
def client(tmp_path, monkeypatch):
    if not DATA_PYTHON.is_file():
        pytest.skip('Existing data worker runtime not available')
    monkeypatch.setenv('WORKBENCH_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('WORKBENCH_DATA_PYTHON', str(DATA_PYTHON))
    with TestClient(app, base_url='http://127.0.0.1:8766') as value:
        yield value


def workspace(client):
    response = client.post(PREFIX + '/workspaces', headers=LOCAL,
                           json={'schema_version': '1.0', 'name': '自动对账'})
    assert response.status_code == 201, response.text
    return response.json()['workspace_id']


def upload(client, workspace_id, name, text):
    response = client.post(f'{PREFIX}/workspaces/{workspace_id}/assets/uploads', headers=LOCAL,
                           params={'name': name}, content=text.encode('utf-8'))
    assert response.status_code == 201, response.text
    value = response.json()
    return {'asset_id': value['asset_id'], 'asset_revision': value['revision'], 'blob_hash': value['blob_hash']}


def template(client, workspace_id):
    role = {'formats': ['csv'], 'columns': [{'header': '编号', 'kinds': ['text']},
            {'header': '金额', 'kinds': ['text']}], 'import_options': {},
            'steps': [{'operator': 'trim', 'parameters': {'column': '编号'}},
                      {'operator': 'cast', 'parameters': {'column': '金额', 'target': 'decimal'}}]}
    response = client.post(PREFIX + '/batch-templates', headers=LOCAL, json={'schema_version': '1.0',
        'workspace_id': workspace_id, 'name': '已确认对账', 'scenario': 'data_compare',
        'config': {'roles': {'left': role, 'right': role}, 'comparison': {
            'key_pairs': [['编号', '编号']], 'compare_pairs': [['金额', '金额']], 'amount_pairs': [['金额', '金额']],
            'amount_tolerance': '0.01', 'type_policy': 'strict', 'null_policy': 'never_match'}}})
    assert response.status_code == 201, response.text
    return response.json()


def request(workspace_id, saved, inputs, key='runA'):
    return {'schema_version': '1.0', 'request_key': key, 'workspace_id': workspace_id,
            'template_id': saved['template_id'], 'template_version': saved['version'], 'inputs': inputs}


def terminal(client, batch_id, timeout=40, poll_interval=0.05, stall_timeout=None):
    started = time.monotonic()
    last_progress, last_changed = None, started
    while time.monotonic() - started < timeout:
        batch = client.get(f'{PREFIX}/batches/{batch_id}').json()
        if batch['status'] not in {'queued', 'running', 'cancelling'}:
            return batch
        progress = (batch['stage'], batch['coverage']['next_ordinal'],
                    tuple((key, value.get('status')) for key, value in batch['receipts'].items()
                          if isinstance(value, dict) and 'status' in value))
        if progress != last_progress:
            last_progress, last_changed = progress, time.monotonic()
        elif stall_timeout is not None and time.monotonic() - last_changed > stall_timeout:
            pytest.fail(f'Batch progress stalled at {progress}')
        time.sleep(poll_interval)
    pytest.fail('Batch did not reach a terminal/waiting state')


def test_real_data_closure_replay_reordered_reuse_and_disposition(client):
    ws = workspace(client)
    saved = template(client, ws)
    number = '00000000000000000001'
    left = upload(client, ws, '左.csv', f'编号,金额\n{number},10.00\nB,20.00\nD,1\nD,2\n')
    right = upload(client, ws, '右.csv', f'金额,编号\n10.01,{number}\n21.00,B\n3,D\n')
    payload = request(ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}])
    assert client.post(PREFIX + '/batches/preflight', headers=LOCAL, json=payload).json()['ok'] is True
    response = client.post(PREFIX + '/batches', headers=LOCAL, json=payload)
    assert response.status_code == 202, response.text
    batch = terminal(client, response.json()['batch_id'])
    assert batch['status'] == 'succeeded', batch
    ProcessingBatch.model_validate(batch)
    assert batch['artifact_state'] == 'collected'
    assert batch['summary']['coverage']['same_pairs'] == 1
    assert batch['summary']['coverage']['different_pairs'] == 1
    assert batch['summary']['pending'] > 1 and batch['summary']['review_complete'] is False
    assert any(item['category'].startswith('ambiguous') for item in batch['issues'])
    assert batch['summary']['coverage']['left_total'] == 4
    assert batch['coverage']['left_rows'] == 4 and batch['coverage']['right_rows'] == 3
    assert batch['coverage']['result_rows'] == 5 and batch['coverage']['required_items'] == 4
    assert batch['coverage']['index_complete'] is True
    assert batch['coverage']['pending'] == batch['coverage']['indexed_items'] == 4
    found = []
    for offset in (0, 2):
        found += client.get(f"{PREFIX}/batches/{batch['batch_id']}/issues",
                            params={'offset': offset, 'limit': 2}).json()['items']
    assert len({item['item_id'] for item in found}) == 4
    bundle = batch['artifacts'][0]
    downloaded = client.get(f"{PREFIX}/workspaces/{ws}/artifacts/{bundle['bundle_id']}/files/results.csv")
    assert downloaded.status_code == 200 and number in downloaded.content.decode('utf-8-sig')
    assert client.post(PREFIX + '/batches', headers=LOCAL, json=payload).json()['batch_id'] == batch['batch_id']
    listed = client.get(PREFIX + '/batches', params={'workspace_id': ws, 'request_key': payload['request_key']}).json()
    assert listed['total'] == 1
    assert client.get(f'{PREFIX}/workspaces/{ws}/artifacts').json()['total'] == 1
    item = batch['issues'][0]
    url = f"{PREFIX}/batches/{batch['batch_id']}/dispositions"
    actor = {'id': 'localUser', 'display_name': '本机编制人', 'source': 'local'}
    def disposal(seen, request_id, action, **extra):
        return {'request_id': request_id, 'expected_revision': seen.get('revision', 1),
                'actor': actor, 'item_ids': [item['item_id']], 'action': action, **extra}
    confirmed = client.post(url, headers=LOCAL, json=disposal(batch, 'confirmA', 'confirm', note='已核对原件')).json()
    assert confirmed['summary']['confirmed'] == 1
    assert confirmed['revision'] == batch['revision'] + 1
    replay = client.post(url, headers=LOCAL, json=disposal(batch, 'confirmA', 'confirm', note='已核对原件'))
    assert replay.status_code == 200 and replay.json() == confirmed
    conflict = client.post(url, headers=LOCAL, json=disposal(batch, 'otherWindow', 'defer'))
    assert conflict.status_code == 409 and '版本冲突' in conflict.text
    assert client.post(url, headers=LOCAL, json=disposal(batch, 'confirmA', 'confirm', note='不同说明')).status_code == 409
    deferred = client.post(url, headers=LOCAL, json=disposal(confirmed, 'deferA', 'defer')).json()
    assert deferred['summary']['confirmed'] == 0 and deferred['summary']['deferred'] == 1
    assert client.post(url, headers=LOCAL, json=disposal(batch, 'confirmA', 'confirm', note='已核对原件')).json() == confirmed
    reopened = client.post(url, headers=LOCAL, json=disposal(deferred, 'reopenA', 'reopen')).json()
    assert reopened['summary']['deferred'] == 0 and reopened['issues'][0]['raw_output'] == item['raw_output']
    history = app.state.batch_service.store.history(batch['batch_id'])
    assert history['total'] == 3
    assert [event['revision'] for event in history['items']] == [2, 3, 4]
    assert all(event['actor'] == actor and event['legacy'] is False for event in history['items'])
    assert client.post(url, headers=LOCAL, json=disposal(reopened, 'invalidA', 'confirm', item_ids=['otherIssue'])).status_code == 409
    from_session = client.post(PREFIX + '/batch-templates/from-data-session', headers=LOCAL, json={
        'schema_version': '1.0', 'workspace_id': ws, 'name': '复用已确认规则', 'session_id': batch['summary']['session_id']})
    assert from_session.status_code == 201, from_session.text
    assert from_session.json()['config']['comparison']['key_pairs'] == [['编号', '编号']]
    session = app.state.data_workbench.get(batch['summary']['session_id'])
    assert len(session['uploads']) == 2 and len(session['state']['exports']) == 1
    # Read-back persisted runner recovery must reuse all child job and export identities.
    app.state.batch_service.store.update(batch['batch_id'], status='queued')
    app.state.batch_service.schedule(batch['batch_id'])
    recovered = terminal(client, batch['batch_id'])
    assert recovered['status'] == 'succeeded', recovered
    assert recovered['artifacts'][0]['bundle_id'] == bundle['bundle_id']
    assert len(app.state.data_workbench.get(session['id'])['state']['exports']) == 1


def test_real_data_review_preserves_original_formats_and_same_rows(client):
    import csv
    import hashlib
    import io
    from openpyxl import load_workbook
    from packages.contracts.automation import ReviewBundle
    ws, saved = workspace(client), None
    saved = template(client, ws)
    number = '00123456789012345678'
    left = upload(client, ws, 'left-review.csv', f'编号,金额\n{number},1\nB,2\n')
    right = upload(client, ws, 'right-review.csv', f'编号,金额\n{number},1\nB,3\n')
    submitted = client.post(PREFIX + '/batches', headers=LOCAL,
        json=request(ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}])).json()
    batch = terminal(client, submitted['batch_id'])
    assert batch['status'] == 'succeeded' and batch['coverage']['result_rows'] == 2
    original = batch['artifacts'][0]
    base = f'{PREFIX}/workspaces/{ws}/artifacts'
    old = {file['name']: client.get(f"{base}/{original['bundle_id']}/files/{file['name']}").content for file in original['files']}
    actor = {'id': 'localReviewer', 'display_name': '本机复核人', 'source': 'local'}
    confirmed = client.post(f"{PREFIX}/batches/{batch['batch_id']}/dispositions", headers=LOCAL,
        json={'request_id': 'reviewConfirm', 'expected_revision': batch['revision'], 'actor': actor,
              'item_ids': [item['item_id'] for item in batch['issues']], 'action': 'confirm', 'note': '差异已核对'}).json()
    created = client.post(f"{PREFIX}/batches/{batch['batch_id']}/review-artifacts", headers=LOCAL,
        json={'request_id': 'actualDataReview', 'expected_revision': confirmed['revision'], 'actor': actor, 'kind': 'final'})
    assert created.status_code == 201, created.text
    response = created.json()
    ProcessingBatch.model_validate(response)
    bundle = response['artifacts'][-1]
    ReviewBundle.model_validate(bundle)
    def download(name):
        reply = client.get(f"{base}/{bundle['bundle_id']}/files/{name}")
        assert reply.status_code == 200, reply.text
        return reply.content
    records = list(csv.DictReader(io.StringIO(download('results.csv').decode('utf-8-sig'))))
    assert len(records) == 2 and {row['disposition'] for row in records} == {'confirmed', 'not_required'}
    assert number in download('results.csv').decode('utf-8-sig')
    workbook = load_workbook(io.BytesIO(download('results.xlsx')), read_only=True)
    try:
        assert len(list(workbook.active.iter_rows())) == 3
    finally:
        workbook.close()
    provenance = json.loads(download('provenance.json'))
    assert provenance['coverage']['index_complete'] is True and provenance['coverage']['result_rows'] == 2
    assert len(provenance['rows']) == 2 and len(provenance['original_bundles'][0]['files']) == 5
    for name, content in old.items():
        assert download('original-1-' + name) == content
        assert client.get(f"{base}/{original['bundle_id']}/files/{name}").content == content
        descriptor = next(file for file in bundle['files'] if file['name'] == 'original-1-' + name)
        assert descriptor['blob_hash'] == hashlib.sha256(content).hexdigest()
        assert descriptor['artifact_id'] != next(file['artifact_id'] for file in original['files'] if file['name'] == name)


def test_preflight_fail_closed_and_template_versions(client):
    ws = workspace(client)
    saved = template(client, ws)
    left = upload(client, ws, 'left.csv', '编号,金额\nA,1\n')
    bad = upload(client, ws, 'bad.csv', '编号,编号\nA,B\n')
    payload = request(ws, saved, [{'role': 'left', **left}, {'role': 'right', **bad}])
    checked = client.post(PREFIX + '/batches/preflight', headers=LOCAL, json=payload).json()
    assert checked['ok'] is False
    assert {'missing_column', 'duplicate_column'} <= {item['category'] for item in checked['issues']}
    posted = client.post(PREFIX + '/batches', headers=LOCAL, json=payload).json()
    batch = terminal(client, posted['batch_id'])
    assert batch['status'] == 'waiting_input' and batch['receipts'] == {}
    assert app.state.data_workbench.list() == []
    modified = {key: saved[key] for key in ('schema_version', 'workspace_id', 'template_id', 'name', 'scenario', 'config')}
    modified['name'] = '第二版本'
    version2 = client.post(PREFIX + '/batch-templates', headers=LOCAL, json=modified).json()
    assert version2['version'] == 2
    assert app.state.batch_service.store.template(ws, saved['template_id'], 1)['name'] == '已确认对账'
    conflicting = {**payload, 'template_version': 2}
    assert client.post(PREFIX + '/batches', headers=LOCAL, json=conflicting).status_code == 409


def test_cancel_waits_for_live_child_terminal_and_stops_dispatch(client, monkeypatch):
    ws = workspace(client)
    saved = template(client, ws)
    left = upload(client, ws, 'left.csv', '编号,金额\n' + ''.join(f'{i:020d},1\n' for i in range(50_000)))
    right = upload(client, ws, 'right.csv', '编号,金额\nA,1\n')
    payload = request(ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}])
    # Structure is already known for this fixture; keep cancellation focused on the real subprocess.
    monkeypatch.setattr(app.state.batch_service.data, 'preflight', lambda *_: [])
    posted = client.post(PREFIX + '/batches', headers=LOCAL, json=payload).json()
    started = time.monotonic()
    while time.monotonic() - started < 30:
        batch = client.get(f"{PREFIX}/batches/{posted['batch_id']}").json()
        receipt = batch['receipts'].get('import_left')
        if receipt and receipt['status'] in {'running', 'queued'}:
            break
        time.sleep(0.01)
    else:
        pytest.fail('Real import did not start')
    assert client.post(f"{PREFIX}/batches/{posted['batch_id']}/control", headers=LOCAL,
                       json={'operation': 'cancel'}).status_code == 200
    cancelled = terminal(client, posted['batch_id'])
    assert cancelled['status'] == 'cancelled'
    assert 'import_right' not in cancelled['receipts'] and 'export' not in cancelled['receipts']
    child = app.state.data_workbench.job(receipt['session_id'], receipt['job_id'])
    assert child['status'] not in {'running', 'queued', 'cancelling'}
    assert receipt['job_id'] not in app.state.data_workbench.processes


def test_durable_worker_reply_recovers_after_response_loss(tmp_path):
    service = DataWorkbenchService(tmp_path, DATA_PYTHON)
    session = service.create('回执丢失')
    session_id, job_id = session['id'], 'jLostResponse'
    request = {'job_id': job_id, 'session_dir': str(service.session_dir(session_id)),
               'action': 'compare', 'payload': {}, 'state': session['state']}
    with closing_connection(service.db()) as db:
        db.execute('INSERT INTO jobs VALUES(?,?,?,?,?,?,?,?)',
                   (job_id, session_id, 'compare', 'running', None, None, now(), now()))
        db.execute('INSERT INTO operation_requests VALUES(?,?,?,?,?,?)',
                   ('replyLost', session_id, job_id, fingerprint({}), 1, encoded(request)))
        db.commit()
    folder = service.session_dir(session_id) / 'jobs' / job_id
    folder.mkdir()
    state = {**session['state'], 'result_id': 'fixedResult'}
    (folder / 'response.json').write_text(encoded({'ok': True, 'state': state, 'result': {'id': 'fixedResult'}}), encoding='utf-8')
    service.close()
    reopened = DataWorkbenchService(tmp_path, DATA_PYTHON)
    try:
        assert reopened.operation('replyLost')['id'] == job_id
        assert reopened.operation('replyLost')['status'] == 'succeeded'
        assert reopened.get(session_id)['state']['result_id'] == 'fixedResult'
        assert reopened.get(session_id)['version'] == 2
    finally:
        reopened.close()


def test_service_restart_retries_collection_without_new_results(tmp_path, monkeypatch):
    monkeypatch.setenv('WORKBENCH_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('WORKBENCH_DATA_PYTHON', str(DATA_PYTHON))
    with TestClient(app, base_url='http://127.0.0.1:8766') as first:
        ws = workspace(first)
        saved = template(first, ws)
        left = upload(first, ws, 'left.csv', '编号,金额\n00000000000000000001,10.00\n')
        right = upload(first, ws, 'right.csv', '编号,金额\n00000000000000000001,10.00\n')
        payload = request(ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}], key='restartCollection')
        original = app.state.artifact_center.publish_data_export
        def unavailable(*args):
            raise ValueError('测试归集暂不可用')
        app.state.artifact_center.publish_data_export = unavailable
        posted = first.post(PREFIX + '/batches', headers=LOCAL, json=payload).json()
        partial = terminal(first, posted['batch_id'])
        assert partial['status'] == 'partial' and partial['artifact_state'] == 'collection_failed'
        export_id = partial['receipts']['export']['job_id']
        session_id = partial['summary']['session_id']
        app.state.artifact_center.publish_data_export = original
    with TestClient(app, base_url='http://127.0.0.1:8766') as restored:
        read_back = restored.get(f"{PREFIX}/batches/{posted['batch_id']}").json()
        assert read_back['artifact_state'] == 'collection_failed'
        assert restored.post(f"{PREFIX}/batches/{posted['batch_id']}/control", headers=LOCAL,
                             json={'operation': 'retry'}).status_code == 200
        completed = terminal(restored, posted['batch_id'])
        assert completed['status'] == 'succeeded', completed
        assert completed['receipts']['export']['job_id'] == export_id
        assert completed['summary']['pending'] == 0
        assert len(app.state.data_workbench.get(session_id)['state']['exports']) == 1
        assert restored.get(f'{PREFIX}/workspaces/{ws}/artifacts').json()['total'] == 1


from contextlib import closing as closing_connection


def test_100002_anomalies_are_fully_indexed_and_json_stays_bounded(client):
    ws = workspace(client)
    saved = template(client, ws)
    # Keep this check focused on true 100,002 result rows, not numeric casts.
    config = saved['config']
    for role in config['roles'].values():
        role['steps'] = []
    saved = client.post(PREFIX + '/batch-templates', headers=LOCAL, json={
        'schema_version': '1.0', 'workspace_id': ws, 'name': '全量不交覆盖',
        'scenario': 'data_compare', 'config': config}).json()
    left = upload(client, ws, 'large-left.csv', '编号,金额\n' + ''.join(f'L{i:06d},1\n' for i in range(50_001)))
    right = upload(client, ws, 'large-right.csv', '编号,金额\n' + ''.join(f'R{i:06d},1\n' for i in range(50_001)))
    posted = client.post(PREFIX + '/batches', headers=LOCAL, json=request(
        ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}], key='fullCoverage')).json()
    batch = terminal(client, posted['batch_id'], timeout=600, poll_interval=1, stall_timeout=120)
    assert batch['status'] == 'succeeded', batch
    ProcessingBatch.model_validate(batch)
    assert batch['coverage'] == {'input_files': 2, 'left_rows': 50_001, 'right_rows': 50_001,
        'result_rows': 100_002, 'required_items': 100_002, 'indexed_items': 100_002,
        'scanned_results': 100_002, 'next_ordinal': 100_002, 'index_complete': True,
        'pending': 100_002, 'confirmed': 0, 'deferred': 0}
    assert batch['summary']['review_complete'] is False and len(batch['issues']) == 30
    # Complete pagination checks every stable identity, independently of the
    # result count. Reverse page traversal must expose the same item universe.
    store = app.state.batch_service.store
    seen = set()
    for offset in range(100_000, -1, -200):
        page = store.issue_page(batch['batch_id'], offset, 200)
        assert page['total'] == 100_002
        ids = {item['item_id'] for item in page['items']}
        assert not ids & seen
        seen.update(ids)
    assert len(seen) == 100_002
    last = client.get(f"{PREFIX}/batches/{batch['batch_id']}/issues", params={'offset': 100_000, 'limit': 200}).json()
    assert len(last['items']) == 2 and last['total'] == 100_002
    assert last['coverage']['index_complete'] is True
    deferred = client.post(f"{PREFIX}/batches/{batch['batch_id']}/dispositions", headers=LOCAL, json={
        'request_id': 'lastPageReview', 'expected_revision': batch['revision'],
        'actor': {'id': 'localUser', 'display_name': '本机编制人', 'source': 'local'},
        'item_ids': [item['item_id'] for item in last['items']], 'action': 'defer'}).json()
    assert deferred['coverage']['deferred'] == 2 and deferred['summary']['pending'] == 100_000
    with closing_connection(store.db()) as db:
        assert db.execute('SELECT count(DISTINCT ordinal) FROM batch_items WHERE batch_id=?',
                          (batch['batch_id'],)).fetchone()[0] == 100_002
        body = db.execute('SELECT body FROM batches WHERE batch_id=?', (batch['batch_id'],)).fetchone()[0]
    assert len(body.encode('utf-8')) < 20_000
    assert json.loads(body)['issues'] == [] and 'issues' not in json.loads(body)['receipts']['results']


def test_mixed_null_duplicate_pair_coverage_across_page_orders(client):
    ws = workspace(client)
    saved = template(client, ws)
    left = upload(client, ws, 'mixed-left.csv', '编号,金额\nA,10\nB,20\nD,1\nD,2\nE,7\n,4\nF,0\n')
    right = upload(client, ws, 'mixed-right.csv', '编号,金额\nA,10.01\nB,21\nD,1\nD,2\nD,3\nC,8\n,4\nF,0\n')
    posted = client.post(PREFIX + '/batches', headers=LOCAL, json=request(
        ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}], key='mixedCoverage')).json()
    batch = terminal(client, posted['batch_id'])
    assert batch['status'] == 'succeeded', batch
    coverage = batch['summary']['coverage']
    assert coverage['left_total'] == coverage['left_categorized'] == 7
    assert coverage['right_total'] == coverage['right_categorized'] == 8
    assert coverage['same_pairs'] == 2 and coverage['different_pairs'] == 1
    assert coverage['only_left'] == coverage['only_right'] == 2
    assert coverage['ambiguous_left_rows'] == 2 and coverage['ambiguous_right_rows'] == 3
    assert batch['coverage']['result_rows'] == 12 and batch['coverage']['required_items'] == 10
    assert batch['coverage']['indexed_items'] == 10 and batch['coverage']['index_complete'] is True
    def collect(order, limit):
        items = []
        for offset in order:
            items += client.get(f"{PREFIX}/batches/{batch['batch_id']}/issues",
                                params={'offset': offset, 'limit': limit}).json()['items']
        assert len(items) == len({item['item_id'] for item in items}) == 10
        return {item['item_id'] for item in items}
    assert collect(range(0, 10, 3), 3) == collect(reversed(range(0, 10, 2)), 2)


def test_index_page_commit_then_restart_resumes_without_duplicates(tmp_path, monkeypatch):
    from workbench_platform.data_api.service import DataError
    monkeypatch.setenv('WORKBENCH_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('WORKBENCH_DATA_PYTHON', str(DATA_PYTHON))
    with TestClient(app, base_url='http://127.0.0.1:8766') as first:
        ws = workspace(first)
        saved = template(first, ws)
        left = upload(first, ws, 'resume-left.csv', '编号,金额\n' + ''.join(f'L{i:04d},1\n' for i in range(1251)))
        right = upload(first, ws, 'resume-right.csv', '编号,金额\n' + ''.join(f'R{i:04d},1\n' for i in range(1251)))
        store = app.state.batch_service.store
        original = store.index_page
        def interrupt_after_commit(*args):
            original(*args)
            raise DataError('测试页提交后中断，已保留恢复点')
        monkeypatch.setattr(store, 'index_page', interrupt_after_commit)
        posted = first.post(PREFIX + '/batches', headers=LOCAL, json=request(
            ws, saved, [{'role': 'left', **left}, {'role': 'right', **right}], key='indexResume')).json()
        paused = terminal(first, posted['batch_id'])
        assert paused['status'] == 'needs_verification'
        assert paused['coverage']['scanned_results'] == paused['coverage']['indexed_items'] == 1000
        assert paused['coverage']['index_complete'] is False and paused['summary']['review_complete'] is False
        compare_job = paused['receipts']['compare']['job_id']
        session_id = paused['receipts']['session']['session_id']
        item = paused['issues'][0]
        response = first.post(f"{PREFIX}/batches/{paused['batch_id']}/dispositions", headers=LOCAL, json={
            'request_id': 'beforeResume', 'expected_revision': paused['revision'],
            'actor': {'id': 'localUser', 'display_name': '本机编制人', 'source': 'local'},
            'item_ids': [item['item_id']], 'action': 'defer', 'note': '中断前人工处置'})
        assert response.status_code == 200, response.text
    with TestClient(app, base_url='http://127.0.0.1:8766') as restored:
        response = restored.post(f"{PREFIX}/batches/{paused['batch_id']}/control", headers=LOCAL, json={'operation': 'retry'})
        assert response.status_code == 200, response.text
        batch = terminal(restored, paused['batch_id'])
        assert batch['status'] == 'succeeded', batch
        assert batch['receipts']['compare']['job_id'] == compare_job
        assert batch['receipts']['session']['session_id'] == session_id
        assert batch['coverage']['index_complete'] is True
        assert batch['coverage']['indexed_items'] == batch['coverage']['scanned_results'] == 2502
        assert batch['coverage']['pending'] == 2501 and batch['coverage']['deferred'] == 1
        assert batch['summary']['review_complete'] is False
        page = restored.get(f"{PREFIX}/batches/{batch['batch_id']}/issues", params={'disposition': 'deferred'}).json()
        assert page['total'] == 1 and page['items'][0]['item_id'] == item['item_id']
        store = app.state.batch_service.store
        with closing_connection(store.db()) as db:
            assert db.execute('SELECT count(*),count(DISTINCT ordinal) FROM batch_items WHERE batch_id=?',
                              (batch['batch_id'],)).fetchone()[:] == (2502, 2502)


def test_index_identity_change_and_unresolved_counts_fail_closed(tmp_path):
    from workbench_platform.batches.store import BatchStore
    from workbench_platform.data_api.service import DataError
    store = BatchStore(tmp_path / 'indexes')
    store.begin_index('batchA', 'fixedA', 2, 2, 1, 1)
    with pytest.raises(DataError, match='身份'):
        store.begin_index('batchA', 'fixedB', 2, 2, 1, 1)
    item = {'item_id': 'itemA', 'category': 'only_left', 'source': {'result_id': 'fixedA'},
            'raw_output': {'ordinal': 0}, 'disposition': 'pending'}
    store.index_page('batchA', 'fixedA', 0, [{'ordinal': 0}], [item])
    store.index_page('batchA', 'fixedA', 0, [{'ordinal': 0}], [item])
    assert store.coverage('batchA', 2)['indexed_items'] == 1
    assert store.coverage('batchA', 2)['index_complete'] is False
    with pytest.raises(DataError, match='内容'):
        store.index_page('batchA', 'fixedA', 0, [{'ordinal': 0}], [{**item, 'title': 'changed'}])


@pytest.mark.parametrize('response_kind', ['page', 'service'])
def test_issue_response_uses_one_snapshot_during_concurrent_index_commit(tmp_path, monkeypatch, response_kind):
    from workbench_platform.batches.store import BatchStore
    from workbench_platform.batches.service import BatchService
    store = BatchStore(tmp_path / 'snapshot')
    template = store.save_template({'workspace_id': 'workspaceA', 'name': '读快照',
                                    'scenario': 'data_compare', 'config': {}})
    batch, _ = store.accept({'schema_version': '1.0', 'workspace_id': 'workspaceA',
        'request_key': 'snapshotA', 'template_id': template['template_id'],
        'template_version': template['version'], 'inputs': []}, template)
    batch_id = batch['batch_id']
    store.begin_index(batch_id, 'fixedA', 2, 2, 1, 1)
    writer = BatchStore(tmp_path / 'snapshot')
    original_connect = store.db
    committed = False
    class InterleavedConnection:
        def __init__(self):
            self.connection = original_connect()
        def execute(self, sql, arguments=()):
            nonlocal committed
            cursor = self.connection.execute(sql, arguments)
            if sql.startswith('SELECT count(*) FROM batch_items i WHERE') and not committed:
                committed = True
                writer.index_page(batch_id, 'fixedA', 0, [{'ordinal': 0}], [{
                    'item_id': 'itemA', 'category': 'only_left', 'source': {'result_id': 'fixedA'},
                    'raw_output': {'ordinal': 0}}])
                writer.update(batch_id, revision=2)
            return cursor
        def close(self):
            self.connection.close()
    monkeypatch.setattr(store, 'db', InterleavedConnection)
    if response_kind == 'page':
        response = store.issue_page(batch_id)
    else:
        service = object.__new__(BatchService)
        service.store = store
        service.workspace = lambda _: None
        response = service.issue_page(batch_id)
        assert response['revision'] == 1
        assert response['coverage']['indexed_items'] == response['coverage']['scanned_results'] == 0
        assert response['summary']['pending'] == 0
    assert committed and response['total'] == 0 and response['items'] == []
    next_page = store.snapshot(batch_id)
    assert next_page[0]['revision'] == 2 and next_page[2]['total'] == 1
    assert next_page[0]['coverage']['indexed_items'] == next_page[0]['coverage']['scanned_results'] == 1


def ledger_fixture(tmp_path, scenario='data_compare', request_key='ledgerRun'):
    from workbench_platform.batches.store import BatchStore
    store = BatchStore(tmp_path / 'ledger')
    template = store.save_template({'schema_version': '1.0', 'workspace_id': 'workspaceA',
        'name': '版本化处置', 'scenario': scenario, 'config': {'link_id': 'linkA'}})
    batch, _ = store.accept({'schema_version': '1.0', 'workspace_id': 'workspaceA',
        'request_key': request_key, 'template_id': template['template_id'], 'template_version': template['version'],
        'inputs': [{'role': 'left' if scenario == 'data_compare' else 'images',
                    'asset_id': 'assetA', 'asset_revision': 1, 'blob_hash': 'a' * 64}]}, template)
    item = {'item_id': 'itemA', 'category': 'different' if scenario == 'data_compare' else 'review',
        'source': {'result_id': 'fixedA', 'asset_id': 'assetA', 'blob_hash': 'a' * 64},
        'raw_output': {'ordinal': 0, 'native_item_id': 'itemA', 'prediction_label_id': 'A',
                       'identifier': '00000000000000000001'}, 'disposition': 'pending', 'note': ''}
    store.begin_index(batch['batch_id'], 'fixedA', 1, 1, 1, 0)
    store.index_page(batch['batch_id'], 'fixedA', 0, [{'ordinal': 0}], [item])
    store.update(batch['batch_id'], status='succeeded', artifacts=[{'bundle_id': 'originalBundleA'}])
    return store, store.get(batch['batch_id']), item


def ledger_payload(batch, request_id, action='confirm', **extra):
    return {'request_id': request_id, 'expected_revision': batch['revision'],
        'actor': {'id': 'localReviewer', 'display_name': '本机复核人', 'source': 'local'},
        'item_ids': ['itemA'], 'action': action, 'note': '已核对原件', **extra}


def test_disposition_revision_and_request_receipt_survive_replay_and_other_windows(tmp_path):
    from workbench_platform.batches.store import BatchStore
    from workbench_platform.data_api.service import DataError
    from packages.contracts.automation import BatchOperation
    store, batch, item = ledger_fixture(tmp_path)
    other_window = BatchStore(tmp_path / 'ledger')
    first = ledger_payload(batch, 'firstWindow')
    receipt = store.disposition(batch['batch_id'], first)
    assert receipt['revision'] == 2 and receipt['batch']['summary']['review_complete'] is True
    with pytest.raises(DataError, match='版本冲突'):
        other_window.disposition(batch['batch_id'], ledger_payload(batch, 'otherWindow', 'defer'))
    with pytest.raises(DataError, match='不同'):
        other_window.disposition(batch['batch_id'], {**first, 'note': '另一说明'})
    assert other_window.disposition(batch['batch_id'], first) == receipt
    assert store.history(batch['batch_id'])['total'] == 1
    later = store.disposition(batch['batch_id'], ledger_payload(receipt['batch'], 'laterReopen', 'reopen'))
    assert later['revision'] == 3 and later['batch']['summary']['review_complete'] is False
    assert other_window.disposition(batch['batch_id'], first) == receipt  # replay precedes stale revision CAS
    operation = store.operation(receipt['operation_id'])
    BatchOperation.model_validate(operation)
    assert operation['state'] == 'committed' and operation['actor'] == first['actor']
    assert operation['receipt'] == receipt
    assert store.history(batch['batch_id'])['total'] == 2
    assert store.selected_items(batch['batch_id'], ['itemA'])[0]['raw_output'] == item['raw_output']
    second_store, second_batch, _ = ledger_fixture(tmp_path, request_key='otherBatch')
    with pytest.raises(DataError, match='不同批次'):
        second_store.disposition(second_batch['batch_id'], first)
    assert second_store.get(second_batch['batch_id'])['revision'] == 1


def test_intent_freezes_origin_and_adoption_and_unknown_operation_blocks_completion(tmp_path):
    from workbench_platform.batches.store import BatchStore, present_batch
    from workbench_platform.data_api.service import DataError
    from packages.contracts.automation import BatchOperation
    store, batch, original = ledger_fixture(tmp_path, 'classification')
    payload = ledger_payload(batch, 'adoptNative', 'adopt', adoption={'native_revision': 2, 'human_label_id': 'B'})
    origin = {'instance_id': 'instanceA', 'project_id': 'projectA', 'original_sha256': 'a' * 64}
    intent = store.begin_operation(batch['batch_id'], payload, origin=origin)
    BatchOperation.model_validate(store.operation(intent['operation_id']))
    assert intent['intent']['origin'] == origin and intent['intent']['native_started'] is False
    assert intent['intent']['targets'][0]['raw_output']['prediction_label_id'] == 'A'
    assert intent['payload']['adoption']['native_revision'] == 2
    with pytest.raises(DataError, match='核对回执'):
        store.commit_operation(intent['operation_id'], adoption={'human_label_id': 'B'})
    with pytest.raises(DataError, match='来源'):
        store.checkpoint_operation(intent['operation_id'], 'intent', intent={
            **intent['intent'], 'origin': {**origin, 'instance_id': 'instanceB'}})
    receipt = {'operation_id': intent['operation_id'], 'state': 'committed',
               'instance_id': 'instanceA', 'project_id': 'projectA', 'revision': 3, 'human_label_id': 'B'}
    store.checkpoint_operation(intent['operation_id'], 'native_committed', native_receipt=receipt)
    adopted = store.commit_operation(intent['operation_id'], adoption={
        'native_revision': 2, 'committed_native_revision': 3, 'human_label_id': 'B', 'native_receipt': receipt})
    item = store.selected_items(batch['batch_id'], ['itemA'])[0]
    assert item['raw_output'] == original['raw_output'] and item['adoption']['human_label_id'] == 'B'
    assert adopted['batch']['coverage']['confirmed'] == 1
    pending = store.begin_operation(batch['batch_id'], ledger_payload(adopted['batch'], 'unknownNative', 'reopen'), origin=origin)
    store.checkpoint_operation(pending['operation_id'], 'needs_verification', reason='原生响应尚不可核实',
        intent={**pending['intent'], 'native_started': True})
    reopened = BatchStore(tmp_path / 'ledger')
    remembered = reopened.begin_operation(batch['batch_id'], pending['payload'], origin={**origin, 'instance_id': 'instanceB'})
    assert remembered['operation_id'] == pending['operation_id'] and remembered['intent']['origin'] == origin
    assert remembered['state'] == 'needs_verification' and remembered['intent']['native_started'] is True
    snapshot, counts, _ = reopened.snapshot(batch['batch_id'])
    assert counts['confirmed'] == 1 and snapshot['coverage']['index_complete'] is True
    displayed = present_batch(snapshot, counts)
    assert displayed['summary']['review_complete'] is False and displayed['summary']['pending_operations'] == 1
    assert displayed['summary']['pending_operation']['operation_id'] == pending['operation_id']
    with pytest.raises(DataError, match='待核实'):
        reopened.begin_operation(batch['batch_id'], ledger_payload(adopted['batch'], 'blindRetry', 'reopen'), origin=origin)
    with pytest.raises(DataError, match='核对回执'):
        reopened.commit_operation(pending['operation_id'])
    assert reopened.get(batch['batch_id'])['revision'] == 2 and reopened.history(batch['batch_id'])['total'] == 1
    assert reopened.history(batch['batch_id'])['items'][0]['adoption']['native_revision'] == 2


def test_legacy_ledger_migration_rolls_back_then_repeats_without_inventing_identity(tmp_path, monkeypatch):
    from workbench_platform.batches.store import BatchStore
    from workbench_platform.data_api.service import DataError
    root = tmp_path / 'legacy'
    root.mkdir()
    item = {'item_id': 'legacyItem', 'category': 'different', 'source': {'result_id': 'oldResult'},
            'raw_output': {'ordinal': 0, 'identifier': '00000000000000000001'},
            'disposition': 'deferred', 'note': '旧人工说明'}
    body = {'schema_version': '1.0', 'batch_id': 'legacyBatch', 'request_key': 'legacyRun',
        'workspace_id': 'workspaceA', 'template_id': 'legacyTemplate', 'template_version': 1,
        'template': {'scenario': 'data_compare'}, 'inputs': [], 'status': 'succeeded', 'stage': 'collect',
        'summary': {}, 'issues': [item], 'receipts': {'results': {'issues': [item], 'result_id': 'oldResult'}},
        'artifact_state': 'collected', 'artifacts': [{'bundle_id': 'oldBundleA'}], 'capabilities': [],
        'cancel_requested': False, 'created_at': '2026-09-30T01:00:00Z', 'updated_at': '2026-09-30T02:00:00Z'}
    raw_body = encoded(body)
    with sqlite3.connect(root / 'batches.sqlite3') as db:
        db.executescript('CREATE TABLE batches(batch_id TEXT PRIMARY KEY,request_key TEXT UNIQUE,workspace_id TEXT,fingerprint TEXT,body TEXT);'
                        'CREATE TABLE disposition_events(event_id TEXT PRIMARY KEY,batch_id TEXT,item_id TEXT,action TEXT,note TEXT,created_at TEXT);')
        db.execute('INSERT INTO batches VALUES(?,?,?,?,?)', ('legacyBatch', 'legacyRun', 'workspaceA', 'a' * 64, raw_body))
        db.execute('INSERT INTO disposition_events VALUES(?,?,?,?,?,?)',
                   ('oldEvent', 'legacyBatch', 'legacyItem', 'defer', '旧人工说明', '2026-09-30T02:00:00Z'))
    original_write = BatchStore._write_items
    def failing_migration(self, db, batch_id, issues):
        original_write(self, db, batch_id, issues)
        raise DataError('迁移注入中断')
    with monkeypatch.context() as patch:
        patch.setattr(BatchStore, '_write_items', failing_migration)
        with pytest.raises(DataError, match='迁移'):
            BatchStore(root)
    with sqlite3.connect(root / 'batches.sqlite3') as db:
        assert db.execute('SELECT body FROM batches').fetchone()[0] == raw_body
        assert db.execute('SELECT count(*) FROM batch_items').fetchone()[0] == 0
        assert db.execute('SELECT count(*) FROM disposition_events').fetchone()[0] == 1
    migrated = BatchStore(root)
    repeated = BatchStore(root)
    read = repeated.get('legacyBatch')
    assert read['revision'] == 1 and read['artifacts'] == body['artifacts']
    assert read['issues'][0]['raw_output'] == item['raw_output'] and read['issues'][0]['disposition'] == 'deferred'
    assert read['coverage']['indexed_items'] == 1 and read['coverage']['index_complete'] is False
    assert migrated.history('legacyBatch') == repeated.history('legacyBatch')
    history = repeated.history('legacyBatch')['items']
    assert len(history) == 1 and history[0]['event_id'] == 'oldEvent'
    assert history[0]['actor'] is None and history[0]['revision'] is None and history[0]['legacy'] is True
    assert history[0]['created_at'] == '2026-09-30T02:00:00Z' and history[0]['note'] == '旧人工说明'


def test_native_receipt_and_dispatched_segment_binding_cannot_be_replaced(tmp_path):
    from workbench_platform.data_api.service import DataError
    store, batch, _ = ledger_fixture(tmp_path, 'classification')
    operation = store.begin_operation(batch['batch_id'], ledger_payload(batch, 'fixedDispatch', 'defer'),
        origin={'instance_id': 'instanceA', 'project_id': 'projectA'})
    segment = {'operation_id': operation['operation_id'] + '_segment0', 'item_ids': ['itemA'],
        'payload_sha256': 'a' * 64, 'payload': {'state': 'deferred', 'native_revision': 1},
        'state': 'intent', 'native_started': True}
    started = store.checkpoint_operation(operation['operation_id'], 'intent',
        intent={**operation['intent'], 'native_started': True, 'segments': [segment]})
    for changed in ({'operation_id': 'replacement'}, {'item_ids': ['differentItem']},
                    {'payload_sha256': 'b' * 64}, {'payload': {'state': 'deferred', 'native_revision': 2}}):
        with pytest.raises(DataError, match='分段'):
            store.checkpoint_operation(operation['operation_id'], 'needs_verification',
                intent={**started['intent'], 'segments': [{**segment, **changed}]})
    with pytest.raises(DataError, match='分段已派发'):
        store.checkpoint_operation(operation['operation_id'], 'needs_verification',
            intent={**started['intent'], 'segments': [{**segment, 'native_started': False}]})
    native_receipt = {'operation_id': segment['operation_id'], 'state': 'committed',
                      'instance_id': 'instanceA', 'payload_hash': 'a' * 64}
    completed_segment = {**segment, 'state': 'native_committed', 'native_receipt': native_receipt}
    verified = store.checkpoint_operation(operation['operation_id'], 'native_committed',
        native_receipt={'segments': [native_receipt]},
        intent={**started['intent'], 'segments': [completed_segment]})
    with pytest.raises(DataError, match='退回'):
        store.checkpoint_operation(operation['operation_id'], 'intent')
    with pytest.raises(DataError, match='回执不能替换'):
        store.checkpoint_operation(operation['operation_id'], 'needs_verification', native_receipt={'segments': []})
    with pytest.raises(DataError, match='分段回执'):
        store.checkpoint_operation(operation['operation_id'], 'needs_verification',
            intent={**verified['intent'], 'segments': [{**completed_segment, 'native_receipt': {}}]})
    unknown = store.checkpoint_operation(operation['operation_id'], 'needs_verification', reason='平台尚待记账')
    assert unknown['native_receipt'] == verified['native_receipt'] and unknown['intent'] == verified['intent']
    replay = store.checkpoint_operation(operation['operation_id'], 'native_committed', native_receipt=verified['native_receipt'])
    assert replay['native_receipt'] == verified['native_receipt'] and replay['intent'] == verified['intent']
    committed = store.commit_operation(operation['operation_id'])
    assert committed['native_receipt'] == verified['native_receipt'] and committed['revision'] == 2


def test_two_initializers_and_disposition_never_rewrite_a_stale_legacy_revision(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from workbench_platform.batches.store import BatchStore
    store, batch, _ = ledger_fixture(tmp_path)
    # Make the existing safe fixture a legacy body while retaining its original
    # indexed row. A's migration must not replay this body over B's new decision.
    with closing_connection(store.db()) as db, db:
        raw = json.loads(db.execute('SELECT body FROM batches WHERE batch_id=?', (batch['batch_id'],)).fetchone()[0])
        raw['issues'] = batch['issues']
        for key in ('revision', 'automation_version', 'delivery'):
            raw.pop(key, None)
        db.execute('UPDATE batches SET body=? WHERE batch_id=?', (encoded(raw), batch['batch_id']))
    old_body_read, release_a, started_b, acquired_b = (threading.Event() for _ in range(4))
    original_write, original_connect = BatchStore._write_items, BatchStore.db
    def paused_write(self, db, batch_id, issues):
        if threading.current_thread().name.startswith('migration-A'):
            old_body_read.set()
            assert release_a.wait(10)
        return original_write(self, db, batch_id, issues)
    class ObservedConnection:
        def __init__(self, connection):
            self.connection = connection
        def execute(self, sql, arguments=()):
            cursor = self.connection.execute(sql, arguments)
            if sql == 'BEGIN IMMEDIATE' and threading.current_thread().name.startswith('migration-B'):
                acquired_b.set()
            return cursor
        def executescript(self, sql):
            return self.connection.executescript(sql)
        def close(self):
            self.connection.close()
        def __enter__(self):
            self.connection.__enter__()
            return self
        def __exit__(self, *args):
            return self.connection.__exit__(*args)
    monkeypatch.setattr(BatchStore, '_write_items', paused_write)
    monkeypatch.setattr(BatchStore, 'db', lambda self: ObservedConnection(original_connect(self)))
    def initialize_and_dispose():
        started_b.set()
        other = BatchStore(tmp_path / 'ledger')
        return other.disposition(batch['batch_id'], ledger_payload(batch, 'migrationWindowB'))
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix='migration-A') as a_pool, \
            ThreadPoolExecutor(max_workers=1, thread_name_prefix='migration-B') as b_pool:
        a = a_pool.submit(BatchStore, tmp_path / 'ledger')
        try:
            assert old_body_read.wait(5)
            b = b_pool.submit(initialize_and_dispose)
            assert started_b.wait(5)
            assert not acquired_b.wait(0.25), 'B obtained the write lock while A held a stale migration read'
        finally:
            release_a.set()
        a.result(timeout=30)
        receipt = b.result(timeout=30)
    final = store.get(batch['batch_id'])
    assert receipt['revision'] == final['revision'] == 2
    assert final['issues'][0]['disposition'] == 'confirmed'
    history = store.history(batch['batch_id'])['items']
    assert len(history) == 1 and history[0]['actor'] == ledger_payload(batch, 'migrationWindowB')['actor']
    assert store.operation(receipt['operation_id'])['receipt'] == receipt


def test_uncertain_native_cancel_remains_unverified_and_retry_keeps_cancel(client, monkeypatch):
    from types import SimpleNamespace
    from workbench_platform.batches.native import NativeBatchError
    ws = workspace(client)
    asset = upload(client, ws, 'input.csv', '编号,金额\nA,1\n')
    saved = client.post(PREFIX + '/batch-templates', headers=LOCAL, json={
        'schema_version': '1.0', 'workspace_id': ws, 'name': '取消回执核实',
        'scenario': 'classification', 'config': {'link_id': 'linkedProject', 'model_id': 'existingModel'}}).json()
    submitted = threading.Event()
    release = threading.Event()
    attempts = []
    def run(batch, checkpoint, cancelled):
        attempts.append({'cancel_requested': batch['cancel_requested'], 'inputs': batch['inputs']})
        if len(attempts) == 1:
            checkpoint('native_submission', {'platform_job_id': 'jExistingNativeJob'})
            submitted.set()
            assert release.wait(5)
            # The native worker may still run; unavailable status cannot prove termination.
            raise NativeBatchError('原生服务不可达，取消终态需核实', 'needs_verification')
        assert cancelled(), 'Retry must continue cancellation, not resume dispatch'
        assert batch['receipts']['native_submission']['platform_job_id'] == 'jExistingNativeJob'
        return {'status': 'cancelled', 'summary': {'total': 1, 'cancelled': 1}, 'issues': [],
                'artifacts': [], 'artifact_state': 'not_generated'}
    monkeypatch.setattr(app.state.batch_service, 'native', SimpleNamespace(
        prepare_runtime=lambda *_: None, preflight=lambda *_: [], run=run))
    payload = request(ws, saved, [{'role': 'images', **asset}], key='cancelUnverified')
    posted = client.post(PREFIX + '/batches', headers=LOCAL, json=payload).json()
    assert submitted.wait(5)
    assert client.post(f"{PREFIX}/batches/{posted['batch_id']}/control", headers=LOCAL,
                       json={'operation': 'cancel'}).status_code == 200
    release.set()
    unknown = terminal(client, posted['batch_id'])
    assert unknown['status'] == 'needs_verification' and unknown['cancel_requested'] is True
    assert unknown['summary']['execution_complete'] is False
    assert client.post(f"{PREFIX}/batches/{posted['batch_id']}/control", headers=LOCAL,
                       json={'operation': 'retry'}).status_code == 200
    reconciled = terminal(client, posted['batch_id'])
    assert reconciled['status'] == 'cancelled' and reconciled['cancel_requested'] is True
    assert len(attempts) == 2 and attempts[1]['cancel_requested'] is True
    assert attempts[0]['inputs'] == attempts[1]['inputs']
