"""Completed data exports become workspace-scoped, replayable artifacts."""

import hashlib
import json
import sqlite3
import csv
import io

import pytest
from openpyxl import load_workbook
from packages.contracts.automation import ProcessingBatch, ReviewBundle

from fastapi.testclient import TestClient

from workbench_platform.app import app

LOCAL = {"host": "127.0.0.1:8766", "origin": "http://127.0.0.1:8766"}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def finished_export(service):
    session = service.create("对账成果")
    session_id = session["id"]
    upload_id = "uSourceA"
    source = service.session_dir(session_id) / "uploads" / "uSourceA.csv"
    source.write_bytes("编号,金额\n甲,10\n".encode())
    source_bytes = source.read_bytes()
    service.upload(session_id, upload_id, "原始表.csv", ".csv", source,
                   sha(source_bytes), len(source_bytes))
    export_id = "jExportA"
    folder = service.session_dir(session_id) / "exports" / export_id
    folder.mkdir(parents=True)
    source_ref = {"asset_id": upload_id, "sha256": sha(source_bytes), "format": "csv"}
    contents = {
        "cleaned.csv": b"source_row,amount\n2,10\n",
        "cleaned.xlsx": b"xlsx-fixture",
        "cleaned.parquet": b"parquet-fixture",
        "summary.json": json.dumps({"revision": "rOriginal", "source": source_ref,
                                    "rows": 1}, ensure_ascii=False).encode(),
        "provenance.json": json.dumps({"source": source_ref, "rows": [2]}).encode(),
    }
    files = []
    for name, data in contents.items():
        path = folder / name
        path.write_bytes(data)
        files.append({"name": name, "relative": path.relative_to(service.session_dir(session_id)).as_posix(),
                      "sha256": sha(data), "bytes": len(data)})
    state = session["state"]
    state["exports"] = [{"id": export_id, "result_id": None, "files": files}]
    with sqlite3.connect(service.db_path) as db:
        db.execute("UPDATE sessions SET state=? WHERE id=?", (json.dumps(state), session_id))
        db.execute("INSERT INTO jobs VALUES (?,?,?,?,?,?,?,?)",
                   (export_id, session_id, "export_table", "succeeded", None, "{}",
                    "2026-09-27T00:00:00Z", "2026-09-27T00:00:00Z"))
    return session_id, export_id, contents, source


def test_publish_replay_download_and_source_links(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKBENCH_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("WORKBENCH_DATA_PYTHON", raising=False)
    with TestClient(app, base_url="http://127.0.0.1:8766") as client:
        first = client.post("/api/platform/v1/workspaces", headers=LOCAL,
                            json={"schema_version": "1.0", "name": "甲"}).json()["workspace_id"]
        second = client.post("/api/platform/v1/workspaces", headers=LOCAL,
                             json={"schema_version": "1.0", "name": "乙"}).json()["workspace_id"]
        session_id, export_id, contents, _ = finished_export(app.state.data_workbench)
        base = f"/api/platform/v1/workspaces/{first}/artifacts"
        other = f"/api/platform/v1/workspaces/{second}/artifacts"
        payload = {"session_id": session_id, "export_id": export_id}
        assert client.post(base + "/data-exports", json=payload).status_code == 403
        created = client.post(base + "/data-exports", headers=LOCAL, json=payload)
        assert created.status_code == 201, created.text
        bundle = created.json()
        assert len(bundle["files"]) == 5
        assert [item["name"] for item in bundle["inputs"]] == ["原始表.csv"]
        assert bundle["input_set_revision"] == 1
        replay = client.post(base + "/data-exports", headers=LOCAL, json=payload)
        assert replay.status_code == 201 and replay.json() == bundle
        page = client.get(base).json()
        assert page["total"] == 1 and page["items"][0] == bundle
        assert client.get(other).json()["total"] == 0
        assert client.get(other + "/" + bundle["bundle_id"]).status_code == 404
        for name, original in contents.items():
            result = client.get(base + f"/{bundle['bundle_id']}/files/{name}")
            assert result.status_code == 200 and result.content == original
        asset_id = bundle["inputs"][0]["asset_id"]
        asset = client.get(f"/api/platform/v1/workspaces/{first}/assets/{asset_id}").json()
        assert asset["original_name"] == "原始表.csv" and asset["source"]["root_id"] == "data"
        assert client.get(f"/api/platform/v1/workspaces/{first}/assets/{asset_id}/content",
                          params={"revision": 1}).content == "编号,金额\n甲,10\n".encode()


def test_source_change_blocks_publication(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKBENCH_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("WORKBENCH_DATA_PYTHON", raising=False)
    with TestClient(app, base_url="http://127.0.0.1:8766") as client:
        workspace_id = client.post("/api/platform/v1/workspaces", headers=LOCAL,
                                   json={"schema_version": "1.0", "name": "甲"}).json()["workspace_id"]
        session_id, export_id, _, source = finished_export(app.state.data_workbench)
        source.write_bytes(b"changed")
        base = f"/api/platform/v1/workspaces/{workspace_id}/artifacts"
        response = client.post(base + "/data-exports", headers=LOCAL,
                               json={"session_id": session_id, "export_id": export_id})
        assert response.status_code == 409
        assert client.get(base).json()["total"] == 0


def test_fixed_asset_set_stages_only_its_pinned_member(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKBENCH_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("WORKBENCH_DATA_PYTHON", raising=False)
    with TestClient(app, base_url="http://127.0.0.1:8766") as client:
        first = client.post("/api/platform/v1/workspaces", headers=LOCAL,
                            json={"schema_version": "1.0", "name": "甲"}).json()["workspace_id"]
        second = client.post("/api/platform/v1/workspaces", headers=LOCAL,
                             json={"schema_version": "1.0", "name": "乙"}).json()["workspace_id"]
        base = f"/api/platform/v1/workspaces/{first}/assets"
        asset = client.post(base + "/uploads?name=原始台账.csv", headers=LOCAL,
                            content="编号,金额\n甲,10\n".encode()).json()
        fixed = client.post(base + "/sets", headers=LOCAL,
                            json={"asset_ids": [asset["asset_id"]]}).json()
        session = client.post("/api/platform/v1/data/sessions", headers=LOCAL,
                              json={"name": "资料集导入"}).json()
        path = f"/api/platform/v1/data/sessions/{session['id']}/asset-uploads"
        payload = {"workspace_id": first, "asset_set_id": fixed["asset_set_id"],
                   "asset_set_revision": fixed["revision"], "item_id": fixed["items"][0]["item_id"]}
        assert client.post(path, json=payload).status_code == 403
        staged = client.post(path, headers=LOCAL, json=payload)
        assert staged.status_code == 201, staged.text
        upload = app.state.data_workbench.upload_source(session["id"], staged.json()["id"])
        assert upload["source_workspace_id"] == first
        assert upload["source_asset_id"] == asset["asset_id"]
        assert upload["source_asset_revision"] == 1
        assert upload["source_blob_hash"] == asset["blob_hash"]
        assert client.post(path, headers=LOCAL, json={**payload, "workspace_id": second}).status_code == 409


@pytest.fixture
def review_client(tmp_path, monkeypatch):
    monkeypatch.setenv('WORKBENCH_DATA_DIR', str(tmp_path))
    monkeypatch.delenv('WORKBENCH_DATA_PYTHON', raising=False)
    with TestClient(app, base_url='http://127.0.0.1:8766') as client:
        yield client


def review_batch(client, count=41, *, complete=True, status='succeeded'):
    ws = client.post('/api/platform/v1/workspaces', headers=LOCAL,
        json={'schema_version': '1.0', 'name': '固定复核'}).json()['workspace_id']
    source = client.post(f'/api/platform/v1/workspaces/{ws}/assets/uploads', headers=LOCAL,
        params={'name': 'source.csv'}, content=b'id,amount\n001,1\n').json()
    store = app.state.batch_service.store
    template = store.save_template({'schema_version': '1.0', 'workspace_id': ws,
        'name': '完整复核', 'scenario': 'data_compare', 'config': {'fixture': 'fixed ledger'}})
    batch, _ = store.accept({'schema_version': '1.0', 'workspace_id': ws, 'request_key': 'reviewRun',
        'template_id': template['template_id'], 'template_version': 1, 'inputs': [{'role': 'left',
        'asset_id': source['asset_id'], 'asset_revision': source['revision'], 'blob_hash': source['blob_hash']}]}, template)
    items = [{'item_id': 'issue' + str(index), 'category': 'different', 'title': '金额差异',
        'source': {'result_id': 'fixedResult', 'asset_id': source['asset_id'], 'blob_hash': source['blob_hash']},
        'raw_output': {'ordinal': index, 'category': 'different', 'key_json': '00123456789012345678',
                       'text': '=formula' if index == 0 else str(index)}, 'disposition': 'pending', 'note': ''}
        for index in range(count)]
    store.begin_index(batch['batch_id'], 'fixedResult', count + (not complete), count + (not complete), count, count)
    store.index_page(batch['batch_id'], 'fixedResult', 0, [{'ordinal': i} for i in range(count)], items)
    store.update(batch['batch_id'], status=status)
    return app.state.batch_service.get(batch['batch_id']), items


def review_request(batch, request_id='reviewA', kind='intermediate'):
    return {'request_id': request_id, 'expected_revision': batch['revision'],
        'actor': {'id': 'localUser', 'display_name': '本机编制人', 'source': 'local'}, 'kind': kind}


def review_create(client, batch, payload):
    return client.post(f"/api/platform/v1/batches/{batch['batch_id']}/review-artifacts", headers=LOCAL, json=payload)


def review_file(client, batch, bundle, name):
    return client.get(f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts/{bundle['bundle_id']}/files/{name}")


def test_review_api_full_snapshot_formats_gate_replay_and_successor(review_client):
    client = review_client
    batch, items = review_batch(client)
    assert review_create(client, batch, review_request(batch, 'prematureFinal', 'final')).status_code == 409
    url = f"/api/platform/v1/batches/{batch['batch_id']}/review-artifacts"
    assert client.post(url, json=review_request(batch)).status_code == 403
    payload = review_request(batch)
    created = review_create(client, batch, payload)
    assert created.status_code == 201, created.text
    response = created.json()
    ProcessingBatch.model_validate(response)
    bundle = response['artifacts'][-1]
    ReviewBundle.model_validate(bundle)
    assert bundle['unresolved_count'] == 41 and len(bundle['unresolved_preview']) == 30
    assert response['delivery']['kind'] == 'intermediate' and bundle['status'] == 'current'
    assert review_create(client, batch, payload).json() == response
    assert review_create(client, batch, {**payload, 'kind': 'final'}).status_code == 409
    assert review_create(client, batch, {**payload, 'actor': {**payload['actor'], 'id': 'anotherUser'}}).status_code == 409
    csv_data = review_file(client, batch, bundle, 'results.csv').content
    rows = list(csv.DictReader(io.StringIO(csv_data.decode('utf-8-sig'))))
    assert len(rows) == 41 and rows[0]['adopted_value'] == "'=formula"
    workbook = load_workbook(io.BytesIO(review_file(client, batch, bundle, 'results.xlsx').content), read_only=True)
    try:
        assert workbook.active.max_row is None or workbook.active.max_row == 42
        xlsx_rows = list(workbook.active.iter_rows())
        assert len(xlsx_rows) == 42 and all(cell.data_type != 'f' for row in xlsx_rows for cell in row)
    finally:
        workbook.close()
    fixed = review_file(client, batch, bundle, 'review.json').json()
    assert len(fixed['items']) == len(fixed['unresolved']) == 41 and fixed['items'][-1]['item_id'] == 'issue40'
    old_hashes = {file['name']: sha(review_file(client, batch, bundle, file['name']).content) for file in bundle['files']}
    confirmed = app.state.batch_service.dispositions(batch['batch_id'], {'request_id': 'confirmFull',
        'expected_revision': batch['revision'], 'actor': payload['actor'], 'action': 'confirm',
        'item_ids': [item['item_id'] for item in items], 'note': '逐项核对原件'})
    final = review_create(client, confirmed, review_request(confirmed, 'reviewFinal', 'final'))
    assert final.status_code == 201, final.text
    result = final.json()
    ProcessingBatch.model_validate(result)
    newer = result['artifacts'][-1]
    ReviewBundle.model_validate(newer)
    assert newer['bundle_id'] != bundle['bundle_id'] and newer['predecessor_bundle_id'] == bundle['bundle_id']
    assert review_file(client, batch, newer, 'history.json').json()['items'][-1]['revision'] == 2
    old_get = client.get(f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts/{bundle['bundle_id']}").json()
    ReviewBundle.model_validate(old_get)
    assert old_get['status'] == 'outdated' and old_get['successor_bundle_ids'] == [newer['bundle_id']]
    assert {name: sha(review_file(client, batch, bundle, name).content) for name in old_hashes} == old_hashes
    reopened = app.state.batch_service.dispositions(batch['batch_id'], {'request_id': 'reopenOne',
        'expected_revision': result['revision'], 'actor': payload['actor'], 'action': 'reopen', 'item_ids': ['issue40']})
    assert reopened['delivery']['status'] == 'outdated'
    latest = client.get(f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts/{newer['bundle_id']}").json()
    assert latest['status'] == 'outdated'
    assert review_create(client, batch, payload).json() == response  # original immutable receipt


@pytest.mark.parametrize('change', ('revision', 'same_revision_status', 'same_revision_index', 'pending_operation'))
def test_review_generation_change_cannot_attach_current(review_client, monkeypatch, change):
    from modules.data_workbench.exporting import results as renderer
    client = review_client
    batch, items = review_batch(client, count=2, complete=change != 'same_revision_index')
    original = renderer.render_review_export
    def during_generation(*args, **kwargs):
        files = original(*args, **kwargs)
        store = app.state.batch_service.store
        if change == 'revision':
            app.state.batch_service.dispositions(batch['batch_id'], {'request_id': 'concurrentConfirm',
                'expected_revision': batch['revision'], 'actor': review_request(batch)['actor'], 'action': 'confirm', 'item_ids': ['issue0']})
        elif change == 'same_revision_status':
            store.update(batch['batch_id'], status='partial')
        elif change == 'same_revision_index':
            extra = {**items[0], 'item_id': 'issue2', 'raw_output': {'ordinal': 2, 'category': 'different'}}
            store.index_page(batch['batch_id'], 'fixedResult', 2, [{'ordinal': 2}], [extra])
        else:
            store.begin_operation(batch['batch_id'], {'request_id': 'unverifiedDuringReview', 'expected_revision': 1,
                'actor': review_request(batch)['actor'], 'action': 'confirm', 'item_ids': ['issue0']})
        return files
    monkeypatch.setattr(renderer, 'render_review_export', during_generation)
    result = review_create(client, batch, review_request(batch))
    assert result.status_code == 201, result.text
    assert result.json()['delivery']['status'] == 'outdated'
    current = app.state.batch_service.get(batch['batch_id'])
    assert current['delivery']['status'] == 'none'
    bundle = result.json()['artifacts'][-1]
    get = client.get(f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts/{bundle['bundle_id']}")
    assert get.status_code == 200 and get.json()['status'] == 'outdated'


def test_review_same_revision_later_execution_invalidates_current(review_client):
    batch, _ = review_batch(review_client, count=2)
    result = review_create(review_client, batch, review_request(batch)).json()
    bundle = result['artifacts'][-1]
    app.state.batch_service.store.update(batch['batch_id'], status='partial')
    current = app.state.batch_service.get(batch['batch_id'])
    assert current['revision'] == batch['revision'] and current['delivery']['status'] == 'outdated'
    get = review_client.get(f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts/{bundle['bundle_id']}")
    assert get.json()['status'] == 'outdated'


def test_review_interrupted_publication_hidden_and_exact_bytes_replay(review_client, monkeypatch):
    batch, _ = review_batch(review_client, count=2)
    center = app.state.artifact_center
    original = center.store.publish_managed
    attempts = []
    def interrupted(*args, **kwargs):
        attempts.append(1)
        if len(attempts) == 4:
            raise OSError('injected after three artifact mappings')
        return original(*args, **kwargs)
    monkeypatch.setattr(center.store, 'publish_managed', interrupted)
    payload = review_request(batch)
    result = review_create(review_client, batch, payload)
    assert result.status_code == 409
    operation = center.batch_store.begin_review(batch['batch_id'], payload)
    bundle_id = operation['bundle_id']
    prefix = f"/api/platform/v1/workspaces/{batch['workspace_id']}/artifacts"
    assert review_client.get(prefix).json()['total'] == 0
    assert review_client.get(prefix + '/' + bundle_id).status_code == 404
    assert review_client.get(prefix + '/' + bundle_id + '/files/results.csv').status_code == 404
    with sqlite3.connect(center.store.db_path) as db:
        mapping = json.loads(db.execute('SELECT files_json FROM batch_review_bundle').fetchone()[0])
        assert db.execute('SELECT count(*) FROM published_artifact').fetchone()[0] == len(mapping) == 3
    first_ids = dict(mapping)
    monkeypatch.setattr(center.store, 'publish_managed', original)
    recovered = review_create(review_client, batch, payload)
    assert recovered.status_code == 201, recovered.text
    bundle = recovered.json()['artifacts'][-1]
    assert bundle['bundle_id'] == bundle_id
    assert all({file['name']: file['artifact_id'] for file in bundle['files']}[name] == identity for name, identity in first_ids.items())
    assert review_create(review_client, batch, payload).json() == recovered.json()


def test_review_ready_before_batch_commit_recovers_original_bundle(review_client, monkeypatch):
    batch, _ = review_batch(review_client, count=2)
    service = app.state.batch_service
    finish = service.store.finish_review
    monkeypatch.setattr(service.store, 'finish_review', lambda *_: (_ for _ in ()).throw(OSError('commit interruption')))
    payload = review_request(batch)
    with pytest.raises(OSError, match='commit interruption'):
        service.review_export(batch['batch_id'], payload)
    visible = app.state.artifact_center.list(batch['workspace_id'], 0, 30)['items'][0]
    assert visible['status'] == 'outdated' and service.get(batch['batch_id'])['delivery']['status'] == 'none'
    with sqlite3.connect(app.state.artifact_center.store.db_path) as db:
        count = db.execute('SELECT count(*) FROM published_artifact').fetchone()[0]
    monkeypatch.setattr(service.store, 'finish_review', finish)
    recovered = service.review_export(batch['batch_id'], payload)
    assert recovered['delivery']['status'] == 'current' and recovered['delivery']['bundle_id'] == visible['bundle_id']
    with sqlite3.connect(app.state.artifact_center.store.db_path) as db:
        assert db.execute('SELECT count(*) FROM published_artifact').fetchone()[0] == count


def test_review_concurrent_publisher_has_one_bundle_and_mapping(review_client, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import threading
    from modules.data_workbench.exporting import results as renderer
    from workbench_platform.data_api.service import DataError
    batch, _ = review_batch(review_client, count=2)
    started, release = threading.Event(), threading.Event()
    original = renderer.render_review_export
    def paused(*args, **kwargs):
        started.set()
        assert release.wait(10)
        return original(*args, **kwargs)
    monkeypatch.setattr(renderer, 'render_review_export', paused)
    service, payload = app.state.batch_service, review_request(batch)
    with ThreadPoolExecutor(max_workers=2) as pool:
        future = pool.submit(service.review_export, batch['batch_id'], payload)
        assert started.wait(10)
        with pytest.raises(DataError, match='正在生成'):
            service.review_export(batch['batch_id'], payload)
        release.set()
        response = future.result(timeout=20)
    assert service.review_export(batch['batch_id'], payload) == response
    with sqlite3.connect(app.state.artifact_center.store.db_path) as db:
        assert db.execute('SELECT count(*) FROM batch_review_bundle').fetchone()[0] == 1
        assert db.execute('SELECT count(*) FROM published_artifact').fetchone()[0] == 7


def test_review_technical_unresolved_and_failed_execution_only_intermediate(review_client):
    batch, items = review_batch(review_client, count=2)
    service, actor = app.state.batch_service, review_request(batch)['actor']
    confirmed = service.dispositions(batch['batch_id'], {'request_id': 'resolvedResults', 'expected_revision': 1,
        'actor': actor, 'action': 'confirm', 'item_ids': [item['item_id'] for item in items]})
    technical = {'item_id': 'technicalCollection', 'category': 'artifact', 'title': '成果归集失败',
        'source': {}, 'raw_output': {}, 'disposition': 'pending', 'note': '保留失败原因'}
    service.store.update(batch['batch_id'], status='partial', issues=[technical])
    current = service.get(batch['batch_id'])
    assert current['coverage']['pending'] == 0 and current['summary']['pending'] == 1
    assert review_create(review_client, current, review_request(current, 'noPartialFinal', 'final')).status_code == 409
    service.store.update(batch['batch_id'], status='failed')
    assert review_create(review_client, current, review_request(current, 'noFailedFinal', 'final')).status_code == 409
    result = review_create(review_client, current, review_request(current, 'technicalIntermediate'))
    assert result.status_code == 201, result.text
    bundle = result.json()['artifacts'][-1]
    ReviewBundle.model_validate(bundle)
    assert bundle['unresolved_count'] == 1 and bundle['unresolved_preview'][0]['category'] == 'artifact'
    document = review_file(review_client, batch, bundle, 'review.json').json()
    assert document['unresolved'][0]['item_id'] == technical['item_id']
    assert document['constraints'][0]['kind'] == 'execution_incomplete'
    rows = list(csv.DictReader(io.StringIO(review_file(review_client, batch, bundle, 'results.csv').content.decode('utf-8-sig'))))
    assert len(rows) == 3 and rows[-1]['category'] == 'artifact'


def test_review_incomplete_coverage_and_pending_operation_cannot_final(review_client):
    batch, items = review_batch(review_client, count=2, complete=False)
    service, actor = app.state.batch_service, review_request(batch)['actor']
    confirmed = service.dispositions(batch['batch_id'], {'request_id': 'resolveIndexed', 'expected_revision': 1,
        'actor': actor, 'action': 'confirm', 'item_ids': [item['item_id'] for item in items]})
    assert confirmed['coverage']['pending'] == 0 and confirmed['coverage']['index_complete'] is False
    assert review_create(review_client, confirmed, review_request(confirmed, 'noIncompleteFinal', 'final')).status_code == 409
    extra = {**items[0], 'item_id': 'issue2', 'disposition': 'confirmed', 'raw_output': {'ordinal': 2, 'category': 'different'}}
    service.store.index_page(batch['batch_id'], 'fixedResult', 2, [{'ordinal': 2}], [extra])
    service.store.begin_operation(batch['batch_id'], {'request_id': 'unverified', 'expected_revision': 2,
        'actor': actor, 'action': 'reopen', 'item_ids': ['issue0']})
    current = service.get(batch['batch_id'])
    assert current['coverage']['index_complete'] is True and current['summary']['pending_operations'] == 1
    assert review_create(review_client, current, review_request(current, 'noUnverifiedFinal', 'final')).status_code == 409
    result = review_create(review_client, current, review_request(current, 'pendingIntermediate'))
    assert result.status_code == 201, result.text
    bundle = result.json()['artifacts'][-1]
    ReviewBundle.model_validate(bundle)
    assert bundle['pending_operations'] == 1 and bundle['unresolved_count'] == 0
    document = review_file(review_client, batch, bundle, 'review.json').json()
    assert document['constraints'][0]['kind'] == 'operation_unverified'
