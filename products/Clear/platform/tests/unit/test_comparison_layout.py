"""Evidence boundaries, independent page requests, rendering and recovery."""
import base64
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import subprocess
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from workbench_platform.comparisons.layout import (
    LAYOUT_SCHEMA, LayoutCancelled, assemble_document, document_markdown,
    normalize_page, parse_page, prepare_pages, file_hash,
)
from workbench_platform.multimodal.runner import worker_environment


@pytest.fixture
def page(tmp_path):
    path = tmp_path / 'page.png'
    # Hash and one-page behavior are independent of bitmap decoding.
    path.write_bytes(b'page-one-pixels')
    return {'image_path': str(path), 'image_sha256': file_hash(path), 'media_type': 'image/png',
            'width': 1000, 'height': 1000, 'page_number': 1, 'page_id': 'page1',
            'source': {'asset_id': 'a', 'asset_revision': 1, 'blob_hash': 'f'*64, 'original_name': 'input.pdf'}}


@pytest.fixture
def local_runtime_inputs(tmp_path, monkeypatch):
    """Mocked runner tests must explicitly own their synthetic installation."""
    program, source, mineru = tmp_path/'program', tmp_path/'ocr-source', tmp_path/'mineru'
    members = [program/'bundles/ocr/runtimes/paddlevl/python.exe',
        program/'bundles/ocr/models/PP-DocLayoutV3/source-manifest.json',
        program/'bundles/ocr/models/PaddleOCR-VL-1.6/source-manifest.json',
        source/'src/ocr_workbench/worker.py', mineru/'source-manifest.json']
    for member in members:
        member.parent.mkdir(parents=True, exist_ok=True)
        member.write_text('synthetic fixture; never executed', encoding='utf-8')
    for name,value in {'WORKBENCH_PROGRAM_ROOT':program, 'WORKBENCH_OCR_SOURCE':source,
        'CLEAR_MINERU_PRO_PYTHON':tmp_path/'python.exe', 'CLEAR_MINERU_PRO_MODEL':mineru}.items():
        monkeypatch.setenv(name,str(value))


def data(text='Phase A complete'):
    return {'blocks': [{'type': 'text', 'text': text, 'bbox': [.1,.2,.8,.3]}],
            'empty_page': False, 'content_complete': True}


def test_each_page_has_one_image_and_fresh_payload_and_verified_cache(page, tmp_path):
    calls = []
    def call(**kwargs):
        calls.append(kwargs)
        return {'data': data(), 'finish_reason': 'stop', 'usage': {'output_tokens': 10}, 'model': 'vision'}
    config = {'mode': 'multimodal', 'provider_id': 'test', 'provider_version': 1}
    first = parse_page(page, config, tmp_path / 'results', call_vlm=call)
    second_page = {**page, 'page_number': 2, 'page_id': 'page2'}
    second = parse_page(second_page, config, tmp_path / 'results', call_vlm=call)
    assert len(calls) == 2
    assert calls[0]['payload'] == calls[1]['payload'] == {}
    assert calls[0]['system'] == calls[1]['system']
    assert not any('messages' in value or 'history' in value for value in calls)
    assert calls[0]['image']['data'] == base64.b64encode(b'page-one-pixels').decode()
    assert first['status'] == second['status'] == 'succeeded'
    assert first['blocks'][0]['source']['coordinate_origin'] == 'model_estimate'
    parse_page(page, config, tmp_path / 'results', call_vlm=call)
    assert len(calls) == 2
    parse_page(page, {**config, 'provider_version': 2}, tmp_path / 'results', call_vlm=call)
    assert len(calls) == 3
    Path(page['image_path']).write_bytes(b'tampered')
    with pytest.raises(ValueError, match='changed'):
        parse_page(page, config, tmp_path / 'results', call_vlm=call)


@pytest.mark.parametrize('box', [[-1,0,1,1], [0,0,float('nan'),1], [1,0,0,1], [0,0,200,100], None])
def test_model_bad_box_falls_back_to_page_reference(page, box):
    raw = data()
    raw['blocks'][0]['bbox'] = box
    normalized = normalize_page(raw, page, {'engine': 'paddlevl'})
    assert normalized['blocks'][0]['source']['precision'] == 'page_only'
    assert normalized['blocks'][0]['bbox'] is None


def test_table_cells_never_inherit_block_coordinates(page):
    raw = {'empty_page': False, 'content_complete': True, 'blocks': [{'type': 'table', 'text': '', 'bbox': [0,0,1,1],
        'table': {'rows': 2, 'columns': 2, 'cells': [
            {'row': 0, 'column': 0, 'text': 'Progress', 'column_span': 2},
            {'row': 1, 'column': 0, 'text': 'A'}, {'row': 1, 'column': 1, 'text': '50%'}]}}]}
    result = normalize_page(raw, page, {'engine': 'paddlevl'})
    assert result['blocks'][0]['text'] == 'Progress\tProgress\nA\t50%'
    assert all(c['source']['precision'] == 'page_only' and c['polygon'] is None for c in result['tables'][0]['cells'])
    invalid = deepcopy(raw)
    invalid['blocks'][0]['table']['cells'].append({'row': 0, 'column': 1, 'text': 'overlap'})
    with pytest.raises(ValueError, match='overlap'):
        normalize_page(invalid, page, {'engine': 'paddlevl'})


def test_ppt_hidden_or_unmatched_text_never_enters_evidence(page):
    page['native_candidates'] = [{'shape_id': 'visible', 'text': 'Phase A complete', 'bbox': [.1,.2,.8,.3]},
                                 {'shape_id': 'hidden', 'text': 'SECRET WRONG 99%', 'bbox': [.1,.2,.8,.3]}]
    result = normalize_page(data(), page, {'engine': 'paddlevl'})
    assert result['blocks'][0]['native_confirmation']['shape_id'] == 'visible'
    assert 'SECRET' not in result['blocks'][0]['text']


def test_sparse_unordered_and_merged_table_cells_keep_column_meaning(page):
    raw = {'empty_page': False, 'content_complete': True, 'blocks': [{'type': 'table', 'text': '',
        'table': {'rows': 3, 'columns': 3, 'cells': [
            {'row': 0, 'column': 2, 'text': 'Progress'}, {'row': 0, 'column': 0, 'text': 'Project'},
            {'row': 2, 'column': 2, 'text': '60%'}, {'row': 1, 'column': 0, 'row_span': 2, 'text': 'Atlas'},
            {'row': 1, 'column': 2, 'text': '50%'}]}}]}
    result = normalize_page(raw, page, {'engine': 'paddlevl'})
    assert result['blocks'][0]['text'] == 'Project\t\tProgress\nAtlas\t\t50%\nAtlas\t\t60%'


def test_no_ocr_text_is_incomplete_without_independent_blank_verification(page):
    result = normalize_page({'blocks': []}, page, {'engine': 'paddlevl'}, local=True)
    assert result['status'] == 'partial' and result['needs_ocr'] and not result['empty_page']
    blank = normalize_page({'blocks': []}, {**page, 'verified_blank': True}, {'engine': 'paddlevl'}, local=True)
    assert blank['status'] == 'succeeded' and blank['empty_page']
    claimed_blank = normalize_page({'blocks': [], 'empty_page': True, 'content_complete': True},
                                   page, {'engine': 'paddlevl'})
    assert claimed_blank['status'] == 'partial'
    model = data()
    model['content_complete'] = False
    assert normalize_page(model, page, {'engine': 'paddlevl'})['status'] == 'partial'
    model['content_complete'] = True
    model['diagnostics'] = ['An unreadable line was omitted']
    assert normalize_page(model, page, {'engine': 'paddlevl'})['status'] == 'partial'


def test_truncation_does_not_become_success_or_repeat_outside_llm_layer(page, tmp_path):
    call = Mock(return_value={'data': data(), 'finish_reason': 'length'})
    result = parse_page(page, {'mode': 'multimodal'}, tmp_path / 'r', call_vlm=call)
    assert result['status'] == 'failed' and result['needs_ocr']
    assert call.call_count == 1
    assert not result['blocks']


def test_unknown_and_cancellation_never_start_fallback(page, tmp_path):
    from workbench_platform.agent.providers import ProviderFault
    fallback = Mock()
    for error in (ProviderFault('response_unknown', 'unknown'), ProviderFault('cancelled', 'cancel'), TimeoutError()):
        with pytest.raises(type(error)):
            parse_page(page, {'mode': 'multimodal', 'fallback_local': True}, tmp_path / 'r',
                       call_vlm=Mock(side_effect=error), call_local=fallback)
    fallback.assert_not_called()
    class Cancelled(Exception):
        pass
    with pytest.raises(Cancelled):
        parse_page(page, {'mode': 'local_ocr'}, tmp_path / 'r', call_local=Mock(side_effect=Cancelled()))


def test_explicit_local_fallback_records_actual_engine(page, tmp_path):
    raw = {'blocks': [{'kind': 'text', 'text': 'local 50%', 'polygon': [[0,0],[100,0],[100,100],[0,100]]}],
           'model_revisions': {'PaddleOCR-VL-1.6': 'pinned'}}
    result = parse_page(page, {'mode': 'multimodal', 'fallback_local': True}, tmp_path / 'r',
                        call_vlm=Mock(side_effect=ValueError('invalid JSON')), call_local=Mock(return_value=raw))
    assert result['status'] == 'succeeded' and result['engine'] == 'paddlevl'
    assert result['fallback_from'] == 'multimodal'
    assert result['blocks'][0]['source']['coordinate_origin'] == 'ocr_region'
    assert result['engine_identity']['model_revisions']['PaddleOCR-VL-1.6'] == 'pinned'


def test_incomplete_vlm_page_uses_opted_in_local_completion_and_keeps_response(page, tmp_path):
    incomplete = {**data(), 'content_complete': False, 'diagnostics': ['Unreadable table']}
    response = {'data': incomplete, 'raw': {'original': 'provider response'}, 'usage': {'output_tokens': 20}}
    local = Mock(return_value={'blocks': [{'kind': 'text', 'text': 'Completed local text'}]})
    result = parse_page(page, {'mode': 'multimodal', 'fallback_local': True}, tmp_path/'fallback-partial',
                        call_vlm=Mock(return_value=response), call_local=local)
    assert result['status'] == 'succeeded' and result['engine'] == 'paddlevl'
    assert result['multimodal_response'] == response
    assert result['multimodal_diagnostics'] == ['Unreadable table']
    assert local.call_count == 1


def test_document_revision_sources_and_markdown_keep_literal_text(page):
    parsed = normalize_page(data('literal <script>'), page, {'engine': 'paddlevl'})
    doc = assemble_document([parsed], page['source'], 'd1', 2)
    assert doc['kind'] == 'clear.document' and doc['coverage']['succeeded'] == 1
    assert doc['pages'][0]['blocks'][0]['source']['document_revision'] == 2
    assert 'literal &lt;script&gt;' in document_markdown(doc)
    assert 'parse_run_id' not in parsed['blocks'][0]['source']


def test_local_worker_uses_explicit_engine_and_shared_gpu_runner(page, tmp_path, local_runtime_inputs):
    from workbench_platform.comparisons.local_ocr import LocalOcrParser
    from workbench_platform.multimodal.runner import Outcome
    coordinator = Mock()
    coordinator.snapshot.return_value = SimpleNamespace(pending=())
    runner = Mock()
    observed = []
    def run(command, output, env, request, **kwargs):
        observed.append((command, env, request, kwargs))
        from workbench_platform.comparisons.layout import fingerprint
        for index, entry in enumerate(request['pages']):
            raw = {'status': 'success', 'blocks': [], 'engine': 'paddlevl'}
            target = Path(entry['checkpoint'])
            target.write_text(json.dumps({'format': 'clear-local-page-v1', 'page_id': entry['page_id'],
                'cache_key': entry['cache_key'], 'image_sha256': entry['image_sha256'], 'result': raw,
                'result_sha256': fingerprint(raw)}), encoding='utf-8')
            kwargs['committed']({'chunk_index': index, 'checkpoint': {'path': os.path.relpath(target, output)}})
        return Outcome('succeeded', 0)
    runner.run.side_effect = run
    manual = SimpleNamespace(coordinator=coordinator, runner=runner, launch_id='testlaunch', instance_id='testinstance')
    local = LocalOcrParser(SimpleNamespace(state=SimpleNamespace(multimodal_service=manual)))
    local(page, {'engine': 'paddlevl'}, tmp_path/'local')
    command, env, request, options = observed[0]
    assert request['engine'] == 'paddlevl' and request['options']['device'] == 'cuda'
    assert request['runner']['protocol'] == 'clear-gpu-worker-v1'
    assert coordinator.request.call_args.args[1] == 'manual'
    assert options['request_id'] == coordinator.request.call_args.args[0]
    assert all(Path(env[key]).drive.upper() == 'F:' for key in ('TEMP','TMP','APPDATA','LOCALAPPDATA','HF_HOME','PADDLE_PDX_CACHE_HOME'))
    assert '--worker' in command


def test_local_batch_reuses_worker_commits_each_page_and_recovers_checkpoints(page, tmp_path, local_runtime_inputs):
    from workbench_platform.comparisons.local_ocr import LocalOcrParser
    from workbench_platform.comparisons.layout import fingerprint
    from workbench_platform.multimodal.runner import Outcome
    coordinator = Mock()
    coordinator.snapshot.return_value = SimpleNamespace(pending=())
    runner = Mock()
    def run(command, output, env, request, **kwargs):
        for index, entry in enumerate(request['pages']):
            raw = {'status': 'success', 'blocks': [{'kind':'text','text':str(index)}], 'engine': 'paddlevl',
                   'load_seconds': 100 if index == 0 else 0}
            target = Path(entry['checkpoint'])
            target.write_text(json.dumps({'format': 'clear-local-page-v1', 'page_id': entry['page_id'],
                'cache_key': entry['cache_key'], 'image_sha256': entry['image_sha256'], 'result': raw,
                'result_sha256': fingerprint(raw)}), encoding='utf-8')
            try:
                kwargs['committed']({'chunk_index': index, 'checkpoint': {'path': os.path.relpath(target, output)}})
            except Exception:
                return Outcome('failed', 1, 'runner performed cleanup')
        return Outcome('succeeded', 0)
    runner.run.side_effect = run
    manual = SimpleNamespace(coordinator=coordinator, runner=runner, launch_id='testlaunch', instance_id='testinstance')
    local = LocalOcrParser(SimpleNamespace(state=SimpleNamespace(multimodal_service=manual)))
    pages = [page, {**page, 'page_id':'page2', 'page_number':2}]
    receipts = []
    result = local.parse_batch(pages, {'engine':'paddlevl'}, tmp_path/'batch', committed=lambda p,r: receipts.append(p['page_id']))
    assert receipts == ['page1','page2'] and list(result) == receipts and runner.run.call_count == 1
    assert result['page2']['load_seconds'] == 0
    receipts.clear()
    local.parse_batch(pages, {'engine':'paddlevl'}, tmp_path/'batch', committed=lambda p,r: receipts.append(p['page_id']))
    assert receipts == ['page1','page2'] and runner.run.call_count == 1
    source_error = ValueError('source changed during durable commit')
    def fail_commit(p,r):
        raise source_error
    with pytest.raises(ValueError) as error:
        local.parse_batch(pages, {'engine':'paddlevl'}, tmp_path/'other-batch', committed=fail_commit)
    assert error.value is source_error
    with pytest.raises(ValueError, match='20'):
        local.parse_batch([{**page,'page_id':f'p{i}'} for i in range(21)], {'engine':'paddlevl'}, tmp_path/'too-large')


def test_local_batch_gpu_yield_only_requeues_unfinished_pages(page, tmp_path, local_runtime_inputs):
    from workbench_platform.comparisons.local_ocr import LocalOcrParser
    from workbench_platform.comparisons.layout import fingerprint
    from workbench_platform.multimodal.runner import Outcome
    coordinator, runner = Mock(), Mock()
    coordinator.snapshot.return_value = SimpleNamespace(pending=(), owner_request_id=None)
    attempts = []
    def run(command, output, env, request, **kwargs):
        attempts.append([p['page_id'] for p in request['pages']])
        entry = request['pages'][0]
        raw = {'status':'success','blocks':[]}
        target = Path(entry['checkpoint'])
        target.write_text(json.dumps({'format':'clear-local-page-v1','page_id':entry['page_id'],
            'cache_key':entry['cache_key'],'image_sha256':entry['image_sha256'],'result':raw,
            'result_sha256':fingerprint(raw)}),encoding='utf-8')
        kwargs['committed']({'chunk_index':0,'checkpoint':{'path':os.path.relpath(target,output)}})
        return Outcome('yielded' if len(attempts)==1 else 'succeeded',75 if len(attempts)==1 else 0)
    runner.run.side_effect = run
    manual=SimpleNamespace(coordinator=coordinator,runner=runner,launch_id='testlaunch',instance_id='testinstance')
    local=LocalOcrParser(SimpleNamespace(state=SimpleNamespace(multimodal_service=manual)))
    results=local.parse_batch([page,{**page,'page_id':'page2'}],{'engine':'paddlevl'},tmp_path/'yield')
    assert attempts == [['page1','page2'],['page2']] and len(results)==2
    assert coordinator.request.call_args_list[0].args[0] != coordinator.request.call_args_list[1].args[0]


def test_ordered_image_source_count_must_match(page, tmp_path):
    with pytest.raises(ValueError, match='identities'):
        prepare_pages([page['image_path']], {'original_name': 'images', 'assets': []}, tmp_path/'render')


def test_renderer_releases_job_handle_even_when_drain_verification_fails(tmp_path, monkeypatch):
    from workbench_platform.comparisons import layout
    job = Mock()
    job.drain.side_effect = RuntimeError('cleanup unconfirmed')
    process = Mock(returncode=0)
    process.poll.return_value = 0
    monkeypatch.setattr(layout, 'WindowsJob', Mock(return_value=job))
    monkeypatch.setattr(layout.subprocess, 'Popen', Mock(return_value=process))
    with pytest.raises(RuntimeError, match='cleanup unconfirmed'):
        layout._run_renderer(['owned-renderer'], tmp_path, {}, None)
    job.close.assert_called_once()
    process.wait.assert_called_once()


@pytest.mark.skipif(not Path(os.environ.get('WORKBENCH_DOCUMENTS_PYTHON', '')).is_file(), reason='Explicit document runtime unavailable')
def test_real_pdf_ppt_hidden_slide_and_ordered_image_rendering(tmp_path):
    # All fixture generation, Office profiles and PDFium outputs stay in pytest's F: base temp.
    fixture_script = '''
from pathlib import Path
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
root=Path(sys.argv[1])
Image.new('RGB',(1200,800),'white').save(root/'one.png')
Image.new('RGB',(800,1200),'gray').save(root/'two.png')
Image.new('RGBA',(800,300),(255,0,0,0)).save(root/'transparent.png')
Image.open(root/'one.png').save(root/'scan.pdf',save_all=True,append_images=[Image.open(root/'two.png')])
p=Presentation()
for number in range(2):
 s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Slide '+str(number+1)
 if number==1:s._element.set('show','0')
 s.notes_slide.notes_text_frame.text='PRIVATE NOTES MUST NOT APPEAR'
p.save(root/'slides.pptx')
'''
    python = Path(os.environ['WORKBENCH_DOCUMENTS_PYTHON'])
    env = worker_environment(tmp_path / 'generator', python)
    subprocess.run([str(python), '-B', '-c', fixture_script, str(tmp_path)], check=True, env=env, cwd=tmp_path)
    base = {'asset_id': 'asset1', 'asset_revision': 1, 'blob_hash': 'a'*64}
    pdf = prepare_pages(tmp_path/'scan.pdf', {**base,'original_name':'scan.pdf'}, tmp_path/'pdf')
    assert len(pdf) == 2 and pdf[0]['width'] > pdf[0]['height'] and pdf[1]['height'] > pdf[1]['width']
    ppt = prepare_pages(tmp_path/'slides.pptx', {**base,'original_name':'slides.pptx'}, tmp_path/'ppt')
    assert len(ppt) == 2 and ppt[1]['hidden'] is True and ppt[1]['slide_number'] == 2
    assert 'PRIVATE NOTES' not in json.dumps([p.get('native_candidates') for p in ppt])
    from pypdf import PdfReader
    visible_pdf = PdfReader(tmp_path/'ppt/office-export/source.pdf')
    assert 'Slide 2' in visible_pdf.pages[1].extract_text()
    assert 'PRIVATE NOTES' not in ''.join(p.extract_text() for p in visible_pdf.pages)
    images = prepare_pages([tmp_path/'two.png',tmp_path/'one.png'],
        {**base, 'original_name': '页面图片组', 'assets': [{**base, 'asset_id':'two'}, {**base, 'asset_id':'one'}]}, tmp_path/'images')
    assert [p['source']['asset_id'] for p in images] == ['two','one']
    assert images[0]['height'] > images[0]['width']
    assert all(Path(p['image_path']).is_relative_to(tmp_path) and file_hash(p['image_path']) == p['image_sha256'] for p in [*pdf,*ppt,*images])
    from workbench_platform.comparisons.layout import resize_page
    smaller = resize_page(images[0], {'mode':'multimodal','max_image_edge':512,'image_max_bytes':65536}, tmp_path/'resized')
    assert max(smaller['width'], smaller['height']) == 512
    assert smaller['source']['asset_id'] == 'two'
    assert smaller['original_image_sha256'] == images[0]['image_sha256']
    assert file_hash(images[0]['image_path']) == images[0]['image_sha256']
    transparent = prepare_pages(tmp_path/'transparent.png', {**base,'original_name':'transparent.png'}, tmp_path/'alpha')
    assert transparent[0]['verified_blank'] is True
    pixel_check = "from PIL import Image;import sys;im=Image.open(sys.argv[1]);assert im.getpixel((5,5))==(255,255,255)"
    subprocess.run([str(python),'-B','-c',pixel_check,transparent[0]['image_path']],check=True,env=env,cwd=tmp_path)
