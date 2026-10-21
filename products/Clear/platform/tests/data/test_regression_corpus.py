"""Independent row-level oracle for the complete tabular workflow."""

import csv
import hashlib
import json
from pathlib import Path
import shutil

import duckdb
from openpyxl import load_workbook

from modules.data_workbench.importing import import_xlsx
from modules.data_workbench.importing.parquet import import_parquet
from modules.data_workbench.interactive_worker import do_operation
from modules.data_workbench.recipes import RecipeStore


FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "data"


def test_cross_file_recipe_comparison_and_export_oracle(tmp_path):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    state = {"tables": {}, "recipes": [], "comparisons": [], "exports": [], "result_id": None}
    counter = 0

    def run(action, payload, upload=None):
        nonlocal state, counter
        counter += 1
        request = {"session_dir": str(tmp_path), "state": state, "payload": payload,
                   "action": action, "job_id": f"j{counter}"}
        if upload:
            source = FIXTURES / "regression" / upload
            destination = uploads / upload
            shutil.copy2(source, destination)
            request["upload"] = {"id": f"u{counter}", "name": upload,
                                 "relative": f"uploads/{upload}", "format": "csv",
                                 "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}
        outcome = do_operation(request)
        state = outcome["state"]
        return outcome

    run("import", {"side": "left", "options": {"header_row": 1}}, "left.csv")
    assert [column["header"] for column in state["tables"]["left"]["columns"]] == [
        "编号", "金额", "日期", "备注"]
    run("step", {"side": "left", "operator": "trim", "parameters": {"column": "c2"}})
    run("step", {"side": "left", "operator": "cast",
                 "parameters": {"column": "c2", "target": "decimal"}})
    recipe_id = run("save_recipe", {"side": "left", "name": "金额标准化"})["result"]["recipe_id"]
    run("import", {"side": "right", "options": {"header_row": 1}}, "right.csv")
    run("replay", {"side": "right", "recipe_id": recipe_id,
                   "mapping": {"c1": "c2", "c2": "c3", "c3": "c1", "c4": "c4"}})
    store = RecipeStore(tmp_path / "recipes.sqlite3")
    left = store.get(state["tables"]["left"]["revision"])["snapshot"]
    right = store.get(state["tables"]["right"]["revision"])["snapshot"]
    assert left["rows"][0]["cells"][0]["value"] == "00123456789012345678"
    assert left["rows"][0]["cells"][1]["value"] == "10.00"
    assert left["rows"][1]["cells"][2]["raw"] == "2026/09/28"
    assert right["rows"][0]["cells"][2]["value"] == "10.00"
    assert right["rows"][1]["cells"][2]["value"] == "21.00"
    assert left["rows"][4]["cells"][0]["kind"] == "null"

    run("compare", {"spec": {"key_pairs": [["c1", "c2"]],
                             "compare_pairs": [["c2", "c3"], ["c4", "c4"]],
                             "amount_pairs": [["c2", "c3"]],
                             "amount_tolerance": "0"}})
    result_id = state["result_id"]
    expected = {"left_total": 5, "right_total": 5,
                "left_categorized": 5, "right_categorized": 5,
                "same_pairs": 1, "different_pairs": 1,
                "only_left": 1, "only_right": 1,
                "ambiguous_keys": 1, "ambiguous_left_rows": 2, "ambiguous_right_rows": 2}
    assert state["comparisons"][0]["coverage"] == expected
    run("export", {"result_id": result_id})
    exported = state["exports"][0]
    files = {item["name"]: tmp_path / item["relative"] for item in exported["files"]}
    assert set(files) == {"results.csv", "results.xlsx", "results.parquet",
                          "summary.json", "provenance.json"}
    for item in exported["files"]:
        with files[item["name"]].open("rb") as stream:
            assert hashlib.file_digest(stream, "sha256").hexdigest() == item["sha256"]
    summary = json.loads(files["summary.json"].read_text(encoding="utf-8"))
    assert summary["coverage"] == expected
    with files["results.csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["category"] for row in rows] == [
        "same", "different", "only_left", "only_right",
        "ambiguous_left", "ambiguous_left", "ambiguous_right", "ambiguous_right"]
    assert rows[0]["L_c1_value"] == rows[0]["R_c2_value"] == "00123456789012345678"
    assert rows[1]["L_c2_value"] == "20.00" and rows[1]["R_c3_value"] == "21.00"
    workbook = load_workbook(files["results.xlsx"], read_only=True, data_only=True)
    try:
        assert sum(1 for _ in workbook.active.values) == 9
    finally:
        workbook.close()
    with duckdb.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM read_parquet(?)",
                          [str(files["results.parquet"])]).fetchone()[0] == 8
    assert len(json.loads(files["provenance.json"].read_text(encoding="utf-8"))["rows"]) == 8


def test_frozen_xlsx_and_parquet_preserve_long_ids_and_loss_markers():
    path = FIXTURES / "tables" / "match.parquet"
    parquet = import_parquet(path, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    assert len(parquet.rows) == 3
    assert parquet.rows[0].cells[1].value == "000123456789012345678901234567"
    assert parquet.rows[0].cells[4].kind == "null"
    assert parquet.rows[1].cells[4].kind == "text"
    assert parquet.rows[1].cells[4].value == ""
    xlsx = FIXTURES / "tables" / "对账.xlsx"
    workbook = load_workbook(xlsx, read_only=True, data_only=False)
    try:
        sheet = workbook.sheetnames[0]
    finally:
        workbook.close()
    table = import_xlsx(xlsx, source_sha256=hashlib.sha256(xlsx.read_bytes()).hexdigest(),
                        sheet_name=sheet)
    assert table.rows[0].cells[1].value == "000123456789012345678901234567"
    assert table.rows[0].cells[4].kind == "unknown_formula"
    assert table.rows[0].cells[4].formula == "=SUM(1,2)"


def test_interactive_old_amount_result_keeps_identity_when_new_comparison_is_exported(tmp_path):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    state = {"tables": {}, "recipes": [], "comparisons": [], "exports": [], "result_id": None}
    counter = 0

    def run(action, payload, *, filename=None, content=None):
        nonlocal state, counter
        counter += 1
        request = {"session_dir": str(tmp_path), "state": state, "payload": payload,
                   "action": action, "job_id": f"legacyExport{counter}"}
        if filename:
            path = uploads / filename
            path.write_text(content, encoding="utf-8")
            request["upload"] = {"id": f"legacyUpload{counter}", "name": filename,
                "relative": "uploads/" + filename, "format": "csv",
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        state = do_operation(request)["state"]

    run("import", {"side": "left"}, filename="left.csv",
        content='id,amount\n001234567890,"1,234.50"\nK,(20.00)\n')
    run("import", {"side": "right"}, filename="right.csv",
        content='id,amount\n001234567890,1234.50\nK,-20.00\n')
    spec = {"key_pairs": [["c1", "c1"]], "compare_pairs": [["c2", "c2"]],
            "amount_pairs": [["c2", "c2"]], "amount_tolerance": "0.01"}
    run("compare", {"spec": {**spec, "amount_format": "plain-decimal-v1"}})
    legacy_id = state["result_id"]
    assert "amount_format" not in state["comparison"]["spec"]
    assert state["comparisons"][0]["coverage"]["different_pairs"] == 2
    run("compare", {"spec": spec})
    current_id = state["result_id"]
    assert current_id != legacy_id
    assert state["comparison"]["spec"]["amount_format"] == "financial-v1"
    assert state["comparisons"][0]["coverage"]["same_pairs"] == 2
    # Exporting the old result reloads its unversioned saved policy, rather
    # than selecting the new default or overwriting the old comparison ID.
    run("export", {"result_id": legacy_id})
    exported = state["exports"][0]
    files = {item["name"]: tmp_path / item["relative"] for item in exported["files"]}
    summary = json.loads(files["summary.json"].read_text("utf-8"))
    assert summary["comparison_id"] == legacy_id
    assert "amount_format" not in summary["policy"]
    assert summary["coverage"]["different_pairs"] == 2
    with files["results.csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["category"] for row in rows] == ["different", "different"]
    assert rows[0]["L_c1_value"] == "001234567890"
    assert rows[0]["L_c2_value"] == "1,234.50"
    assert rows[1]["L_c2_value"] == "(20.00)"


def test_real_financial_producer_marks_data_before_publication_and_blocks_previous_reader(tmp_path, monkeypatch):
    import importlib.util
    import time
    import pytest
    from workbench_platform.data_api.service import DataWorkbenchService
    from workbench_platform.resources import reader_features
    from scripts.product_compatibility import check, DATA_COMPATIBILITY

    python = Path('F:/OCR/release/Clear/versions/0.10.0/program/runtimes/data/python.exe')
    previous = Path('F:/OCR/output/clear-mineru-integration-20261021/continuation-20261021-001/delivery/Clear-0.17.0-documents-mineru-candidate-005')
    if not python.is_file() or not previous.is_dir():
        pytest.skip('Pinned local worker and previous candidate are required')
    data = tmp_path / 'data'
    service = DataWorkbenchService(data / 'data-workbench', python)
    marker = data / 'product-reader-requirements.json'
    session = service.create('金额历史与新格式门禁')
    session_id = session['id']

    def operation(action, payload):
        current = service.get(session_id)
        job = service.submit(session_id, action, payload, current['version'])
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            job = service.job(session_id, job['id'])
            if job['status'] not in {'queued', 'running', 'cancelling'}:
                assert job['status'] == 'succeeded', job
                return service.get(session_id)
            time.sleep(.05)
        pytest.fail('Real data worker did not finish')

    try:
        for side, content in (('left', 'id,amount\n00123,"1,234.50"\nK,(20.00)\n'),
                              ('right', 'id,amount\n00123,1234.50\nK,-20.00\n')):
            path = service.session_dir(session_id) / 'uploads' / (side + '.csv')
            path.write_text(content, encoding='utf-8')
            uploaded = service.upload(session_id, 'u' + side, path.name, '.csv', path,
                hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_size)
            operation('import', {'side': side, 'upload_id': uploaded['id']})
        spec = {'key_pairs': [['c1', 'c1']], 'compare_pairs': [['c2', 'c2']],
                'amount_pairs': [['c2', 'c2']], 'amount_tolerance': '0.01'}
        legacy = operation('compare', {'spec': {**spec, 'amount_format': 'plain-decimal-v1'}})
        legacy_id = legacy['state']['result_id']
        assert not marker.exists()
        operation('export', {'result_id': legacy_id})
        assert not marker.exists(), 'Exporting a saved legacy policy must not upgrade DATA'
        original = reader_features.require_features
        observed = []

        def mark_before_result(root, *features):
            assert service.get(session_id)['state']['result_id'] == legacy_id
            assert 'data-financial-amount-v1' in features
            observed.append(True)
            return original(root, *features)

        monkeypatch.setattr(reader_features, 'require_features', mark_before_result)
        from workbench_platform.data_api import service as service_module
        original_popen = service_module.subprocess.Popen
        launched = []

        def child_after_marker(*args, **kwargs):
            assert marker.is_file(), 'Reader marker must precede child result/response persistence'
            launched.append(True)
            return original_popen(*args, **kwargs)

        monkeypatch.setattr(service_module.subprocess, 'Popen', child_after_marker)
        current = operation('compare', {'spec': spec})
        assert observed and launched and current['state']['result_id'] != legacy_id
        assert current['state']['comparison']['spec']['amount_format'] == 'financial-v1'
        assert 'data-financial-amount-v1' in json.loads(marker.read_text('utf-8'))['reader_features']
        saved_marker = marker.read_bytes()
        loader = importlib.util.spec_from_file_location('previous_financial_reader', previous / 'platform/scripts/product_compatibility.py')
        previous_reader = importlib.util.module_from_spec(loader)
        loader.loader.exec_module(previous_reader)
        previous_descriptor = json.loads((previous / 'candidate.json').read_text('utf-8'))['data_compatibility']
        with pytest.raises(ValueError, match='不能直接降级'):
            previous_reader.check(data, previous_descriptor)
        assert check(data, DATA_COMPATIBILITY)['compatible']
        assert marker.read_bytes() == saved_marker
    finally:
        service.close()
