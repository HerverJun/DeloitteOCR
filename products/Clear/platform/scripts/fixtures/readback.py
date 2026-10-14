"""Independent parser for fixture truth; no generator imports."""

import argparse
import csv
import json
from pathlib import Path


def inspect(path: Path, media_type: str) -> dict:
    if media_type == "image/png":
        from PIL import Image
        with Image.open(path) as image:
            image.load()
            return {"width": image.width, "height": image.height, "mode": image.mode,
                    "pixel_0_0": list(image.getpixel((0, 0))),
                    "pixel_63_31": list(image.getpixel((63, 31)))}
    if media_type == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return {"page_count": len(reader.pages),
                "page_text": [page.extract_text().strip() for page in reader.pages]}
    if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook
        formula_book = load_workbook(path, data_only=False, read_only=True)
        value_book = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = formula_book["对账"]
            rows = [list(row) for row in sheet.iter_rows(min_row=2, max_row=4, min_col=1, max_col=6,
                                                           values_only=True)]
            return {"sheet_names": formula_book.sheetnames, "rows": rows,
                    "formula_cell": sheet["E2"].value, "cached_value": value_book["对账"]["E2"].value,
                    "long_id_type": type(sheet["B2"].value).__name__}
        finally:
            formula_book.close()
            value_book.close()
    if media_type == "text/csv":
        with path.open(encoding="utf-8", newline="") as stream:
            rows = list(csv.DictReader(stream))
        return {"rows": rows, "null_token": "\\N", "empty_string_distinct":
                rows[0]["nullable"] == "\\N" and rows[1]["nullable"] == ""}
    if media_type == "application/vnd.apache.parquet":
        import pyarrow.parquet as pq
        table = pq.read_table(path)
        rows = [{key: str(value) if key == "amount" else value for key, value in row.items()}
                for row in table.to_pylist()]
        return {"rows": rows, "amount_type": str(table.schema.field("amount").type),
                "long_id_type": str(table.schema.field("long_id").type)}
    raise ValueError(f"unsupported media type: {media_type}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--path", type=Path, required=True)
    parser.add_argument("--media-type", required=True)
    args = parser.parse_args()
    print(json.dumps(inspect(args.path, args.media_type), ensure_ascii=False, sort_keys=True))
