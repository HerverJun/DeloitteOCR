"""Real local GPU PDF-to-XLSX workflow with a synthetic controller.

Uses a new workspace, the portable model assets and the normal GPU instance
lock. Does not contact a real controller or use private/sealed documents.
"""
import argparse
import asyncio
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[2]


def checksum(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def save(out, value):
    temporary = out / 'gpu-batch.tmp'
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), 'utf-8')
    temporary.replace(out / 'gpu-batch.json')


async def run(args, report):
    from ocr_workbench.store import Store
    from ocr_workbench.adapter import EngineAdapter
    from ocr_workbench.documents import Documents
    from ocr_workbench.task_queue import TaskQueue
    from ocr_workbench.application_services import ApplicationServices
    from ocr_workbench.agent.runtime import AgentRuntime
    from ocr_workbench.agent.providers import normalize_response
    from ocr_workbench.agent.store import digest
    from ocr_workbench.agent.selection import create_selection, grant_selection
    from openpyxl import load_workbook

    out, bundle = args.output, args.bundle
    store = Store(out / 'workspace')
    project = store.project('Synthetic public GPU batch')['id']
    # Worker code comes from the same immutable snapshot as the service code.
    queue = TaskQueue(store, bundle, adapter_factory=lambda b, engine, session: EngineAdapter(b, engine, session, worker_source=out / 'source/src'))
    documents = Documents(store, bundle, queue)
    doc = await asyncio.to_thread(documents.import_document, project, 'synthetic-invoices.pdf', out / 'input.pdf', 144)
    services = ApplicationServices(store, documents, SimpleNamespace(guard=store.file_lock), queue=queue, bundle=bundle)
    requests = []

    async def model(state, current):
        if state['step'] == 0:
            name, arguments = 'process_pages', {'document_id': doc['id'], 'page_numbers': list(range(1, args.pages + 1)), 'mode': 'ocr', 'engine': args.engine}
        elif state['step'] == 1:
            name, arguments = 'export_results', {'source': 'run_results', 'source_run_id': current['id'], 'format': 'xlsx', 'aggregate': True}
        else:
            name, arguments = None, None
        message = {'role': 'assistant', 'content': '已根据工具实际结果核对批次，并保存工作簿。'}
        if name:
            message = {'role': 'assistant', 'tool_calls': [{'id': 'call-' + str(state['step']), 'type': 'function', 'function': {'name': name, 'arguments': json.dumps(arguments)}}]}
        request, fresh = runtime.agent.begin_model_request(project, current['id'], state['generation'], state['step'], digest(message))
        if not fresh:
            return json.loads(request['normalized_response'])
        requests.append({'step': state['step'], 'tool': name, 'utc': datetime.now(timezone.utc).isoformat(),
                         'unfinished_stages': len(store.rows("SELECT id FROM document_stages WHERE kind='process' AND status NOT IN ('succeeded','failed','cancelled')"))})
        response = normalize_response('openai_chat_completions', {'choices': [{'message': message, 'finish_reason': 'tool_calls' if name else 'stop'}]}, request['request_id'])
        runtime.agent.finish_model_request(project, current['id'], state['generation'], request['request_id'], response)
        return response

    runtime = await AgentRuntime(services, policy={}, model_override=model).open()
    try:
        queue.start()
        documents.start()
        session = runtime.agent.create_session(project, {'client_request_id': 'batch', 'title': 'Real GPU PDF to workbook'})
        selection = create_selection(runtime.agent, project, {'document_ranges': [{'document_id': doc['id'], 'first_page': 1, 'last_page': args.pages}],
            'allow_processing': True, 'allow_export': True, 'engines': [args.engine]}, available_engines=[args.engine])
        current = runtime.agent.create_run(project, session['id'], {'client_request_id': 'process', 'content': '识别明确选择的PDF页并汇总Excel'},
            context={'selection': selection['selection'], 'selection_token': selection['selection_token']}, config={}, limits={})
        grant_selection(runtime.policy, current, selection['selection'])
        started = time.monotonic()
        runtime.schedule(current)
        deadline = started + args.timeout_seconds
        last = None
        while time.monotonic() < deadline:
            current = runtime.agent.run(project, current['id'])
            stage_counts = store.rows("SELECT status,COUNT(*) count FROM document_stages WHERE kind='process' GROUP BY status")
            job_counts = store.rows("SELECT engine,kind,status,COUNT(*) count FROM tasks GROUP BY engine,kind,status")
            report.update(run=current, stages=stage_counts, jobs=job_counts, model_requests=requests,
                          queue=queue.status(), elapsed_seconds=time.monotonic() - started)
            save(out, report)
            summary = [current['status'], stage_counts, job_counts]
            if summary != last:
                print(json.dumps({'run': current['status'], 'stages': stage_counts, 'jobs': job_counts}), flush=True)
                last = summary
            if current['status'] in ('completed', 'failed', 'cancelled', 'waiting_user', 'interrupted'):
                break
            await asyncio.sleep(1)
        assert current['status'] == 'completed', current
        assert current['outcome'] == 'success', current
        assert current['coverage']['completed_count'] == args.pages, current['coverage']
        assert len(requests) == 3 and all(item['unfinished_stages'] == 0 for item in requests), requests
        assert len(store.rows("SELECT id FROM tasks WHERE kind='region_ocr' AND status='succeeded'")) == args.pages
        assert len(store.rows("SELECT id FROM document_stages WHERE kind='process' AND status='succeeded'")) == args.pages
        artifacts = store.rows("SELECT * FROM agent_artifacts WHERE state='ready'")
        assert len(artifacts) == 1
        exported = runtime.artifacts.verify(project, artifacts[0]['id'])
        workbook = load_workbook(exported)
        try:
            values = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
            report['workbook'] = {'path': str(exported), 'sha256': checksum(exported), 'sheet_names': workbook.sheetnames,
                                  'observed_values': values[:100], 'cells_with_values': len(values),
                                  'expected_financial_values_present': {value: value in values for value in ('10.00', '20.00', '30.00')}}
            assert len(values) >= args.pages * 6, 'Workbook lacks the expected batch content'
        finally:
            workbook.close()
        manifest = json.loads(artifacts[0]['manifest'])
        assert len(manifest['results']) == args.pages
        assert [source['page_number'] for source in manifest['sources']] == list(range(1, args.pages + 1))
        assert all(args.engine in source['ocr_engines'] for source in manifest['sources'])
        report.update(status='pass', artifact_manifest=manifest, observer_error=runtime.observer_error,
                      model_accuracy_qualification='not_established', real_controller='not_tested')
        assert not runtime.observer_error
    finally:
        await runtime.close()
        await asyncio.to_thread(documents.stop)
        await asyncio.to_thread(queue.stop)
        report['shutdown'] = {'documents': documents.status(), 'gpu_queue': queue.status()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--bundle', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--pages', type=int, default=3)
    parser.add_argument('--engine', choices=['glm', 'paddlevl'], default='glm')
    parser.add_argument('--timeout-seconds', type=int, default=1800)
    args = parser.parse_args()
    if not 1 <= args.pages <= 100:
        parser.error('Pages must be 1-100')
    args.bundle, args.output = args.bundle.resolve(), args.output.resolve()
    out = args.output
    out.mkdir(parents=True, exist_ok=False)
    shutil.copytree(ROOT / 'src', out / 'source/src', ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
    shutil.copytree(ROOT / 'config', out / 'source/config', ignore=shutil.ignore_patterns('development-machine.json'))
    source = {p.relative_to(out / 'source').as_posix(): checksum(p) for p in (out / 'source').rglob('*') if p.is_file()}
    (out / 'source-manifest.json').write_text(json.dumps(source, indent=2), 'utf-8')
    sys.path.insert(0, str(out / 'source/src'))
    code = """from fpdf import FPDF
import sys
p=FPDF(unit='pt',format=(400,420));p.set_auto_page_break(False)
for n in range(1,int(sys.argv[2])+1):
 p.add_page();p.set_font('Helvetica',size=14);p.text(40,40,f'Invoice {n:04d}')
 rows=[['Item','Amount'],['Alpha','10.00'],['Beta','20.00'],['Total','30.00']]
 for r,row in enumerate(rows):
  for c,value in enumerate(row):
   x=40+c*150;y=65+r*50;p.rect(x,y,150,50);p.text(x+12,y+30,value)
p.output(sys.argv[1])
"""
    subprocess.run([str(args.bundle / 'runtimes/pdf/python.exe'), '-B', '-c', code, str(out / 'input.pdf'), str(args.pages)], check=True)
    gpu = subprocess.run(['nvidia-smi', '--query-gpu=name,memory.total,memory.used,utilization.gpu', '--format=csv,noheader'], capture_output=True, text=True)
    report = {'status': 'running', 'started_utc': datetime.now(timezone.utc).isoformat(), 'engine': args.engine, 'pages': args.pages,
        'gpu_at_start': gpu.stdout.strip(), 'gpu_query_exit': gpu.returncode, 'source_sha256': checksum(out / 'source-manifest.json'),
        'bundle_manifest_sha256': checksum(args.bundle / 'manifest.json'), 'input_sha256': checksum(out / 'input.pdf'),
        'controller': 'synthetic deterministic fixture', 'private_or_sealed_documents_used': False, 'real_controller': 'not_tested'}
    try:
        asyncio.run(run(args, report))
    except BaseException as error:
        report.update(status='failed', error=f'{type(error).__name__}: {error}')
        raise
    finally:
        report['finished_utc'] = datetime.now(timezone.utc).isoformat()
        save(out, report)


if __name__ == '__main__':
    main()
