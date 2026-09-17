"""Smoke-test only the installed package, including its real visual-review GPU queue.

The OCR seed and image are synthetic. A replacement, when offered, is adopted by
this audit script to simulate an explicit human action; the product never adopts
the model response automatically. This is a packaging test, not an accuracy test.
"""
import argparse
import ctypes
from ctypes import wintypes
import hashlib
import json
from pathlib import Path
import re
import secrets
import socket
import subprocess
import sys
import time
import traceback
import urllib.error
import urllib.request
import uuid


SEED_PROGRAM = r'''
import hashlib, json, sys
from pathlib import Path
import ocr_workbench
from PIL import Image, ImageDraw, ImageFont
from ocr_workbench.store import Store
from ocr_workbench.imaging import add_image

bundle, output = map(Path, sys.argv[1:3])
module = Path(ocr_workbench.__file__).resolve()
assert sys.flags.isolated and module.is_relative_to(bundle.resolve()), str(module)
store = Store(output / 'data')
project = store.project('Installed multimodal smoke test')
image = Image.new('RGB', (1200, 180), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 64)
truth, before = 'TOTAL 001.05', 'TOTAL 001.00'
draw.text((40, 45), truth, font=font, fill='black')
path = output / 'synthetic-line.png'
image.save(path)
photo = add_image(store, project['id'], 'synthetic-line.png', path)
version = store.one('versions', photo['active_version'])
task = store.enqueue(project['id'], [version['id']], ['ppocr'])[0]
assert store.claim()['id'] == task
store.complete(task, {'engine': 'ppocr', 'text': before, 'tables': [],
    'blocks': [{'kind': 'text', 'text': before, 'confidence': .5,
                'polygon': [[20,20],[1160,20],[1160,160],[20,160]]}],
    'image': {'width': 1200, 'height': 180}, 'project_image_version': version['id']})
result_id = store.one('tasks', task)['result_id']
seed = {'project_id': project['id'], 'image_id': photo['id'], 'version_id': version['id'],
        'result_id': result_id, 'before': before, 'printed_truth': truth,
        'ocr_seed_is_synthetic': True, 'module_path': str(module), 'version': ocr_workbench.__version__,
        'module_sha256': hashlib.sha256(module.read_bytes()).hexdigest(), 'isolated': bool(sys.flags.isolated)}
(output / 'seed.json').write_text(json.dumps(seed, ensure_ascii=False, indent=2), 'utf-8')
print(json.dumps(seed, ensure_ascii=False))
'''

EXPORT_READBACK_PROGRAM = r'''
import json, sys
from pathlib import Path
from openpyxl import load_workbook
root = Path(sys.argv[1])
report = json.loads((root / 'review-report.json').read_text('utf-8'))
book = load_workbook(root / 'review-report.xlsx', read_only=True, data_only=False)
try:
    cells = [cell for sheet in book for row in sheet.iter_rows() for cell in row]
    rows = list(book['校验清单'].iter_rows(values_only=True))
    assert len(rows) == len(report['proposals']) + 1
    assert not any(cell.data_type == 'f' for cell in cells)
    assert report['automatic_adoption'] is False
    assert all(any(row[0] == item['id'] and row[4] == item['before'] and row[5] == item['after'] for row in rows[1:]) for item in report['proposals'])
    print(json.dumps({'passed': True, 'sheets': book.sheetnames, 'review_rows': len(rows),
        'formula_cells': 0, 'proposals': len(report['proposals']), 'revision': report['revision']}, ensure_ascii=False))
finally:
    book.close()
'''


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), "utf-8")


def process_snapshot():
    """Read parentage with Win32 only; do not depend on host/site packages."""
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)

    class Entry(ctypes.Structure):
        _fields_ = [("dwSize", wintypes.DWORD), ("cntUsage", wintypes.DWORD),
                    ("pid", wintypes.DWORD), ("heap", ctypes.c_size_t),
                    ("module", wintypes.DWORD), ("threads", wintypes.DWORD),
                    ("parent", wintypes.DWORD), ("priority", wintypes.LONG),
                    ("flags", wintypes.DWORD), ("exe", wintypes.WCHAR * 260)]

    kernel.CreateToolhelp32Snapshot.argtypes = [wintypes.DWORD, wintypes.DWORD]
    kernel.CreateToolhelp32Snapshot.restype = wintypes.HANDLE
    kernel.Process32FirstW.argtypes = [wintypes.HANDLE, ctypes.POINTER(Entry)]
    kernel.Process32NextW.argtypes = [wintypes.HANDLE, ctypes.POINTER(Entry)]
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = kernel.CreateToolhelp32Snapshot(2, 0)
    if handle == ctypes.c_void_p(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())
    values = []
    try:
        entry = Entry()
        entry.dwSize = ctypes.sizeof(Entry)
        found = kernel.Process32FirstW(handle, ctypes.byref(entry))
        while found:
            values.append({"pid": int(entry.pid), "parent": int(entry.parent), "exe": entry.exe})
            found = kernel.Process32NextW(handle, ctypes.byref(entry))
    finally:
        kernel.CloseHandle(handle)
    return values


def process_identity(pid):
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel.OpenProcess.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel.GetProcessTimes.argtypes = [wintypes.HANDLE] + [ctypes.POINTER(wintypes.FILETIME)] * 4
    kernel.GetExitCodeProcess.argtypes = [wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD)]
    kernel.QueryFullProcessImageNameW.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
    handle = kernel.OpenProcess(0x1000, False, pid)
    if not handle:
        return None
    try:
        code = wintypes.DWORD()
        if not kernel.GetExitCodeProcess(handle, ctypes.byref(code)) or code.value != 259:
            return None
        created, exited, cpu_kernel, cpu_user = (wintypes.FILETIME() for _ in range(4))
        if not kernel.GetProcessTimes(handle, ctypes.byref(created), ctypes.byref(exited), ctypes.byref(cpu_kernel), ctypes.byref(cpu_user)):
            return None
        buffer, length = ctypes.create_unicode_buffer(32768), wintypes.DWORD(32768)
        executable = buffer.value if kernel.QueryFullProcessImageNameW(handle, 0, buffer, ctypes.byref(length)) else ""
        return {"pid": pid, "created": (created.dwHighDateTime << 32) | created.dwLowDateTime, "executable": executable}
    finally:
        kernel.CloseHandle(handle)


class InstalledAudit:
    def __init__(self, bundle, output):
        self.bundle, self.output = bundle, output
        self.runtime = bundle / "runtimes/service/python.exe"
        self.process = None
        self.known = {}
        self.opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        self.token = secrets.token_urlsafe(48)
        self.token_file = output / "session-token.txt"
        self.token_file.write_text(self.token, "utf-8")
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
            listener.bind(("127.0.0.1", 0))
            self.port = listener.getsockname()[1]
        self.base = f"http://127.0.0.1:{self.port}"

    def embedded(self, program, *args):
        result = subprocess.run([str(self.runtime), "-I", "-B", "-X", "utf8", "-c", program, *map(str, args)],
            cwd=self.bundle, capture_output=True, text=True, encoding="utf-8", timeout=90,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if result.returncode:
            raise RuntimeError(f"Installed Python failed ({result.returncode}):\n{result.stderr}\n{result.stdout}")
        return json.loads(result.stdout)

    def seed(self):
        return self.embedded(SEED_PROGRAM, self.bundle, self.output)

    def track(self):
        if not self.process:
            return
        records = process_snapshot()
        parents = {record["pid"] for record in self.alive_tracked()}
        if self.process.poll() is None:
            parents.add(self.process.pid)
        changed = True
        while changed:
            changed = False
            for record in records:
                if record["pid"] not in parents and record["parent"] in parents:
                    parents.add(record["pid"])
                    changed = True
        for record in records:
            if record["pid"] in parents and record["pid"] not in self.known:
                identity = process_identity(record["pid"])
                if identity:
                    self.known[record["pid"]] = {**record, **identity}

    def alive_tracked(self):
        return [record for record in self.known.values()
                if (current := process_identity(record["pid"])) and current["created"] == record["created"]]

    def http(self, path, method="GET", body=None, raw=False, timeout=30):
        headers = {"Authorization": "Bearer " + self.token}
        data = None
        if body is not None:
            data = json.dumps(body, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        request = urllib.request.Request(self.base + path, data=data, method=method, headers=headers)
        try:
            with self.opener.open(request, timeout=timeout) as response:
                content = response.read()
                metadata = {"status": response.status, "content_type": response.headers.get("Content-Type"),
                            "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        except urllib.error.HTTPError as error:
            raise RuntimeError(f"HTTP {error.code} {path}: {error.read().decode('utf-8', errors='replace')}") from error
        return (content, metadata) if raw else json.loads(content)

    def start(self):
        self.log = (self.output / "service.log").open("wb")
        self.command = [str(self.runtime), "-I", "-B", "-X", "utf8", "-m", "ocr_workbench.service",
                        "--bundle", str(self.bundle), "--data", str(self.output / "data"),
                        "--port", str(self.port), "--token-file", str(self.token_file)]
        self.process = subprocess.Popen(self.command, cwd=self.bundle, stdout=self.log, stderr=subprocess.STDOUT,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        deadline = time.monotonic() + 120
        while time.monotonic() < deadline:
            self.track()
            if self.process.poll() is not None:
                raise RuntimeError(f"Installed service exited during startup: {self.process.returncode}")
            try:
                health = self.http("/api/health", timeout=2)
                if health["status"] == "ready":
                    return health
            except (OSError, urllib.error.URLError):
                pass
            time.sleep(.3)
        raise TimeoutError("Installed service did not become ready within 120 seconds")

    def stop(self):
        evidence = {"shutdown_requested": False, "forced_cleanup": [], "surviving_processes": []}
        try:
            if self.process and self.process.poll() is None:
                self.track()
                try:
                    self.http("/api/shutdown", "POST", {}, timeout=10)
                    evidence["shutdown_requested"] = True
                except Exception as error:
                    evidence["shutdown_error"] = str(error)
                deadline = time.monotonic() + 45
                while time.monotonic() < deadline and (self.process.poll() is None or self.alive_tracked()):
                    self.track()
                    time.sleep(.25)
                if self.process.poll() is None:
                    self.process.terminate()
                    evidence["forced_cleanup"].append(self.process.pid)
                    self.process.wait(timeout=10)
            # Emergency cleanup is restricted to identities captured from this service tree.
            for record in self.alive_tracked():
                kernel = ctypes.WinDLL("kernel32", use_last_error=True)
                kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
                kernel.OpenProcess.restype = wintypes.HANDLE
                kernel.TerminateProcess.argtypes = [wintypes.HANDLE, wintypes.UINT]
                kernel.CloseHandle.argtypes = [wintypes.HANDLE]
                current = process_identity(record["pid"])
                if not current or current["created"] != record["created"]:
                    continue
                handle = kernel.OpenProcess(1, False, record["pid"])
                if handle:
                    try:
                        if kernel.TerminateProcess(handle, 1):
                            evidence["forced_cleanup"].append(record["pid"])
                    finally:
                        kernel.CloseHandle(handle)
            deadline = time.monotonic() + 5
            while time.monotonic() < deadline and self.alive_tracked():
                time.sleep(.1)
            evidence["surviving_processes"] = self.alive_tracked()
            evidence["service_returncode"] = self.process.poll() if self.process else None
            evidence["observed_processes"] = list(self.known.values())
        finally:
            if hasattr(self, "log"):
                self.log.close()
            self.token_file.unlink(missing_ok=True)
            evidence["token_file_removed"] = not self.token_file.exists()
            write_json(self.output / "shutdown.json", evidence)
        return evidence


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bundle", type=Path, default=Path("D:/OCR-multimodal-workbench-20260917/bundle"))
    parser.add_argument("--output", type=Path)
    parser.add_argument("--model", help="Configured profile; defaults to the installed catalog default")
    parser.add_argument("--timeout", type=int, default=600, help="Maximum seconds for real model inference")
    args = parser.parse_args()
    bundle = args.bundle.resolve()
    output = (args.output or bundle.parent / "installed-smoke").resolve()
    if output == bundle or output.is_relative_to(bundle):
        raise ValueError("The isolated smoke output must be outside the installed bundle")
    if not (bundle / "runtimes/service/python.exe").is_file():
        raise FileNotFoundError("Installed service runtime is missing")
    output.mkdir(parents=True, exist_ok=False)
    report = {"passed": False, "bundle": str(bundle), "output": str(output), "checks": [], "errors": [],
              "scope": "Installed-package real inference and API smoke test; synthetic OCR input; no accuracy benchmark.",
              "decision_driver": "Audit script simulates explicit human acceptance of a synthetic fixture suggestion; product automatic acceptance remains disabled."}
    audit = InstalledAudit(bundle, output)

    def check(name, condition=True):
        if not condition:
            raise AssertionError(name)
        report["checks"].append(name)
        print(name, flush=True)

    try:
        seed = audit.seed()
        report["seed"] = seed
        check("seed created exclusively by isolated installed Python and package", seed["isolated"] and Path(seed["module_path"]).is_relative_to(bundle))
        health = audit.start()
        report["health"] = health
        report["service_command"] = audit.command
        report["loopback_port"] = audit.port
        check("real installed service and queues are ready with matching package version", health["version"] == seed["version"] and not health.get("review_only"))
        html, metadata = audit.http("/", raw=True)
        scripts = re.findall(r'<script\b[^>]*\bsrc=["\']([^"\']+\.js)["\']', html.decode("utf-8"))
        assets = []
        for source in scripts:
            if not source.startswith("/") or source.startswith("//"):
                raise ValueError("Unexpected non-local frontend asset")
            content, asset = audit.http(source, raw=True)
            installed_file = (bundle / "web" / source.lstrip("/")).resolve()
            if not installed_file.is_relative_to(bundle / "web"):
                raise ValueError("Frontend asset is outside installed web directory")
            check("served frontend asset matches installed bytes", content == installed_file.read_bytes())
            assets.append({"path": source, **asset, "contains_multimodal_ui": "多模态二轮审校" in content.decode("utf-8")})
        report["frontend"] = {"index": metadata, "assets": assets}
        check("installed frontend contains the visual-review interface", bool(assets) and any(item["contains_multimodal_ui"] for item in assets))
        catalog = audit.http("/api/multimodal/models")
        write_json(output / "models.json", catalog)
        model_id = args.model or catalog.get("default_model")
        model = next((item for item in catalog["models"] if item["id"] == model_id), None)
        check("selected installed visual model is available", bool(model and model["available"]))
        report["model"] = model
        original_result = audit.http(f"/api/results/{seed['result_id']}")
        body = {"revision": original_result["revision"], "version_id": seed["version_id"], "model_id": model_id,
                "scope": "target", "target": {"kind": "text", "start": 0, "end": len(seed["before"])}, "request_id": str(uuid.uuid4())}
        queued = audit.http(f"/api/results/{seed['result_id']}/multimodal", "POST", body)
        task_id = queued["task_id"]
        report["submission"] = {"body": body, "task_id": task_id}
        changes, previous = [], None
        deadline = time.monotonic() + args.timeout
        started = time.monotonic()
        while time.monotonic() < deadline:
            audit.track()
            if audit.process.poll() is not None:
                raise RuntimeError("Installed service exited during inference")
            view = audit.http(f"/api/results/{seed['result_id']}/multimodal")
            task = next(item for item in view["requests"] if item["task_id"] == task_id)
            state = (task["status"], task.get("phase"), task.get("error"))
            if state != previous:
                changes.append({"seconds": round(time.monotonic() - started, 3), "status": state[0], "phase": state[1], "error": state[2]})
                write_json(output / "task-transitions.json", changes)
                print(f"review task: {state[0]} / {state[1]}", flush=True)
                previous = state
            if task["status"] in {"succeeded", "failed", "cancelled", "paused", "interrupted"}:
                break
            time.sleep(.5)
        else:
            raise TimeoutError("Installed real multimodal task exceeded inference timeout")
        write_json(output / "review-before-decision.json", view)
        report["task"] = task
        report["task_elapsed_seconds"] = round(time.monotonic() - started, 3)
        check("real GPU queue completed visual review", task["status"] == "succeeded")
        check("real llama server was observed in this service process tree", any(item["exe"].lower() == "llama-server.exe" for item in audit.known.values()))
        proposals = [item for item in view["proposals"] if item["task_id"] == task_id]
        check("model returned traceable proposals for the requested scope", bool(proposals) and all(item["evidence"]["version_id"] == seed["version_id"] for item in proposals))
        before_decision = audit.http(f"/api/results/{seed['result_id']}")
        check("model completion preserves the saved OCR draft until explicit adoption", before_decision["edited"] == original_result["edited"] and before_decision["revision"] == original_result["revision"])
        replacement = next((item for item in proposals if item["decision"] == "replace" and item["state"] == "pending"), None)
        if replacement:
            decision = {"action": "accept", "revision": before_decision["revision"], "version_id": seed["version_id"], "request_id": str(uuid.uuid4())}
            saved = audit.http(f"/api/results/{seed['result_id']}/multimodal/{replacement['id']}/decision", "POST", decision)
            report["scripted_human_decision"] = {"proposal_id": replacement["id"], "body": decision, "before": replacement["before"], "after": replacement["after"]}
            check("scripted explicit acceptance records one editable revision", saved["revision"] == before_decision["revision"] + 1 and saved["edited"]["text"] == replacement["after"])
        else:
            saved = before_decision
            report["scripted_human_decision"] = {"skipped": True, "reason": "The model returned no replace proposal; no artificial replacement was invented."}
        check("original OCR output is preserved", saved["original"] == original_result["original"])
        project = audit.http(f"/api/projects/{seed['project_id']}")
        check("visual review does not auto-confirm the page", next(item for item in project["images"] if item["id"] == seed["image_id"]).get("review_status") != "confirmed")
        exports = {}
        for format in ("json", "xlsx"):
            content, metadata = audit.http(f"/api/results/{seed['result_id']}/multimodal/report?format={format}", raw=True)
            path = output / f"review-report.{format}"
            path.write_bytes(content)
            exports[format] = {**metadata, "path": str(path)}
        report["exports"] = exports
        readback = audit.embedded(EXPORT_READBACK_PROGRAM, output)
        write_json(output / "export-readback.json", readback)
        check("installed JSON and XLSX exports preserve proposal strings and provenance", readback["passed"])
        report["passed"] = True
    except BaseException as error:
        report["errors"].append({"message": str(error), "traceback": traceback.format_exc()})
    finally:
        try:
            shutdown = audit.stop()
            report["shutdown"] = shutdown
            if shutdown["forced_cleanup"] or shutdown["surviving_processes"] or shutdown["service_returncode"] not in (0, None):
                report["passed"] = False
                report["errors"].append({"message": "Service/process tree did not shut down cleanly"})
            elif audit.process:
                report["checks"].append("shutdown endpoint closes service and observed model processes")
        except BaseException as error:
            report["passed"] = False
            report["errors"].append({"message": "Shutdown verification failed: " + str(error), "traceback": traceback.format_exc()})
        write_json(output / "installed-smoke-report.json", report)
    print(json.dumps({"passed": report["passed"], "checks": len(report["checks"]), "report": str(output / 'installed-smoke-report.json')}, ensure_ascii=False), flush=True)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
