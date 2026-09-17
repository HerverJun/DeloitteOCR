"""Scripted source choice on a real public page, with independent PDF/XLSX readback."""
import argparse
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time
import zipfile


def main():
    p=argparse.ArgumentParser()
    for name in ('bundle','workspace','output'):p.add_argument('--'+name,type=Path,required=True)
    p.add_argument('--result',required=True);p.add_argument('--installed',action='store_true')
    p.add_argument('--provider-prefix',default='',help='Explicit engineering source choice, not an automatic product preference')
    p.add_argument('--refresh-tool',action='store_true',help='Refresh auxiliary PDF candidates from an older saved workspace')
    a=p.parse_args();root=Path(__file__).resolve().parents[1]
    sys.path.insert(0,str(a.bundle/'app' if a.installed else root/'src'))
    from ocr_workbench.store import Store
    from ocr_workbench.structure_store import refresh_proposals,decide_structure
    from ocr_workbench.document_conflicts import acknowledge_conflict
    from ocr_workbench.exporting import build_export
    from ocr_workbench.pdf_export import build_pdf_export
    from ocr_workbench.documents import Documents
    from ocr_workbench.geometry import geometry_view,bind_manual
    from ocr_workbench.coordinates import bounds,box_polygon
    from ocr_workbench import __version__
    from openpyxl import load_workbook
    a.output.mkdir(parents=True,exist_ok=False)
    shutil.copytree(a.workspace,a.output/'workspace')
    store=Store(a.output/'workspace');manager=Documents(store,a.bundle)
    result=store.result(a.result);original=result['original']
    result=store.save(a.result,{'text':original['text'],'tables':original['tables']},result['revision'])
    if a.refresh_tool:
        from ocr_workbench.table_tool import enqueue
        stage=enqueue(manager,a.result,result['revision'])['stage_id']
        while store.one('document_stages',stage)['status']=='queued':
            if not manager.step():raise RuntimeError('Tool retry did not run')
        if store.one('document_stages',stage)['status']!='succeeded':raise RuntimeError(store.one('document_stages',stage)['error'])
    view=refresh_proposals(store,a.result,result['revision'])
    proposal=next(p for p in view['proposals'] if p['kind']=='native_table' and p['can_apply'] and p['provider'].startswith(a.provider_prefix))
    result=decide_structure(store,a.result,proposal['id'],{'action':'accept','request_id':'public-native-export-followup',
        'revision':proposal['revision'],'basis':proposal['basis'],'version_id':proposal['version_id'],'acknowledge_unverified_empty':True})
    # Explicit scripted native-source preference; no human quality/efficiency claim.
    edit=result['edited'];choices=[]
    for conflict in original['document']['conflicts']:
        text=conflict['ocr']['text']
        if edit['text'].count(text)!=1:raise ValueError('Ambiguous scripted OCR removal')
        edit['text']=edit['text'].replace(text,'',1)
        choices.append({'conflict_id':conflict['id'],'removed_ocr':text,'retained_native':[u['text'] for u in conflict['native']]})
    result=store.save(a.result,edit,result['revision'])
    for choice in choices:acknowledge_conflict(store,a.result,choice['conflict_id'],result['revision'])
    timing=[]
    for _ in range(50):
        tick=time.perf_counter();geometry_view(store,a.result);timing.append((time.perf_counter()-tick)*1000)
    workbook=build_export(store,[a.result],'xlsx')
    with zipfile.ZipFile(workbook) as z:
        data=z.read('OCR-result.xlsx');lineage=json.loads(z.read('sources/structure-'+a.result+'.json'))
    (a.output/'public-structure.xlsx').write_bytes(data)
    wb=load_workbook(io.BytesIO(data));checked=0
    for ti,table in enumerate(result['edited']['tables']):
        sheet=wb['Table '+str(ti+1)]
        for cell in table['cells']:
            value=sheet.cell(cell['row']+1,cell['column']+1)
            assert (value.value or '')==cell['text']
            assert not cell['text'] or value.data_type=='s'
            checked+=1
    preflight=build_pdf_export(store,manager,{'result_ids':[a.result]},preflight=True)
    bindings=[]
    if not preflight['ready']:
        for failure in preflight['failures']:
            for item in failure.get('items',[]):
                target=item.get('target',{})
                if target.get('kind')!='text':continue
                value=result['edited']['text'][target['start']:target['end']]
                matches=[c for c in original['document']['conflicts'] if ''.join(value.split())==''.join(u['text'] for u in c['native']).replace(' ','')]
                if len(matches)!=1:continue
                boxes=[bounds(u['polygon']) for u in matches[0]['native']]
                poly=box_polygon([min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)])
                bind_manual(store,a.result,{'revision':result['revision'],'version_id':proposal['version_id'],'target':target,'polygon':poly})
                bindings.append({'target':target,'polygon':poly,'source':'scripted explicit binding to native title units'})
        preflight=build_pdf_export(store,manager,{'result_ids':[a.result]},preflight=True)
    if not preflight['ready']:raise ValueError(json.dumps(preflight,ensure_ascii=False))
    pdf=build_pdf_export(store,manager,{'result_ids':[a.result]})
    shutil.copy2(pdf,a.output/'public-structure.pdf')
    shutil.copy2(pdf.with_suffix('.sources.json'),a.output/'public-structure.sources.json')
    code='''import json,sys,pypdfium2 as f,pikepdf
from PIL import ImageChops
p=f.PdfDocument(sys.argv[1]);page=p[0];text=page.get_textpage().get_text_range()
page.render(scale=1.5).to_pil().save(sys.argv[2])
with pikepdf.open(sys.argv[1]) as doc:attachments=list(doc.attachments)
print(json.dumps({'pages':len(p),'text':text,'attachments':attachments},ensure_ascii=False))'''
    read=subprocess.run([str(a.bundle/'runtimes/pdf/python.exe'),'-X','utf8','-c',code,str(pdf),str(a.output/'public-structure.png')],check=True,capture_output=True,text=True,encoding='utf-8')
    pdf_read=json.loads(read.stdout);assert pdf_read['pages']==1 and pdf_read['attachments']
    normalized=''.join(pdf_read['text'].split())
    for table in result['edited']['tables']:
        for cell in table['cells']:
            assert ''.join(cell['text'].split()) in normalized,cell['text']
    receipt={'version':__version__,'installed_app':a.installed,'module_source':sys.modules['ocr_workbench.structure_store'].__file__,
        'result_id':a.result,'revision':result['revision'],'tables':len(result['edited']['tables']),'xlsx_cells_read':checked,
        'pdf_all_cell_text_found':True,'pdf_attachments':pdf_read['attachments'],'pdf_preflight':preflight,
        'cached_geometry_p95_ms':sorted(timing)[47],'scripted_choices':choices,'scripted_bindings':bindings,'human_review_time':'unmeasured',
        'quality_acceptance':False,'source_revision':lineage['revision'],'structure_provider':proposal['provider']}
    (a.output/'readback.json').write_text(json.dumps(pdf_read,ensure_ascii=False,indent=2),encoding='utf-8')
    (a.output/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    manager.stop();print(json.dumps(receipt,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
