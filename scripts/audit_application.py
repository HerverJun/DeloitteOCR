"""Real packaged application acceptance. Never mocks OCR or edits model output."""

import argparse
import copy
import ctypes
import hashlib
import io
import json
from pathlib import Path
import subprocess
import time
import urllib.request
import urllib.error
import uuid
from PIL import Image, ImageDraw
from pillow_heif import register_heif_opener
from openpyxl import load_workbook


def alive(pid):
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.OpenProcess.restype = ctypes.c_void_p
    kernel.CloseHandle.argtypes = [ctypes.c_void_p]
    handle = kernel.OpenProcess(0x1000, False, pid)
    if not handle:
        return False
    code = ctypes.c_ulong()
    kernel.GetExitCodeProcess.argtypes = [
        ctypes.c_void_p,
        ctypes.POINTER(ctypes.c_ulong),
    ]
    kernel.GetExitCodeProcess(handle, ctypes.byref(code))
    kernel.CloseHandle(handle)
    return code.value == 259


def until(fn, timeout=300):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = fn()
        if value:
            return value
        time.sleep(0.15)
    raise TimeoutError("Acceptance condition timed out")


class Application:
    def __init__(self, bundle, output):
        self.bundle, self.output = bundle, output
        self.data = output / "项目 中文 空格"
        self.process = None
        self.opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))

    def start(self):
        self.process = subprocess.Popen(
            [
                str(self.bundle / "launcher/OfflineOCRLauncher.exe"),
                "--data",
                str(self.data),
                "--no-browser",
            ],
            creationflags=subprocess.CREATE_NO_WINDOW,
        )

        def ready():
            try:
                self.state = json.loads(
                    (self.data / "launcher/launcher-state.json").read_text("utf-8")
                )
                if self.state["pid"] != self.process.pid:
                    return False
                self.token = (
                    (self.data / "launcher/session-token.txt")
                    .read_text("utf-8")
                    .strip()
                )
                self.base = "http://127.0.0.1:" + str(self.state["port"])
                return self.api("/health")["status"] == "ready"
            except (OSError, ValueError, urllib.error.URLError):
                return False

        until(ready, 60)

    def api(self, path, method="GET", body=None, raw=False, headers=None):
        request_headers = {"Authorization": "Bearer " + self.token}
        if headers is not None:
            request_headers = headers
        if isinstance(body, dict):
            body = json.dumps(body, ensure_ascii=False).encode("utf-8")
            request_headers["Content-Type"] = "application/json"
        request = urllib.request.Request(
            self.base + "/api" + path, data=body, method=method, headers=request_headers
        )
        with self.opener.open(request, timeout=120) as response:
            value = response.read()
        return value if raw else json.loads(value)

    def import_file(self, project, path):
        boundary = uuid.uuid4().hex
        body = (
            (
                f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="{path.name}"\r\nContent-Type: application/octet-stream\r\n\r\n'
            ).encode("utf-8")
            + path.read_bytes()
            + f"\r\n--{boundary}--\r\n".encode()
        )
        return self.api(
            "/projects/" + project + "/images",
            "POST",
            body,
            headers={
                "Authorization": "Bearer " + self.token,
                "Content-Type": "multipart/form-data; boundary=" + boundary,
            },
        )

    def tasks(self, project):
        return self.api("/projects/" + project)["tasks"]

    def enqueue(self, project, versions, engines):
        return self.api(
            "/projects/" + project + "/tasks",
            "POST",
            {"version_ids": versions, "engines": engines},
        )["task_ids"]

    def action(self, project, action, ids):
        return self.api(
            "/projects/" + project + "/queue/" + action, "POST", {"task_ids": ids}
        )

    def completed(self, project, ids):
        tasks = [t for t in self.tasks(project) if t["id"] in ids]
        if all(t["status"] in {"succeeded", "failed", "cancelled"} for t in tasks):
            return tasks

    def stop(self):
        if self.process and self.process.poll() is None:
            self.api("/shutdown", "POST", {})
            self.process.wait(50)
            until(lambda: not alive(self.state["service_pid"]), 20)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--bundle", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--quick",
        action="store_true",
        help="Four engines once; omit crash and interaction cases",
    )
    args = parser.parse_args()
    if args.output.exists() and any(args.output.iterdir()):
        raise ValueError("Use a fresh evidence directory")
    args.output.mkdir(parents=True, exist_ok=True)
    app = Application(args.bundle.resolve(), args.output.resolve())
    report = {"passed": False, "checks": [], "engines": [], "errors": []}

    def check(name, condition=True):
        if not condition:
            raise AssertionError(name)
        report["checks"].append(name)
        print(name, flush=True)

    def expect_status(call, status):
        try:
            call()
        except urllib.error.HTTPError as error:
            assert error.code == status, (error.code, error.read())
        else:
            raise AssertionError("Expected HTTP " + str(status))

    try:
        app.start()
        check("packaged launcher starts embedded service in Chinese data path")
        expect_status(lambda: app.api("/state", headers={}), 401)
        expect_status(
            lambda: app.api("/state", headers={"Origin": "https://example.com"}), 403
        )
        check("unauthenticated and cross-origin requests rejected")
        project = app.api("/projects", "POST", {"name": "四引擎验收"})["id"]
        photos = [
            app.import_file(project, args.bundle / "fixtures" / name)["images"][0]
            for name in ["printed.png", "table.png"]
        ]
        versions = [p["active_version"] for p in photos]
        ids = app.enqueue(
            project,
            versions[1:] if args.quick else versions,
            ["ppocr", "paddlevl", "glm", "hunyuan"],
        )
        tasks = until(lambda: app.completed(project, ids), 900)
        check(
            "all real engine tasks succeeded",
            all(t["status"] == "succeeded" for t in tasks),
        )
        results = [app.api("/results/" + t["result_id"]) for t in tasks]
        for engine in ["ppocr", "paddlevl", "glm", "hunyuan"]:
            group = [r for r in results if r["original"]["engine"] == engine]
            check(
                engine + " retains exact long identifier",
                all("00123456789012345678" in r["original"]["text"] for r in group),
            )
            check(
                engine + " resident session reused",
                len({r["original"]["engine_session"]["id"] for r in group}) == 1,
            )
            table = next(
                r
                for r in group
                if r["original"]["project_image_version"] == versions[1]
            )
            check(
                engine + " truthful table capability",
                bool(table["original"]["tables"]) == (engine != "ppocr"),
            )
            report["engines"].extend(
                [
                    {
                        "engine": engine,
                        "result_id": r["id"],
                        "elapsed_seconds": r["original"]["elapsed_seconds"],
                        "session": r["original"]["engine_session"],
                    }
                    for r in group
                ]
            )
        ordered = sorted(tasks, key=lambda t: t["started"])
        check(
            "GPU tasks serialized",
            all(a["finished"] <= b["started"] for a, b in zip(ordered, ordered[1:])),
        )
        until(lambda: app.api("/state")["queue"]["engine"] is None, 30)
        until(
            lambda: all(
                not alive(r["original"]["engine_session"]["pid"]) for r in results
            ),
            30,
        )
        check("engine processes released after queue idle")
        if not args.quick:
            register_heif_opener()
            for suffix, fmt in [
                ("jpg", "JPEG"),
                ("bmp", "BMP"),
                ("tiff", "TIFF"),
                ("webp", "WEBP"),
                ("heic", "HEIF"),
            ]:
                target = args.output / ("decode." + suffix)
                Image.new("RGB", (90, 60), "red").save(target, format=fmt)
                imported = app.import_file(project, target)
                check(
                    "decode " + suffix,
                    len(imported["images"]) == 1 and not imported["errors"],
                )
            corrupt = args.output / "corrupt.png"
            corrupt.write_bytes(b"not an image")
            check(
                "corrupt import reported without adding image",
                bool(app.import_file(project, corrupt)["errors"]),
            )
            parent = versions[1]
            source = next(
                v
                for v in app.api("/projects/" + project)["versions"]
                if v["id"] == parent
            )
            w, h = source["width"], source["height"]
            for op in [
                {"kind": "rotate", "degrees": 90},
                {"kind": "crop", "box": [0, 0, w - 10, h - 10]},
                {
                    "kind": "perspective",
                    "points": [[0, 0], [w - 1, 0], [w - 1, h - 1], [0, h - 1]],
                },
                {"kind": "contrast", "factor": 1.3},
            ]:
                transformed = app.api("/versions/" + parent + "/transform", "POST", op)
                check(
                    "immutable " + op["kind"],
                    transformed["parent_id"] == parent and transformed["id"] != parent,
                )
            color = args.output / "dewarp-color.png"
            with Image.open(args.bundle / "fixtures/table.png") as im:
                im = im.convert("RGB")
                ImageDraw.Draw(im).rectangle((20, 20, im.width - 20, 100), fill="red")
                im.save(color)
            colored = app.import_file(project, color)["images"][0]
            dewarp = app.api(
                "/versions/" + colored["active_version"] + "/transform",
                "POST",
                {"kind": "dewarp"},
            )["task_id"]
            dewarped = until(lambda: app.completed(project, [dewarp]), 300)[0]
            check(
                "actual UVDoc creates a version, not an OCR result",
                dewarped["status"] == "succeeded"
                and dewarped["result_version_id"]
                and not dewarped["result_id"],
            )
            data = app.api(
                "/versions/" + dewarped["result_version_id"] + "/image", raw=True
            )
            (args.output / "dewarped.png").write_bytes(data)
            with Image.open(io.BytesIO(data)) as im:
                pixels = list(im.resize((150, 150)).getdata())
                red = sum(r > 150 and r > g * 1.5 and r > b * 1.5 for r, g, b in pixels)
                blue = sum(
                    b > 150 and b > r * 1.5 and b > g * 1.5 for r, g, b in pixels
                )
                check("UVDoc RGB colors preserved", red > 20 and red > blue * 2)
            blank = args.output / "blank.png"
            Image.new("RGB", (640, 480), "white").save(blank)
            blank_photo = app.import_file(project, blank)["images"][0]
            failed = app.enqueue(project, [blank_photo["active_version"]], ["ppocr"])
            check(
                "blank OCR fails without fabricated success",
                until(lambda: app.completed(project, failed))[0]["status"] == "failed",
            )
            app.action(project, "retry", failed)
            check(
                "failed item can retry and remains truthful",
                until(lambda: app.completed(project, failed))[0]["status"] == "failed",
            )
            cancel_ids = app.enqueue(project, versions, ["paddlevl"])
            running = until(
                lambda: next(
                    (
                        t
                        for t in app.tasks(project)
                        if t["id"] in cancel_ids and t["status"] == "running"
                    ),
                    None,
                )
            )
            queued = [i for i in cancel_ids if i != running["id"]]
            app.action(project, "pause", queued)
            check(
                "pause queued tasks",
                all(
                    t["status"] == "paused"
                    for t in app.tasks(project)
                    if t["id"] in queued
                ),
            )
            app.action(project, "cancel", [running["id"]])
            until(lambda: app.api("/state")["queue"]["engine"] is None, 30)
            check(
                "cancel running model load",
                next(t for t in app.tasks(project) if t["id"] == running["id"])[
                    "status"
                ]
                == "cancelled",
            )
            app.action(project, "retry", [running["id"]])
            app.action(project, "resume", queued)
            check(
                "retry cancelled and resume paused succeed",
                all(
                    t["status"] == "succeeded"
                    for t in until(lambda: app.completed(project, cancel_ids), 400)
                ),
            )
            result = next(
                r
                for r in results
                if r["original"]["engine"] == "paddlevl" and r["original"]["tables"]
            )
            original = copy.deepcopy(result["original"])
            edit = copy.deepcopy(result["edited"])
            cell = next(c for c in edit["tables"][0]["cells"] if c["text"] == "扫描仪")
            cell["text"] = "扫描仪（审计校对）"
            edit["text"] += "\n校对备注 00001234567890123456"
            saved = app.api(
                "/results/" + result["id"],
                "PUT",
                {"edited": edit, "revision": result["revision"]},
            )
            expect_status(
                lambda: app.api(
                    "/results/" + result["id"],
                    "PUT",
                    {"edited": edit, "revision": result["revision"]},
                ),
                409,
            )
            undone = app.api(
                "/results/" + result["id"] + "/history",
                "POST",
                {"direction": -1, "revision": saved["revision"]},
            )
            check("undo restores original edit", undone["edited"] == result["edited"])
            redone = app.api(
                "/results/" + result["id"] + "/history",
                "POST",
                {"direction": 1, "revision": undone["revision"]},
            )
            check(
                "redo and immutable original",
                redone["edited"] == edit and redone["original"] == original,
            )
            for fmt in ["txt", "md", "json", "xlsx"]:
                data = app.api(
                    "/export",
                    "POST",
                    {"result_ids": [result["id"]], "format": fmt},
                    raw=True,
                )
                (args.output / ("edited." + fmt)).write_bytes(data)
                if fmt == "md":
                    check(
                        "Markdown exports edited cells without stale duplicate",
                        data.decode().count("扫描仪（审计校对）") == 1,
                    )
                if fmt == "xlsx":
                    wb = load_workbook(io.BytesIO(data))
                    cells = [c for ws in wb for row in ws for c in row]
                    check(
                        "Excel preserves edited cell and long string",
                        any(c.value == "扫描仪（审计校对）" for c in cells)
                        and any(
                            c.value == "00123456789012345678" and c.data_type == "s"
                            for c in cells
                        ),
                    )
                    check(
                        "Excel preserves merged heading",
                        bool(wb.active.merged_cells.ranges),
                    )
            data = app.api(
                "/export",
                "POST",
                {
                    "result_ids": [r["id"] for r in results if r["edited"]["tables"]],
                    "format": "xlsx",
                },
                raw=True,
            )
            check(
                "aggregate exports tables to separate sheets",
                len(load_workbook(io.BytesIO(data)).sheetnames)
                == sum(len(r["edited"]["tables"]) for r in results),
            )
            crash_ids = app.enqueue(project, versions, ["paddlevl"])
            crashed = until(
                lambda: next(
                    (
                        t
                        for t in app.tasks(project)
                        if t["id"] in crash_ids and t["status"] == "running"
                    ),
                    None,
                )
            )
            old_service = app.state["service_pid"]
            app.process.kill()
            app.process.wait(10)
            until(lambda: not alive(old_service), 30)
            app.start()
            recovered = [t for t in app.tasks(project) if t["id"] in crash_ids]
            check(
                "crash recovery requires deliberate resume",
                {t["status"] for t in recovered} == {"interrupted", "paused"},
            )
            check(
                "edits survive application crash",
                app.api("/results/" + result["id"])["edited"] == edit,
            )
            check(
                "completed results not duplicated",
                all(
                    next(t for t in app.tasks(project) if t["id"] == old["id"])[
                        "result_id"
                    ]
                    == old["result_id"]
                    for old in tasks
                ),
            )
            app.action(project, "resume", crash_ids)
            check(
                "resume after crash succeeds",
                all(
                    t["status"] == "succeeded"
                    for t in until(lambda: app.completed(project, crash_ids), 400)
                ),
            )
            for photo in photos:
                check(
                    "original file hash unchanged " + photo["name"],
                    hashlib.sha256(
                        (app.data / photo["original_path"]).read_bytes()
                    ).hexdigest()
                    == photo["sha256"],
                )
        app.stop()
        check(
            "graceful launcher/service exit and token cleanup",
            not (app.data / "launcher/session-token.txt").exists(),
        )
        report["passed"] = True
    except BaseException as error:
        report["errors"].append(repr(error))
        raise
    finally:
        try:
            app.stop()
        finally:
            (args.output / "application-audit.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), "utf-8"
            )


if __name__ == "__main__":
    main()
