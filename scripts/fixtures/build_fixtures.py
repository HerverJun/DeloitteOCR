"""Generate small, fully synthetic, deterministic v1 fixtures.

Usage: python build_fixtures.py --output-dir <isolated/data> --manifest <isolated/manifest.json>
No input file, business directory or external service is read.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pyarrow as pa
import pyarrow.parquet as pq
from PIL import Image, ImageDraw
from openpyxl import Workbook
from reportlab.pdfgen import canvas

LONG_ID = "000123456789012345678901234567"
PDF_PAGES = ["SYNTHETIC INVOICE 000123456789012345678901234567", "SECOND PAGE TABLE A 12.50"]
CSV_ROWS = [
    {"row_id": "r1", "long_id": LONG_ID, "match_key": "dup", "amount": "12.500000", "nullable": "\\N"},
    {"row_id": "r2", "long_id": "000000000000000000000000000002", "match_key": "dup", "amount": "0.000001", "nullable": ""},
    {"row_id": "r3", "long_id": "000000000000000000000000000003", "match_key": "unique", "amount": "-3.250000", "nullable": "值"},
]
PARQUET_ROWS = [{**row, "nullable": None if row["nullable"] == "\\N" else row["nullable"]}
                for row in CSV_ROWS]
ORIGIN = {"kind": "synthetic", "creator": "Workbench fixture generator", "license": "CC0-1.0",
          "upstream_url": None}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical_zip(path: Path) -> None:
    """Remove ZIP member clock/host variation without altering workbook XML."""
    with ZipFile(path) as source:
        contents = {name: source.read(name) for name in source.namelist()}
    core = contents["docProps/core.xml"]
    core, replacements = re.subn(
        rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
        rb"\g<1>2020-01-01T00:00:00Z\g<2>", core,
    )
    if replacements != 1:
        raise ValueError("Expected exactly one XLSX modified timestamp")
    contents["docProps/core.xml"] = core
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temporary:
        temp_path = Path(temporary.name)
    with ZipFile(temp_path, "w") as output:
        for name in sorted(contents):
            info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            output.writestr(info, contents[name], compress_type=ZIP_DEFLATED, compresslevel=9)
    temp_path.replace(path)


def image(path: Path, changed: bool = False) -> None:
    im = Image.new("RGB", (64, 32), (248, 248, 246))
    draw = ImageDraw.Draw(im)
    draw.rectangle((4, 4, 59, 27), outline=(18, 54, 85), width=2)
    draw.text((9, 10), "ID 0123", fill=(20, 20, 20))
    if changed:
        im.putpixel((63, 31), (255, 0, 0))
    im.save(path, format="PNG", optimize=False)


def pdf(path: Path) -> None:
    doc = canvas.Canvas(str(path), pagesize=(320, 180), invariant=1, pageCompression=0)
    doc.setTitle("Synthetic two-page invoice")
    doc.setAuthor("Workbench fixture generator")
    for text in PDF_PAGES:
        doc.setFont("Helvetica", 10)
        doc.drawString(20, 140, text)
        doc.showPage()
    doc.save()


def xlsx(path: Path) -> None:
    wb = Workbook()
    wb.properties.creator = "Workbench fixture generator"
    wb.properties.created = datetime(2020, 1, 1)
    wb.properties.modified = datetime(2020, 1, 1)
    ws = wb.active
    ws.title = "对账"
    ws.append(["row_id", "long_id", "match_key", "amount", "formula", "nullable"])
    ws.append(["r1", LONG_ID, "dup", "12.500000", "=SUM(1,2)", None])
    ws.append(["r2", "000000000000000000000000000002", "dup", "0.000001", None, ""])
    ws.append(["r3", "000000000000000000000000000003", "unique", "-3.250000", None, "值"])
    for row in range(2, 5):
        ws[f"B{row}"].number_format = "@"
        ws[f"D{row}"].number_format = "@"
    wb.save(path)
    canonical_zip(path)


def csv_file(path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=list(CSV_ROWS[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(CSV_ROWS)


def parquet(path: Path) -> None:
    schema = pa.schema([
        pa.field("row_id", pa.string(), nullable=False),
        pa.field("long_id", pa.string(), nullable=False),
        pa.field("match_key", pa.string(), nullable=False),
        pa.field("amount", pa.decimal128(20, 6), nullable=False),
        pa.field("nullable", pa.string()),
    ])
    columns = {name: [Decimal(row[name]) if name == "amount" else row[name] for row in PARQUET_ROWS]
               for name in schema.names}
    table = pa.Table.from_pydict(columns, schema=schema)
    pq.write_table(table, path, version="2.6", compression="NONE", use_dictionary=False,
                   row_group_size=3, write_statistics=True)


def build(output_dir: Path, manifest_path: Path) -> dict:
    output_dir = output_dir.resolve()
    manifest_path = manifest_path.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    expected_image = {"width": 64, "height": 32, "mode": "RGB", "pixel_0_0": [248, 248, 246],
                      "pixel_63_31": [248, 248, 246]}
    definitions = [
        ("images/same-source-a.png", "image/png", "sourceAlpha", 1, expected_image),
        ("images/same-source-b.png", "image/png", "sourceBeta", 1, expected_image),
        ("images/票据-甲-v2.png", "image/png", "sourceAlpha", 2,
         expected_image | {"pixel_63_31": [255, 0, 0]}),
        ("docs/invoice.pdf", "application/pdf", "sourcePdf", 1,
         {"page_count": 2, "page_text": PDF_PAGES}),
        ("tables/对账.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
         "sourceXlsx", 1, {"sheet_names": ["对账"], "rows": [
             ["r1", LONG_ID, "dup", "12.500000", "=SUM(1,2)", None],
             ["r2", "000000000000000000000000000002", "dup", "0.000001", None, None],
             ["r3", "000000000000000000000000000003", "unique", "-3.250000", None, "值"]],
             "formula_cell": "=SUM(1,2)", "cached_value": None, "long_id_type": "str"}),
        ("tables/match.csv", "text/csv", "sourceCsv", 1,
         {"rows": CSV_ROWS, "null_token": "\\N", "empty_string_distinct": True}),
        ("tables/match.parquet", "application/vnd.apache.parquet", "sourceParquet", 1,
         {"rows": PARQUET_ROWS, "amount_type": "decimal128(20, 6)", "long_id_type": "string"}),
    ]
    fixtures = []
    for relative, media_type, source_id, revision, expected in definitions:
        target = output_dir / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        if relative.endswith("same-source-b.png"):
            shutil.copyfile(output_dir / "images/same-source-a.png", target)
        elif relative.endswith(".png"):
            image(target, changed=revision == 2)
        elif relative.endswith(".pdf"):
            pdf(target)
        elif relative.endswith(".xlsx"):
            xlsx(target)
        elif relative.endswith(".csv"):
            csv_file(target)
        else:
            parquet(target)
        fixtures.append({"path": f"data/{relative}", "sha256": sha(target), "size_bytes": target.stat().st_size,
                         "media_type": media_type, "origin": ORIGIN, "source_id": source_id,
                         "source_revision": revision, "expected": expected,
                         "byte_deterministic": True})
    manifest = {"schema_version": 1, "generator": "scripts/fixtures/build_fixtures.py",
                "policy": "Self-authored synthetic CC0-1.0 material; no private data or external source assets",
                "real_public_sources": [], "fixtures": fixtures,
                "relationships": {"same_bytes_distinct_sources": ["data/images/same-source-a.png",
                                                                "data/images/same-source-b.png"],
                                  "changed_source_revisions": ["data/images/same-source-a.png",
                                                               "data/images/票据-甲-v2.png"]}}
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    manifest = build(args.output_dir, args.manifest)
    print(f"generated {len(manifest['fixtures'])} fixtures")
