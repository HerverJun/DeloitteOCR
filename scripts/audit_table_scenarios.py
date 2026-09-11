"""Real inference on deterministic multi-table and perspective engineering images."""

import argparse
import copy
import hashlib
import json
from pathlib import Path
import sys
import time
import cv2
import numpy as np
from PIL import Image
from openpyxl import load_workbook
from ocr_workbench.adapter import EngineAdapter
from ocr_workbench.tables import export_xlsx, parse_tables

sys.path.insert(0, str(Path(__file__).resolve().parent))
from benchmark_metrics import table_metrics


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--bundle", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--real-case", type=Path)
    a = p.parse_args()
    if a.output.exists() and any(a.output.iterdir()):
        raise ValueError("Fresh evidence directory required")
    a.output.mkdir(parents=True, exist_ok=True)
    source = a.bundle / "fixtures/table.png"
    expected = json.loads((a.bundle / "fixtures/expectations.json").read_text("utf-8"))["table"]
    cells = [{"row": 0, "column": 0, "row_span": 1, "column_span": 4, "text": "采购记录（合并标题）"}]
    cells += [{"row": r+1, "column": c, "row_span": 1, "column_span": 1, "text": text} for r, values in enumerate(expected["data"]) for c, text in enumerate(values)]
    table = {"rows": 5, "columns": 4, "cells": cells}
    with Image.open(source) as original:
        original = original.convert("RGB")
        crop = original.crop((50, 150, 1350, 700))
        multi = Image.new("RGB", (1400, 1400), "white")
        multi.paste(crop, (50, 60))
        multi.paste(crop, (50, 790))
        multi.save(a.output / "multiple.png")
        src = np.float32([[0, 0], [1399, 0], [1399, 999], [0, 999]])
        dst = np.float32([[220, 70], [1380, 180], [1490, 1080], [70, 990]])
        matrix = cv2.getPerspectiveTransform(src, dst)
        projected = cv2.warpPerspective(np.asarray(original), matrix, (1560, 1160), borderValue=(245, 245, 245))
        Image.fromarray(projected).save(a.output / "perspective.png")
        restored = cv2.warpPerspective(projected, np.linalg.inv(matrix), (1400, 1000), borderValue=(255, 255, 255))
        Image.fromarray(restored).save(a.output / "perspective-corrected.png")
    scenarios = [
        ("multiple", [table, table]),
        ("perspective", [table]),
        ("perspective-corrected", [table]),
    ]
    real_case = None
    if a.real_case:
        real_case = json.loads(a.real_case.read_text("utf-8"))
        source_image = Path(real_case["image"])
        assert hashlib.sha256(source_image.read_bytes()).hexdigest() == real_case["image_sha256"]
        with Image.open(source_image) as image:
            image.convert("RGB").save(a.output / "real-multiple.png")
        truth = [t for block in real_case["annotations"] for t in parse_tables(block["html"])]
        scenarios.append(("real-multiple", truth))
    report = {
        "complete": False,
        "scope": "Synthetic multi-table and perspective reprojection of known engineering fixture. Corrected uses known geometry; not a genuine smartphone photo or a user-drawn GUI correction. Recognition errors remain in metrics; engineering export checks do not assert perfect model accuracy.",
        "bundle_manifest_sha256": hashlib.sha256((a.bundle / "manifest.json").read_bytes()).hexdigest(),
        "audit_script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "perspective_matrix": matrix.tolist(),
        "real_case": ({"sha256": hashlib.sha256(a.real_case.read_bytes()).hexdigest(), "image_sha256": real_case["image_sha256"], "expected_tables": len(truth), "scope": real_case["scope"]} if real_case else None),
        "records": [],
    }
    try:
        for engine in ["paddlevl", "glm", "hunyuan"]:
            adapter = EngineAdapter(a.bundle, engine, a.output / "sessions")
            try:
                for name, truth in scenarios:
                    output = a.output / engine / name
                    output.mkdir(parents=True)
                    image = a.output / (name + ".png")
                    record = {"engine": engine, "scenario": name, "image_sha256": hashlib.sha256(image.read_bytes()).hexdigest(), "expected_tables": len(truth), "succeeded": False}
                    started = time.monotonic()
                    result = {}
                    try:
                        if not adapter.ready:
                            adapter.load()
                        result = adapter.recognize(image, output)
                        record["succeeded"] = True
                        predicted = result["tables"]
                        record["predicted_tables"] = len(predicted)
                        if predicted:
                            edited = copy.deepcopy(predicted)
                            first = edited[0]["cells"][0]
                            first["text"] = "00123456789012345678"
                            workbook = output / "edited.xlsx"
                            export_xlsx(edited, workbook)
                            book = load_workbook(workbook)
                            assert len(book.worksheets) == len(predicted)
                            record["exported_sheets"] = len(book.worksheets)
                            for sheet, t in zip(book.worksheets, edited):
                                offset = 1 if t.get("caption") else 0
                                for cell in t["cells"]:
                                    r, c = cell["row"] + 1 + offset, cell["column"] + 1
                                    value = sheet.cell(r, c)
                                    assert (value.value or "") == cell["text"]
                                    if cell["text"]:
                                        assert value.data_type == "s"
                                    if cell["row_span"] > 1 or cell["column_span"] > 1:
                                        assert any(m.min_row == r and m.min_col == c and m.max_row == r+cell["row_span"]-1 and m.max_col == c+cell["column_span"]-1 for m in sheet.merged_cells.ranges)
                            book.close()
                            record["edited_export_verified"] = True
                            record["workbook_sha256"] = hashlib.sha256(workbook.read_bytes()).hexdigest()
                        else:
                            record["edited_export_verified"] = False
                    except Exception as error:
                        record["error"] = str(error)
                        adapter.unload()
                        adapter = EngineAdapter(a.bundle, engine, a.output / "sessions")
                    record["seconds"] = time.monotonic() - started
                    record["metrics"] = table_metrics(truth, result.get("tables", []))
                    report["records"].append(record)
                    (a.output / "table-scenarios.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")
                    print(json.dumps(record, ensure_ascii=False), flush=True)
            finally:
                adapter.unload()
        report["complete"] = True
        report["engineering_passed"] = all(r["succeeded"] and r.get("edited_export_verified") and "error" not in r for r in report["records"])
    finally:
        (a.output / "table-scenarios.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")
    return 0 if report.get("engineering_passed") else 2


if __name__ == "__main__":
    raise SystemExit(main())
