"""Public engineering evidence, deliberately separate from annotated acceptance."""
import argparse
from collections import Counter
from copy import deepcopy
import io
import json
from pathlib import Path
import shutil
import statistics
import sys
import time
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'src'),str(ROOT/'scripts')]
from geometry_eval_common import file_lock,sha,write_json
from evaluate_public_pdf_geometry import artifact
from ocr_workbench.structure_diagnostics import prepare_candidates
from ocr_workbench.geometry_contract import fingerprint


def controlled(args,samples):
    rows=[]
    for sample in samples:
        native=json.loads((args.data/'pages'/sample['id']/'native.json').read_text('utf-8'))['native']
        context=artifact(args.inference/'context',sample,'geometry.json')
        pools={'ocr':context['ocr_blocks'],'native':[dict(u,id='native-'+str(i),source_kind='native',granularity='word')
            for i,u in enumerate(native['units']) if u.get('reliable')]}
        for pool,blocks in pools.items():
            for provider in ('paddle','tableformer-raw','rapidtable','paddlevl'):
                row={'id':sample['id'],'document_id':sample['original_document_id'],'pool':pool,'provider':provider,
                     'shared_input_sha256':fingerprint(blocks),'references_read':False}
                try:
                    if provider=='paddlevl':
                        adopted=artifact(args.inference/'adopted',sample,'result.json')
                        if not adopted['tables'] or not all(c.get('polygon') for t in adopted['tables'] for c in t['cells']):
                            raise ValueError('VL has no spatial cell skeleton; fixed-token spatial filling is not measurable')
                        raise ValueError('VL cell geometry contract not qualified')
                    pred=artifact(args.inference/provider,sample,'geometry.json')
                    value=prepare_candidates(pred,blocks,source_result=pool+'-fixed',image_version=sample['sha256'],width=sample['width'],height=sample['height'])
                    write_json(args.output/'B'/pool/provider/(sample['id']+'.json'),value)
                    row.update(status='success',tables=len(value['tables']),cells=sum(len(t['skeleton']['cells']) for t in value['tables']),
                        unassigned=sum(len(t['unassigned_token_ids']) for t in value['tables']),
                        invalid_tables=sum(bool(t['reason_codes']) for t in value['tables']),
                        maximum_table_ms=max((t['timing_ms'] for t in value['tables']),default=0),
                        reasons=dict(Counter(r for t in value['tables'] for r in t['reason_codes'])))
                except Exception as error:
                    row.update(status='unavailable' if provider=='paddlevl' else 'failed',error=str(error))
                rows.append(row)
    write_json(args.output/'B/summary.json',{'rows':rows,'quality_scores':'unmeasured; no independent annotations',
        'comparison':'same detected regions, fixed token pool per cohort, local-v3 ownership; native and OCR never combined'})


def product(args,samples):
    from ocr_workbench.store import Store
    from ocr_workbench.documents import Documents
    from ocr_workbench.task_queue import TaskQueue
    from ocr_workbench.adapter import EngineAdapter
    from ocr_workbench.structure_store import refresh_proposals,decide_structure
    from ocr_workbench.geometry import geometry_view
    from ocr_workbench.document_review import document_review_queue
    from ocr_workbench.exporting import build_export
    from ocr_workbench.pdf_export import build_pdf_export
    from openpyxl import load_workbook
    root=args.output/'C';root.mkdir(parents=True)
    store=Store(root/'workspace')
    def factory(bundle,engine,sessions):
        return EngineAdapter(bundle,engine,sessions,worker_source=ROOT/'src')
    gpu=TaskQueue(store,args.bundle,adapter_factory=factory)
    manager=Documents(store,args.bundle,gpu_queue=gpu)
    project=store.project('Public structure workflow engineering run')
    documents={};rows=[]
    try:
        for sample in samples:
            row={'id':sample['id'],'scripted_review':True,'human_review_time':'unmeasured'}
            start=time.perf_counter()
            try:
                doc_id=sample['original_document_id']
                if doc_id not in documents:
                    imported=manager.import_document(project['id'],doc_id+'.pdf',args.data/'sources'/doc_id/'source.pdf',dpi=150)
                    documents[doc_id]=imported['id']
                page=store.rows('SELECT * FROM pages WHERE document_id=? AND page_number=?',(documents[doc_id],sample['page_number']))[0]
                stage=manager.process(page['id'],'auto')
                manager.action(documents[doc_id],'pause')
                assert store.one('document_stages',stage)['status']=='paused'
                manager.action(documents[doc_id],'resume')
                while True:
                    cpu_work=manager.step();gpu_work=gpu.step()
                    if not cpu_work and not gpu_work:break
                    if time.perf_counter()-start>900:raise TimeoutError('Product page budget exceeded')
                finished=store.one('document_stages',stage)
                if finished['status']!='succeeded':raise RuntimeError(finished.get('error') or finished['status'])
                result_id=json.loads(finished['output'])['result_id']
                result=store.result(result_id)
                row['baseline_revision']=result['revision'];row['baseline_tables']=len(result['edited']['tables'])
                write_json(root/(sample['id']+'-before.json'),result)
                view=refresh_proposals(store,result_id,result['revision'])
                write_json(root/(sample['id']+'-proposals.json'),view)
                options=[p for p in view['proposals'] if p['can_apply']]
                row['suggestions']=len(view['proposals']);row['applicable']=len(options)
                if options:
                    # Predeclared engineering action, never a model-quality label.
                    p=options[0]
                    saved=decide_structure(store,result_id,p['id'],{'request_id':'engineering-'+sample['id'],'action':'accept',
                        'revision':p['revision'],'basis':p['basis'],'version_id':p['version_id'],'acknowledge_unverified_empty':True})
                    undo=store.history(result_id,-1,saved['revision'])
                    result=store.history(result_id,1,undo['revision'])
                    assert result['edited']==saved['edited']
                    row['accepted_kind']=p['kind']
                write_json(root/(sample['id']+'-after.json'),result)
                times=[]
                for _ in range(20):
                    tick=time.perf_counter();geometry_view(store,result_id);times.append((time.perf_counter()-tick)*1000)
                row['cached_geometry_p95_ms']=sorted(times)[18]
                if result['edited']['tables']:
                    tick=time.perf_counter();export=build_export(store,[result_id],'xlsx');row['xlsx_seconds']=time.perf_counter()-tick
                    shutil.copy2(export,root/(sample['id']+export.suffix))
                    if export.suffix=='.zip':
                        with zipfile.ZipFile(export) as z: content=io.BytesIO(z.read('OCR-result.xlsx'))
                    else:content=export
                    wb=load_workbook(content)
                    for ti,table in enumerate(result['edited']['tables']):
                        sheet=wb['Table '+str(ti+1)]
                        for cell in table['cells']:
                            value=sheet.cell(cell['row']+1,cell['column']+1)
                            assert (value.value or '')==cell['text'],(sample['id'],value.coordinate)
                            assert not cell['text'] or value.data_type=='s'
                            if cell['row_span']*cell['column_span']>1:
                                assert any(r.min_row==cell['row']+1 and r.min_col==cell['column']+1 and r.max_row==cell['row']+cell['row_span'] and r.max_col==cell['column']+cell['column_span'] for r in sheet.merged_cells.ranges)
                    row['xlsx_readback']='passed'
                tick=time.perf_counter()
                preflight=build_pdf_export(store,manager,{'page_ids':[page['id']]},preflight=True)
                row['pdf_preflight']=preflight
                if preflight['ready']:
                    export=build_pdf_export(store,manager,{'page_ids':[page['id']]})
                    shutil.copy2(export,root/(sample['id']+'.pdf'))
                    shutil.copy2(export.with_suffix('.sources.json'),root/(sample['id']+'.sources.json'))
                    check=manager.cpu.call({'operation':'inspect','path':str(export),'dpi':72})
                    assert len(check['pages'])==1
                    row['pdf_readback']='one_page_verified'
                row['pdf_seconds']=time.perf_counter()-tick
                row.update(status='success',result_id=result_id,page_class=result['original'].get('document',{}).get('page_class'),
                    result_tables=len(result['edited']['tables']),cache_reused=manager.process(page['id'],'auto')==stage)
                write_json(root/(sample['id']+'-queue.json'),document_review_queue(store,documents[doc_id]))
            except Exception as error:
                row.update(status='failed',error=repr(error))
            row['wall_seconds']=time.perf_counter()-start;rows.append(row)
            write_json(root/'summary.json',{'rows':rows,'independent_quality':'unmeasured','route':'auto + ppocr fallback + default paddle/local-v2 structure detection',
                'baseline':'same product result before structure review; no independent baseline-quality claim'})
            print('C',sample['id'],row['status'],round(row['wall_seconds'],2),row.get('error',''),flush=True)
    finally:
        manager.stop();gpu.unload()


def main():
    p=argparse.ArgumentParser()
    for name in ('data','inference','output','bundle'):p.add_argument('--'+name,type=Path,required=True)
    p.add_argument('--phase',choices=['B','C','all'],default='all')
    p.add_argument('--samples',help='Explicit comma-separated follow-up IDs; original receipts are retained')
    args=p.parse_args()
    args.output.mkdir(parents=True,exist_ok=False)
    manifest=args.data/'inputs/development.json';samples=json.loads(manifest.read_text('utf-8'))['samples']
    if args.samples:
        names=set(args.samples.split(','));samples=[s for s in samples if s['id'] in names]
        if {s['id'] for s in samples}!=names:raise ValueError('Unknown follow-up sample')
    write_json(args.output/'run-lock.json',{'files':file_lock([*ROOT.glob('src/ocr_workbench/*.py'),Path(__file__),manifest]),
        'input_artifacts':file_lock([p for p in args.inference.rglob('*.json') if 'sessions' not in p.parts]),
        'scope':'public engineering; no ground truth; no human participants','defaults_changed':False,'sample_ids':[s['id'] for s in samples]})
    if args.phase in ('B','all'):controlled(args,samples)
    if args.phase in ('C','all'):product(args,samples)


if __name__=='__main__':main()
