"""Independent workbook/ZIP oracle from authoritative result snapshots, without truth."""
from __future__ import annotations
import argparse
from copy import deepcopy
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import PurePosixPath
import re
import zipfile
from common import FORMAT, digest, fpath, frozen_json, indexed, read_json
from quality import tables, raw

SOURCE_FIELDS = ["工作簿", "工作表", "表格序号", "原图", "图片 ID", "图像版本", "输入 SHA-256", "引擎", "结果 ID", "校对 revision"]
METADATA_LEGACY_FIELDS = ["类别", "识别原文", "来源", "原图坐标", "核对状态"]
METADATA_CURRENT_FIELDS = METADATA_LEGACY_FIELDS + ["另一识别候选（未人工确认）", "处理依据", "另一候选来源"]


def current_metadata(entry):
    """Resolve metadata from authoritative current candidate and review identity.

    New eight-column metadata requires its saved original evidence and preserves
    alternatives. A confirmation is current only when the human review binds
    this input and the revision produced by that review. Legacy five-column
    artifacts remain a separate explicit contract, never a header-based escape.
    """
    candidate = entry['candidate']
    original = raw(candidate)
    original_screen = original.get('screen') or {}
    edited_screen = (candidate.get('edited') or {}).get('screen') or {}
    originals = original_screen.get('business_metadata', [])
    metadata = deepcopy(edited_screen['business_metadata'] if 'business_metadata' in edited_screen else originals)
    if (not isinstance(metadata, list) or any(not isinstance(p, dict) or not isinstance(p.get('text'), str) for p in metadata)):
        raise ValueError('Authoritative metadata is not a literal part list')
    evidence = original_screen.get('business_metadata_evidence') or {}
    version = evidence.get('policy_version')
    current = (version == 'screen-business-context-v4' or bool(edited_screen.get('metadata_review'))
               or any('transcription_evidence' in p or 'alternative_text' in p for p in originals))
    errors = []
    if current and evidence.get('items') != originals:
        errors.append({'code': 'metadata_original_evidence_binding'})
    review = edited_screen.get('metadata_review') or {}
    review_bound = (review.get('state') == 'accepted' and review.get('actor') == 'human'
                    and review.get('input_version_id') == entry['input_version_id']
                    and type(review.get('base_revision')) is int
                    and entry['result_revision'] == review['base_revision'] + 1
                    and entry.get('human_confirmed') is True and entry.get('review_state') == 'accepted')
    stale_confirmation = (review.get('state') == 'accepted'
                          and entry['result_revision'] != review.get('base_revision', -2) + 1)
    logs = edited_screen.get('metadata_edit_log') or []
    sources = original_screen.get('sources') or []
    for index, part in enumerate(metadata):
        source_part = originals[index] if index < len(originals) else None
        if current and source_part is None:
            errors.append({'code': 'metadata_unbound_added_part', 'item': index})
            continue
        if source_part is not None and current:
            for key in ('kind', 'box', 'evidence_refs', 'transcription_evidence'):
                if part.get(key) != source_part.get(key):
                    errors.append({'code': 'metadata_geometry_or_evidence_changed', 'item': index, 'field': key})
            if any(part.get(key) != source_part.get(key) for key in
                   ('text', 'source', 'selected_text_basis', 'alternative_text', 'alternative_source')):
                bound_logs = [log for log in logs if log.get('item_index') == index and log.get('actor') == 'human'
                              and log.get('after') == part['text'] and log.get('image_version') == entry['input_version_id']
                              and type(log.get('base_revision')) is int and log['base_revision'] < entry['result_revision']
                              and log.get('original_selected_text') == source_part.get('text')
                              and log.get('original_evidence_refs') == source_part.get('evidence_refs', [])
                              and log.get('original_transcription_evidence') == source_part.get('transcription_evidence', {})]
                if not bound_logs:
                    errors.append({'code': 'metadata_unbound_human_edit', 'item': index})
                else:
                    log = bound_logs[-1]
                    choice = log.get('choice')
                    transcription = source_part.get('transcription_evidence') or {}
                    chosen = (transcription.get('primary_text', source_part['text']) if choice == 'primary'
                              else transcription.get('auxiliary_text') if choice == 'auxiliary' else log['after'])
                    alternative = (transcription.get('auxiliary_text', source_part['text']) if choice == 'primary'
                                   else transcription.get('primary_text', source_part['text']))
                    alternative_source = ((transcription.get('auxiliary_source') or {}).get('source') if choice == 'primary'
                                          else source_part.get('source'))
                    if (choice not in {'manual', 'primary', 'auxiliary'} or chosen != part['text']
                            or part.get('source') != ('human_edit' if choice == 'manual' else 'human_selected_' + choice)
                            or part.get('selected_text_basis') != 'human_metadata_' + choice
                            or part.get('alternative_text') != alternative
                            or part.get('alternative_source') != alternative_source):
                        errors.append({'code': 'metadata_human_choice_source_binding', 'item': index})
            transcription = source_part.get('transcription_evidence') or {}
            if source_part.get('alternative_text'):
                primary = transcription.get('primary_source') or {}
                auxiliary = transcription.get('auxiliary_source') or {}
                matching = [source for source in sources if all(source.get(key) == primary.get(key)
                             for key in ('result_id', 'task_id', 'engine', 'revision'))]
                if (source_part.get('alternative_text') != transcription.get('auxiliary_text')
                        or source_part.get('alternative_source') != auxiliary.get('source')
                        or source_part.get('text') != transcription.get('primary_text')
                        or source_part.get('primary_source') != primary or len(matching) != 1
                        or primary.get('input_sha256') != entry['image_sha256']
                        or primary.get('image_version') != entry['input_version_id']):
                    errors.append({'code': 'metadata_alternative_source_binding', 'item': index})
        if stale_confirmation:
            part.update(human_confirmed=False, review_required=True, review_reason='metadata_confirmation_revision_changed')
        elif part.get('human_confirmed') and not review_bound:
            errors.append({'code': 'metadata_false_human_confirmation', 'item': index})
    if current and len(metadata) != len(originals):
        errors.append({'code': 'metadata_current_scope_changed'})
    return metadata, METADATA_CURRENT_FIELDS if current else METADATA_LEGACY_FIELDS, errors


def number_value(value, kind):
    literal = str(value)
    if re.fullmatch(r"0\d+", literal) or re.fullmatch(r"\d{16,}", literal) or literal.startswith(("=", "+")):
        raise ValueError("Protected identifier/formula-like text cannot become numeric")
    literal = literal.replace(",", "").replace("−", "-")
    if literal.startswith("(") and literal.endswith(")"):
        literal = "-" + literal[1:-1]
    if kind == "percent":
        if not literal.endswith("%"):
            raise ValueError("Percent authorization requires percent source")
        return Decimal(literal[:-1]) / 100
    if kind != "number" or "%" in literal:
        raise ValueError("Wrong numeric authorization kind")
    return Decimal(literal)


def check_book(book, filename, entry):
    from openpyxl.utils import get_column_letter
    candidate = entry["candidate"]
    expected_tables, original = tables(candidate), raw(candidate)
    metadata, metadata_fields, metadata_errors = current_metadata(entry)
    expected_sheets = [f"Table {i + 1}" for i in range(len(expected_tables))] + ["来源索引"]
    if metadata:
        expected_sheets.append("业务说明")
    errors, count, numeric_count = list(metadata_errors), 0, 0
    if book.sheetnames != expected_sheets:
        errors.append({"code": "sheet_mapping", "expected": expected_sheets, "actual": book.sheetnames})
    authorized = {}
    authorization = entry.get("numeric_authorization")
    if authorization:
        identity = {"input_version_id": entry["input_version_id"], "result_id": entry["result_id"], "result_revision": entry["result_revision"]}
        if not authorization.get("human_confirmed") or any(authorization.get(k) != v for k, v in identity.items()):
            raise ValueError("Numeric authorization not confirmed/current")
        for cell in authorization["cells"]:
            position = cell["table_index"], cell["row"], cell["column"]
            if position in authorized:
                raise ValueError("Duplicate numeric cell authorization")
            authorized[position] = cell["kind"]
    consumed = set()
    for ti, table in enumerate(expected_tables):
        name = f"Table {ti + 1}"
        if name not in book.sheetnames:
            continue
        sheet = book[name]
        offset = int(bool(table.get("caption")))
        content = {(1, 1): table["caption"]} if offset else {}
        merges, occupied = set(), set()
        for cell in table["cells"]:
            r, c = cell["row"] + offset + 1, cell["column"] + 1
            rs, cs = cell.get("row_span", 1), cell.get("column_span", 1)
            slots = {(y, x) for y in range(r, r + rs) for x in range(c, c + cs)}
            if occupied & slots or not isinstance(cell["text"], str):
                raise ValueError("Authoritative candidate grid/type invalid")
            occupied |= slots
            expected = None if cell["text"] == "" else cell["text"]
            got = sheet.cell(r, c)
            position = ti, cell["row"], cell["column"]
            if position in authorized:
                consumed.add(position)
                want_number = number_value(cell["text"], authorized[position])
                try:
                    correct = got.data_type == "n" and Decimal(str(got.value)) == want_number
                except (InvalidOperation, ValueError):
                    correct = False
                if not correct:
                    errors.append({"code": "authorized_numeric_value", "sheet": name, "row": r, "column": c})
                if authorized[position] == "percent" and "%" not in got.number_format:
                    errors.append({"code": "authorized_percent_format", "sheet": name, "row": r, "column": c})
                content[r, c] = got.value if correct else expected
                numeric_count += 1
            else:
                if got.value != expected:
                    errors.append({"code": "literal_value", "sheet": name, "row": r, "column": c, "expected": expected, "actual": got.value})
                if expected is not None and (got.data_type != "s" or got.number_format != "@"):
                    errors.append({"code": "protected_text_type", "sheet": name, "row": r, "column": c})
                content[r, c] = expected
            if got.data_type == "f":
                errors.append({"code": "invented_formula", "sheet": name, "row": r, "column": c})
            if rs > 1 or cs > 1:
                merges.add(f"{get_column_letter(c)}{r}:{get_column_letter(c + cs - 1)}{r + rs - 1}")
            count += 1
        if offset and sheet.cell(1, 1).value != table["caption"]:
            errors.append({"code": "caption"})
        if {str(m) for m in sheet.merged_cells.ranges} != merges:
            errors.append({"code": "merged_spans", "sheet": name})
        observed = {(cell.row, cell.column): cell.value for row in sheet for cell in row if cell.value is not None}
        if observed != {p: v for p, v in content.items() if v is not None}:
            errors.append({"code": "extra_or_missing_cell", "sheet": name})
    if set(authorized) != consumed:
        raise ValueError("Numeric authorization references missing grid position")
    if "来源索引" in book.sheetnames:
        sheet = book["来源索引"]
        headers = [c.value for c in sheet[1]]
        if sheet.max_row != len(expected_tables) + 1 or any(headers.count(label) != 1 for label in SOURCE_FIELDS):
            errors.append({"code": "source_index_shape_or_headers"})
        else:
            common = {"工作簿": filename, "原图": entry["source_name"], "图片 ID": entry["image_id"],
                      "图像版本": entry["input_version_id"], "输入 SHA-256": entry["image_sha256"],
                      "引擎": original["engine"], "结果 ID": entry["result_id"], "校对 revision": str(entry["result_revision"])}
            for ti in range(len(expected_tables)):
                source = {**common, "工作表": f"Table {ti + 1}", "表格序号": str(ti + 1)}
                for label, want in source.items():
                    got = sheet.cell(ti + 2, headers.index(label) + 1)
                    if got.value != want or got.data_type != "s" or got.number_format != "@":
                        errors.append({"code": "source_binding", "field": label, "table": ti})
    if metadata and "业务说明" in book.sheetnames:
        sheet = book["业务说明"]
        if sheet.max_row != len(metadata) + 1 or [c.value for c in sheet[1]] != metadata_fields:
            errors.append({"code": "metadata_shape"})
        for row, item in enumerate(metadata, 2):
            if metadata_fields == METADATA_CURRENT_FIELDS:
                status = ('人工已确认' if item.get('human_confirmed') else
                          '需要人工核对' if item.get('review_required') else '识别候选，未逐项人工确认')
                expected = [item['kind'], item['text'], item.get('source', ''),
                            json.dumps(item.get('box'), ensure_ascii=False, separators=(',', ':')), status,
                            item.get('alternative_text', ''), item.get('review_reason', item.get('selected_text_basis', '')),
                            item.get('alternative_source', '')]
            else:
                expected = [item["kind"], item["text"], item.get("source", ""),
                            json.dumps(item.get("box"), ensure_ascii=False, separators=(",", ":")),
                            "独立文字识别，未逐项人工确认"]
            for column, value in enumerate(expected, 1):
                if sheet.cell(row, column).value != (None if value == "" else value):
                    errors.append({"code": "metadata_value", "row": row, "column": column})
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    errors.append({"code": "formula_in_any_sheet", "sheet": sheet.title, "row": cell.row, "column": cell.column})
    return {"status": "passed" if not errors else "failed", "errors": errors,
            "checked_business_cells": count, "authorized_numeric_cells": numeric_count,
            "source_fields": SOURCE_FIELDS, "metadata_rows": len(metadata),
            "metadata_contract": 'source_bound_eight_columns' if metadata_fields == METADATA_CURRENT_FIELDS else 'legacy_five_columns'}


def check_bundle(path, snapshot):
    from openpyxl import load_workbook
    path = fpath(path)
    items = indexed(snapshot["items"], "input_version_id", "authoritative current input")
    if len({i["id"] for i in items.values()}) != len(items):
        raise ValueError("Duplicate sample mapping")
    if snapshot.get("zip_sha256") and digest(path) != snapshot["zip_sha256"]:
        raise ValueError("ZIP differs from authoritative receipt")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or archive.testzip() is not None:
            raise ValueError("ZIP corrupt/duplicate entries")
        manifest = json.loads(archive.read("manifest.json"))
        entries = indexed(manifest["items"], "input_version_id", "ZIP input")
        if set(items) != set(entries) or manifest["batch_id"] != snapshot["batch_id"] or manifest["policy"] != snapshot["policy"]:
            raise ValueError("ZIP input range/batch/policy differs")
        freshness_checked = False
        if snapshot.get("version_vector") is not None:
            if (snapshot.get("current_version_vector") != snapshot["version_vector"]
                    or manifest.get("version_vector") != snapshot["version_vector"]):
                raise ValueError("Current selection/review/export-option vector changed or absent")
            freshness_checked = True
        filenames = [i["file_name"] for i in entries.values() if i.get("file_name")]
        if len(set(filenames)) != len(filenames) or set(filenames) != {n for n in names if n.endswith(".xlsx")} or len(filenames) != manifest["success_files"]:
            raise ValueError("ZIP merges/omits/unmaps workbook")
        checks = []
        for vid, current in items.items():
            exported = entries[vid]
            for field in ["image_id", "source_name", "status", "result_id", "result_revision", "review_state", "human_confirmed"]:
                if current.get(field) != exported.get(field):
                    raise ValueError(f"Stale/different manifest identity: {vid}/{field}")
            name = exported.get("file_name")
            should_export = current["status"] in {"completed", "needs_review"} and bool(tables(current.get("candidate") or {}))
            if bool(name) != should_export:
                raise ValueError("Fake success/omitted table workbook")
            if not name:
                checks.append({"id": current["id"], "status": "no_book", "native_status": current["status"], "reason": exported.get("export_error") or exported.get("error")})
                continue
            if PurePosixPath(name).name != name or "\\" in name:
                raise ValueError("Unsafe workbook ZIP path")
            payload = archive.read(name)
            book = load_workbook(io.BytesIO(payload), data_only=False)
            try:
                checks.append({"id": current["id"], "file_name": name, "xlsx_sha256": hashlib.sha256(payload).hexdigest(), **check_book(book, name, current)})
            finally:
                book.close()
    return {"format": FORMAT, "scope": "Actual workbook vs authoritative current candidate and identities; no image truth, no exporter called",
            "status": "passed" if all(c["status"] in {"passed", "no_book"} for c in checks) else "failed",
            "inputs": len(items), "workbooks": len(filenames), "no_book": sum(c["status"] == "no_book" for c in checks),
            "records": checks, "zip_sha256": digest(path),
            "fidelity_status": "passed" if all(c["status"] in {"passed", "no_book"} for c in checks) else "failed",
            "freshness_status": "passed" if freshness_checked else "not_checked",
            "current_delivery_verified": freshness_checked and all(c["status"] in {"passed", "no_book"} for c in checks),
            "limitations": ["No-book accounting is not OCR correctness or proof that a table is absent.",
                            "Caller must capture a live current-version snapshot at both start and finish; old snapshot verifies history only.",
                            "Selected-candidate/export-option freshness requires the version vector added by S05, separately checked against live state."]}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", required=True)
    parser.add_argument("--snapshot", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    report = check_bundle(args.zip, read_json(args.snapshot))
    frozen_json(args.output, report)
    print(report["status"], report["inputs"], report["workbooks"])
