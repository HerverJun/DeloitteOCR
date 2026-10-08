"""Explicit actual queue/Agent continuation after immutable recovery preparation."""
from __future__ import annotations
import argparse
import asyncio
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import http.client
import importlib
import json
import os
from pathlib import Path
import secrets
import shutil
import socket
import sqlite3
import subprocess
import sys
import threading
import time
from types import SimpleNamespace
from urllib.parse import urlsplit
import zipfile

from common import FORMAT, fingerprint, fpath, frozen_json, read_json
import controlled_recovery as recovery

GPU_SLOT = 'h60-controlled-recovery-20261008-002'


def resume_request_id(value):
    identity={'action':'resume','batch_id':value['original_batch_id'],
              'execution_scope':str(recovery.plain_path(value['output'])),
              'gpu_slot':value['gpu_slot'],'prepared_sha256':value['prepared']['sha256']}
    return 'controlledResume'+fingerprint(identity)[:32]


def save(path,value):
    frozen_json(path,value)


def append(path,value):
    path=fpath(path)
    with path.open('a',encoding='utf-8') as stream:
        stream.write(json.dumps(value,ensure_ascii=False,allow_nan=False)+'\n')


def now():return datetime.now(timezone.utc).isoformat()


def gate(prepared,binding,slot):
    if slot!=GPU_SLOT:raise ValueError('Explicit root-owned H60 GPU serial slot is required')
    validated=recovery.validate_launch(prepared,binding)
    original=read_json(prepared)
    if original.get('continuation_kind') == 'failed-recovery-followup':
        lineage=original['lineage']
        if (len(lineage['base_primary_task_ids'])!=26 or len(lineage['added_primary_task_ids'])!=10
                or len(original['preserved_primary_tasks'])!=36
                or set(original['preserved_primary_tasks'])!=set(lineage['base_primary_task_ids'])|set(lineage['added_primary_task_ids'])
                or len(lineage.get('closed_execution_segments',[]))!=1):
            raise ValueError('Actual H60 continuation requires the exact H26 to H36 stopped lineage')
    elif len(original['preserved_primary_tasks'])!=26:
        raise ValueError('Actual continuation is limited to the frozen original H60 development scope')
    if original['input_count']!=60 or len(original['successful_context_tasks'])!=60:
        raise ValueError('Actual continuation is limited to the frozen original H60 development scope')
    state=recovery.native_state(Path(original['runtime_root'])/recovery.NATIVE_DB)
    if any(row['status']!='paused' for row in state['screen_items']) or state['screen_batches'][0]['status']!='paused':
        raise ValueError('Original scope must be paused before the actual service starts')
    policy=json.loads(state['screen_batches'][0]['policy'])
    if policy.get('primary_engine')!='hunyuan' or policy.get('request_timeout_seconds')!=180 or policy.get('budget_seconds')!=360:
        raise ValueError('Frozen H60 original execution policy/budget must not change')
    return validated,original,state


def asset_inventory(bundle,policy):
    config=read_json(bundle/'config/engines.json')
    names=set(name for engine in [policy['primary_engine'],*policy['fallback_engines']] for name in config[engine]['models'])
    paths={}
    for name in sorted(names):
        for relative,entry in recovery.file_inventory(bundle/'models'/name).items():
            paths['models/'+name+'/'+relative]=entry
    for engine in [policy['primary_engine'],*policy['fallback_engines']]:
        runtime=bundle/'runtimes'/engine
        for relative in ('python.exe','python._pth','python311._pth','python312._pth','python313._pth'):
            if (runtime/relative).is_file():paths['runtimes/'+engine+'/'+relative]=recovery.bound(runtime/relative)
    paths['config/engines.json']=recovery.bound(bundle/'config/engines.json')
    return paths


def readiness(prepared,binding,original):
    if original.get('continuation_kind')!='failed-recovery-followup':return None
    scope=recovery.plain_path(original['recovery_root'])
    path=scope/'evidence/independent-readback-001/result.json'
    evidence=read_json(path)
    if (evidence.get('success') is not True or evidence['prepared']!=recovery.bound(prepared)
            or evidence['source_binding']!=recovery.bound(binding) or evidence.get('actual_recovery_proven') is not False
            or evidence.get('service_or_model_started') is not False):
        raise ValueError('Exact follow-up preparation requires CPU independent readiness evidence')
    binding=evidence['restore_cpu'];cpu=read_json(binding['path'])
    if recovery.bound(binding['path'])!=binding or cpu.get('success') is not True:
        raise ValueError('Actual platform restoration CPU qualification changed')
    return {'independent_readback':recovery.bound(path),'restore_cpu':binding,
            'cpu_only':True,'actual_recovery_proven':False,'fresh_root_gpu_slot_confirmation_required':True}


def prepare(prepared,binding,output,slot):
    recovery.require_environment()
    validated,original,state=gate(prepared,binding,slot)
    ready=readiness(prepared,binding,original)
    output=recovery.plain_path(output)
    if output.exists() or not output.is_relative_to(Path(original['recovery_root'])):
        raise ValueError('Use one new execution directory inside the prepared F recovery scope')
    policy=json.loads(state['screen_batches'][0]['policy'])
    bundle=fpath(read_json(Path(original['runtime_root'])/'prepared.json')['bundle'])
    models=asset_inventory(bundle,policy)
    output.mkdir(parents=True,exist_ok=False)
    value={'format':'controlled-native-recovery-actual-execution-v1','stage':'prepared',
        'prepared':recovery.bound(prepared),'source_binding':recovery.bound(binding),
        'output':str(output),'runtime_root':original['runtime_root'],'gpu_slot':slot,
        'sources':validated['sources'],'bundle':str(bundle),'model_asset_files':models,
        'helper':recovery.bound(Path(__file__)),'original_batch_id':original['batch_id'],
        'original_task_ids':[row['id'] for row in state['tasks']],
        'original_context_rows':state['screen_contexts'],
        'original_context_task_rows':[row for row in state['tasks'] if row['kind']=='screen_context'],
        'original_primary_task_ids':original['preserved_primary_tasks'],
        'input_version_ids':[row['input_version_id'] for row in state['screen_items']],
        'policy':policy,'deadline':state['screen_batches'][0]['deadline'],
        'readiness':ready,
        'agent_layer':'F:/OCR/output/platform-agent-dev-20261004/dependencies/agent',
        'pdf_layer':'F:/OCR/output/clear-automation-dev-20261002/worktrees/ocr/build/public-quality-20260916/tool/site-packages/pdfplumber',
        'performance_baseline_qualified':False,'truth_read':False,'model_or_service_started':False,
        'producer_note':'Original identities and frozen task metadata remain historical; actual continuation producer is separately bound to the new complete source snapshot.'}
    value['resume_request_id']=resume_request_id(value)
    save(output/'prepared-execution.json',value)
    return value


def configure(value):
    output=fpath(value['output']);runtime=fpath(value['runtime_root'])
    sources=value['sources']['sources'];platform=fpath(sources['platform']['path']);ocr=fpath(sources['ocr']['path'])
    sys.path[:0]=[str(platform),str(platform/'src'),str(ocr/'src'),str(fpath(value['pdf_layer']).parent)]
    os.environ['WORKBENCH_AGENT_LAYER']=str(fpath(value['agent_layer']))
    from workbench_platform.supervisor.processes import storage_environment
    os.environ.update(storage_environment(runtime))
    os.environ['PYTHONDONTWRITEBYTECODE']='1'
    recovery.require_environment()
    return output,runtime,platform,ocr


def native_http(service,method,path,body=None):
    parsed=urlsplit(service.base_url)
    connection=http.client.HTTPConnection('127.0.0.1',parsed.port,timeout=10)
    started=time.perf_counter()
    try:
        payload=json.dumps(body,ensure_ascii=False).encode() if body is not None else None
        connection.request(method,path,body=payload,headers={'Authorization':'Bearer '+service.token,'Content-Type':'application/json'})
        response=connection.getresponse();content=response.read(16*1024*1024+1)
        if response.status!=200 or len(content)>16*1024*1024:raise ValueError('Native actual HTTP contract failed')
        result=json.loads(content)
        append(service.trace,{'utc':now(),'method':method,'path':path,'response_seconds':time.perf_counter()-started,
            'http_status':response.status,'status':result.get('status'),'counts':result.get('counts')})
        return result
    except Exception as error:
        append(service.trace,{'utc':now(),'method':method,'path':path,'response_seconds':time.perf_counter()-started,'error_type':type(error).__name__})
        raise
    finally:connection.close()


def serve(execution_path,port):
    value=read_json(execution_path);output,runtime,platform,ocr=configure(value)
    cutoff=time.monotonic()+30
    while not (output/'native-start-allowed.json').is_file():
        if time.monotonic()>=cutoff:raise ValueError('Native service was not assigned to its owning Job')
        time.sleep(.1)
    allowed=read_json(output/'native-start-allowed.json')
    if allowed['native_pid']!=os.getpid():raise ValueError('Native ownership gate belongs to another process')
    from ocr_workbench.service import create_app
    import uvicorn
    token=(output/'native-token.txt').read_text('utf-8')
    app=create_app(fpath(value['bundle']),runtime/'native/ocr',token,start_queue=True,agent_enabled=False)
    server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=port,log_level='warning'))
    stopped=threading.Event();workers={};jobs={}
    def watch():
        while not stopped.wait(.3):
            adapter=app.state.queue.adapter
            if adapter and adapter.process:
                process=adapter.process
                if process.pid not in workers:
                    workers[process.pid]=process;jobs[process.pid]=adapter.job
                    append(output/'worker-argv.jsonl',{'utc':now(),'pid':process.pid,'engine':adapter.engine,
                        'actual_argv':process.args,'worker_source':str(adapter.worker_source),'ipc':str(adapter.ipc)})
            if (output/'stop-native.json').exists():server.should_exit=True
    watcher=threading.Thread(target=watch,daemon=True);watcher.start()
    try:server.run()
    finally:
        stopped.set();watcher.join(timeout=2)
        save(output/'native-own-cleanup.json',{'pid':os.getpid(),'workers':[
            {'pid':pid,'exit_code':process.poll(),'worker_exited':process.poll() is not None,
             'job_closed':jobs[pid] is None or jobs[pid].handle is None,
             'job_drained':jobs[pid] is None or jobs[pid].drained}
            for pid,process in workers.items()],
            'queue_alive':app.state.queue.status()['alive'],
            'module_origins':{name:str(module.__file__) for name,module in sys.modules.items()
                if name.startswith('ocr_workbench') and getattr(module,'__file__',None)}})


def restore_services(root,service,mapping):
    from workbench_platform.assets.registry import AssetRegistry
    from workbench_platform.assets.sets import AssetSetStore
    from workbench_platform.artifacts.center import ArtifactCenter
    from workbench_platform.bridge.context import ResolvedContext
    from workbench_platform.jobs.native_ocr import OcrAdapter
    from workbench_platform.jobs.ocr_handoff import OcrHandoff
    from workbench_platform.jobs.screen_batches import ScreenBatches
    from workbench_platform.store import WorkspaceStore
    workspace=WorkspaceStore(root/'workspaces.sqlite3')
    links=workspace.list_links('wreplay')
    if len(links)!=1 or links[0].module_project_id!=service.project_id:raise ValueError('Original workspace/module identity changed')
    link=links[0]
    assets=AssetRegistry(root/'assets.sqlite3',root/'blobs',{'replay':root/'inputs'})
    fixed=AssetSetStore(assets);keeper=assets._connect()
    context=ResolvedContext('wreplay','ocr','replayInstance',link.link_id,service.project_id,'owner','/apps/ocr/workflow',service.base_url)
    receipt=SimpleNamespace(instance_id='replayInstance',base_url=service.base_url)
    launches=SimpleNamespace(current=lambda family:receipt)
    jobs=SimpleNamespace(workspaces=workspace,launches=launches,_token=lambda:service.token,_live=lambda supplied:supplied==context)
    def bound(w,family,link_id):
        if (w,family,link_id)!=('wreplay','ocr',link.link_id):raise ValueError('Continuation scope mismatch')
        return context
    jobs._context=bound;jobs._adapter=lambda supplied:OcrAdapter(supplied,token=jobs._token,live=jobs._live,timeout=10)
    handoff=OcrHandoff(root/'ocr-handoff.sqlite3',jobs,assets,fixed)
    center=ArtifactCenter(root/'artifacts',assets,fixed,None,ocr_handoff=handoff)
    batches=ScreenBatches(root/'screen-batches.sqlite3',handoff,center)
    return SimpleNamespace(workspace_store=workspace,asset_library=(assets,fixed,None),job_service=jobs,launch_service=launches,
        supervisor=SimpleNamespace(registry=SimpleNamespace(installations={'ocr':SimpleNamespace(installed=True)})),
        ocr_handoff=handoff,artifact_center=center,screen_batches=batches,forensic_service=None,
        data_workbench=SimpleNamespace(list_page=lambda *a,**k:{'items':[]}),batch_store=None),keeper


def restored_mapping(root,batch_id):
    with closing(recovery.connect(root/'screen-batches.sqlite3',readonly=True)) as db:
        mapping=json.loads(db.execute('SELECT source_json FROM screen_batch_refs WHERE batch_id=?',(batch_id,)).fetchone()[0])
    with closing(recovery.connect(root/recovery.NATIVE_DB,readonly=True)) as db:
        positions={row['input_version_id']:row['position'] for row in db.execute(
            'SELECT input_version_id,position FROM screen_items WHERE batch_id=?',(batch_id,))}
    if len(mapping)!=60 or {row['input_version_id'] for row in mapping}!=set(positions):
        raise ValueError('Original platform/native input mapping differs')
    for source in mapping:source['position']=positions[source['input_version_id']]
    return mapping


def verify_restored_assets(services,mapping,root):
    assets,fixed,_=services.asset_library
    if assets.roots!={'replay':recovery.plain_path(root/'inputs')}:
        raise ValueError('Original asset root identity changed')
    db=assets._connect()
    try:
        before=[dict(row) for row in db.execute('SELECT * FROM registered_asset ORDER BY asset_id')]
        versions=[dict(row) for row in db.execute('SELECT * FROM asset_set_versions ORDER BY asset_set_id,revision')]
    finally:db.close()
    if len(before)!=60 or {row['asset_id'] for row in before}!={row['asset_id'] for row in mapping}:
        raise ValueError('Original 60 asset identities changed')
    reads=[]
    for source in mapping:
        asset=assets.get('wreplay',source['asset_id'])
        with assets.stage_verified('wreplay',asset.asset_id,revision=asset.revision,blob_hash=asset.blob_hash) as path:
            reads.append({'asset_id':asset.asset_id,'revision':asset.revision,'blob_hash':asset.blob_hash,
                          'pixel_sha256':recovery.file_sha(path),'input_version_id':source['input_version_id']})
    for row in versions:fixed.get(row['workspace_id'],row['asset_set_id'],row['revision'])
    db=assets._connect()
    try:
        after=[dict(row) for row in db.execute('SELECT * FROM registered_asset ORDER BY asset_id')]
        after_versions=[dict(row) for row in db.execute('SELECT * FROM asset_set_versions ORDER BY asset_set_id,revision')]
    finally:db.close()
    if before!=after or versions!=after_versions:
        raise ValueError('Asset restoration cannot create registrations or advance revisions')
    if any(root.joinpath('blobs').iterdir()):
        raise ValueError('Verified original referenced assets left new blob state')
    return {'asset_count':len(reads),'asset_reads':reads,'registrations_unchanged':True,
            'fixed_set_versions_unchanged':True,'assets_created':0,'revision_advanced':False}


def observe_first_downloadable_excel(services,batch_id,output):
    from workbench_platform.jobs.screen_batches import ScreenExport
    from openpyxl import load_workbook
    from io import BytesIO
    bundle=services.screen_batches.export('wreplay',batch_id,ScreenExport(include_pending_review=True))
    artifact=services.artifact_center.file('wreplay',bundle['bundle_id'],'ocr-results.zip')
    destination=output/'first-available-results.zip'
    with services.artifact_center.store.stage_download('wreplay',artifact.artifact_id,
            expected_revision=artifact.revision,expected_hash=artifact.blob_hash) as path:
        shutil.copyfile(path,destination)
    with zipfile.ZipFile(destination) as archive:
        names=[name for name in archive.namelist() if name.lower().endswith('.xlsx')]
        if not names:raise ValueError('Observed partial export has no downloadable Excel')
        for name in names:
            book=load_workbook(BytesIO(archive.read(name)),read_only=True,data_only=False)
            if not book.sheetnames:raise ValueError('Observed Excel has no sheet')
            book.close()
    return {'bundle_id':bundle['bundle_id'],'zip':recovery.bound(destination),'xlsx_files':len(names),
            'observed_utc':now(),'xlsx_readable':True,'content_quality_proven':False,
            'scope':'first successful partial ZIP download/readback; may become historical as batch continues'}


async def collect(output,runtime,value,batch_id,mapping,inputs_sha):
    from workbench_platform.jobs.screen_batches import ScreenExport
    summary=await asyncio.to_thread(value.screen_batches.status,'wreplay',batch_id,limit=1)
    items=[];records=[]
    refs=read_json(runtime/'input-map.json')['items']
    for offset in range(0,summary['counts']['total'],30):
        page=await asyncio.to_thread(value.screen_batches.status,'wreplay',batch_id,offset=offset,limit=30)
        for item in page['items']:
            detail=await asyncio.to_thread(value.screen_batches.detail,'wreplay',batch_id,item['input_version_id'])
            source=next(source for source in mapping if source['input_version_id']==item['input_version_id'])
            sample=next(row for row in refs if row['position']==source['position'])
            detail.update(sample_id=sample['sample_id'],source_image_sha256=sample['source_image_sha256'])
            save(output/'native-receipts'/(sample['sample_id']+'.json'),detail)
            result=detail.get('result') or {};candidate={'original':result.get('original') or {},'edited':result.get('edited') or {}}
            entry={'id':sample['sample_id'],'input_version_id':item['input_version_id'],'image_id':item['image_id'],
                'source_name':detail['source_name'],'image_sha256':sample['source_image_sha256'],
                'result_id':result.get('id'),'result_revision':result.get('revision'),
                'status':detail['status'],'review_state':detail['review_state'],'human_confirmed':detail.get('human_confirmed',False),'candidate':candidate}
            items.append(entry)
            records.append({**entry,'status':'success' if result else detail['status'],
                'review_image':detail['status']=='needs_review','review_regions':[],
                'automatic_delivery':False,'issues':detail.get('issues',[])})
    save(output/'observations.json',{'format':FORMAT,'evidence_kind':'native_model_receipts','inputs_sha256':inputs_sha,
        'records':records,'truth_read':False,'controlled_continuation':True,
        'partial_run':summary['counts']['pending']>0,'performance_baseline_qualified':False})
    export=None
    if summary['counts']['completed']+summary['counts']['needs_review']:
        before=await asyncio.to_thread(value.screen_batches.delivery_snapshot,'wreplay',batch_id)
        bundle=await asyncio.to_thread(value.screen_batches.export,'wreplay',batch_id,ScreenExport(include_pending_review=True))
        artifact=value.artifact_center.file('wreplay',bundle['bundle_id'],'ocr-results.zip')
        target=output/'actual-results.zip'
        with value.artifact_center.store.stage_download('wreplay',artifact.artifact_id,expected_revision=artifact.revision,expected_hash=artifact.blob_hash) as path:
            shutil.copyfile(path,target)
        after=await asyncio.to_thread(value.screen_batches.delivery_snapshot,'wreplay',batch_id)
        options={'include_pending_review':True,'numeric_columns':None}
        vector=lambda snapshot:{**{key:snapshot[key] for key in ('snapshot_id','batch_id','project_id','policy_sha256','items')},'export_options':options}
        snapshot={'batch_id':batch_id,'policy':summary['policy'],'zip_sha256':artifact.blob_hash,'items':items,
            'version_vector':vector(before),'current_version_vector':vector(after),'truth_read':False,'manifest_used_as_oracle':False}
        save(output/'readback-snapshot.json',snapshot)
        from readback import check_bundle
        readback=check_bundle(target,snapshot)
        save(output/'xlsx-zip-readback.json',readback)
        export={'bundle_id':bundle['bundle_id'],'zip':recovery.bound(target),'xlsx_files':readback['workbooks'],
            'no_book':readback['no_book'],'fidelity':readback['fidelity_status'],'freshness':readback['freshness_status'],
            'current_delivery_verified':readback['current_delivery_verified'],'post_agent_independent_export':True}
    return {'summary':summary,'export':export,'observations':recovery.bound(output/'observations.json')}


async def run(execution_path,slot):
    recovery.require_environment()
    value=read_json(execution_path)
    if value['stage']!='prepared' or value['gpu_slot']!=slot or slot!=GPU_SLOT:raise ValueError('Execution slot/preparation changed')
    if recovery.bound(Path(__file__))!=value['helper']:raise ValueError('Actual runner changed after CPU preparation')
    if value.get('resume_request_id')!=resume_request_id(value):raise ValueError('Execution-specific original batch resume identity changed')
    checked,original,_=gate(value['prepared']['path'],value['source_binding']['path'],slot)
    if readiness(value['prepared']['path'],value['source_binding']['path'],original)!=value.get('readiness'):
        raise ValueError('Independent preparation readiness changed before dispatch')
    if checked['sources']!=value['sources']:raise ValueError('Actual source binding changed')
    output=fpath(value['output'])
    if (output/'started.json').exists():raise ValueError('Actual execution cannot restart an existing attempt')
    policy=value['policy'];bundle=fpath(value['bundle'])
    if asset_inventory(bundle,policy)!=value['model_asset_files']:raise ValueError('Original model assets changed before actual dispatch')
    output,runtime,platform,ocr=configure(value)
    from ocr_workbench.processes import ProcessJob
    from workbench_platform.agent.runtime import AgentRuntime
    import workbench_platform.agent.runtime as runtime_module
    from scripts.evaluate_screen_agent_native import native_timing_receipts
    token=secrets.token_urlsafe(36);(output/'native-token.txt').write_text(token,'utf-8')
    with socket.socket() as probe:probe.bind(('127.0.0.1',0));port=probe.getsockname()[1]
    service=SimpleNamespace(base_url=f'http://127.0.0.1:{port}',token=token,project_id=original['native_project_id'],trace=output/'native-http-timings.jsonl')
    python=bundle/'runtimes/control/python.exe'
    bootstrap="import sys,runpy;sys.path.insert(0,sys.argv.pop(1));runpy.run_path(sys.argv.pop(1),run_name='__main__')"
    argv=[str(python),'-B','-u','-X','utf8','-c',bootstrap,str(Path(__file__).parent),str(Path(__file__)),'serve',
          '--execution',str(execution_path),'--port',str(port)]
    save(output/'started.json',{'format':value['format'],'utc':now(),'runner_pid':os.getpid(),'runner_actual_argv':list(sys.orig_argv),
        'native_service_actual_argv':argv,'cwd':str(runtime),'gpu_slot':slot,'truth_read':False,'performance_baseline_qualified':False})
    started=time.perf_counter();process=None;job=None;agent=None;keeper=None;run_id=None;final=None;collected=None;error=None
    requests=[];waiting=[];baseline_requests=0
    first_model_result=None;first_excel=None;partial_export_at_count=-1;partial_export_failures=[]
    def mark(stage,**details):append(output/'stage-timestamps.jsonl',{'stage':stage,'utc':now(),'elapsed_seconds':time.perf_counter()-started,**details})
    original_request=runtime_module.request
    async def measured(config,key,payload,request_id,**kwargs):
        save(output/'controller-inputs'/(hashlib.sha256(request_id.encode()).hexdigest()+'.json'),{'request_id':request_id,'payload':payload,'truth_read':False})
        start=time.perf_counter()
        try:return await original_request(config,key,payload,request_id,**kwargs)
        finally:requests.append({'request_id':request_id,'wall_seconds':time.perf_counter()-start})
    with (output/'native-service.log').open('wb') as log:
        try:
            mapping=restored_mapping(runtime,original['batch_id'])
            services,keeper=restore_services(runtime,service,mapping)
            save(output/'restored-assets-pre-dispatch.json',verify_restored_assets(services,mapping,runtime))
            agent=AgentRuntime(runtime/'agent',services);services.agent=agent
            runs=agent.store.rows('SELECT id,status FROM runs')
            if len(runs)!=1 or runs[0]['status']!='needs_verification':raise ValueError('Original Agent identity/state differs')
            run_id=runs[0]['id'];baseline_requests=len(agent.store.rows('SELECT id FROM model_requests WHERE run_id=?',(run_id,)))
            mark('platform_services_and_original_agent_restored_before_native_start',run_id=run_id)
            job=ProcessJob();process=subprocess.Popen(argv,cwd=runtime,env=os.environ.copy(),stdout=log,stderr=subprocess.STDOUT,creationflags=subprocess.CREATE_NO_WINDOW)
            job.assign(process);save(output/'native-start-allowed.json',{'native_pid':process.pid,'assigned_owned_job':True})
            mark('native_started',pid=process.pid)
            print(json.dumps({'stage':'native_started','runner_pid':os.getpid(),'native_service_pid':process.pid,'gpu_slot':slot}),flush=True)
            ready=time.perf_counter()+90
            while time.perf_counter()<ready:
                if process.poll() is not None:raise ValueError('Owned native service exited before readiness')
                try:
                    health=await asyncio.to_thread(native_http,service,'GET','/api/health')
                    if health.get('status')=='ready' and health.get('queue',{}).get('healthy'):break
                except (OSError,ValueError):pass
                await asyncio.sleep(.5)
            else:raise ValueError('Owned native queue did not become ready')
            initial=await asyncio.to_thread(native_http,service,'GET',f'/api/projects/{service.project_id}/screen-batches/{original["batch_id"]}?limit=1')
            if initial['status']!='paused' or initial['policy']!=policy:raise ValueError('Original paused policy differs after service start')
            resume=await asyncio.to_thread(native_http,service,'POST',f'/api/projects/{service.project_id}/screen-batches/{original["batch_id"]}/action',
                {'action':'resume','request_id':value['resume_request_id']})
            mark('original_batch_resumed',batch_id=resume['batch_id'],counts=resume['counts'])
            runtime_module.request=measured;await agent.start()
            mark('agent_reconcile_start',run_id=run_id)
            restored=await agent.reconcile(run_id)
            if restored['id']!=run_id:raise ValueError('Agent reconcile replaced the durable original run')
            mark('agent_reconciled',run_id=run_id,status=restored['status'])
            last_native=0
            while time.time()<value['deadline']:
                current=agent.store.run(run_id)
                count=len(agent.store.rows('SELECT id FROM model_requests WHERE run_id=?',(run_id,)))
                if current['status']=='waiting_job':waiting.append({'elapsed_seconds':time.perf_counter()-started,'model_requests':count})
                if time.perf_counter()-last_native>=5:
                    state=await asyncio.to_thread(native_http,service,'GET',f'/api/projects/{service.project_id}/screen-batches/{original["batch_id"]}?limit=1')
                    append(output/'progress.jsonl',{'utc':now(),'elapsed_seconds':time.perf_counter()-started,'agent_status':current['status'],
                        'model_requests':count,'native_status':state['status'],'counts':state['counts']})
                    print(json.dumps({'stage':'progress','agent_status':current['status'],'native_status':state['status'],'counts':state['counts'],'model_requests':count}),flush=True)
                    if first_model_result is None:
                        with closing(recovery.connect(runtime/recovery.NATIVE_DB,readonly=True)) as observation:
                            result=observation.execute('SELECT id FROM results ORDER BY rowid LIMIT 1').fetchone()
                        if result:
                            first_model_result={'result_id':result['id'],'observed_utc':now(),
                                                'elapsed_seconds':time.perf_counter()-started}
                            mark('first_native_model_result_observed',**first_model_result)
                    available=state['counts']['completed']+state['counts']['needs_review']
                    if first_excel is None and available>partial_export_at_count and available>0:
                        partial_export_at_count=available
                        try:
                            first_excel=await asyncio.to_thread(observe_first_downloadable_excel,services,original['batch_id'],output)
                            first_excel['elapsed_seconds']=time.perf_counter()-started
                            mark('first_downloadable_excel_readback',**first_excel)
                        except Exception as problem:
                            fault={'type':type(problem).__name__,'code':getattr(problem,'code',None),'available_count':available}
                            partial_export_failures.append(fault);mark('partial_export_observation_failed',**fault)
                    last_native=time.perf_counter()
                if current['status'] not in {'running','waiting_job'}:break
                await asyncio.sleep(1)
            final=agent.store.run(run_id);mark('agent_terminal',run_id=run_id,status=final['status'])
            collected=await collect(output,runtime,services,original['batch_id'],mapping,read_json(runtime/'prepared.json')['inputs_sha256'])
            operations=agent.store.rows('SELECT id,run_id,call_id,tool,state,result,created,updated FROM operations WHERE run_id=? ORDER BY created',(run_id,))
            save(output/'agent-operations.json',operations)
            save(output/'agent-model-requests.json',[{key:row[key] for key in ('id','run_id','step','state','fingerprint','created')}|
                {'response_sha256':hashlib.sha256(str(row['response'] or '').encode()).hexdigest()}
                for row in agent.store.rows('SELECT * FROM model_requests WHERE run_id=?',(run_id,))])
            save(output/'agent-events.json',agent.store.events(final['session_id'],limit=10000)['items'])
        except BaseException as failure:
            import traceback
            error={'type':type(failure).__name__,'code':getattr(failure,'code',None),
                   'locations':[{'file':frame.filename,'line':frame.lineno,'function':frame.name}
                                for frame in traceback.extract_tb(failure.__traceback__)]}
            mark('failure',**error)
            if agent and run_id:final=agent.store.run(run_id)
        finally:
            mark('cleanup_start')
            if agent:await agent.close()
            runtime_module.request=original_request
            if keeper:keeper.close()
            save(output/'stop-native.json',{'utc':now(),'owned_native_pid':process.pid if process else None})
            if process:
                try:await asyncio.to_thread(process.wait,45)
                except subprocess.TimeoutExpired:
                    if job:await asyncio.to_thread(job.drain,15)
                    await asyncio.to_thread(process.wait,15)
            if job:
                await asyncio.to_thread(job.drain,15);job.close()
            mark('cleanup_finished',native_exit_code=process.poll() if process else None,owned_job_drained=bool(job and job.drained))
    tables={}
    with closing(sqlite3.connect((runtime/recovery.NATIVE_DB).as_uri()+'?mode=ro',uri=True)) as db:
        db.row_factory=sqlite3.Row
        for name in ('tasks','results','versions','images','screen_batches','screen_items','screen_tasks','screen_contexts','ocr_submissions','selections','reviews'):
            tables[name]=[dict(row) for row in db.execute('SELECT * FROM '+name+' ORDER BY rowid')]
            save(output/'native-tables'/(name+'.json'),tables[name])
    save(output/'native-task-timings.json',native_timing_receipts(tables))
    context_ids=[row['id'] for row in tables['tasks'] if row['kind']=='screen_context']
    context_preserved=(context_ids==[row['id'] for row in value['original_context_task_rows']]
        and [row for row in tables['tasks'] if row['kind']=='screen_context']==value['original_context_task_rows']
        and tables['screen_contexts']==value['original_context_rows'])
    primary_ids={row['task_id'] for row in tables['screen_tasks']}
    identity_ok=set(value['original_primary_task_ids'])<=primary_ids and len(tables['screen_batches'])==1 and tables['screen_batches'][0]['deadline']==value['deadline']
    source_after=recovery.frozen_sources(value['sources']['freeze']['path'])
    assets_after=asset_inventory(bundle,policy)
    save(output/'model-assets-after.json',assets_after)
    save(output/'controller-timings.json',requests)
    proof=read_json(original['original_inventory']['path'])
    original_unchanged=recovery.file_inventory(Path(original['original_run']))==proof
    byte_copy_unchanged=recovery.file_inventory(Path(original['original_byte_copy']))==proof
    summary={'format':value['format'],'status':final['status'] if final else 'failed','error':error,'run_id':run_id,
        'native_batch_id':original['batch_id'],'scope_inputs':60,'wall_seconds':time.perf_counter()-started,
        'wall_scope':'actual continuation through export/readback/cleanup/assets checks; original failed work reported separately',
        'agent_reconciled':any(json.loads(line)['stage']=='agent_reconciled' for line in (output/'stage-timestamps.jsonl').read_text('utf-8').splitlines()),
        'collected':collected,'original_context_tasks_unchanged_no_recomputation':context_preserved,
        'original_primary_task_ids_preserved':identity_ok,'original_primary_count':len(value['original_primary_task_ids']),'actual_primary_task_count':len(tables['screen_tasks']),
        'new_tasks_are_remaining_inputs_of_original_batch':primary_ids-set(value['original_primary_task_ids']),
        'new_native_batches':len(tables['screen_batches'])-1,'frozen_source_unchanged':source_after==value['sources'],
        'model_assets_unchanged':assets_after==value['model_asset_files'],'original_run_unchanged':original_unchanged,
        'original_byte_copy_unchanged':byte_copy_unchanged,'owned_service_exited':process is not None and process.poll() is not None,
        'owned_job_drained_and_closed':bool(job and job.drained and job.handle is None),
        'baseline_model_requests':baseline_requests,'new_controller_requests':len(requests),'waiting_request_counts':waiting,
        'first_native_model_result':first_model_result,'first_downloadable_excel':first_excel,
        'partial_export_observation_failures':partial_export_failures,
        'performance_baseline_qualified':False,'truth_read':False,'new_producer_source':value['sources'],
        'original_task_provenance':original['termination_evidence'],
        'helper_unchanged':recovery.bound(Path(__file__))==value['helper']}
    summary['new_tasks_are_remaining_inputs_of_original_batch']=sorted(summary['new_tasks_are_remaining_inputs_of_original_batch'])
    save(output/'result.json',summary)
    print(json.dumps({key:summary[key] for key in ('status','error','agent_reconciled','original_context_tasks_unchanged_no_recomputation','original_primary_task_ids_preserved','actual_primary_task_count','owned_service_exited','owned_job_drained_and_closed')},ensure_ascii=False),flush=True)
    if error or not all(summary[key] for key in ('agent_reconciled','original_context_tasks_unchanged_no_recomputation','original_primary_task_ids_preserved','frozen_source_unchanged','model_assets_unchanged','original_run_unchanged','original_byte_copy_unchanged','owned_service_exited','owned_job_drained_and_closed','helper_unchanged')):
        raise SystemExit(1)


def main():
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    item=sub.add_parser('prepare')
    for name in ('prepared','source-binding','output','gpu-slot'):item.add_argument('--'+name,required=True)
    item=sub.add_parser('run')
    item.add_argument('--execution',required=True);item.add_argument('--gpu-slot',required=True)
    item=sub.add_parser('serve');item.add_argument('--execution',required=True);item.add_argument('--port',required=True,type=int)
    args=parser.parse_args()
    if args.command=='prepare':
        value=prepare(args.prepared,args.source_binding,args.output,args.gpu_slot)
        print(json.dumps({'stage':'prepared','output':value['output'],'gpu_slot':value['gpu_slot'],'model_or_service_started':False}),flush=True)
    elif args.command=='run':asyncio.run(run(fpath(args.execution),args.gpu_slot))
    else:serve(fpath(args.execution),args.port)


if __name__=='__main__':main()
