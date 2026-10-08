"""Meaningful CPU positive/negative gates; fixtures are never OCR evaluation evidence."""
from __future__ import annotations
import argparse
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from openpyxl import Workbook
from common import FORMAT, digest, fpath, frozen_json
from quality import evaluate, match_tables, paired_report, score_sample
from readback import check_book, check_bundle, SOURCE_FIELDS
from native_bridge import normalize_native_index
from raw_bridge import collect as collect_raw, source_hash

OUTPUT = None


def truth(sid="s1"):
    return {"id": sid, "image_sha256": "a" * 64, "cohort": "chinese", "group_id": "fixture-family", "source_id": sid,
            "level": "clean", "challenge": None,
            "tables": [{"rows": 2, "columns": 2, "caption": "", "cells": [
                {"row": 0, "column": 0, "text": "编号", "row_span": 1, "column_span": 1, "visibility": "visible"},
                {"row": 0, "column": 1, "text": "金额", "row_span": 1, "column_span": 1, "visibility": "visible"},
                {"row": 1, "column": 0, "text": "000123", "row_span": 1, "column_span": 1, "visibility": "visible"},
                {"row": 1, "column": 1, "text": "−20.00", "row_span": 1, "column_span": 1, "visibility": "visible"}]}],
            "context_targets": [{"kind": "title", "text": "往来核验", "visibility": "visible"},
                                {"kind": "unit", "text": "单位：元", "visibility": "visible"},
                                {"kind": "note", "text": "B 是业务代码", "visibility": "visible"}],
            "ui_markers": ["粘贴", "Excel"], "ui_regions": []}


def observation(target):
    return {"id": target["id"], "image_sha256": target["image_sha256"], "status": "success", "result_id": "res-" + target["id"],
            "result_revision": 0, "candidate": {"engine": "fixture", "tables": deepcopy(target["tables"]),
                "screen": {"business_metadata": [{"kind": p["kind"], "text": p["text"]} for p in target["context_targets"]]}},
            "review_image": True, "review_regions": [], "automatic_delivery": False, "human_confirmed": False}


def entry():
    source = truth()
    record = observation(source)
    record["candidate"]["tables"][0]["cells"] += [
        {"row": 2, "column": 0, "text": "202610070000000000005700", "row_span": 1, "column_span": 1},
        {"row": 2, "column": 1, "text": "=0.8", "row_span": 1, "column_span": 1}]
    record["candidate"]["tables"][0]["rows"] = 3
    for p in record["candidate"]["screen"]["business_metadata"]:
        p.update(source="fixture", box=[1, 2, 3, 4])
    return {**record, "status": "needs_review", "input_version_id": "v1", "image_id": "im1", "source_name": "重复.png",
            "review_state": "automatic", "file_name": "0001-重复.xlsx"}


def workbook(source):
    book = Workbook()
    book.remove(book.active)
    table = source["candidate"]["tables"][0]
    sheet = book.create_sheet("Table 1")
    for cell in table["cells"]:
        value = sheet.cell(cell["row"] + 1, cell["column"] + 1)
        value.value = cell["text"] or None
        if cell["text"]:
            value.data_type = "s"
            value.number_format = "@"
    index = book.create_sheet("来源索引")
    index.append(SOURCE_FIELDS)
    mapping = {"工作簿": source["file_name"], "工作表": "Table 1", "表格序号": "1", "原图": source["source_name"],
               "图片 ID": source["image_id"], "图像版本": source["input_version_id"], "输入 SHA-256": source["image_sha256"],
               "引擎": "fixture", "结果 ID": source["result_id"], "校对 revision": "0"}
    index.append([mapping[field] for field in SOURCE_FIELDS])
    for row in index:
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
    notes = book.create_sheet("业务说明")
    notes.append(["类别", "识别原文", "来源", "原图坐标", "核对状态"])
    for part in source["candidate"]["screen"]["business_metadata"]:
        notes.append([part["kind"], part["text"], part["source"], json.dumps(part["box"], separators=(",", ":")), "独立文字识别，未逐项人工确认"])
    return book


class QualityTests(unittest.TestCase):
    def test_missing_and_failed_remain_in_denominator(self):
        a, b = truth(), truth("s2")
        obs = observation(a)
        obs["status"] = "failed"
        report = evaluate([a, b], [obs])
        total = report["summary"]["totals"]
        self.assertEqual(total["images"], 2)
        self.assertEqual(total["visible_critical_cells"], 4)
        self.assertEqual(total["failed_images"], 1)
        self.assertEqual(total["missing_images"], 1)

    def test_duplicate_extra_hash_and_result_identity_rejected(self):
        t, o = truth(), observation(truth())
        for rows in [[o, o], [{**o, "id": "unknown"}], [{**o, "image_sha256": "b" * 64}], [{**o, "result_id": None}]]:
            with self.assertRaises(ValueError):
                evaluate([t], rows)

    def test_all_human_review_is_not_localization_or_automation(self):
        t = truth()
        o = observation(t)
        o["candidate"]["tables"][0]["cells"][2]["text"] = "123"
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["review_reference_cells"], 4)
        self.assertEqual(m["local_error_hits"], 0)
        self.assertEqual(m["automatic_images"], 0)

    def test_failed_missing_output_is_not_silent_automatic_release(self):
        t, o = truth(), observation(truth())
        o.update(status="failed", review_image=False)
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["visible_critical_error_cells"], 2)
        self.assertEqual(m["silent_critical_error_cells"], 0)

    def test_native_bridge_does_not_promote_generic_flag_to_local_evidence(self):
        index = {"format": "actual-native-offline-receipt-index-v1", "truth_opened": False,
                 "inputs_sha256": "a", "native_table_identity_sha256": "b", "native_database": "fixture-only",
                 "batch_id": "b1", "actual_policy": {},
                 "rows": [{"id": "s1", "input_sha256": "a", "scoring_status": "success", "native_status": "needs_review",
                           "issues": [{"code": "quality_review_required"}], "human_accepted": False,
                           "result_id": "r1", "result_revision": 0, "selected_candidate": {"quality": {"automatic_delivery": False}}}]}
        record = normalize_native_index(index)["records"][0]
        self.assertTrue(record["review_image"])
        self.assertEqual(record["review_regions"], [])
        self.assertFalse(record["automatic_delivery"])
        index["truth_opened"] = True
        with self.assertRaises(ValueError):
            normalize_native_index(index)

    def test_existing_uncorrected_error_released_is_new_silent_error(self):
        t = truth()
        before = observation(t)
        before["candidate"]["tables"][0]["cells"][2]["text"] = "123"
        before["candidate"]["tables"][0]["cells"][3]["text"] = "20"
        after = deepcopy(before)
        after["candidate"]["tables"][0]["cells"][3]["text"] = "−20.00"
        after["review_image"] = False
        after["automatic_delivery"] = True
        paired = paired_report(evaluate([t], [before]), evaluate([t], [after]))
        self.assertEqual(paired["totals"]["net_visible_cell_repairs"], 1)
        self.assertEqual(paired["totals"]["new_silent_critical_release_errors"], 1)
        self.assertFalse(paired["content_gate_passes"])

    def test_unreliable_or_unproven_index_is_not_local_detection(self):
        t, o = truth(), observation(truth())
        o["review_image"] = False
        o["review_regions"] = [{"scope": "cell", "table_index": 0, "row": 1, "column": 0, "reliable": False, "evidence_ref": "image"},
                               {"scope": "cell", "table_index": 0, "row": 1, "column": 1, "reliable": True}]
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["local_review_cells"], 0)
        self.assertEqual(m["review_reference_cells"], 0)
        self.assertEqual(m["invalid_review_regions"], 2)

    def test_hidden_references_not_in_visible_error_denominator(self):
        t = truth()
        t["tables"][0]["cells"][2]["visibility"] = "occluded"
        o = observation(t)
        o["candidate"]["tables"][0]["cells"][2]["text"] = "wrong"
        o["review_image"] = False
        o["automatic_delivery"] = True
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["visible_critical_cells"], 1)
        self.assertEqual(m["visible_error_cells"], 0)
        self.assertEqual(m["uncertain_nonempty_outputs"], 1)
        self.assertEqual(m["automatic_uncertain_images"], 1)
        self.assertEqual(m["automatic_error_images"], 1)

    def test_ui_and_metadata_duplicate_detected_without_deleting_business_B(self):
        t, o = truth(), observation(truth())
        o["candidate"]["screen"]["business_metadata"].extend([{"kind": "note", "text": "000123"}, {"kind": "note", "text": "粘贴"}])
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["metadata_exact_targets"], 3)
        self.assertEqual(m["metadata_extra_parts"], 2)
        self.assertEqual(m["metadata_duplicate_table_text_parts"], 1)
        self.assertEqual(m["ui_marker_leaks"], 1)

    def test_metadata_kind_sequence_and_missing_checked(self):
        t, o = truth(), observation(truth())
        o["candidate"]["screen"]["business_metadata"] = [{"kind": "note", "text": "单位：元"}, {"kind": "title", "text": "往来核验"}]
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["metadata_missing_targets"], 2)
        self.assertEqual(m["metadata_sequence_exact_images"], 0)

    def test_metadata_order_is_part_of_annotated_exactness(self):
        t, o = truth(), observation(truth())
        o["candidate"]["screen"]["business_metadata"].reverse()
        m = score_sample(t, o)["metrics"]
        self.assertEqual(m["metadata_exact_targets"], 3)
        self.assertEqual(m["metadata_sequence_exact_images"], 0)
        self.assertFalse(m["annotated_content_exact_images"])

    def test_existing_automatic_residuals_cannot_pass_gate_with_one_fix(self):
        t, before = truth(), observation(truth())
        before.update(review_image=False, automatic_delivery=True)
        before["candidate"]["tables"][0]["cells"][2]["text"] = "123"
        before["candidate"]["tables"][0]["cells"][3]["text"] = "20"
        after = deepcopy(before)
        after["candidate"]["tables"][0]["cells"][3]["text"] = "−20.00"
        result = paired_report(evaluate([t], [before]), evaluate([t], [after]))
        self.assertEqual(result["totals"]["new_silent_critical_release_errors"], 0)
        self.assertFalse(result["content_gate_passes"])

    def test_sign_glyph_differences_not_repaired(self):
        t, o = truth(), observation(truth())
        o["candidate"]["tables"][0]["cells"][3]["text"] = "-20.00"
        self.assertEqual(score_sample(t, o)["metrics"]["visible_critical_error_cells"], 1)

    def test_missing_first_table_does_not_misalign_second(self):
        a = truth()["tables"][0]
        b = deepcopy(a)
        b["cells"][2]["text"] = "987654"
        b["cells"][3]["text"] = "99.00"
        pairs, unmatched = match_tables([a, b], [b])
        self.assertIsNone(pairs[0][2])
        self.assertEqual(pairs[1][2], 0)
        self.assertEqual(unmatched, [])


class ReadbackTests(unittest.TestCase):
    def test_literal_ids_and_formula_like_are_preserved(self):
        e = entry()
        self.assertEqual(check_book(workbook(e), e["file_name"], e)["status"], "passed")

    def test_changed_value_source_merge_and_formula_detected(self):
        e = entry()
        for mutation in ["value", "source", "merge", "formula"]:
            b = workbook(e)
            if mutation == "value":
                b["Table 1"].cell(2, 1).value = "123"
            elif mutation == "source":
                b["来源索引"].cell(2, 9).value = "different-result"
            elif mutation == "merge":
                b["Table 1"].merge_cells("A2:B2")
            else:
                b["业务说明"].cell(2, 2).value = "=0.8"
            self.assertEqual(check_book(b, e["file_name"], e)["status"], "failed", mutation)

    def test_numeric_confirmation_and_revision_required(self):
        e = entry()
        authorization = {"human_confirmed": True, "input_version_id": "v1", "result_id": e["result_id"], "result_revision": 0,
                         "cells": [{"table_index": 0, "row": 1, "column": 1, "kind": "number"}]}
        e["numeric_authorization"] = authorization
        b = workbook(e)
        b["Table 1"].cell(2, 2).value = -20
        self.assertEqual(check_book(b, e["file_name"], e)["status"], "passed")
        for key, value in [("human_confirmed", False), ("result_revision", 2)]:
            changed = deepcopy(e)
            changed["numeric_authorization"][key] = value
            with self.assertRaises(ValueError):
                check_book(b, e["file_name"], changed)

    def test_authorized_identifier_and_formula_conversion_rejected(self):
        e = entry()
        for row, column in [(1, 0), (2, 0), (2, 1)]:
            changed = deepcopy(e)
            changed["numeric_authorization"] = {"human_confirmed": True, "input_version_id": "v1", "result_id": e["result_id"], "result_revision": 0,
                                                "cells": [{"table_index": 0, "row": row, "column": column, "kind": "number"}]}
            with self.assertRaises(ValueError):
                check_book(workbook(changed), e["file_name"], changed)

    def test_same_count_old_zip_and_selection_vector_rejected(self):
        e = entry()
        data = io.BytesIO()
        workbook(e).save(data)
        vector = {"v1": {"result_id": e["result_id"], "revision": 0, "selected_candidate_id": "candidate1", "review_epoch": 0, "options": {}}}
        manifest = {"batch_id": "b1", "policy": {}, "success_files": 1, "items": [{k: v for k, v in e.items() if k != "candidate"}], "version_vector": vector}
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            path = Path(temporary) / "actual-fixture.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("manifest.json", json.dumps(manifest))
                archive.writestr(e["file_name"], data.getvalue())
            snapshot = {"batch_id": "b1", "policy": {}, "items": [e], "version_vector": vector, "current_version_vector": deepcopy(vector)}
            self.assertTrue(check_bundle(path, snapshot)["current_delivery_verified"])
            snapshot["current_version_vector"]["v1"]["selected_candidate_id"] = "candidate2"
            with self.assertRaises(ValueError):
                check_bundle(path, snapshot)
            snapshot.pop("version_vector")
            snapshot.pop("current_version_vector")
            self.assertFalse(check_bundle(path, snapshot)["current_delivery_verified"])
            snapshot["items"][0]["result_revision"] = 1
            with self.assertRaises(ValueError):
                check_bundle(path, snapshot)

    def test_off_drive_target_rejected_before_write(self):
        with self.assertRaises(ValueError):
            fpath("C:/Users/A/Desktop/eval-test.json")


def raw_fixture(temporary):
    root = Path(temporary)
    worker = root / "worker"
    worker.mkdir()
    (worker / "fixture.py").write_text("# identity-only CPU fixture; no model\n", "utf-8")
    bundle = root / "bundle"
    (bundle / "config").mkdir(parents=True)
    frozen_json(bundle / "config/engines.json", {"fixture": True})
    image = root / "opaque-fixture-input.bin"
    image.write_bytes(b"identity-only CPU fixture, not OCR input")
    inputs = root / "inputs.json"
    frozen_json(inputs, {"truth_available_to_inference": False, "samples": [{"id": "s1", "image": str(image), "image_sha256": digest(image)}]})
    run = root / "run"
    (run / "s1").mkdir(parents=True)
    frozen_json(run / "binding.json", {"inputs_sha256": digest(inputs), "worker_source_sha256": source_hash(worker),
                "engine": "fixture", "engine_config_sha256": digest(bundle / "config/engines.json"), "bundle": str(bundle),
                "labels_or_ideal_regions_provided": False, "offset": 0, "limit": None})
    result = run / "s1/result.json"
    frozen_json(result, {"engine": "fixture", "tables": []})
    frozen_json(run / "s1/record.json", {"id": "s1", "image_sha256": digest(image), "engine": "fixture", "status": "success",
                "result_sha256": digest(result), "prepared_images": [], "quality": {}, "attempt": 1})
    return run, inputs, worker


def partial_fixture(temporary, include_success=False):
    """Synthetic termination protocol only, never claims an actual stopped process."""
    run, inputs, worker = raw_fixture(temporary)
    manifest = json.loads(inputs.read_text("utf-8"))
    first = manifest["samples"][0]
    manifest["samples"] = [{**first, "id": f"s{i + 1}"} for i in range(12)]
    inputs.write_text(json.dumps(manifest), "utf-8")
    binding_path = run / "binding.json"
    binding = json.loads(binding_path.read_text("utf-8"))
    binding["inputs_sha256"] = digest(inputs)
    binding_path.write_text(json.dumps(binding), "utf-8")
    states = ["success", "failed", "running"] if include_success else ["failed", "failed", "failed", "running"]
    for index, status in enumerate(states):
        folder = run / f"s{index + 1}"
        folder.mkdir(exist_ok=True)
        if status == "success":
            continue
        record = {"id": f"s{index + 1}", "image_sha256": first["image_sha256"], "engine": "fixture", "status": status,
                  "attempt": 1, "output": str(folder / "attempt-0001")}
        if status == "failed":
            record.update(error_type="TimeoutError", error="generic fixture timeout", failure_phase="fixture_explicit_load")
        (folder / "record.json").write_text(json.dumps(record), "utf-8")
    executable = str(Path(temporary) / "fixture-python.exe")
    command = f'"{executable}" -B fixture-run.py --inputs "{inputs}" --bundle "{binding["bundle"]}" --worker-source "{worker}" --output "{run}" --engine fixture'
    original = {"format": "screen-benchmark-external-termination-v1", "run": str(run), "binding_sha256": digest(binding_path),
                "process_id": 1234, "executable": executable, "command_line": command,
                "created_at": "2026-10-07T12:00:00Z", "stop_requested_at": "2026-10-07T12:01:00Z",
                "stopped_at": "2026-10-07T12:01:01Z", "process_running_after_stop": False,
                "completed_at_stop": sum(status in {"success", "failed", "cancelled"} for status in states), "total": 12,
                "fixture_only": True}
    original_path = run / "fixture-external.json"
    frozen_json(original_path, original)
    receipt = {"format": "owned-benchmark-termination-v1", "run_root": str(run), "engine": "fixture",
               "run_binding_sha256": digest(binding_path), "inputs_sha256": digest(inputs),
               "benchmark_process": {"pid": 1234, "creation_time_utc": original["created_at"],
                                     "command_line_sha256": hashlib.sha256(command.encode("utf-8")).hexdigest()},
               "ownership_verified": True, "exit_observed": True,
               "stop_requested_at_utc": original["stop_requested_at"], "stopped_at_utc": original["stopped_at"],
               "original_evidence_file": str(original_path), "original_evidence_sha256": digest(original_path),
               "completed_at_stop": original["completed_at_stop"], "total": 12, "fixture_only": True}
    receipt_path = run / "fixture-owned.json"
    frozen_json(receipt_path, receipt)
    return run, inputs, worker, receipt_path


class RawBridgeTests(unittest.TestCase):
    def test_immutable_raw_identity_and_zero_delivery_scope(self):
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker = raw_fixture(temporary)
            result = collect_raw(run, inputs, worker)
            self.assertTrue(result["complete_snapshot"])
            self.assertTrue(result["records"][0]["result_id"].startswith("raw-output-sha256:"))
            self.assertFalse(result["records"][0]["automatic_delivery"])
            self.assertIsNone(result["performance"]["full_batch_wall_seconds"])

    def test_result_tamper_and_worker_changes_rejected(self):
        for kind in ["result", "worker"]:
            with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
                run, inputs, worker = raw_fixture(temporary)
                if kind == "result":
                    (run / "s1/result.json").write_text("{}", "utf-8")
                else:
                    (worker / "fixture.py").write_text("# changed\n", "utf-8")
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker)

    def test_partial_missing_record_is_preserved(self):
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker = raw_fixture(temporary)
            (run / "s1/record.json").unlink()
            result = collect_raw(run, inputs, worker)
            self.assertFalse(result["complete_snapshot"])
            self.assertEqual(result["records"][0]["status"], "missing")


class PartialRunTests(unittest.TestCase):
    def test_full_denominator_and_unknown_not_failure(self):
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker, receipt = partial_fixture(temporary)
            result = collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)
            self.assertEqual(len(result["records"]), 12)
            self.assertEqual([r["status"] for r in result["records"][:4]], ["failed", "failed", "failed", "interrupted_unknown"])
            self.assertEqual(sum(r["status"] == "not_attempted" for r in result["records"]), 8)
            self.assertEqual(result["records"][0]["failure_phase"], "fixture_explicit_load")
            self.assertEqual(result["records"][0]["error"], "generic fixture timeout")
            self.assertEqual(result["records"][3]["source_status"], "running")
            self.assertIsNone(result["records"][3]["result_id"])
            self.assertFalse(result["complete_snapshot"])
            self.assertFalse(result["complete_engine_comparison"])
            self.assertFalse(result["default_route_selection_eligible"])
            targets = [truth(r["id"]) for r in result["records"]]
            for target, record in zip(targets, result["records"]):
                target["image_sha256"] = record["image_sha256"]
            totals = evaluate(targets, result["records"])["summary"]["totals"]
            self.assertEqual((totals["images"], totals["visible_reference_cells"], totals["visible_critical_cells"]), (12, 48, 24))
            self.assertEqual((totals["failed_images"], totals["interrupted_unknown_images"], totals["not_attempted_images"]), (3, 1, 8))
            self.assertEqual((totals["missing_images"], totals["pending_images"], totals["automatic_images"], totals["no_table_correct_images"]), (0, 0, 0, 0))

    def test_partial_success_still_requires_actual_output_sha(self):
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker, receipt = partial_fixture(temporary, include_success=True)
            result = collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)
            self.assertTrue(result["records"][0]["result_id"].startswith("raw-output-sha256:"))
            (run / "s1/result.json").write_text("{}", "utf-8")
            with self.assertRaises(ValueError):
                collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)

    def test_partial_requires_explicit_receipt_and_mode(self):
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker, receipt = partial_fixture(temporary)
            for kwargs in ({"partial_run": True}, {"termination_receipt": receipt}):
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker, **kwargs)

    def test_wrong_termination_bindings_and_unobserved_exit_rejected(self):
        for field, value in [("run_root", "F:/OCR/wrong-fixture-run"), ("engine", "other"),
                             ("inputs_sha256", "0" * 64), ("run_binding_sha256", "0" * 64),
                             ("worker_source_sha256", "0" * 64), ("engine_config_sha256", "0" * 64),
                             ("ownership_verified", False), ("exit_observed", False), ("total", 11), ("completed_at_stop", 0)]:
            with self.subTest(field=field), tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
                run, inputs, worker, receipt = partial_fixture(temporary)
                changed = json.loads(receipt.read_text("utf-8"))
                changed[field] = value
                receipt.write_text(json.dumps(changed), "utf-8")
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)

    def test_process_and_stop_time_identity_rejected(self):
        mutations = [("pid", 9999), ("command_line_sha256", "0" * 64), ("creation_time_utc", "2026-10-07T12:02:00Z"),
                     ("stopped_at_utc", "2026-10-07T12:00:30Z"), ("stop_requested_at_utc", "2026-10-07T12:01:00"),
                     ("stopped_at_utc", "nonsense"), ("stopped_at_utc", "2026-10-07T12:01:01.0000001Z")]
        for field, value in mutations:
            with self.subTest(field=field), tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
                run, inputs, worker, receipt = partial_fixture(temporary)
                changed = json.loads(receipt.read_text("utf-8"))
                target = changed["benchmark_process"] if field in {"pid", "command_line_sha256", "creation_time_utc"} else changed
                target[field] = value
                receipt.write_text(json.dumps(changed), "utf-8")
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)

    def test_original_receipt_and_command_target_tamper_rejected(self):
        for kind in ("bytes", "run_target", "still_running"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
                run, inputs, worker, receipt = partial_fixture(temporary)
                owned = json.loads(receipt.read_text("utf-8"))
                original_path = Path(owned["original_evidence_file"])
                original = json.loads(original_path.read_text("utf-8"))
                if kind == "run_target":
                    original["command_line"] = original["command_line"].replace(f'--output "{run}"', '--output "F:/OCR/wrong-fixture-run"')
                    owned["benchmark_process"]["command_line_sha256"] = hashlib.sha256(original["command_line"].encode("utf-8")).hexdigest()
                elif kind == "still_running":
                    original["process_running_after_stop"] = True
                original_path.write_text(json.dumps(original), "utf-8")
                if kind != "bytes":
                    owned["original_evidence_sha256"] = digest(original_path)
                    receipt.write_text(json.dumps(owned), "utf-8")
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)

    def test_receipt_change_during_collection_rejected(self):
        import raw_bridge
        with tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
            run, inputs, worker, receipt = partial_fixture(temporary)
            reader = raw_bridge.read_json
            def mutate_after_read(path):
                value = reader(path)
                if Path(path) == receipt:
                    receipt.write_text(json.dumps(value) + "\n ", "utf-8")
                return value
            with patch("raw_bridge.read_json", side_effect=mutate_after_read), self.assertRaises(ValueError):
                collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)

    def test_partial_source_input_config_failed_record_and_prepared_tamper_rejected(self):
        for kind in ("source", "input", "config", "failed_identity", "prepared"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(dir=OUTPUT) as temporary:
                run, inputs, worker, receipt = partial_fixture(temporary)
                if kind == "source":
                    (worker / "fixture.py").write_text("# mutated\n", "utf-8")
                elif kind == "input":
                    image = Path(json.loads(inputs.read_text("utf-8"))["samples"][0]["image"])
                    image.write_bytes(b"changed original input")
                elif kind == "config":
                    bundle = Path(json.loads((run / "binding.json").read_text("utf-8"))["bundle"])
                    (bundle / "config/engines.json").write_text("{}", "utf-8")
                else:
                    record_path = run / "s1/record.json"
                    record = json.loads(record_path.read_text("utf-8"))
                    if kind == "failed_identity":
                        record["image_sha256"] = "0" * 64
                    else:
                        prepared = run / "s1/prepared.bin"
                        prepared.write_bytes(b"fixture prepared image")
                        record["prepared_images"] = [{"path": str(prepared), "sha256": "0" * 64}]
                    record_path.write_text(json.dumps(record), "utf-8")
                with self.assertRaises(ValueError):
                    collect_raw(run, inputs, worker, partial_run=True, termination_receipt=receipt)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--receipt", required=True)
    args = parser.parse_args()
    OUTPUT = fpath(args.output_root)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    suite = unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(QualityTests), unittest.defaultTestLoader.loadTestsFromTestCase(ReadbackTests),
                               unittest.defaultTestLoader.loadTestsFromTestCase(RawBridgeTests), unittest.defaultTestLoader.loadTestsFromTestCase(PartialRunTests)])
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    frozen_json(args.receipt, {"format": FORMAT, "status": "passed" if result.wasSuccessful() else "failed", "tests_run": result.testsRun,
                "errors": len(result.errors), "failures": len(result.failures), "scope": "CPU contract fixtures only; no model, Agent, native export or human review evidence",
                "failure_details": [{"test": str(case), "traceback": trace} for case, trace in result.errors + result.failures]})
    raise SystemExit(0 if result.wasSuccessful() else 1)
