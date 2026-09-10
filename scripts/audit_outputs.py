"""Check actual exported workbooks and failure paths without loading a GPU."""
import argparse
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from openpyxl import load_workbook

p=argparse.ArgumentParser()
p.add_argument('--bundle',type=Path,required=True)
p.add_argument('--inference',type=Path,required=True)
p.add_argument('--output',type=Path,required=True)
a=p.parse_args()
checks=[]
for path in sorted(a.inference.glob('*/result.xlsx')):
    workbook=load_workbook(path)
    strings=[]
    sheets=[]
    for sheet in workbook:
        values=[c for row in sheet for c in row if c.value is not None]
        assert all(c.data_type=='s' for c in values), path
        strings.extend(c.value for c in values)
        sheets.append({'name':sheet.title,'merges':sorted(str(r) for r in sheet.merged_cells.ranges),
                       'rows':sheet.max_row,'columns':sheet.max_column,
                       'values':[[c.value for c in row] for row in sheet]})
    assert '00123456789012345678' in strings, path
    checks.append({'kind':'real_xlsx','case':path.parent.name,'passed':True,'sheets':sheets})
assert len(checks)>=3, 'Expected all three structural engines to export a real workbook'
with tempfile.TemporaryDirectory(prefix='ocr-negative-') as temporary:
    temp=Path(temporary)
    corrupt=temp/'corrupt.png'
    corrupt.write_bytes(b'not an image')
    fake=temp/'missing-assets'
    (fake/'config').mkdir(parents=True)
    shutil.copy2(a.bundle/'config/engines.json',fake/'config/engines.json')
    for engine in ['ppocr','paddlevl','glm','hunyuan']:
        for case in ['corrupt_image','missing_model_manifest']:
            output=temp/f'{engine}-{case}'
            cmd=[str(a.bundle/'runtimes'/engine/'python.exe'),'-X','utf8','-I','-m','ocr_workbench.worker',
                 '--bundle',str(fake if case=='missing_model_manifest' else a.bundle),
                 '--engine',engine,'--image',str(corrupt if case=='corrupt_image' else a.bundle/'fixtures/printed.png'),
                 '--output',str(output)]
            result=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',timeout=30)
            error=json.loads((output/'error.json').read_text(encoding='utf-8'))
            assert result.returncode!=0 and not (output/'result.json').exists()
            assert error['error_type']==('UnidentifiedImageError' if case=='corrupt_image' else 'FileNotFoundError')
            checks.append({'kind':case,'engine':engine,'passed':True,'exit_code':result.returncode,
                           'error_type':error['error_type'],'message':error['message'],'success_file_absent':True})
a.output.write_text(json.dumps({'passed':True,'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(checks)}))
