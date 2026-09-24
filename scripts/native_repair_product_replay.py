"""Exercise native research candidates through actual PDF/store/product gates."""
from copy import deepcopy
import argparse
import json
from pathlib import Path
import shutil
import sys
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'src'),str(ROOT/'scripts')]
from native_repair_pilot import BUILD,AUDIT,read,save,sha
from ocr_workbench.store import Store
from ocr_workbench.documents import Documents
from ocr_workbench.page_processing import merge_page
from ocr_workbench.native_tables import native_table_preview
from ocr_workbench.structure_repair import content_baseline,native_header_patches
from ocr_workbench.structure_store import refresh_proposals,context,decide_structure,structure_view
from ocr_workbench.structure_arbitration import build_snapshot
from ocr_workbench.geometry_contract import fingerprint
from ocr_workbench import table_tool


def shape_literal(table):
    return (table['rows'],table['columns'],sorted((c['row'],c['column'],c['row_span'],c['column_span'],c['text']) for c in table['cells']))


def header_signature(table):
    return (shape_literal(table), sorted((c['row'],c['column'],c.get('is_header',False),c.get('header_role')) for c in table['cells']))


def reject(store,rid,proposal,version,reason):
    # 'keep' retires the entire table scope; reject is the ordinary per-proposal action.
    decide_structure(store,rid,proposal['id'],{'request_id':'research-filter-'+proposal['id'],
        'action':'reject','revision':proposal['revision'],'basis':proposal['basis'],'version_id':version['id']})
    return {'id':proposal['id'],'reason':reason}


def main(name):
    if not name.startswith('product-replay') or not name.replace('-','').isalnum():raise ValueError('Invalid replay name')
    folder=BUILD/name;folder.mkdir(exist_ok=False)
    store=Store(folder/'workspace');project=store.project('Native content evidence research')
    manager=Documents(store,Path('E:/DeloitteOCR-ComplexTables-20260919/bundle'),review_only=True)
    rows=read(BUILD/'pages.json');docs={};records=[]
    for document in dict.fromkeys(r['document'] for r in rows):
        source=BUILD/'sources'/f'{document}.pdf';temp=folder/f'{document}.import.pdf';shutil.copy2(source,temp)
        docs[document]=manager.import_document(project['id'],source.name,temp,dpi=150)
        assert sha(source)==docs[document]['sha256']
    for row in rows:
        started=time.monotonic();doc=docs[row['document']]
        page=store.document_pages(doc['id'],offset=row['page']-1,limit=1)[0]
        assert page['page_number']==row['page']
        image=manager.ensure_rendered(page['id']);page=store.one('pages',page['id'])
        version=store.one('versions',image['active_version'])
        rendered=read(store.file(page['native_result']));native=rendered['native']
        table_tool.prepare(manager,page,doc,version,native,'native',lambda:False)
        if native['table_tool']['state']!='ready':raise ValueError(native['table_tool'])
        raw=merge_page(native,[],version,doc,page,'native')
        prediction=native['table_structure']
        preview=native_table_preview(raw,prediction,version['width'],version['height'])
        if preview is None:raise ValueError('Native preview missing')
        assert [shape_literal(t) for t in preview['tables']]==[shape_literal(t) for t in row['preview']['tables']], 'Product extraction drift'
        # Add the frozen research variants to the ordinary tool candidate run.
        # This is experiment injection, not a claim of UI/default integration.
        variants=[];ranked={}
        for ti in row['selected_indices']:
            current=preview['tables'][ti]
            baseline=content_baseline(current,result_id=row['id'],revision=0,image_version=version['id'],image_sha256=version['sha256'])
            patches,_=native_header_patches(baseline)
            ranked[ti]=patches
            for item in patches:
                table=deepcopy(prediction['pdfplumber_tables'][ti]);headers={(c['row'],c['column']):c.get('is_header',False) for c in item['patch']['after']['cells']}
                for c in table['cells']:c['is_header']=headers[c['row'],c['column']]
                variants.append({'variant':item['variant']+f'-table-{ti}','pdfplumber_tables':[table]})
        prediction=deepcopy(prediction);prediction.setdefault('experimental_alternatives',[]).extend(variants)
        native['table_tool']['run_id']=fingerprint(prediction)
        preview['document']['table_tool']=native['table_tool'];preview['document']['table_structure']=prediction
        task=store.enqueue(project['id'],[version['id']],['pdf-native'],engine_packages={'pdf-native':'builtin'})[0]
        store.claim();store.complete(task,preview);rid=store.one('tasks',task)['result_id']
        with store.transaction() as db:table_tool.record(db,rid,version,prediction,raw['blocks'])
        view=refresh_proposals(store,rid,0)
        record={'page_id':row['id'],'result_id':rid,'version_id':version['id'],'image_sha256':version['sha256'],
                'source_unchanged':sha(BUILD/'sources'/f"{row['document']}.pdf")==row['source_sha256'],
                'same_baseline_literals':True,'tool_state':view['table_tool']['state'],'proposals':view['proposals'],'tables':[]}
        for ti in row['selected_indices']:
            options=[p for p in view['proposals'] if p['table_indices']==[ti] and p['can_apply']]
            info={'table':ti,'applicable_proposals':len(options),'snapshot':None,'reason':None,'filter_decisions':[]}
            retained=[];used=set()
            for patch in ranked[ti]:
                matches=sorted((p for p in options if p['kind']=='replace_table'
                    and header_signature(p['proposed_tables'][0])==header_signature(patch['patch']['after'])),key=lambda p:p['id'])
                if matches and matches[0]['id'] not in used:
                    retained.append(matches[0]);used.add(matches[0]['id'])
            info['ranked_variants']=[{'variant':patch['variant'],'signature_sha256':fingerprint(header_signature(patch['patch']['after']))} for patch in ranked[ti]]
            for option in options:
                if option['id'] not in used:info['filter_decisions'].append(reject(store,rid,option,version,'not a unique exact bounded full-table header variant'))
            while retained:
                try:
                    with store.transaction() as db:info['snapshot']=build_snapshot(db,context(db,rid),{'table':ti},version)
                    info['reason']=None;break
                except ValueError as error:
                    info['reason']=str(error)
                    if '超过上限' not in str(error) and '候选过多' not in str(error):break
                    if len(retained)==1:break
                    info['filter_decisions'].append(reject(store,rid,retained.pop(),version,'lowest fixed rank removed to meet unchanged product budget'))
            if not retained:info['reason']='No exact bounded full-table research patch passed ordinary proposal checks'
            info['retained_proposal_ids']=[p['id'] for p in retained]
            # Diagnose the unchanged gate using the exact candidate envelope.
            info['retained_candidate_characters']=[]
            with store.transaction() as db:
                from ocr_workbench.structure_store import structure_snapshot
                from ocr_workbench.complex_table_contract import validate_grid
                live=structure_snapshot(db,context(db,rid))
                sets={p['id']:json.loads(p['payload']) for p in live['candidate_rows']}
                for p in retained:
                    pool=sets[p['candidate_set_id']]['tokens'];table=p['proposed_tables'][0]
                    ids=validate_grid(table,tokens=pool)['token_ids'];byid={t['id']:t for t in pool}
                    envelope={'id':p['id'],'basis':p['basis'],'candidate_set_id':p['candidate_set_id'],
                        'token_pool_sha256':p['token_pool_sha256'],'table':table,'token_ids':ids,
                        'tokens':[byid[i] for i in ids],'kind':p['kind']}
                    info['retained_candidate_characters'].append({'id':p['id'],'characters':len(json.dumps([envelope],ensure_ascii=False)),
                        'cells':len(table['cells']),'tokens':len(ids)})
            info['final_proposals']=[p for p in structure_view(store,rid)['proposals'] if p['table_indices']==[ti]]
            record['tables'].append(info)
        # Local adoption is independent of external eligibility. Exercise only
        # the fixed first bounded variant, never use reference scores to choose.
        record['roundtrips']=[]
        for ti in row['selected_indices']:
            before=store.result(rid);fresh=refresh_proposals(store,rid,before['revision'])
            desired=header_signature(ranked[ti][0]['patch']['after'])
            matches=sorted((p for p in fresh['proposals'] if p['can_apply'] and p['kind']=='replace_table'
                and p['table_indices']==[ti] and header_signature(p['proposed_tables'][0])==desired),key=lambda p:p['id'])
            # For the first table the retained first choice is still pending;
            # subsequent tables get fresh revision-bound candidates after edits.
            if not matches:raise ValueError('Fixed local repair unavailable for roundtrip')
            p=matches[0]
            saved=decide_structure(store,rid,p['id'],{'request_id':'research-adopt-'+p['id'],'action':'accept',
                'revision':before['revision'],'basis':p['basis'],'version_id':version['id'],
                'acknowledge_unverified_empty':True})
            assert [shape_literal(t) for t in before['edited']['tables']]==[shape_literal(t) for t in saved['edited']['tables']]
            provenance=lambda table:{(c['row'],c['column']):c.get('native_content') for c in table['cells']}
            assert provenance(before['edited']['tables'][ti])==provenance(saved['edited']['tables'][ti])
            assert before['original']==saved['original']
            undo=store.history(rid,-1,saved['revision']);assert undo['edited']==before['edited']
            redo=store.history(rid,1,undo['revision']);assert redo['edited']==saved['edited']
            reopened=Store(store.root).result(rid);assert reopened['edited']==saved['edited']
            from ocr_workbench.tables import export_xlsx
            from openpyxl import load_workbook
            export=folder/f"{row['id']}-table-{ti}-adopted.xlsx";export_xlsx(reopened['edited']['tables'],export)
            book=load_workbook(export)
            for number,table in enumerate(reopened['edited']['tables']):
                sheet=book.worksheets[number];offset=1 if table.get('caption') else 0
                for cell in table['cells']:
                    out=sheet.cell(cell['row']+1+offset,cell['column']+1)
                    assert (out.value or '')==cell['text']
                    assert bool(out.font.bold)==bool(cell.get('is_header'))
                expected={str(__import__('openpyxl').worksheet.cell_range.CellRange(min_row=c['row']+1+offset,min_col=c['column']+1,
                    max_row=c['row']+c['row_span']+offset,max_col=c['column']+c['column_span'])) for c in table['cells'] if c['row_span']>1 or c['column_span']>1}
                assert {str(m) for m in sheet.merged_cells.ranges}==expected
            book.close()
            record['roundtrips'].append({'table':ti,'variant':ranked[ti][0]['variant'],'literal_preserved':True,
                'native_provenance_preserved':True,'original_unchanged':True,'undo_redo_reopen':True,'xlsx_values_headers_spans':True,
                'xlsx_sha256':sha(export),'research_adoption_only':True})
        record['seconds']=time.monotonic()-started
        records.append(record);save(folder/(row['id']+'.json'),record)
        print(row['id'],'applicable',sum(t['applicable_proposals'] for t in record['tables']),
              'eligible',sum(t['snapshot'] is not None for t in record['tables']),flush=True)
    save(folder/'rows.json',records)
    summary={'pages':len(records),'tables':sum(len(r['tables']) for r in records),
        'applicable_proposals':sum(t['applicable_proposals'] for r in records for t in r['tables']),
        'snapshot_eligible_tables':sum(t['snapshot'] is not None for r in records for t in r['tables']),
        'rejections':[{'page':r['page_id'],'table':t['table'],'reason':t['reason']} for r in records for t in r['tables'] if t['reason']],
        'source_unchanged':all(r['source_unchanged'] for r in records),
        'scope':'actual Documents import/render and ordinary native preview/store/product gates; frozen research candidates injected as explicit tool alternatives; not default integration'}
    save(AUDIT/(name+'.json'),summary);print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--name',default='product-replay');args=parser.parse_args();main(args.name)
