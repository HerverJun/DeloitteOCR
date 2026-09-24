"""Replay frozen model replies through real queue/completion/adoption/export paths."""
from copy import deepcopy
from pathlib import Path
import shutil
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'src'),str(ROOT/'scripts')]
from native_repair_pilot import read,save,sha
from native_repair_product_replay import shape_literal,header_signature
from ocr_workbench.store import Store
from ocr_workbench.multimodal_store import enqueue_review,prepare_review,complete_review
from ocr_workbench.structure_store import structure_view,decide_structure
from ocr_workbench.structure_wire import decode_response,encode,restore,evidence
from ocr_workbench.geometry_contract import fingerprint
from ocr_workbench.tables import export_xlsx


def main():
    from openpyxl import load_workbook
    records=[]
    for run in ('luna-local-review-20260919-06','luna-local-review-20260919-07'):
        root=ROOT/'build'/run
        for item in read(root/'prepared.json'):
            if not item['snapshot']:continue
            case=item['case'];folder=root/'response-replay'/case;folder.mkdir(parents=True,exist_ok=False)
            shutil.copytree(Path(item['store']),folder/'workspace')
            store=Store(folder/'workspace');rid=item['result_id'];before=store.result(rid)
            config={'backend':'external','protocol':'openai','base_url':'http://127.0.0.1:9/v1',
                'model':'simulation:gpt-5.6-luna','revision':1,'config_sha256':'isolated-model-response-replay',
                'prompt_version':'structure-evidence-v2','limits':{'request_timeout_seconds':1}}
            task=enqueue_review(store,rid,{'request_id':'frozen-luna-replay','model_id':'external:simulation-luna',
                'revision':before['revision'],'version_id':item['version']['id'],'scope':'table',
                'review_kind':'structure','table':item['table']},config)
            claimed=store.claim(external=True);assert claimed and claimed['id']==task
            snapshot=prepare_review(store,task)
            assert snapshot['structure']==item['snapshot']
            wire=read(root/'inbox'/case/'wire.json')
            assert fingerprint(wire)==snapshot['structure']['wire_sha256']
            assert restore(wire['archive'])==evidence(snapshot['structure'])
            raw=read(root/'inbox'/case/'response.json');canonical=decode_response(raw,snapshot['structure'])
            completed=complete_review(store,task,{'structure_response':canonical,'simulation':True,'network_calls':0})
            assert completed and store.result(rid)['edited']==before['edited']
            view=structure_view(store,rid)
            recommendation=next(r for r in view['arbitrations'] if r['task_id']==task)
            assert recommendation['response']==canonical and not recommendation['stale']
            if canonical['decision']=='select':
                p=next(p for p in view['proposals'] if p['id']==canonical['candidate_id'])
                saved=decide_structure(store,rid,p['id'],{'request_id':'research-luna-adopt','action':'accept',
                    'revision':before['revision'],'basis':p['basis'],'version_id':item['version']['id'],
                    'acknowledge_unverified_empty':True})
                assert [shape_literal(t) for t in before['edited']['tables']]==[shape_literal(t) for t in saved['edited']['tables']]
                assert saved['original']==before['original']
                target=saved['edited']['tables'][item['table']]
                assert header_signature(target)==header_signature(p['proposed_tables'][0])
                assert any(r['task_id']==task for r in target['structure_review']['external_recommendations'])
                provenance=lambda t:{(c['row'],c['column']):c.get('native_content') for c in t['cells']}
                assert provenance(target)==provenance(before['edited']['tables'][item['table']])
                undo=store.history(rid,-1,saved['revision']);assert undo['edited']==before['edited']
                redo=store.history(rid,1,undo['revision']);assert redo['edited']==saved['edited']
                assert Store(store.root).result(rid)['edited']==saved['edited']
                path=folder/'adopted.xlsx';export_xlsx(saved['edited']['tables'],path)
                book=load_workbook(path)
                for i,table in enumerate(saved['edited']['tables']):
                    sheet=book.worksheets[i];offset=1 if table.get('caption') else 0
                    for cell in table['cells']:
                        out=sheet.cell(cell['row']+1+offset,cell['column']+1)
                        assert (out.value or '')==cell['text'] and bool(out.font.bold)==bool(cell.get('is_header'))
                book.close()
                exported=sha(path)
            else:exported=None
            records.append({'run':run,'case':case,'decision':canonical['decision'],'wire_exactly_restored':True,
                'real_queue_completion':True,'recommendation_no_automatic_adoption':True,'literal_and_provenance_preserved':True,
                'undo_redo_reopen':True,'xlsx_sha256':exported,'response_sha256':sha(root/'inbox'/case/'response.json'),
                'network_calls':0,'scope':'frozen actual agent response replay in isolated Stores; not real endpoint dispatch'})
            print(run,case,'passed',flush=True)
    save(ROOT/'audit/luna-local-review-20260919-06/response-roundtrips.json',{'tables':len(records),'passed':True,'records':records})


if __name__=='__main__':main()
