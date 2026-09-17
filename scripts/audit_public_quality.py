"""Fixed-tool public development runs; unresolved regions are never accuracy."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import time
import urllib.request
import shutil

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))


def write(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fetch(url):
    with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'OCR-public-tool-validation'}),timeout=45) as response:
        return response.read()


def prepare(args):
    from ocr_workbench.documents import PdfCPU
    names=['nics-background-checks-2015-11.pdf','nics-background-checks-2015-11-rotated.pdf',
           'WARN-Report-for-7-1-2015-to-03-25-2016.pdf','senate-expenditures.pdf',
           'federal-register-2020-17221.pdf','cupertino_usd_4-6-16.pdf','pdffill-demo.pdf']
    commit=json.loads(fetch('https://api.github.com/repos/jsvine/pdfplumber/commits/v0.11.10'))['sha']
    write(args.output/'selection-before-download.json',{'commit':commit,'filenames':names,'pages':'first page of every document',
        'tool_config_sha256':sha(ROOT/'config/pdf-table-tool.json'),'quality_acceptance':False,
        'scope':'upstream public regression material, not an independent test'})
    samples=[];attempts=[];cpu=PdfCPU(args.bundle,args.output)
    for name in names:
        sid='tool-public-'+Path(name).stem.lower()
        url=f'https://raw.githubusercontent.com/jsvine/pdfplumber/{commit}/tests/pdfs/{name}'
        row={'id':sid,'source_url':url}
        try:
            source=args.output/'sources'/sid/'source.pdf';source.parent.mkdir(parents=True)
            source.write_bytes(fetch(url));row['source_pdf_sha256']=sha(source)
            inspected=cpu.call({'operation':'inspect','path':str(source.resolve()),'dpi':150})
            folder=args.output/'pages'/(sid+'-p1');folder.mkdir(parents=True)
            value=cpu.call({'operation':'render','path':str(source.resolve()),'page_number':1,'dpi':150,'image_output':str((folder/'input.png').resolve())})
            write(folder/'native.json',value)
            samples.append({'id':sid+'-p1','original_document_id':sid,'group_id':'nics' if name.startswith('nics-') else sid,
                'page_number':1,'image':'../pages/'+sid+'-p1/input.png','sha256':sha(folder/'input.png'),
                'width':value['metadata']['width'],'height':value['metadata']['height'],
                'source_pdf_sha256':sha(source),'native_result_sha256':sha(folder/'native.json'),
                'independent_cell_annotations':False,'physical_scan_verified':False,'source_url':url})
            row.update(status='success',page_count=len(inspected['pages']))
        except Exception as error:row.update(status='failed',error=repr(error))
        attempts.append(row);write(args.output/'preparation.json',{'rows':attempts})
        print(json.dumps(row,ensure_ascii=False),flush=True)
    write(args.output/'inputs/development.json',{'samples':samples,'scope':'upstream public development regressions; no independent quality acceptance'})


def product(args):
    from ocr_workbench.store import Store
    from ocr_workbench.documents import Documents
    from ocr_workbench.task_queue import TaskQueue
    from ocr_workbench.adapter import EngineAdapter
    from ocr_workbench.structure_store import refresh_proposals
    from ocr_workbench.pdf_export import build_pdf_export
    from ocr_workbench.document_review import document_review_queue
    samples=json.loads((args.data/'inputs/development.json').read_text('utf-8'))['samples']
    source_lock={p.relative_to(ROOT).as_posix():sha(p) for p in (ROOT/'src/ocr_workbench').glob('*.py')}
    write(args.output/'run-lock.json',{'source':source_lock,'tool_config_sha256':sha(ROOT/'config/pdf-table-tool.json'),
        'manifest_sha256':sha(args.data/'inputs/development.json'),'route':'auto/ppocr; paddle/local-v2 unchanged',
        'adoption':'no automatic candidate adoption; no automatic empty-region acknowledgement'})
    store=Store(args.output/'workspace')
    gpu=TaskQueue(store,args.bundle,adapter_factory=lambda b,e,s:EngineAdapter(b,e,s,worker_source=ROOT/'src'))
    manager=Documents(store,args.bundle,gpu_queue=gpu)
    project=store.project('通用工具公开材料验证');documents={};rows=[]
    try:
        for sample in samples:
            start=time.monotonic();row={'id':sample['id'],'source_pdf_sha256':sample['source_pdf_sha256'],'image_sha256':sample['sha256']}
            try:
                sid=sample['original_document_id']
                if sid not in documents:
                    documents[sid]=manager.import_document(project['id'],sid+'.pdf',args.data/'sources'/sid/'source.pdf',dpi=150)['id']
                page=store.rows('SELECT * FROM pages WHERE document_id=? AND page_number=?',(documents[sid],sample['page_number']))[0]
                stage=manager.process(page['id'],'auto')
                while True:
                    cpu_work=manager.step();gpu_work=gpu.step()
                    if not cpu_work and not gpu_work:break
                    if time.monotonic()-start>900:raise TimeoutError('Page processing budget exceeded')
                final=store.one('document_stages',stage)
                if final['status']!='succeeded':raise RuntimeError(final.get('error') or final['status'])
                result=store.result(json.loads(final['output'])['result_id'])
                write(args.output/(sample['id']+'-result.json'),result)
                view=refresh_proposals(store,result['id'],result['revision'])
                write(args.output/(sample['id']+'-proposals.json'),view)
                preflight=build_pdf_export(store,manager,{'page_ids':[page['id']]},preflight=True)
                write(args.output/(sample['id']+'-preflight.json'),preflight)
                raw=result['original'];conflicts=raw['document']['conflicts']
                empty=sum(c.get('reason')=='region_no_text' for c in conflicts)
                applicable=[p for p in view['proposals'] if p['can_apply']]
                row.update(status='processed_with_review' if conflicts else 'processed',result_id=result['id'],
                    revision=result['revision'],native_units=sum(b.get('source')=='pdf-native' for b in raw['blocks']),
                    empty_regions_requiring_review=empty,overlap_conflicts=len(conflicts)-empty,
                    pdf_preflight_ready=preflight['ready'],candidate_sets=len(view['candidates']),
                    applicable_proposals=len(applicable),applicable_tool_proposals=sum(p['provider'].startswith('pdfplumber/') for p in applicable),
                    tool_error=raw['document'].get('table_structure_error'),
                    pending_review_tasks=document_review_queue(store,documents[sid])['total'],
                    cache_reused=manager.process(page['id'],'auto')==stage)
            except Exception as error:row.update(status='failed',error=repr(error))
            row['seconds']=time.monotonic()-start;rows.append(row)
            write(args.output/'summary.json',{'rows':rows,'human_efficiency':'unmeasured','quality_acceptance':False,
                'empty_region_is_correct':False,'source_unchanged':all(sha(ROOT/p)==h for p,h in source_lock.items())})
            print(json.dumps(row,ensure_ascii=False),flush=True)
    finally:
        manager.stop();gpu.unload()


def review(args):
    """Replay final review code on immutable OCR/tool outputs, without inference."""
    from ocr_workbench.store import Store
    from ocr_workbench.structure_store import refresh_proposals,decide_structure
    from ocr_workbench.exporting import build_export
    from openpyxl import load_workbook
    source=json.loads((args.data/'summary.json').read_text('utf-8'))
    shutil.copytree(args.data/'workspace',args.output/'workspace',ignore=shutil.ignore_patterns('engine-sessions','pdf-ipc','__pycache__'))
    store=Store(args.output/'workspace');rows=[]
    lock={p.relative_to(ROOT).as_posix():sha(p) for p in (ROOT/'src/ocr_workbench').glob('*.py')}
    write(args.output/'run-lock.json',{'source':lock,'source_run':str(args.data),'tool_config_sha256':sha(ROOT/'config/pdf-table-tool.json'),
        'inference':'unchanged saved OCR and upstream table outputs','scripted_acceptance':'first applicable tool alternative in upstream numeric order; not a quality label'})
    for old in source['rows']:
        row={'id':old['id'],'source_status':old['status']}
        if 'result_id' not in old:
            row.update(status='unavailable',error=old.get('error'));rows.append(row);continue
        try:
            result=store.result(old['result_id'])
            # Recompute only this isolated audit copy's derived pending review.
            with store.transaction() as db:
                db.execute("UPDATE structure_proposals SET state='stale' WHERE result_id=? AND state IN ('pending','deferred')",(result['id'],))
            view=refresh_proposals(store,result['id'],result['revision'])
            write(args.output/(old['id']+'-proposals.json'),view)
            options=sorted((p for p in view['proposals'] if p['provider'].startswith('pdfplumber/') and p['can_apply']),
                           key=lambda p:int(p['provider'].rsplit(':',1)[-1]))
            row.update(status='checked',applicable_tool_proposals=len(options))
            if options:
                p=options[0]
                saved=decide_structure(store,result['id'],p['id'],{'request_id':'tool-audit-'+old['id'],'action':'accept',
                    'revision':p['revision'],'basis':p['basis'],'version_id':p['version_id'],'acknowledge_unverified_empty':True})
                before_conflicts=result['original']['document']['conflicts']
                assert saved['original']['document']['conflicts']==before_conflicts
                for conflict in before_conflicts:
                    if conflict['ocr']['text']:assert conflict['ocr']['text'] in saved['edited']['text']
                undone=store.history(result['id'],-1,saved['revision']);assert undone['edited']==result['edited']
                redone=store.history(result['id'],1,undone['revision']);assert redone['edited']==saved['edited']
                path=build_export(store,[result['id']],'xlsx');shutil.copy2(path,args.output/(old['id']+path.suffix))
                import io,zipfile
                if path.suffix=='.zip':
                    with zipfile.ZipFile(path) as z:content=io.BytesIO(z.read('OCR-result.xlsx'))
                else:content=path
                book=load_workbook(content);count=0
                for ti,table in enumerate(redone['edited']['tables']):
                    sheet=book['Table '+str(ti+1)]
                    for cell in table['cells']:
                        value=sheet.cell(cell['row']+1,cell['column']+1)
                        assert (value.value or '')==cell['text']
                        assert not cell['text'] or value.data_type=='s'
                        count+=1
                write(args.output/(old['id']+'-adopted.json'),redone)
                row.update(scripted_provider=p['provider'],xlsx_literal_cells=count,undo_redo='passed',
                    conflicts_preserved=len(before_conflicts),tables=[{'rows':t['rows'],'columns':t['columns'],'cells':len(t['cells']),
                    'spans':[{'row':c['row'],'column':c['column'],'row_span':c['row_span'],'column_span':c['column_span'],'text':c['text']} for c in t['cells'] if max(c['row_span'],c['column_span'])>1]} for t in redone['edited']['tables']])
        except Exception as error:row.update(status='failed',error=repr(error))
        rows.append(row);print(json.dumps(row,ensure_ascii=False),flush=True)
        write(args.output/'summary.json',{'rows':rows,'quality_acceptance':False,'human_efficiency':'unmeasured',
            'source_unchanged':all(sha(ROOT/p)==h for p,h in lock.items())})


def main():
    p=argparse.ArgumentParser();p.add_argument('--phase',choices=['prepare','product','review'],required=True)
    p.add_argument('--data',type=Path)
    p.add_argument('--output',type=Path,required=True);p.add_argument('--bundle',type=Path,required=True)
    args=p.parse_args();args.output.mkdir(parents=True,exist_ok=False)
    {'prepare':prepare,'product':product,'review':review}[args.phase](args)


if __name__=='__main__':main()
