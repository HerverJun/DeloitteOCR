"""Prepare or run an eight-input production native queue recovery check.

Default: freeze source/config and eight real input copies on F:, then stop.
--run: use actual create_app, its lifespan, TaskQueue and EngineAdapter. After
at least two table-bearing final items, interrupt this app's own live OCR work
through screen_batches.stop()/queue.stop(), reopen the same workspace, resume
the same batch with an explicit idempotency key, and export exactly once.

No controller, truth, frozen OCR receipt, mock adapter or prefilled cache is
loaded. This is a controlled service/queue lifecycle restart, not an OS crash.
The root task must schedule --run after the other model queue has stopped.
"""
from __future__ import annotations

import argparse
import asyncio
from collections import Counter
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import secrets
import shutil
import sys
import time
import zipfile

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
# Reuse the existing real-native harness's F-only paths, environment and source
# inventory. Importing it performs no preparation, OCR, server or model call.
import evaluate_screen_agent_native as native_harness

f_path = native_harness.f_path
sha = native_harness.sha
inventory = native_harness.inventory
TERMINAL = {"completed", "needs_review", "failed", "cancelled"}
SUBMIT_ID = "nativeRecoveryEightSubmit20261007"
PAUSE_ID = "nativeRecoveryEightPause20261007"
RESUME_ID = "nativeRecoveryEightResume20261007"


def now():
    return datetime.now(timezone.utc).isoformat()


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()


def write_new(path, value):
    path = f_path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")


def read_json(path):
    return json.loads(f_path(path).read_text(encoding="utf-8"))


def configure(args):
    root = f_path(args.output)
    # Existing preparation is reusable only for one real run, not overwritten.
    if (root / "run-started.json").exists():
        raise ValueError("This evidence already started a run; use a fresh output")
    if root.exists() and any(root.iterdir()) and not (root / "prepared.json").is_file():
        raise ValueError("Output must be empty or an existing intact preparation")
    root.mkdir(parents=True, exist_ok=True)
    # The existing helper also calls production storage_environment(root),
    # validating all current/subprocess cache, temp and application data paths.
    return native_harness.configure(args)


def prepare(args, root):
    if (root / "prepared.json").is_file():
        binding = read_json(root / "prepared.json")
        mapping = read_json(root / "input-map.json")["items"]
        verify_prepared(args, root, binding, mapping)
        return binding, mapping
    inputs = f_path(args.inputs)
    source = f_path(args.native_source)
    bundle = f_path(args.bundle)
    data = read_json(inputs)
    if data.get("truth_available_to_inference") is not False:
        raise ValueError("The input manifest must explicitly isolate offline truth")
    all_samples = data["samples"]
    indexed = {sample["id"]: sample for sample in all_samples}
    if len(indexed) != len(all_samples):
        raise ValueError("Frozen sample IDs are not unique")
    selected = args.sample_id or [sample["id"] for sample in all_samples[:8]]
    if len(selected) != 8 or len(set(selected)) != 8 or not set(selected) <= set(indexed):
        raise ValueError("Select exactly eight unique frozen input IDs")
    if any(not re.fullmatch(r"[A-Za-z0-9_.-]{1,180}", ident)
           or ident in {".", ".."} for ident in selected):
        raise ValueError("Sample IDs must be safe filenames")
    original_files = inventory(source)
    frozen = root / "frozen-native"
    for relative in original_files:
        original = f_path(source / relative)
        target = f_path(frozen / relative)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(original, target)
    files = inventory(frozen)
    if files != original_files or inventory(source) != original_files:
        raise ValueError("Native source changed during freezing")
    policy_source = f_path(args.policy_file) if args.policy_file else frozen / "config/screen-policy.json"
    policy = read_json(policy_source)
    if not isinstance(policy, dict) or policy.get("visual_review"):
        raise ValueError("Use an explicit production policy without unsupported automatic visual review")
    write_new(root / "run-policy.json", policy)
    engines = bundle / "config/engines.json"
    if not engines.is_file() or not (bundle / "runtimes/control/python.exe").is_file():
        raise ValueError("Read-only shared model bundle is incomplete")
    folder = root / "inputs"
    folder.mkdir()
    mapping = []
    for position, ident in enumerate(selected):
        sample = indexed[ident]
        original = f_path(sample["image"])
        if sha(original) != sample["image_sha256"]:
            raise ValueError("Frozen input bytes changed: " + ident)
        target = folder / f"{position + 1:04d}-{ident}{original.suffix.lower()}"
        shutil.copyfile(original, target)
        if sha(target) != sample["image_sha256"]:
            raise ValueError("Input copy differs: " + ident)
        mapping.append({"position": position, "sample_id": ident,
            "source_image": str(original), "image_sha256": sample["image_sha256"],
            "input_copy": str(target)})
    binding = {"format": "screen-native-eight-recovery-prepared-v1", "created": now(),
        "inputs": str(inputs), "inputs_sha256": sha(inputs), "selected_ids": selected,
        "selection_rule": "explicit --sample-id order, else first eight manifest entries; no OCR/truth-based selection",
        "scope_inputs": 8, "native_source_original": str(source), "native_source": str(frozen),
        "native_source_files": files, "native_source_inventory_sha256": fingerprint(files),
        "bundle": str(bundle), "bundle_engines_sha256": sha(engines),
        "run_policy": policy, "run_policy_sha256": sha(root / "run-policy.json"),
        "source_default_policy_sha256": sha(frozen / "config/screen-policy.json"),
        "policy_source": str(policy_source), "driver_sha256": sha(__file__),
        "reused_harness_sha256": sha(native_harness.__file__),
        "truth_read": False, "frozen_results_loaded": False,
        "prefilled_cache": False, "controller_calls": 0, "gpu_calls": 0,
        "prepared_only": True, "workspace": str(root / "native/ocr")}
    write_new(root / "input-map.json", {"items": mapping})
    write_new(root / "prepared.json", binding)
    return binding, mapping


def verify_prepared(args, root, binding, mapping):
    expected = {"inputs": str(f_path(args.inputs)), "bundle": str(f_path(args.bundle)),
                "native_source_original": str(f_path(args.native_source))}
    if any(binding.get(key) != value for key, value in expected.items()):
        raise ValueError("Preparation belongs to different inputs/source/bundle")
    if args.sample_id and args.sample_id != binding["selected_ids"]:
        raise ValueError("Explicit selected inputs changed")
    if args.policy_file and read_json(args.policy_file) != binding["run_policy"]:
        raise ValueError("Explicit run policy changed")
    if (len(mapping) != 8 or [row["sample_id"] for row in mapping] != binding["selected_ids"]
        or inventory(binding["native_source"]) != binding["native_source_files"]
        or sha(args.inputs) != binding["inputs_sha256"]
        or sha(root / "run-policy.json") != binding["run_policy_sha256"]
        or read_json(root / "run-policy.json") != binding["run_policy"]
        or sha(f_path(args.bundle) / "config/engines.json") != binding["bundle_engines_sha256"]
        or sha(__file__) != binding["driver_sha256"]
        or sha(native_harness.__file__) != binding["reused_harness_sha256"]):
        raise ValueError("Prepared source/input/config/driver identity changed")
    for row in mapping:
        if (sha(row["source_image"]) != row["image_sha256"]
                or sha(row["input_copy"]) != row["image_sha256"]
                or not f_path(row["input_copy"]).is_relative_to(root / "inputs")):
            raise ValueError("An original/input copy changed")


def table_snapshot(store):
    wanted = {"tasks", "results", "images", "versions", "selections", "ocr_submissions"}
    with store.transaction() as db:
        names = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        wanted |= {name for name in names if name.startswith("screen_")}
        return {name: [dict(row) for row in db.execute('SELECT * FROM "' + name + '"')]
            for name in sorted(wanted & names) if re.fullmatch(r"[A-Za-z_]+", name)}


def attempt_inventory(workspace):
    root = f_path(workspace) / "task-results"
    output = {}
    if root.exists():
        for task in sorted(root.iterdir()):
            if not task.is_dir():
                continue
            attempts = []
            for folder in sorted(task.iterdir()):
                if not folder.is_dir():
                    continue
                files = []
                for path in sorted(folder.rglob("*")):
                    path = f_path(path)
                    if path.is_file():
                        files.append({"path": path.relative_to(folder).as_posix(),
                            "bytes": path.stat().st_size,
                            "sha256": sha(path) if path.suffix.lower() == ".json" else None})
                attempts.append({"path": str(f_path(folder)), "files": files})
            output[task.name] = attempts
    return output


class Observer:
    """Observe production-owned workers and IPC; never replace their behavior."""
    def __init__(self):
        self.workers = {}
        self.requests = {}

    def capture(self, app, phase):
        queue = app.state.queue
        with queue.guard:
            adapter, task_id = queue.adapter, queue.current
        if adapter is None:
            return None
        key = str(f_path(adapter.ipc))
        value = self.workers.setdefault(key, {"adapter": adapter, "process": None,
            "job": None, "pid": None, "engine": adapter.engine, "ipc": key})
        if adapter.process is not None:
            value.update(process=adapter.process, pid=adapter.process.pid)
        if adapter.job is not None:
            value["job"] = adapter.job
        request_path = adapter.ipc / "request.json"
        try:
            request = read_json(request_path)
        except (OSError, ValueError):
            return None
        output = f_path(request["output"])
        expected_root = app.state.store.root / "task-results"
        if not output.is_relative_to(expected_root):
            raise ValueError("Observed worker request escaped this isolated workspace")
        request_task = output.relative_to(expected_root).parts[0]
        ident = (key, request["id"])
        observed = self.requests.setdefault(ident, {"phase_first_observed": phase,
            "observed_at": now(), "task_id": request_task, "pid": value["pid"],
            "engine": adapter.engine, "ipc": key, "request": request,
            "request_sha256": fingerprint(request), "output": str(output)})
        # A persisted previous request is evidence, not proof that it is active.
        if (task_id != request_task or not adapter.ready or value["process"] is None
                or value["process"].poll() is not None):
            return None
        task = app.state.store.one("tasks", task_id)
        if task["status"] != "running" or task.get("kind") in {"screen_cached", "screen_context_cached"}:
            return None
        try:
            response = read_json(adapter.ipc / "response.json")
        except (OSError, ValueError):
            response = None
        if response and response.get("id") == request["id"]:
            return None
        return {**observed, "task": task, "active_request_verified": True}

    def release_evidence(self):
        rows = []
        for value in self.workers.values():
            process, job, adapter = value["process"], value["job"], value["adapter"]
            rows.append({"ipc": value["ipc"], "engine": value["engine"], "pid": value["pid"],
                "process_exit_code": process.poll() if process else None,
                "job_drained": bool(job and job.drained),
                "job_closed": bool(job and job.handle is None),
                "adapter_reservation_closed": adapter.reservation is None,
                "adapter_gpu_lock_closed": adapter.gpu_lock is None,
                "released": bool(process and process.poll() is not None and job
                    and job.drained and job.handle is None and adapter.reservation is None
                    and adapter.gpu_lock is None)})
        return rows


async def api(client, method, path, body=None, files=None):
    response = await client.request(method, path, json=body if files is None else None, files=files)
    if response.status_code != 200:
        raise ValueError(f"Production API {method} {path} returned {response.status_code}: {response.text[:1200]}")
    return response.json()


async def stop_own_queue(app, observer, phase):
    observer.capture(app, phase)
    await asyncio.to_thread(app.state.screen_batches.stop)
    stopping = asyncio.create_task(asyncio.to_thread(app.state.queue.stop))
    while not stopping.done():
        observer.capture(app, phase)
        await asyncio.sleep(0.05)
    await stopping
    if app.state.queue.status()["alive"]:
        raise ValueError("Owned GPU queue did not exit")


async def wait_for(app, observer, phase, timeout, predicate):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        active = observer.capture(app, phase)
        answer = predicate(active)
        if answer is not None:
            return answer
        if app.state.screen_batches.worker_error or app.state.queue.status()["state"] == "faulted":
            raise ValueError("Production coordinator/queue faulted; retained native logs describe it")
        await asyncio.sleep(0.05)
    raise TimeoutError("Timed out waiting for production " + phase)


def pin_successes(store, batch_id):
    items = store.rows("SELECT * FROM screen_items WHERE batch_id=? ORDER BY position", (batch_id,))
    result = {}
    for item in items:
        if item["status"] not in {"completed", "needs_review"} or not item["result_id"]:
            continue
        record = store.one("results", item["result_id"])
        if not json.loads(record["edited"]).get("tables"):
            continue
        version = item["input_version_id"]
        tasks = store.rows("SELECT * FROM tasks WHERE input_version_id=? ORDER BY id", (version,))
        finals = store.rows("SELECT * FROM screen_final_tasks WHERE batch_id=? AND input_version_id=? ORDER BY content_hash", (batch_id, version))
        selected = store.rows("SELECT * FROM selections WHERE image_id=?", (item["image_id"],))
        result[version] = {"item": item, "result": record, "tasks": tasks,
            "finals": finals, "selection": selected}
    return result


def validate_success_retention(store, batch_id, pins, old_attempts, workspace):
    after = pin_successes(store, batch_id)
    attempts = attempt_inventory(workspace)
    rows = []
    for version, before in pins.items():
        current = after.get(version)
        if current is None or any(current[key] != before[key] for key in ("result", "tasks", "finals", "selection")):
            raise ValueError("An already successful item was recalculated/replaced: " + version)
        if any(current["item"][key] != before["item"][key]
               for key in ("result_id", "result_revision", "candidates", "review_state")):
            raise ValueError("Successful final candidate identity changed: " + version)
        for task in before["tasks"]:
            if attempts.get(task["id"], []) != old_attempts.get(task["id"], []):
                raise ValueError("A successful item has a new or changed compute attempt")
        rows.append({"input_version_id": version, "result_id": before["result"]["id"],
            "result_original_sha256": fingerprint(before["result"]["original"]),
            "result_edited_sha256": fingerprint(before["result"]["edited"]),
            "task_ids": [task["id"] for task in before["tasks"]],
            "no_new_tasks_results_or_attempts": True, "selected_final_unchanged": True})
    return rows


def check_export(path, manifest, store, batch_id, evidence_root):
    from openpyxl import load_workbook
    from ocr_workbench.tables import export_xlsx
    expected_items = {item["input_version_id"]: item for item in store.rows(
        "SELECT * FROM screen_items WHERE batch_id=?", (batch_id,))}
    if len(manifest["items"]) != 8 or {item["input_version_id"] for item in manifest["items"]} != set(expected_items):
        raise ValueError("Export manifest omits/duplicates an input")
    names = [item["file_name"] for item in manifest["items"] if item.get("file_name")]
    if len(set(names)) != len(names):
        raise ValueError("Export duplicates a workbook filename")
    # Compare exported cell values/types/spans with a fresh standard exporter
    # rendering of the same actual selected result. This is fidelity checking,
    # not OCR accuracy: no offline truth is opened by this harness.
    checks = []
    with zipfile.ZipFile(f_path(path)) as archive:
        if archive.testzip() is not None or len(archive.namelist()) != len(set(archive.namelist())):
            raise ValueError("ZIP is corrupt or contains duplicate members")
        inside = json.loads(archive.read("manifest.json"))
        if inside != manifest or set(name for name in archive.namelist() if name.endswith(".xlsx")) != set(names):
            raise ValueError("ZIP content differs from the authoritative manifest")
        if len(names) != manifest["success_files"]:
            raise ValueError("Successful workbook count differs")
        for entry in manifest["items"]:
            item = expected_items[entry["input_version_id"]]
            name = entry.get("file_name")
            if item["status"] in {"failed", "cancelled"} and name:
                raise ValueError("Failed input generated a fake workbook")
            if not name:
                continue
            actual = store.result(item["result_id"])
            actual_tables = actual["edited"].get("tables", [])
            if not actual_tables or entry["result_id"] != actual["id"]:
                raise ValueError("Workbook belongs to a nonfinal/no-table result")
            exported = load_workbook(io.BytesIO(archive.read(name)), data_only=False)
            reference_root = f_path(evidence_root) / "export-fidelity"
            reference_root.mkdir(exist_ok=True)
            reference = reference_root / name
            export_xlsx(actual_tables, reference, source_rows=[{
                "image": item["source_name"], "result_id": actual["id"],
                "engine": actual["original"].get("engine", ""),
                "version_id": item["input_version_id"], "revision": actual["revision"]}])
            expected = load_workbook(reference, data_only=False)
            try:
                metadata = actual["original"].get("screen", {}).get("business_metadata", [])
                if len(exported.worksheets) != len(expected.worksheets) + bool(metadata):
                    raise ValueError("Workbook missed/duplicated a table/source/metadata sheet")
                table_checks = []
                for index in range(len(expected.worksheets)):
                    left, right = exported.worksheets[index], expected.worksheets[index]
                    def cells(sheet):
                        return {(cell.row, cell.column): (cell.value, cell.data_type)
                            for row in sheet for cell in row if cell.value is not None}
                    if cells(left) != cells(right) or set(map(str, left.merged_cells.ranges)) != set(map(str, right.merged_cells.ranges)):
                        raise ValueError("XLSX changed selected values/types/merged spans")
                    if left.title != right.title:
                        raise ValueError("Selected table/source sheet identity changed")
                    table_checks.append({"sheet_index": index, "sheet": left.title,
                        "rows": left.max_row, "columns": left.max_column,
                        "values_types_and_merged_ranges_match": True})
                if metadata:
                    sheet = exported.worksheets[-1]
                    if sheet.title != "业务说明" or sheet.max_row != len(metadata) + 1:
                        raise ValueError("Business metadata sheet identity differs")
                    for row_index, part in enumerate(metadata, 2):
                        if (sheet.cell(row_index, 1).value != part["kind"]
                                or sheet.cell(row_index, 2).value != part["text"]
                                or sheet.cell(row_index, 2).data_type != "s"):
                            raise ValueError("Business metadata text/type changed")
                checks.append({"input_version_id": item["input_version_id"], "result_id": actual["id"],
                    "file_name": name, "sheets": table_checks,
                    "business_metadata_rows": len(metadata), "human_review": "not_performed"})
            finally:
                exported.close()
                expected.close()
    return {"zip_path": str(f_path(path)), "zip_sha256": sha(path), "manifest_items": 8,
        "xlsx_files": len(names), "workbooks": checks, "truth_read": False,
        "scope": "Readback fidelity to actual selected OCR result; no OCR correctness claim"}


async def execute(args, root, binding, mapping):
    verify_prepared(args, root, binding, mapping)
    workspace = f_path(binding["workspace"])
    if workspace.exists():
        raise ValueError("Real run requires a fresh native workspace with no prefilled cache")
    source = f_path(binding["native_source"])
    sys.path[:0] = [str(source / "src"), str(f_path(args.pdf_layer).parent), str(f_path(args.agent_layer))]
    import httpx
    from ocr_workbench.service import create_app
    from ocr_workbench.adapter import EngineAdapter
    from ocr_workbench.screen_batches import validate_policy
    import ocr_workbench.service as service_module
    if not f_path(service_module.__file__).is_relative_to(source / "src"):
        raise ValueError("Loaded service belongs to another source snapshot")
    resolved_policy = validate_policy(binding["run_policy"], read_json(source / "config/screen-policy.json"))
    write_new(root / "run-started.json", {"created": now(), "prepared_sha256": sha(root / "prepared.json"),
        "resolved_policy": resolved_policy, "resolved_policy_sha256": fingerprint(resolved_policy),
        "truth_read": False, "controller_calls": 0, "real_create_app_taskqueue_engineadapter": True})
    observer = Observer()
    summary = {"format": "screen-native-eight-recovery-result-v1", "status": "running",
        "scope_inputs": 8, "prepared_sha256": sha(root / "prepared.json"),
        "truth_read": False, "frozen_results_loaded": False, "prefilled_cache": False,
        "controller_calls": 0, "restart_type": "controlled queue.stop and actual create_app lifecycle reopen",
        "os_crash_tested": False, "native_source_inventory_sha256": binding["native_source_inventory_sha256"],
        "run_policy_sha256": binding["run_policy_sha256"], "human_review": "not_performed"}
    token = secrets.token_urlsafe(36)
    started = time.perf_counter()
    app = None
    try:
        app = create_app(f_path(args.bundle), workspace, token, start_queue=True, agent_enabled=False)
        if app.state.queue.factory is not EngineAdapter:
            raise ValueError("Actual default EngineAdapter was replaced")
        cache_names = [name for name in table_snapshot(app.state.store) if "cache" in name]
        initial = table_snapshot(app.state.store)
        if any(initial[name] for name in cache_names) or initial.get("results"):
            raise ValueError("Fresh run contains seeded OCR/cache results")
        async with app.router.lifespan_context(app):
            transport = httpx.ASGITransport(app=app, raise_app_exceptions=True)
            async with httpx.AsyncClient(transport=transport, base_url="http://127.0.0.1",
                headers={"Authorization": "Bearer " + token}, timeout=90) as client:
                project = await api(client, "POST", "/api/projects",
                    {"name": "Actual eight-input native recovery", "request_id": "nativeRecoveryProject20261007"})
                project_id = project["id"]
                for row in mapping:
                    # Multipart sends a dedicated byte copy; frozen source paths
                    # are never passed to add_image's consumable temporary input.
                    content = f_path(row["input_copy"]).read_bytes()
                    uploaded = await api(client, "POST", f"/api/projects/{project_id}/images",
                        files=[("files", (Path(row["input_copy"]).name, content, "image/png"))])
                    if uploaded["errors"] or len(uploaded["images"]) != 1:
                        raise ValueError("Actual native import failed")
                    image = uploaded["images"][0]
                    row.update(image_id=image["id"], input_version_id=image["active_version"])
                write_new(root / "native-input-map.json", {"items": mapping})
                submit_body = {"version_ids": [row["input_version_id"] for row in mapping],
                    "request_id": SUBMIT_ID, "policy": resolved_policy}
                endpoint = f"/api/projects/{project_id}/screen-batches"
                submitted = await api(client, "POST", endpoint, submit_body)
                batch_id = submitted["batch_id"]
                summary.update(project_id=project_id, batch_id=batch_id,
                    resolved_policy=submitted["policy"], schema_version=app.state.store.schema_version)
                if submitted["policy"] != resolved_policy:
                    raise ValueError("Submitted policy differs from the frozen resolved policy")
                repeated = await api(client, "POST", endpoint, submit_body)
                if repeated["batch_id"] != batch_id:
                    raise ValueError("Repeated submission created another batch")
                write_new(root / "submit-receipts.json", {"request": submit_body,
                    "first": submitted, "repeat": repeated})
                def threshold(active):
                    pins = pin_successes(app.state.store, batch_id)
                    if len(pins) >= 2:
                        return pins
                    items = app.state.store.rows("SELECT status FROM screen_items WHERE batch_id=?", (batch_id,))
                    if all(item["status"] in TERMINAL for item in items):
                        raise ValueError("Fewer than two table-bearing successful final items")
                    return None
                await wait_for(app, observer, "first-two-successes", args.timeout, threshold)
                # Stop scheduling first. Then pause queued items durably while
                # leaving the currently running real recognition interruptible.
                await asyncio.to_thread(app.state.screen_batches.stop)
                def live_ocr(active):
                    if active and app.state.store.rows("SELECT 1 FROM screen_tasks WHERE task_id=? AND batch_id=?", (active["task_id"], batch_id)):
                        return active
                    if not app.state.store.rows("SELECT 1 FROM tasks WHERE status IN ('queued','running')"):
                        raise ValueError("No live real OCR task remains to interrupt")
                    return None
                active = await wait_for(app, observer, "live-inference-to-interrupt", args.timeout, live_ocr)
                pins = pin_successes(app.state.store, batch_id)
                if len(pins) < 2:
                    raise ValueError("Successful item threshold was lost")
                paused = await api(client, "POST", endpoint + f"/{batch_id}/action",
                    {"action": "pause", "request_id": PAUSE_ID})
                await stop_own_queue(app, observer, "controlled-interruption")
                interrupted = app.state.store.one("tasks", active["task_id"])
                before = table_snapshot(app.state.store)
                old_attempts = attempt_inventory(workspace)
                write_new(root / "before-restart-tables.json", before)
                write_new(root / "before-restart-attempts.json", old_attempts)
                write_new(root / "interruption.json", {"active_request": active,
                    "pause_receipt": paused, "interrupted_task": interrupted,
                    "successful_pins": pins, "worker_release": observer.release_evidence(),
                    "queue_status": app.state.queue.status()})
                if interrupted["status"] != "interrupted":
                    raise ValueError("Live request raced to a different outcome; interruption was not proven")
                if not f_path(active["output"]).is_dir():
                    raise ValueError("Interrupted attempt directory was not preserved")
                if not all(row["released"] for row in observer.release_evidence()):
                    raise ValueError("Owned worker Job/process/GPU lock release was not proven")
        verify_prepared(args, root, binding, mapping)
        # Recreate the actual application/store, not a synthetic result replay.
        app = create_app(f_path(args.bundle), workspace, token, start_queue=True, agent_enabled=False)
        if app.state.queue.factory is not EngineAdapter:
            raise ValueError("Reopened app changed its EngineAdapter factory")
        async with app.router.lifespan_context(app):
            transport = httpx.ASGITransport(app=app, raise_app_exceptions=True)
            async with httpx.AsyncClient(transport=transport, base_url="http://127.0.0.1",
                headers={"Authorization": "Bearer " + token}, timeout=90) as client:
                endpoint = f"/api/projects/{project_id}/screen-batches"
                lookup = await api(client, "GET", f"/api/projects/{project_id}/screen-batch-requests/{SUBMIT_ID}")
                if not lookup["submitted"] or lookup["batch_id"] != batch_id:
                    raise ValueError("Original batch request identity was not retained")
                resubmitted = await api(client, "POST", endpoint, submit_body)
                if resubmitted["batch_id"] != batch_id or resubmitted["policy"] != resolved_policy:
                    raise ValueError("Recovery duplicate submission changed batch/policy identity")
                resume_body = {"action": "resume", "request_id": RESUME_ID}
                resumed = await api(client, "POST", endpoint + f"/{batch_id}/action", resume_body)
                repeated_resume = await api(client, "POST", endpoint + f"/{batch_id}/action", resume_body)
                if resumed != repeated_resume:
                    raise ValueError("Repeated resume idempotency key did not return its original receipt")
                write_new(root / "resume-receipts.json", {"lookup": lookup,
                    "resubmitted": resubmitted, "first": resumed, "repeat": repeated_resume})
                def finished(active):
                    rows = app.state.store.rows("SELECT * FROM screen_items WHERE batch_id=?", (batch_id,))
                    return rows if len(rows) == 8 and all(row["status"] in TERMINAL for row in rows) else None
                items = await wait_for(app, observer, "resumed-final-terminal-state", args.timeout, finished)
                retention = validate_success_retention(app.state.store, batch_id, pins, old_attempts, workspace)
                if app.state.store.one("tasks", active["task_id"])["status"] != "succeeded":
                    raise ValueError("Interrupted task did not finish through real resumed computation")
                after_attempts = attempt_inventory(workspace)
                for old in old_attempts.get(active["task_id"], []):
                    match = next((entry for entry in after_attempts.get(active["task_id"], []) if entry["path"] == old["path"]), None)
                    if match != old:
                        raise ValueError("Interrupted old compute attempt was changed or deleted")
                if len(after_attempts.get(active["task_id"], [])) <= len(old_attempts.get(active["task_id"], [])):
                    raise ValueError("Resumed interrupted task has no new real attempt")
                tables = table_snapshot(app.state.store)
                if len(tables["screen_batches"]) != 1:
                    raise ValueError("More than one submitted batch exists")
                final_keys = [(row["batch_id"], row["input_version_id"], row["content_hash"]) for row in tables["screen_final_tasks"]]
                if len(final_keys) != len(set(final_keys)):
                    raise ValueError("Duplicate final content adoption exists")
                result_tasks = [row["task_id"] for row in tables["results"]]
                if len(result_tasks) != len(set(result_tasks)):
                    raise ValueError("One actual task produced duplicate stored results")
                # No export is created before recovery; export exactly once now.
                exports = workspace / "screen-batches" / batch_id / "exports"
                if exports.exists() and list(exports.iterdir()):
                    raise ValueError("Unexpected pre-recovery export exists")
                archive, manifest = await asyncio.to_thread(app.state.screen_batches.export,
                    project_id, batch_id, {"include_pending_review": True})
                if len(list(exports.iterdir())) != 1:
                    raise ValueError("Final export created duplicate export folders")
                exported = await asyncio.to_thread(check_export, archive, manifest,
                    app.state.store, batch_id, root)
                shutil.copyfile(f_path(archive), root / "native-eight-final.zip")
                write_new(root / "export-readback.json", exported)
                write_new(root / "after-restart-tables.json", tables)
                write_new(root / "after-restart-attempts.json", after_attempts)
                summary.update(status="passed", counts=dict(Counter(item["status"] for item in items)),
                    successful_retention=retention, interrupted_task_id=active["task_id"],
                    interrupted_attempt_preserved=True, same_batch_and_resume_idempotent=True,
                    no_duplicate_batch_or_final_content=True, final_export_calls=1,
                    export=exported, delivery_completeness="partial" if any(item["status"] != "completed" for item in items) else "complete")
                await stop_own_queue(app, observer, "final-shutdown")
        summary["worker_release"] = observer.release_evidence()
        if not summary["worker_release"] or not all(row["released"] for row in summary["worker_release"]):
            raise ValueError("Final owned worker resource release was not proven")
        summary["service_exit"] = {"queue_alive": app.state.queue.status()["alive"],
            "coordinator_alive": bool(app.state.screen_batches.thread and app.state.screen_batches.thread.is_alive()),
            "fusion_queue_alive": app.state.fusion_queue.status()["alive"],
            "external_queue_alive": app.state.external_queue.status()["alive"],
            "document_threads_alive": app.state.documents.status()["alive"],
            "actual_app_lifespans_closed": 2}
        if any(value for key, value in summary["service_exit"].items() if key.endswith("alive")):
            raise ValueError("Service threads survived lifecycle exit")
    except Exception as error:
        summary.update(status="failed", error_type=type(error).__name__, error=str(error))
        raise
    finally:
        if app is not None:
            try:
                await stop_own_queue(app, observer, "cleanup")
            except Exception as cleanup_error:
                summary.update(status="failed", cleanup_error=str(cleanup_error))
        summary.update(finished=now(), wall_seconds=time.perf_counter() - started,
            observed_real_worker_requests=list(observer.requests.values()), worker_release=observer.release_evidence())
        summary["observed_engine_request_count"] = len(observer.requests)
        summary["engine_request_count_scope"] = "Observed production IPC request identities, not an imputed complete GPU inference count; authoritative task/attempt/cache tables are saved separately"
        write_new(root / "native-recovery-summary.json", summary)
    if summary["status"] != "passed":
        raise RuntimeError("Recovery check failed; inspect preserved native-recovery-summary.json")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("output", "inputs", "native-source", "bundle"):
        parser.add_argument("--" + name, type=Path, required=True)
    parser.add_argument("--sample-id", action="append", help="Exactly eight IDs, preserving supplied order")
    parser.add_argument("--policy-file", type=Path, help="Optional explicit bounded production policy; frozen on preparation")
    parser.add_argument("--agent-layer", type=Path, default=Path("F:/OCR/output/platform-agent-dev-20261004/dependencies/agent"))
    parser.add_argument("--pdf-layer", type=Path, default=Path("F:/OCR/output/clear-automation-dev-20261002/worktrees/ocr/build/public-quality-20260916/tool/site-packages/pdfplumber"))
    parser.add_argument("--timeout", type=float, default=7200)
    parser.add_argument("--run", action="store_true", help="Starts real GPU OCR; root must schedule the exclusive evaluation window")
    args = parser.parse_args()
    if args.timeout <= 0:
        parser.error("timeout must be positive")
    root = configure(args)
    binding, mapping = prepare(args, root)
    if args.run:
        asyncio.run(execute(args, root, binding, mapping))
        print(json.dumps({"status": read_json(root / "native-recovery-summary.json")["status"],
            "summary": str(root / "native-recovery-summary.json")}, ensure_ascii=False))
    else:
        print(json.dumps({"status": "prepared_only", "scope_inputs": 8,
            "gpu_calls": 0, "controller_calls": 0, "truth_read": False,
            "prepared": str(root / "prepared.json")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
