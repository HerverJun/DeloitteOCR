"""Compare frozen real OCR runs and independently read their real XLSX artifacts."""
from __future__ import annotations
import argparse
import io
import json
from pathlib import Path
import statistics
import sys
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parent))
from score import score_run, result_tables, text

def readback(path, result):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book=load_workbook(path,data_only=False);tables=result_tables(result);errors=[];checked=0
    sheets=[s for s in book.worksheets if s.title!='来源索引']
    if len(sheets)!=len(tables):errors.append({'kind':'sheet_count','expected':len(tables),'actual':len(sheets)})
    for index,table in enumerate(tables):
        if index>=len(sheets):break
        sheet=sheets[index];offset=1 if table.get('caption') else 0;expected_positions=set();expected_merges=set()
        if offset:
            expected_positions.add((1,1))
            if str(sheet.cell(1,1).value or '')!=table['caption']:errors.append({'kind':'caption','sheet':sheet.title})
        for cell in table.get('cells',[]):
            r,c=cell['row']+1+offset,cell['column']+1;expected_positions.add((r,c));actual=sheet.cell(r,c);value='' if actual.value is None else str(actual.value);checked+=1
            if value!=cell['text']:errors.append({'kind':'literal_value','sheet':sheet.title,'row':r,'column':c,'expected':cell['text'],'actual':value})
            if cell['text'] and actual.data_type!='s':errors.append({'kind':'unsafe_type','sheet':sheet.title,'row':r,'column':c,'data_type':actual.data_type})
            if cell.get('row_span',1)>1 or cell.get('column_span',1)>1:
                expected_merges.add(f'{get_column_letter(c)}{r}:{get_column_letter(c+cell.get("column_span",1)-1)}{r+cell.get("row_span",1)-1}')
        actual_merges={str(merged) for merged in sheet.merged_cells.ranges}
        if actual_merges!=expected_merges:errors.append({'kind':'merged_spans','sheet':sheet.title,'expected':sorted(expected_merges),'actual':sorted(actual_merges)})
        for row in sheet:
            for cell in row:
                if cell.value not in (None,'') and (cell.row,cell.column) not in expected_positions:errors.append({'kind':'extra_value','sheet':sheet.title,'row':cell.row,'column':cell.column,'value':str(cell.value)})
    book.close()
    return {'checked_cells':checked,'passes':not errors,'errors':errors}

def run_readbacks(runroot):
    records=[]
    for record_path in sorted(Path(runroot).glob('*/record.json')):
        record=json.loads(record_path.read_text('utf-8'))
        if record['status']!='success':continue
        attempt=Path(record['output']);files=list(attempt.rglob('result.xlsx'))
        if not files:records.append({'id':record['id'],'status':'no_xlsx','note':'No table or export failure; kept visible, never replaced with truth workbook'});continue
        for file in files:
            raw=file.with_name('result.json')
            try:
                result=json.loads(raw.read_text('utf-8'));check=readback(file,result)
                records.append({'id':record['id'],'path':str(file),'status':'passed' if check['passes'] else 'failed',**check})
            except Exception as e:records.append({'id':record['id'],'path':str(file),'status':'failed','error':str(e)})
    return {'scope':'Actual engine XLSX vs that engine raw result, independent of truth. Product batch ZIP mapping is a separate check.',
            'workbooks':len(records),'passed':sum(r['status']=='passed' for r in records),'failed':sum(r['status']=='failed' for r in records),'no_xlsx':sum(r['status']=='no_xlsx' for r in records),'records':records}

def quantile(values,q):
    if not values:return None
    values=sorted(values);position=(len(values)-1)*q;low=int(position);high=min(len(values)-1,low+1)
    return values[low]+(values[high]-values[low])*(position-low)

def compare(truth,runs):
    reports={};rows=[]
    for name,path in runs.items():
        score=score_run(truth,path);files=list(path.glob('*/record.json'));records=[json.loads(p.read_text('utf-8')) for p in files]
        complete=[r for r in records if r['status'] in {'success','failed'}]
        durations=[r['inference_seconds'] for r in complete if r.get('inference_seconds') is not None]
        summary_path=path/'run-summary.json';summary=json.loads(summary_path.read_text('utf-8')) if summary_path.exists() else None
        binding=json.loads((path/'binding.json').read_text('utf-8'))
        report={'binding':binding,'score':score,'performance':{'is_finished':score['completed_snapshot_is_final'],'run_summary':summary,'completed_records':len(complete),'ocr_p50_seconds':quantile(durations,.5),'ocr_p95_seconds':quantile(durations,.95),'ocr_sum_seconds':sum(durations),'preprocessing_sum_seconds':sum(r.get('preprocessing_seconds',0) for r in complete),'model_load_seconds':sum(r.get('model_load_seconds') or 0 for r in complete),'note':'Stage sums are diagnostics; batch wall time comes from run-summary, not sum of potentially overlapping stages'},'xlsx_readback':run_readbacks(path)}
        reports[name]=report
        metric=score['summary'];rows.append({'name':name,'complete':score['completed_snapshot_is_final'],'images':metric['totals']['images'],'failed':metric['totals'].get('failed_images',0),'missing':metric['totals'].get('missing_images',0),'whole_image_exact':metric['ratios']['whole_image_correct_rate'],'critical_exact':metric['ratios']['critical_exact_rate'],'ocr_p95_seconds':report['performance']['ocr_p95_seconds'],'xlsx_passed':report['xlsx_readback']['passed'],'xlsx_failed':report['xlsx_readback']['failed']})
    return {'scope':'Real model outputs on frozen screen simulations; engineering readback never proves recognition accuracy. No truth injected into inference or export.', 'comparison':rows,'runs':reports}

def main():
    p=argparse.ArgumentParser();p.add_argument('--truth',type=Path,required=True);p.add_argument('--run',action='append',required=True,help='label=F:/actual/run');p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    if a.output.resolve().drive.upper()!='F:':raise ValueError('F-only output')
    runs={}
    for value in a.run:
        name,path=value.split('=',1)
        if name in runs:raise ValueError('Duplicate run label')
        runs[name]=Path(path)
    report=compare(a.truth,runs);a.output.parent.mkdir(parents=True,exist_ok=True);a.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n','utf-8');print(json.dumps(report['comparison'],ensure_ascii=False))

if __name__=='__main__':main()
