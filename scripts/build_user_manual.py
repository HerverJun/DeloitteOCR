"""Build the editable Chinese handbook using the documented workspace runtime."""
import json
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'docs/manual-content.json').read_text('utf-8'))
OUT = ROOT / 'output/documents'
ASSETS = ROOT / 'docs/manual-assets'
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)
sources = {
    'overview': ('ui-structure/01-structure-comparison.png', None),
    'quick': ('ui-document/02-review-whole-table.png', (821, 409, 1417, 681)),
    'structure': ('ui-structure/01-structure-comparison.png', (825, 593, 1414, 934)),
    'visual': ('ui-multimodal/03b-evidence-and-decision.png', (826, 341, 1413, 499)),
}
for key, (source, crop) in sources.items():
    image = Image.open(ROOT / 'audit/bugfixes-20260917' / source).convert('RGB')
    if crop:
        image = image.crop(crop)
    image.save(ASSETS / f'{key}.png')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(215.9), Mm(279.4)
section.top_margin, section.bottom_margin = Mm(18), Mm(17)
section.left_margin = section.right_margin = Mm(26.45)
section.header_distance = Mm(8)
section.footer_distance = Mm(8)


def font_style(style, size, bold=False, color='000000'):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(rf.attrib):
        if attr.lower().endswith('theme'):
            del rf.attrib[attr]
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + name), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.2


for name, size, bold in [('Normal', 11, False), ('Title', 27, True),
                         ('Subtitle', 12, False), ('Heading 1', 20, True),
                         ('Heading 2', 12, True), ('Caption', 8.5, False)]:
    font_style(doc.styles[name], size, bold)
doc.styles['Normal'].paragraph_format.widow_control = True
doc.styles['Title'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Caption'].paragraph_format.space_after = Pt(9)
doc.styles['Caption'].font.color.rgb = RGBColor.from_string('575F59')
for name in ('Heading 1', 'Heading 2'):
    doc.styles[name].paragraph_format.keep_with_next = True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

hp = section.header.paragraphs[0]
hp.text = 'DeloitteOCR 使用手册'
hp.runs[0].font.size = Pt(8)
hp.runs[0].font.color.rgb = RGBColor(0, 0, 0)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('0.11.0rc1 修复版     ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
fp._p.append(field)
for run in fp.runs:
    run.font.size = Pt(8)


def para(text, style=None):
    return doc.add_paragraph(text, style)


def new_list(kind):
    numbering = doc.part.numbering_part.element
    abstract_id = max((int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))), default=-1) + 1
    num_id = max((int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))), default=0) + 1
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for tag, val in [('start', '1'), ('numFmt', 'decimal' if kind == 'steps' else 'bullet'), ('lvlText', '%1.' if kind == 'steps' else '•')]:
        node = OxmlElement('w:' + tag); node.set(qn('w:val'), val); level.append(node)
    abstract.append(level); numbering.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(abstract_id)); num.append(ref); numbering.append(num)
    return num_id


def table(rows, widths):
    t = doc.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Mm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + edge)
        for k, v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): b.set(qn('w:' + k), v)
        borders.append(b)
    pr.append(borders)
    margin = OxmlElement('w:tblCellMar')
    for edge, val in [('top','85'),('bottom','85'),('left','110'),('right','110')]:
        m = OxmlElement('w:' + edge); m.set(qn('w:w'),val); m.set(qn('w:type'),'dxa'); margin.append(m)
    pr.append(margin)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        no_split = OxmlElement('w:cantSplit'); t.rows[i]._tr.get_or_add_trPr().append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); t.rows[i]._tr.get_or_add_trPr().append(repeat)
        for j, value in enumerate(row):
            cell = cells[j]; cell.width = Mm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '3A3A3A' if i == 0 else ('F3F3F3' if i % 2 == 0 else 'FFFFFF')); cell._tc.get_or_add_tcPr().append(shade)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            if widths[j] <= 34: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value); run.font.size = Pt(9.5); run.bold = i == 0
            run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '202020')
    spacer = para(''); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = 0.3


markdown = [f'# {DATA["title"]}', '', DATA['subtitle'], '', DATA['version'], '']
for index, page in enumerate(DATA['pages']):
    if index:
        doc.add_page_break()
        para(page['title'], 'Heading 1')
    else:
        para(DATA['title'], 'Title')
        para(DATA['subtitle'], 'Subtitle')
        p = para(DATA['version']); p.runs[0].font.size = Pt(9); p.paragraph_format.space_after = Pt(14)
        para(page['title'], 'Heading 2')
    markdown.extend([f'## {page["title"]}', ''])
    for block in page['blocks']:
        kind = block['type']
        if kind in ('p','h2','note'):
            p = para(block['text'], 'Heading 2' if kind == 'h2' else None)
            if kind == 'note':
                p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(8)
                for run in p.runs: run.font.size = Pt(10)
            markdown.extend([('### ' if kind == 'h2' else '> ' if kind == 'note' else '') + block['text'], ''])
        elif kind in ('steps','bullets'):
            num_id = new_list(kind)
            for n, text in enumerate(block['items'],1):
                p = para(text)
                props = p._p.get_or_add_pPr()
                num_pr = OxmlElement('w:numPr')
                level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); num_pr.append(level)
                ident = OxmlElement('w:numId'); ident.set(qn('w:val'), str(num_id)); num_pr.append(ident)
                props.append(num_pr)
                p.paragraph_format.left_indent = Mm(5)
                p.paragraph_format.first_line_indent = Mm(-5)
                p.paragraph_format.space_after = Pt(6)
                markdown.append((f'{n}. ' if kind == 'steps' else '- ') + text)
            markdown.append('')
        elif kind == 'table':
            table(block['rows'], block['widths'])
            for n, row in enumerate(block['rows']):
                markdown.append('| ' + ' | '.join(row) + ' |')
                if n == 0: markdown.append('| ' + ' | '.join('---' for _ in row) + ' |')
            markdown.append('')
        elif kind == 'image':
            p = para(''); p.paragraph_format.keep_with_next = True
            run = p.add_run(); pic = run.add_picture(str(ASSETS / (block['key'] + '.png')), width=Mm(block['width']))
            pic._inline.docPr.set('descr', block['caption'])
            para(block['caption'], 'Caption')
            markdown.extend([f'![{block["caption"]}](manual-assets/{block["key"]}.png)', ''])

doc.core_properties.title = DATA['title']
doc.core_properties.subject = '离线识别 校对 复核与导出操作'
doc.core_properties.author = 'DeloitteOCR 项目'
doc.core_properties.keywords = 'OCR,使用手册,PDF,表格,校对,视觉审校'
doc.core_properties.comments = ''
doc.save(OUT / 'DeloitteOCR使用手册.docx')
(ROOT / 'docs/DeloitteOCR使用手册.md').write_text('\n'.join(markdown) + '\n', 'utf-8')
print(json.dumps({'docx':str(OUT / 'DeloitteOCR使用手册.docx'), 'sections':len(DATA['pages'])}, ensure_ascii=False))
