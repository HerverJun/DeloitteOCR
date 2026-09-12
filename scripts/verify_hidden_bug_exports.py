"""Verify actual HTTP exports, stale writes and partial import against an audit seed."""
import argparse
import csv
import io
import json
from pathlib import Path
import urllib.error
import urllib.request
import zipfile
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    seed = json.loads(args.seed.read_text("utf-8"))
    args.output.mkdir(parents=True, exist_ok=True)
    checks = []

    def request(path, method="GET", body=None, content_type="application/json"):
        payload = json.dumps(body).encode() if isinstance(body, dict) else body
        req = urllib.request.Request(seed["base"] + "/api" + path, payload, method=method,
            headers={"Authorization": "Bearer " + seed["token"], "Content-Type": content_type})
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                return response.status, response.read()
        except urllib.error.HTTPError as error:
            return error.code, error.read()

    def check(name, passed):
        assert passed, name
        checks.append(name)
        print("PASS", name, flush=True)

    key = seed["seeds"]["真实GLM表格"]["result"]
    status, payload = request("/results/" + key)
    result = json.loads(payload)
    check("GET supplies table source spans without changing original evidence", status == 200 and
          result["text_sources"][0]["table_index"] == 0 and "文字回写" not in result["original"]["text"])
    expected = result["edited"]["tables"][0]["cells"]
    expected_cell = next(c["text"] for c in expected if c["row"] == 2 and c["column"] == 1)
    for format in ("txt", "md", "json", "xlsx"):
        status, data = request("/export", "POST", {"result_ids": [key], "format": format})
        check(format + " export succeeds", status == 200)
        (args.output / ("real-glm." + format)).write_bytes(data)
        if format == "txt":
            check("TXT contains latest cell and literal long identifier without markup", expected_cell in data.decode() and
                  "00123456789012345678" in data.decode() and "<table" not in data.decode())
        elif format == "md":
            check("Markdown renders latest saved cells", expected_cell in data.decode() and data.decode().count("<table>") == 1)
        elif format == "json":
            check("JSON keeps immutable original and edited table values", json.loads(data)["edited"]["tables"][0]["cells"] == expected)
        else:
            book = load_workbook(io.BytesIO(data), data_only=False)
            check("XLSX retains merged cells and long identifiers as strings", book["Table 1"]["A3"].value == "00123456789012345678" and
                  book["Table 1"]["B3"].value == expected_cell and "A1:D1" in str(book["Table 1"].merged_cells) and "来源索引" in book.sheetnames)
    quotes = seed["seeds"]["引号换行"]["result"]
    status, data = request("/export", "POST", {"result_ids": [quotes], "format": "txt"})
    values = list(csv.reader(io.StringIO(data.decode(), newline=None), delimiter="\t"))
    check("HTTP TXT round-trips multiline and quoted cells", values == [["编号", "备注"], ["00001", "第一行\n第二行"], ["00002", '"quoted"']])
    mixed = seed["seeds"]["混排双表"]["result"]
    status, data = request("/export", "POST", {"result_ids": [key, mixed], "format": "xlsx"})
    archive = zipfile.ZipFile(io.BytesIO(data))
    check("Batch XLSX exports both images with source index", status == 200 and set(archive.namelist()) == {"0001.xlsx", "0002.xlsx", "sources.json"})
    (args.output / "batch.zip").write_bytes(data)
    status, _ = request("/results/" + key, "PUT", {"revision": result["revision"] - 1, "edited": result["edited"]})
    check("Stale save is rejected and cannot overwrite current edit", status == 409 and json.loads(request("/results/" + key)[1])["edited"] == result["edited"])
    boundary = "hidden-audit-boundary"
    image = Path(__file__).resolve().parents[1] / "fixtures/table.png"
    parts = []
    for name, data in [("import-ok.png", image.read_bytes()), ("broken.png", b"intentional bad image")]:
        parts.extend([f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="{name}"\r\nContent-Type: image/png\r\n\r\n'.encode(), data, b"\r\n"])
    parts.append(f"--{boundary}--\r\n".encode())
    status, data = request("/projects/" + seed["project"] + "/images", "POST", b"".join(parts), "multipart/form-data; boundary=" + boundary)
    imported = json.loads(data)
    check("Real HTTP partial image import keeps good file and reports bad file", status == 200 and len(imported["images"]) == 1 and len(imported["errors"]) == 1)
    (args.output / "verification.json").write_text(json.dumps({"checks": checks, "passed": len(checks)}, ensure_ascii=False, indent=2), "utf-8")


if __name__ == "__main__":
    main()
