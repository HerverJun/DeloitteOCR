"""Independently read actual GUI downloads and preserve their content hashes."""

import argparse
import hashlib
import io
import json
from pathlib import Path
import zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--evidence", type=Path, required=True)
    p.add_argument("--features", action="store_true")
    a = p.parse_args()
    root = a.evidence
    receipt = json.loads((root / "result.json").read_text("utf-8"))
    assert receipt["passed"] is True
    identifier = "00123456789012345678"
    report = {"passed": False, "scope": "Independent GUI download readback; edited text and original output remain separate", "bundle_manifest_sha256": receipt["bundle_manifest_sha256"], "sha256": {}}
    if a.features:
        paths = [root / "batch-separate.zip", root / "batch-aggregate.xlsx"]
        with zipfile.ZipFile(paths[0]) as archive:
            assert archive.namelist() == ["0001.xlsx", "0002.xlsx"]
            for name in archive.namelist():
                book = load_workbook(io.BytesIO(archive.read(name)))
                try:
                    assert len(book.worksheets) == 1
                    assert str(next(iter(book.active.merged_cells.ranges))) == "A1:D1"
                    assert any(c.value == identifier and c.data_type == "s" for row in book.active for c in row)
                finally:
                    book.close()
        book = load_workbook(paths[1])
        try:
            assert len(book.worksheets) == 2
            for sheet in book:
                assert any(c.value == identifier and c.data_type == "s" for row in sheet for c in row)
                assert "A1:D1" in sheet.merged_cells
        finally:
            book.close()
        report["workbooks"] = 2
        report["aggregate_sheets"] = 2
    else:
        paths = [root / ("edited." + extension) for extension in ["json", "xlsx", "txt", "md"]]
        result = json.loads(paths[0].read_text("utf-8"))
        edited = "扫描仪（浏览器校对）"
        assert any(c["text"] == edited for t in result["edited"]["tables"] for c in t["cells"])
        assert not any(c["text"] == edited for t in result["original"]["tables"] for c in t["cells"])
        assert "浏览器文字校对 00001234567890123456" in result["edited"]["text"]
        assert "浏览器文字校对" not in result["original"]["text"]
        for path in paths[2:]:
            text = path.read_text("utf-8")
            assert edited in text and identifier in text
        book = load_workbook(paths[1])
        try:
            tables = result["edited"]["tables"]
            assert len(book.worksheets) == len(tables)
            for sheet, table in zip(book.worksheets, tables):
                offset = int(bool(table.get("caption")))
                expected_merges = set()
                for c in table["cells"]:
                    row, col = c["row"] + 1 + offset, c["column"] + 1
                    actual = sheet.cell(row, col)
                    assert (actual.value or "") == c["text"]
                    if c["text"]:
                        assert actual.data_type == "s"
                    if c["row_span"] > 1 or c["column_span"] > 1:
                        expected_merges.add(f'{get_column_letter(col)}{row}:{get_column_letter(col+c["column_span"]-1)}{row+c["row_span"]-1}')
                actual_merges = {str(m) for m in sheet.merged_cells.ranges if m.min_row > offset}
                assert actual_merges == expected_merges
        finally:
            book.close()
    report["sha256"] = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    report["passed"] = True
    with (root / "export-verification.json").open("x", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
