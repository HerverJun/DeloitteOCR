"""Exact readback, safe spreadsheet text and version-bound artifact publication."""

from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import asdict, replace
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from modules.data_workbench.comparison import ComparisonSpec, compare_exact
from modules.data_workbench.exporting import ExportError, SourceDescriptor, export_and_publish
from modules.data_workbench.importing import import_csv
from modules.data_workbench.recipes import RecipeStore
from packages.contracts.models import JobRef
from workbench_platform.artifacts import ArtifactStore
import modules.data_workbench.exporting.results as export_module


EXAMPLES = json.loads((Path(__file__).parents[2] / "packages/contracts/examples/positive.json").read_text("utf-8"))


def setup(tmp_path: Path, *, formatted_amounts=False, amount_format="financial-v1"):
    assets = []
    left_amount = '"1,234.50"' if formatted_amounts else '-2.50'
    right_amount = '1234.50' if formatted_amounts else '-2.50'
    left_second = '(20.00)' if formatted_amounts else '0'
    right_second = '-20.00' if formatted_amounts else '0'
    for name, content in (
        ("left", f'id,amount,note\n00123456789012345678,{left_amount},"=HYPERLINK(\"\"x\"\")"\n'
         f'K,{left_second},""\n'),
        ("right", f'id,amount,note\n00123456789012345678,{right_amount},"=HYPERLINK(\"\"x\"\")"\n'
         f'K,{right_second},\n'),
    ):
        path = tmp_path / f"{name}.csv"
        path.write_text(content, encoding="utf-8")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        table = import_csv(path, source_sha256=digest)
        recipes = RecipeStore(tmp_path / f"{name}.sqlite3")
        revision = recipes.import_table(table)
        assets.append((revision, recipes.get(revision)["snapshot"],
                       SourceDescriptor(f"asset{name.capitalize()}", 1, digest, "csv")))
    left, right = assets
    result = compare_exact(left[1], right[1], left_revision=left[0], right_revision=right[0],
                           spec=ComparisonSpec((('c1', 'c1'),), (('c2', 'c2'), ('c3', 'c3')),
                                               amount_pairs=(('c2', 'c2'),), amount_tolerance="0",
                                               amount_format=amount_format))
    blobs, exports = tmp_path / "blobs", tmp_path / "exports"
    blobs.mkdir()
    exports.mkdir()
    store = ArtifactStore(tmp_path / "artifacts.sqlite3", blobs)
    identity = {"comparison_id": result.comparison_id,
                "left_source": asdict(left[2]), "right_source": asdict(right[2])}
    manifest_hash = hashlib.sha256(export_module._canonical(identity)).hexdigest()
    job_data = {**EXAMPLES["JobRef"], "app_id": "data", "instance_id": "dataInstance",
                "native_status": "complete", "status": "succeeded", "stage": "export",
                "counts": {"schema_version": "1.0", "total": 1, "succeeded": 1,
                           "failed": 0, "cancelled": 0, "pending": 0},
                "capabilities": [], "input_manifest_hash": manifest_hash}
    job = JobRef.model_validate(job_data)
    return left, right, result, store, exports, job


def run_export(tmp_path: Path):
    left, right, result, store, exports, job = setup(tmp_path)
    bundle = export_and_publish(result, left[1], right[1], left_source=left[2], right_source=right[2],
                                output_root=exports, artifact_store=store, workspace_id="workspaceA",
                                job_ref=job, application_version="1.0.0", rule_version="exact-v1")
    return left, right, result, store, job, bundle


def test_all_formats_read_back_with_literals_and_provenance(tmp_path: Path):
    left, right, result, store, job, bundle = run_export(tmp_path)
    assert set(bundle.files) == set(bundle.artifacts) == {
        "results.csv", "results.xlsx", "results.parquet", "summary.json", "provenance.json"}
    with bundle.files["results.csv"].open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    assert len(csv_rows) == 2
    assert csv_rows[0]["L_c1_value"] == "00123456789012345678"
    assert csv_rows[0]["L_c2_value"] == "'-2.50"  # safe presentation, numeric kind retained
    assert csv_rows[0]["L_c3_value"].startswith("'=HYPERLINK")
    assert csv_rows[1]["L_c3_kind"] == "text"
    assert csv_rows[1]["R_c3_kind"] == "null"
    with duckdb.connect() as db:
        parquet = db.execute("SELECT L_c1_value,L_c2_value,L_c3_value FROM read_parquet(?)",
                             [str(bundle.files["results.parquet"])]).fetchall()
    assert parquet[0] == ("00123456789012345678", "-2.50", '=HYPERLINK("x")')
    workbook = load_workbook(bundle.files["results.xlsx"], read_only=True, data_only=False)
    try:
        assert all(cell.data_type != "f" for row in workbook.active.iter_rows() for cell in row)
        assert workbook.active.cell(2, 1).value == "same"
    finally:
        workbook.close()
    provenance = json.loads(bundle.files["provenance.json"].read_text("utf-8"))
    assert provenance["rows"][0]["left"]["row_id"] == result.same[0].left.row_id
    assert provenance["rows"][0]["left"]["source_row"] == 2
    assert provenance["rows"][0]["left"]["asset_id"] == left[2].asset_id
    assert provenance["rows"][0]["right"]["asset_id"] == right[2].asset_id
    summary = json.loads(bundle.files["summary.json"].read_text("utf-8"))
    assert summary["coverage"] == result.coverage
    assert summary["csv_text_policy"] == "prefix_apostrophe_for_formula_prefix"
    assert bundle.summary_hash == hashlib.sha256(bundle.files["summary.json"].read_bytes()).hexdigest()
    for name, artifact_id in bundle.artifacts.items():
        artifact = store.get("workspaceA", artifact_id)
        assert artifact.input_manifest_hash == job.input_manifest_hash
        assert artifact.rule_or_model_version == "exact-v1"
        assert len(artifact.source_refs) == 2
        with store.stage_download("workspaceA", artifact_id, expected_revision=artifact.revision,
                                  expected_hash=artifact.blob_hash) as path:
            assert path.read_bytes() == bundle.files[name].read_bytes()


def test_stale_comparison_or_job_identity_is_rejected_before_publication(tmp_path: Path):
    left, right, result, store, exports, job = setup(tmp_path)
    changed = replace(result, coverage={**result.coverage, "left_total": 999})
    with pytest.raises(ExportError, match="changed"):
        export_and_publish(changed, left[1], right[1], left_source=left[2], right_source=right[2],
                           output_root=exports, artifact_store=store, workspace_id="workspaceA",
                           job_ref=job, application_version="1", rule_version="rule")
    wrong = job.model_copy(update={"input_manifest_hash": "0" * 64})
    with pytest.raises(ExportError, match="input identity"):
        export_and_publish(result, left[1], right[1], left_source=left[2], right_source=right[2],
                           output_root=exports, artifact_store=store, workspace_id="workspaceA",
                           job_ref=wrong, application_version="1", rule_version="rule")
    assert not list(exports.iterdir())


def test_independent_readback_detects_changed_result(tmp_path: Path):
    left, right, result, _, _, bundle = run_export(tmp_path)
    headers = export_module._headers(left[1], right[1])
    rows, _ = export_module._rows(result, left[1], right[1], left[2], right[2])
    with bundle.files["results.csv"].open("a", encoding="utf-8") as stream:
        stream.write("tampered\n")
    with pytest.raises(ExportError, match="CSV readback"):
        export_module._readback(bundle.files, headers, rows)


def test_xlsx_locator_and_formula_readback_guard(tmp_path: Path):
    source = SourceDescriptor("assetWorkbook", 2, "a" * 64, "xlsx", "工资")
    locator = source.locator()
    assert locator.kind == "table" and locator.sheet_name == "工资"
    assert locator.asset_revision == 2
    left, right, result, _, _, bundle = run_export(tmp_path)
    workbook = load_workbook(bundle.files["results.xlsx"])
    workbook.active.cell(2, 1, "=1+1")
    workbook.save(bundle.files["results.xlsx"])
    workbook.close()
    headers = export_module._headers(left[1], right[1])
    rows, _ = export_module._rows(result, left[1], right[1], left[2], right[2])
    with pytest.raises(ExportError, match="formula"):
        export_module._readback(bundle.files, headers, rows)


def test_source_descriptor_and_csv_formula_prefix_validation():
    for sheet in ("", 7, [], "Sheet"):
        with pytest.raises(ExportError):
            SourceDescriptor("assetA", 1, "a" * 64, "csv", sheet)
    for sheet in (None, "", 7, []):
        with pytest.raises(ExportError):
            SourceDescriptor("assetA", 1, "a" * 64, "xlsx", sheet)
    for dangerous in ("=1+1", " +1+1", "\ufeff=1+1", "\t=1+1", "\r=1+1"):
        assert export_module._safe_csv(dangerous) == "'" + dangerous
    assert export_module._safe_csv("ordinary") == "ordinary"


def test_saved_legacy_amount_policy_exports_without_reinterpreting_values(tmp_path: Path):
    left, right, legacy, store, exports, job = setup(
        tmp_path, formatted_amounts=True, amount_format="plain-decimal-v1")
    # This is the persisted policy shape from before amount-format versioning.
    saved = replace(legacy, policy=json.loads(json.dumps(legacy.policy)))
    assert "amount_format" not in saved.policy
    assert len(saved.different) == 2
    assert all("c2:c2" in pair.different_columns for pair in saved.different)
    current = compare_exact(left[1], right[1], left_revision=left[0], right_revision=right[0],
        spec=ComparisonSpec((('c1', 'c1'),), (('c2', 'c2'), ('c3', 'c3')),
                            amount_pairs=(('c2', 'c2'),)))
    assert current.policy["amount_format"] == "financial-v1"
    assert current.comparison_id != saved.comparison_id
    assert len(current.same) == 1
    before = json.dumps([left[1], right[1]], sort_keys=True)
    bundle = export_and_publish(saved, left[1], right[1], left_source=left[2], right_source=right[2],
        output_root=exports, artifact_store=store, workspace_id="workspaceA", job_ref=job,
        application_version="1.0.0", rule_version="exact-v1")
    summary = json.loads(bundle.files["summary.json"].read_text("utf-8"))
    assert summary["comparison_id"] == saved.comparison_id
    assert "amount_format" not in summary["policy"]
    assert summary["coverage"]["different_pairs"] == 2
    with bundle.files["results.csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["category"] for row in rows] == ["different", "different"]
    assert rows[0]["L_c2_value"] == "1,234.50"
    assert rows[1]["L_c2_value"] == "(20.00)"
    assert json.dumps([left[1], right[1]], sort_keys=True) == before
