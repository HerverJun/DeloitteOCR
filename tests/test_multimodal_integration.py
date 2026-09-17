"""Queue, API and actual exported-file tests with an injected local review session."""

import asyncio
import hashlib
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image
from openpyxl import load_workbook

from ocr_workbench.editing import tables_html
from ocr_workbench.exporting import build_export
from ocr_workbench.imaging import add_image
from ocr_workbench.multimodal_export import build_review_report, capture_report
from ocr_workbench.multimodal_runtime import ReviewCancelled
from ocr_workbench.multimodal_store import (decide_review, enqueue_review, prepare_review, view_review)
from ocr_workbench.pdf_export import capture_pdf
from ocr_workbench.service import create_app
from ocr_workbench.store import Conflict, Store
from ocr_workbench.tables import parse_tables
from ocr_workbench.task_queue import TaskQueue


class MultimodalIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.bundle = self.root / "bundle"
        (self.bundle / "config").mkdir(parents=True)
        (self.bundle / "config/engines.json").write_text(json.dumps({"ppocr":{"name":"Test OCR","models":[]}}), "utf-8")
        self.store = Store(self.root / "data")
        self.project = self.store.project("通用集成")
        image = self.root / "image.png"
        Image.new("RGB", (300,200), "white").save(image)
        self.image = add_image(self.store,self.project["id"],"image.png",image)
        self.version_id = self.image["active_version"]
        table = {"rows":1,"columns":2,"cells":[{"row":0,"column":c,"row_span":1,"column_span":1,"text":v}
            for c,v in enumerate(("00001","=SUM(A1:A2)"))]}
        text = "标题\n"+tables_html([table])
        table["source"] = parse_tables(text)[0]["source"]
        self.original = {"text":text,"tables":[table],"blocks":[],"engine":"ppocr"}
        self.result_id = self.completed(self.original)
        self.queue = TaskQueue(self.store,self.bundle)
        self.addCleanup(self.queue.stop)
        self.replacements = {"00001":"000001","=SUM(A1:A2)":"=SUM(A1:A3)"}
        self.entered = threading.Event()
        self.closed = 0
        self.behavior = None
        self.loaded_config = None
        owner = self

        class FakeReviewSession:
            def __init__(self,bundle,config,output,cancel_event):
                self.cancel_event = cancel_event
                owner.loaded_config = config

            def __enter__(self):
                owner.entered.set()
                return self

            def __exit__(self,*args):
                owner.closed += 1

            def review(self,image_path,snapshot,progress_callback=None):
                if owner.behavior:
                    owner.behavior(self.cancel_event)
                response = {"items":[{"target_id":t["id"],"decision":"replace" if t["before"] in owner.replacements else "keep",
                    "after":owner.replacements.get(t["before"],t["before"]),"reason":"原图可见内容"} for t in snapshot["targets"]],
                    "summary":"逐项检查", "identity":{"model":"fake-local","quantization":"Q4"},
                    "evidence":{"image_sha256":snapshot["image_sha256"]},"timing":{"elapsed_ms":4}}
                if progress_callback:
                    progress_callback(len(snapshot["targets"]),len(snapshot["targets"]))
                return response

        self.queue.review_factory = FakeReviewSession

    def completed(self,raw):
        task = self.store.enqueue(self.project["id"],[self.version_id],["ppocr"])[0]
        self.store.claim()
        self.store.complete(task,raw)
        return self.store.one("tasks",task)["result_id"]

    def request(self,name="review-001"):
        return {"revision":self.store.result(self.result_id)["revision"],"version_id":self.version_id,
                "model_id":"small-local","scope":"page","request_id":name}

    def enqueue(self,name="review-001"):
        return enqueue_review(self.store,self.result_id,self.request(name),{"profile_id":"small-local","config_sha256":"fixed"})

    def generate(self):
        task = self.enqueue()
        self.assertTrue(self.queue.step())
        self.assertEqual(self.store.one("tasks",task)["status"],"succeeded",self.store.one("tasks",task)["error"])
        return task

    def adopt(self):
        for index,proposal in enumerate(view_review(self.store,self.result_id)["proposals"]):
            if proposal["decision"] == "replace":
                decide_review(self.store,self.result_id,proposal["id"],{"action":"accept",
                    "revision":self.store.result(self.result_id)["revision"],"version_id":self.version_id,"request_id":f"accept-{index}"})

    @staticmethod
    def route(app,path,method="GET"):
        return next(r.endpoint for r in app.routes if r.path == path and method in r.methods)

    def app(self, *, review_only=False, mocked_models=False):
        if mocked_models:
            load = patch("ocr_workbench.multimodal_runtime.load_config",side_effect=lambda bundle,profile=None:
                {"profile_id":profile or "small-local","prompt_version":"test"})
            ready = patch("ocr_workbench.multimodal_runtime.review_readiness",return_value={"ready":True,"label":"Local small model",
                "profiles":[{"id":"small-local"}]})
            with load, ready:
                return create_app(self.bundle,self.store.root,"t"*32,start_queue=False,review_only=review_only)
        return create_app(self.bundle,self.store.root,"t"*32,start_queue=False,review_only=review_only)

    def test_queue_only_publishes_proposals_and_closes_session(self):
        before = self.store.result(self.result_id)
        task = self.generate()
        self.assertEqual(self.store.result(self.result_id),before)
        self.assertEqual(len(self.store.rows("SELECT * FROM results")),1)
        self.assertIsNone(self.store.one("tasks",task)["result_id"])
        self.assertEqual(self.closed,1)
        self.assertEqual(self.loaded_config["config_sha256"],"fixed")
        artifacts = list((self.store.root / "task-results" / task).glob("*/review.json"))
        self.assertEqual(len(artifacts),1)
        self.assertEqual(json.loads(artifacts[0].read_text("utf-8"))["identity"]["model"],"fake-local")

    def test_cancellation_and_expired_inference_publish_no_proposal(self):
        task = self.enqueue()
        self.behavior = lambda cancel: self.queue.action(self.project["id"],"cancel",[task])
        self.queue.step()
        self.assertEqual(self.store.one("tasks",task)["status"],"cancelled")
        self.assertFalse(view_review(self.store,self.result_id)["proposals"])
        self.behavior = lambda cancel: self.store.save(self.result_id,{"text":"人工内容","tables":[]},0)
        late = self.enqueue("late")
        self.queue.step()
        self.assertEqual(self.store.one("tasks",late)["status"],"cancelled")
        self.assertIn("过期",self.store.one("tasks",late)["phase"])
        self.assertFalse(view_review(self.store,self.result_id)["proposals"])
        self.assertEqual(self.store.result(self.result_id)["edited"]["text"],"人工内容")

    def test_deciding_earlier_task_expires_running_other_task_without_losing_siblings(self):
        first_task = self.generate()
        proposals = view_review(self.store,self.result_id)["proposals"]
        chosen = next(p for p in proposals if p["before"] == "00001")
        second_task = self.enqueue("second-task")
        self.behavior = lambda cancel: decide_review(self.store,self.result_id,chosen["id"],
            {"action":"accept","revision":0,"version_id":self.version_id,"request_id":"during-other-review"})
        self.queue.step()
        view = view_review(self.store,self.result_id)
        self.assertEqual(self.store.one("tasks",second_task)["status"],"cancelled")
        self.assertEqual(self.store.one("tasks",first_task)["status"],"succeeded")
        self.assertEqual({p["task_id"] for p in view["proposals"]},{first_task})
        self.assertEqual(view["counts"]["accepted"],1)
        self.assertEqual(view["counts"]["pending"],2)
        self.assertEqual(self.store.result(self.result_id)["original"],self.original)

    def test_failure_retry_and_launch_recovery_keep_same_snapshot(self):
        task = self.enqueue()
        def fail(cancel):
            raise RuntimeError("synthetic decode failure")
        self.behavior = fail
        self.queue.step()
        self.assertEqual(self.store.one("tasks",task)["status"],"failed")
        original_snapshot = prepare_review(self.store,task,require_running=False)
        self.behavior = None
        self.queue.action(self.project["id"],"retry",[task])
        self.store.recover(fusion=False)
        self.assertEqual(self.store.one("tasks",task)["status"],"paused")
        self.queue.action(self.project["id"],"resume",[task])
        self.queue.step()
        self.assertEqual(self.store.one("tasks",task)["status"],"succeeded")
        self.assertEqual(prepare_review(self.store,task,require_running=False),original_snapshot)
        self.assertEqual(len(view_review(self.store,self.result_id)["requests"]),1)

    def test_runtime_cancel_on_shutdown_is_interrupted_and_resumable(self):
        def block_until_cancel(cancel):
            if not cancel.wait(3):
                raise AssertionError("queue shutdown did not signal cancellation")
            raise ReviewCancelled("visual review cancelled")
        self.behavior = block_until_cancel
        self.queue.start()
        task = self.enqueue()
        self.queue.wake.set()
        self.assertTrue(self.entered.wait(3))
        self.queue.stop()
        self.assertEqual(self.store.one("tasks",task)["status"],"interrupted")
        self.assertFalse(view_review(self.store,self.result_id)["proposals"])
        self.assertEqual(self.closed,1)
        self.assertEqual(prepare_review(self.store,task,require_running=False)["task_id"],task)

    def test_reports_preserve_literals_and_include_failed_tasks(self):
        self.generate()
        self.adopt()
        failed = self.enqueue("failed-later")
        self.behavior = lambda cancel: (_ for _ in ()).throw(RuntimeError("invalid JSON response"))
        self.queue.step()
        json_path = build_review_report(self.store,self.result_id,"json")
        report = json.loads(json_path.read_text("utf-8"))
        self.assertEqual(next(r for r in report["requests"] if r["task_id"] == failed)["task"]["status"],"failed")
        accepted = [p for p in report["proposals"] if p["state"] == "accepted"]
        self.assertEqual({p["after"] for p in accepted},{"000001","=SUM(A1:A3)"})
        markdown = build_review_report(self.store,self.result_id,"md").read_text("utf-8")
        self.assertIn("000001",markdown)
        self.assertIn("=SUM(A1:A3)",markdown)
        self.assertIn("invalid JSON response",markdown)
        book = load_workbook(build_review_report(self.store,self.result_id,"xlsx"),data_only=False)
        try:
            cells = [cell for row in book["校验清单"] for cell in row]
            self.assertIn("000001",[cell.value for cell in cells])
            self.assertIn("=SUM(A1:A3)",[cell.value for cell in cells])
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        finally:
            book.close()

    def test_report_pins_content_and_provenance_while_another_store_edits(self):
        from ocr_workbench.multimodal_store import review_sources
        self.generate()
        other = Store(self.store.root)
        changed = []
        def concurrent_edit(db,result):
            if not changed:
                changed.append(other.save(self.result_id,{"text":"concurrent human edit","tables":[]},0))
            return review_sources(db,result)
        with patch("ocr_workbench.multimodal_store.review_sources",side_effect=concurrent_edit):
            report = capture_report(self.store,self.result_id)
        self.assertEqual(report["revision"],0)
        self.assertTrue(all(p["state"] == "pending" and p["revision"] == 0 for p in report["proposals"]))
        self.assertEqual(self.store.result(self.result_id)["revision"],1)
        self.assertTrue(all(p["state"] == "stale" for p in view_review(self.store,self.result_id)["proposals"]))

    def test_standard_exports_carry_same_revision_sources_and_literal_content(self):
        self.generate()
        self.adopt()
        current = self.store.result(self.result_id)
        for format in ("txt","json","xlsx"):
            with self.subTest(format=format):
                path = build_export(self.store,[self.result_id],format)
                self.assertEqual(path.suffix,".zip")
                with zipfile.ZipFile(path) as archive:
                    sources = json.loads(archive.read(f"sources/multimodal-{self.result_id}.json"))
                    self.assertEqual(sources["revision"],current["revision"])
                    self.assertEqual(sources["result_id"],self.result_id)
                    self.assertFalse(sources["automatic_adoption"])
                    payload = archive.read(next(n for n in archive.namelist() if n.endswith("."+format) and not n.startswith("sources/")))
                    if format == "txt":
                        self.assertIn("000001",payload.decode("utf-8"))
                        self.assertIn("=SUM(A1:A3)",payload.decode("utf-8"))
                    elif format == "json":
                        saved = json.loads(payload)
                        self.assertEqual(saved["original"],self.original)
                        self.assertEqual(saved["edited"],current["edited"])
                    else:
                        book = load_workbook(io.BytesIO(payload),data_only=False)
                        try:
                            cells = [cell for sheet in book for row in sheet for cell in row]
                            self.assertIn("000001",[c.value for c in cells])
                            self.assertIn("=SUM(A1:A3)",[c.value for c in cells])
                            self.assertFalse(any(c.data_type == "f" for c in cells))
                        finally:
                            book.close()

    def test_pdf_capture_contains_review_provenance_without_changing_ocr_units(self):
        raw = {"text":"00001","tables":[],"blocks":[{"text":"00001","polygon":[[10,10],[100,10],[100,30],[10,30]]}],"engine":"ppocr"}
        self.result_id = self.completed(raw)
        with self.store.transaction() as db:
            db.execute("UPDATE selections SET result_id=? WHERE image_id=?",(self.result_id,self.image["id"]))
        self.generate()
        folder = self.root / "pdf-capture"
        folder.mkdir()
        captured = capture_pdf(self.store,{"result_ids":[self.result_id]},folder)
        self.assertFalse(captured["failures"],captured["failures"])
        page = json.loads(Path(captured["documents"][0]["pages"][0]).read_text("utf-8"))
        self.assertEqual([u["text"] for u in page["units"]],["00001"])
        self.assertEqual(page["multimodal_sources"]["result_id"],self.result_id)
        self.assertEqual(page["multimodal_sources"]["proposals"][0]["after"],"000001")
        self.assertEqual(page["multimodal_sources"]["proposals"][0]["state"],"pending")

    def test_models_endpoint_handles_missing_config_and_advertises_ready_profile(self):
        app = self.app()
        value = self.route(app,"/api/multimodal/models")()
        self.assertEqual(value["models"],[])
        self.assertTrue(value["reason"])
        app = self.app(mocked_models=True)
        value = self.route(app,"/api/multimodal/models")()
        self.assertEqual(value["default_model"],"small-local")
        self.assertEqual(value["models"][0]["label"],"Local small model")
        self.assertTrue(value["models"][0]["available"])

    def test_one_broken_model_profile_does_not_hide_available_models(self):
        def config(bundle,profile=None):
            if profile == "broken-profile":
                raise ValueError("备用模型缺少 SHA256")
            return {"profile_id":profile or "small-local","prompt_version":"test"}
        with patch("ocr_workbench.multimodal_runtime.load_config",side_effect=config), \
                patch("ocr_workbench.multimodal_runtime.review_readiness",return_value={"ready":True,"label":"Local small model",
                    "profiles":[{"id":"small-local"},{"id":"broken-profile","label":"Broken alternative"}]}):
            app = create_app(self.bundle,self.store.root,"t"*32,start_queue=False)
        value = self.route(app,"/api/multimodal/models")()
        self.assertTrue(next(m for m in value["models"] if m["id"] == "small-local")["available"])
        broken = next(m for m in value["models"] if m["id"] == "broken-profile")
        self.assertFalse(broken["available"])
        self.assertIn("SHA256",broken["reason"])

    def test_real_catalog_isolates_malformed_alternative_metadata(self):
        config_file = Path(__file__).resolve().parents[1] / "config/multimodal-review.json"
        config = json.loads(config_file.read_text("utf-8"))
        default = config["default_profile"]
        alternate = next(key for key in config["profiles"] if key != default)
        config["profiles"][alternate].pop("label")
        # Small, correctly pinned presence fixtures; no executable or model is launched.
        assets = list(config["runtime"]["assets"])
        assets.extend(asset for profile in config["profiles"].values() for key,asset in profile.items()
                      if key in ("model","projector","license_asset"))
        for asset in assets:
            asset["bytes"], asset["sha256"] = 1, hashlib.sha256(b"x").hexdigest()
            path = self.bundle / asset["path"]
            path.parent.mkdir(parents=True,exist_ok=True)
            path.write_bytes(b"x")
        (self.bundle / "config/multimodal-review.json").write_text(json.dumps(config),"utf-8")
        app = self.app()
        catalog = self.route(app,"/api/multimodal/models")()
        self.assertTrue(next(p for p in catalog["models"] if p["id"] == default)["available"])
        self.assertFalse(next(p for p in catalog["models"] if p["id"] == alternate)["available"])

    def test_review_only_routes_allow_decision_and_report_but_deny_inference(self):
        task = self.generate()
        proposal = next(p for p in view_review(self.store,self.result_id)["proposals"] if p["decision"] == "replace")
        app = self.app(review_only=True,mocked_models=True)
        self.assertFalse(self.route(app,"/api/multimodal/models")()["models"][0]["available"])
        with self.assertRaisesRegex(ValueError,"仅校对"):
            self.route(app,"/api/results/{key}/multimodal","POST")(self.result_id,self.request("new"))
        for action in ("retry","resume"):
            with self.assertRaisesRegex(ValueError,"仅校对"):
                self.route(app,"/api/results/{key}/multimodal/tasks/{task_id}/{action}","POST")(self.result_id,task,action,{})
        saved = self.route(app,"/api/results/{key}/multimodal/{proposal_id}/decision","POST")(self.result_id,proposal["id"],
            {"action":"accept","revision":0,"version_id":self.version_id,"request_id":"review-only-accept"})
        self.assertEqual(saved["revision"],1)
        self.assertIn("text_sources",saved)
        response = self.route(app,"/api/results/{key}/multimodal/report")(self.result_id,"json")
        report_path = Path(response.path)
        self.assertEqual(json.loads(report_path.read_text("utf-8"))["revision"],1)
        asyncio.run(response.background())
        self.assertFalse(report_path.parent.exists())

    def test_task_routes_validate_ownership_and_reject_stale_retry(self):
        task = self.enqueue()
        app = self.app(mocked_models=True)
        action = self.route(app,"/api/results/{key}/multimodal/tasks/{task_id}/{action}","POST")
        with self.assertRaisesRegex(ValueError,"不属于"):
            action("unrelated-result",task,"cancel",{})
        self.assertEqual(action(self.result_id,task,"cancel",{})["task_ids"],[task])
        self.assertEqual(action(self.result_id,task,"retry",{})["task_ids"],[task])
        action(self.result_id,task,"cancel",{})
        self.store.save(self.result_id,{"text":"edited after cancellation","tables":[]},0)
        with self.assertRaises(Conflict):
            action(self.result_id,task,"retry",{})

    def test_committed_submission_replay_survives_model_removal_and_review_only_mode(self):
        body = self.request()
        original_app = self.app(mocked_models=True)
        submit = self.route(original_app,"/api/results/{key}/multimodal","POST")
        first = submit(self.result_id,body)
        # There is no model config in this bundle. A genuine new request cannot load it.
        restarted = self.app()
        resumed = self.route(restarted,"/api/results/{key}/multimodal","POST")
        with patch.object(restarted.state.queue.wake,"set") as wake:
            self.assertEqual(resumed(self.result_id,body),first)
            wake.assert_not_called()
        self.assertEqual(len(self.store.rows("SELECT * FROM multimodal_requests")),1)
        with self.assertRaises((ValueError,OSError)):
            resumed(self.result_id,self.request("new-without-model"))
        readonly = self.app(review_only=True)
        replay = self.route(readonly,"/api/results/{key}/multimodal","POST")
        self.assertEqual(replay(self.result_id,body),first)
        with self.assertRaisesRegex(ValueError,"仅校对"):
            replay(self.result_id,self.request("new-review-only"))
        with self.assertRaises(Conflict):
            replay(self.result_id,{**body,"scope":"target","target":{"kind":"text","start":0,"end":1}})


if __name__ == "__main__":
    unittest.main()
