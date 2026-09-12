from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from ocr_workbench.fusion import default_policy
from ocr_workbench.imaging import add_image
from ocr_workbench.store import Store, Conflict
from ocr_workbench.task_queue import TaskQueue, FusionQueue
from test_fusion import table


class FusionStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root/"workspace")
        self.project = self.store.project("融合回归")
        path = self.root/"image.png"
        Image.new("RGB", (600, 400), "white").save(path)
        self.photo = add_image(self.store, self.project["id"], "image.png", path)
        self.policy = default_policy("table")
        self.policy["baseline"] = "glm"
        self.cpu = FusionQueue(self.store, self.root)

    def batch(self, fusion=False, values=None):
        from ocr_workbench.tables import parse_tables
        ids = self.store.enqueue(self.project["id"], [self.photo["active_version"]], ["glm", "paddlevl", "hunyuan"],
                                 fusion_policy=self.policy if fusion else None, request_id="batch-request-1" if fusion else None)
        while task := self.store.claim():
            text = (values or {}).get(task["engine"], table(value="00123" if task["engine"] == "glm" else "00124"))
            raw = {"engine": task["engine"], "text": text, "tables": parse_tables(text), "blocks": [],
                   "project_image_version": task["version_id"], "image": {"width": 600, "height": 400}}
            self.store.complete(task["id"], raw)
        return ids

    def create(self, request_id="reuse-request-1"):
        ids = self.batch()
        results = [self.store.one("tasks", key)["result_id"] for key in ids]
        fusion_id = self.store.enqueue_fusion(self.project["id"], results, self.policy, request_id)[0]
        self.assertTrue(self.cpu.step())
        result_id = self.store.one("tasks", fusion_id)["result_id"]
        return result_id, results

    def adopt(self, result_id):
        with self.store.transaction() as db:
            db.execute("UPDATE selections SET result_id=? WHERE image_id=?", (result_id, self.photo["id"]))

    def decision(self, result_id, issue, action="keep", request_id="decision-request-1", **extra):
        body = {"action": action, "request_id": request_id, "revision": self.store.result(result_id)["revision"],
                "basis": issue["basis"], "current_fingerprint": issue["current_fingerprint"],
                "version_id": self.photo["active_version"], **extra}
        return body

    def test_cpu_routes_never_construct_an_adapter_and_fusion_stays_preview(self):
        with patch("ocr_workbench.task_queue.EngineAdapter", side_effect=AssertionError("GPU forbidden")):
            result_id, originals = self.create()
        self.assertIn(self.store.project_snapshot(self.project["id"])["images"][0]["selected_result"], originals)
        self.assertNotEqual(self.store.project_snapshot(self.project["id"])["images"][0]["selected_result"], result_id)
        self.assertIsNone(self.store.claim())
        self.assertEqual(self.store.result(result_id)["original"]["origin"], "fusion")
        self.assertNotIn("units", self.store.result(result_id)["original"]["fusion"])
        self.assertGreater(len(self.store.rows("SELECT * FROM fusion_evidence")), 0)

    def test_dependency_waits_for_all_terminal_and_records_failure(self):
        ids = self.store.enqueue(self.project["id"], [self.photo["active_version"]], ["glm", "paddlevl"],
                                 fusion_policy=self.policy, request_id="dependency-request-1")
        self.assertIsNone(self.store.claim(fusion=True))
        task = self.store.claim()
        self.store.complete(task["id"], {"engine": task["engine"], "text": "原文", "tables": [], "blocks": []})
        self.assertIsNone(self.store.claim(fusion=True))
        other = self.store.claim()
        with self.store.transaction() as db:
            db.execute("UPDATE tasks SET status='failed',error='expected failure' WHERE id=?", (other["id"],))
        self.assertTrue(self.cpu.step())
        fused = self.store.one("tasks", ids[-1])
        self.assertEqual(fused["status"], "succeeded")
        self.assertEqual(self.store.result(fused["result_id"])["original"]["fusion"]["coverage"], .5)

    def test_request_retry_does_not_create_extra_tasks_and_sources_are_frozen(self):
        result_id, originals = self.create()
        before = len(self.store.rows("SELECT * FROM tasks"))
        self.store.save(originals[0], {"text": "人工改过", "tables": []}, 0)
        again = self.store.enqueue_fusion(self.project["id"], originals, self.policy, "reuse-request-1")
        self.assertEqual(again, [self.store.result(result_id)["task_id"]])
        self.assertEqual(len(self.store.rows("SELECT * FROM tasks")), before)
        self.assertNotIn("人工改过", json.dumps(self.store.fusion_sources(again[0]), ensure_ascii=False))
        with self.assertRaisesRegex(Conflict, "请求标识"):
            self.store.enqueue_fusion(self.project["id"], originals[:1], self.policy, "reuse-request-1")
        with self.assertRaisesRegex(ValueError, "原始 OCR"):
            self.store.enqueue_fusion(self.project["id"], [result_id], self.policy, "reject-fused-request")

    def test_atomic_candidate_response_retry_and_confirmation_invalidation(self):
        result_id, originals = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)["issues"][0]
        chosen = next(c for c in issue["candidates"] if c["value"] == "00124")
        body = self.decision(result_id, issue, "candidate", candidate_id=chosen["id"])
        self.store.set_review(self.photo["id"], "confirmed", result_id, 0, self.photo["active_version"])
        saved = self.store.decide_issue(result_id, issue["id"], body)
        self.assertEqual(saved["edited"]["tables"][0]["cells"][-1]["text"], "00124")
        self.assertIn("00124", saved["edited"]["text"])
        retry = self.store.decide_issue(result_id, issue["id"], body)
        self.assertEqual(saved, retry)
        self.assertEqual(len(self.store.rows("SELECT * FROM fusion_decisions")), 1)
        self.assertEqual(self.store.review_issues(result_id)["counts"]["resolved"], 1)
        self.assertEqual(self.store.review_states(self.project["id"])[self.photo["id"]]["status"], "pending")
        self.assertEqual(Store(self.store.root).review_issues(result_id)["position"], issue["id"])
        self.assertNotIn("00124", self.store.result(originals[0])["original"]["text"])

    def test_conflict_and_transaction_failure_do_not_finish_issue(self):
        result_id, _ = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)["issues"][0]
        body = self.decision(result_id, issue)
        body["revision"] = 100
        with self.assertRaises(Conflict):
            self.store.decide_issue(result_id, issue["id"], body)
        self.assertEqual(self.store.review_issues(result_id)["counts"]["pending"], 1)
        body["revision"] = 0
        with patch("ocr_workbench.fusion_store.reconcile", side_effect=OSError("disk unavailable")):
            with self.assertRaises(OSError):
                self.store.decide_issue(result_id, issue["id"], body)
        self.assertEqual(self.store.result(result_id)["revision"], 0)
        self.assertEqual(self.store.review_issues(result_id)["counts"]["pending"], 1)
        self.assertEqual(self.store.rows("SELECT * FROM fusion_decisions"), [])

    def test_unrelated_edit_preserves_decision_but_structure_and_undo_invalidate(self):
        result_id, _ = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)["issues"][0]
        saved = self.store.decide_issue(result_id, issue["id"], self.decision(result_id, issue))
        edit = deepcopy(saved["edited"])
        edit["tables"][0]["cells"][0]["text"] = "表头修改"
        saved = self.store.save(result_id, edit, saved["revision"])
        self.assertEqual(self.store.review_issues(result_id)["counts"]["resolved"], 1)
        edit = deepcopy(saved["edited"])
        edit["tables"][0]["rows"] += 1
        self.store.save(result_id, edit, saved["revision"])
        self.assertEqual(self.store.review_issues(result_id)["counts"]["stale"], 1)
        result_id2, _ = self.create("reuse-request-2")
        self.adopt(result_id2)
        issue2 = self.store.review_issues(result_id2)["issues"][0]
        saved2 = self.store.decide_issue(result_id2, issue2["id"], self.decision(result_id2, issue2, request_id="decision-request-2"))
        self.store.history(result_id2, -1, saved2["revision"])
        self.assertEqual(self.store.review_issues(result_id2)["counts"]["stale"], 1)

    def test_preview_cannot_submit_decision_and_pagination_is_bounded(self):
        result_id, _ = self.create()
        issue = self.store.review_issues(result_id)["issues"][0]
        self.assertFalse(issue["context_current"])
        with self.assertRaises(Conflict):
            self.store.decide_issue(result_id, issue["id"], self.decision(result_id, issue))
        with self.assertRaises(ValueError):
            self.store.review_issues(result_id, limit=1000)

    def test_mixed_handwriting_review_uses_saved_document_and_supports_legacy_targets(self):
        self.policy = default_policy("handwriting")
        self.policy["baseline"] = "glm"
        result_id, _ = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)["issues"][0]
        self.assertEqual(issue["target"]["kind"], "document")
        self.assertEqual(issue["current_value"], self.store.result(result_id)["edited"]["text"])
        # Existing v5 workspaces used pre-render HTML offsets and cached values.
        with self.store.transaction() as db:
            db.execute("UPDATE fusion_issues SET target=?,current_value=? WHERE id=?",
                       (json.dumps({"kind": "text", "start": 0, "end": len(issue["baseline"])}),
                        json.dumps(issue["baseline"]), issue["id"]))
        issue = self.store.review_issues(result_id)["issues"][0]
        self.assertEqual(issue["target"]["kind"], "document")
        with self.assertRaisesRegex(ValueError, "含表格"):
            self.store.decide_issue(result_id, issue["id"], self.decision(
                result_id, issue, "candidate", candidate_id=issue["candidates"][0]["id"]))
        saved = self.store.decide_issue(result_id, issue["id"], self.decision(result_id, issue, "question"))
        self.assertEqual(saved["review_decision"]["state"], "question")
        issue = self.store.review_issues(result_id)["issues"][0]
        saved = self.store.decide_issue(result_id, issue["id"], self.decision(result_id, issue, request_id="keep-mixed-document"))
        self.assertEqual(saved["review_decision"]["state"], "resolved")
        self.assertEqual(saved["edited"]["tables"][0]["cells"][-1]["text"], "00123")

    def test_all_exports_include_same_saved_fusion_evidence_and_literal_numbers(self):
        import zipfile
        from ocr_workbench.exporting import build_export
        from openpyxl import load_workbook
        from io import BytesIO
        result_id, _ = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)["issues"][0]
        saved = self.store.decide_issue(result_id, issue["id"], self.decision(result_id, issue, "manual", value="000098765432109876"))
        for format in ("txt", "md", "json", "xlsx"):
            path = build_export(self.store, [result_id], format)
            self.assertEqual(path.suffix, ".zip")
            with zipfile.ZipFile(path) as archive:
                manifest = json.loads(archive.read("fusion-sources.json"))
                metadata = manifest["results"][0]
                evidence = json.loads(archive.read(metadata["file"]))
                self.assertEqual(evidence["revision"], saved["revision"])
                self.assertEqual(evidence["review_summary"]["resolved"], 1)
                self.assertFalse(evidence["review_summary"]["human_confirmed"])
                self.assertEqual(len(evidence["sources"]), 3)
                content = archive.read("OCR-result." + format)
                if format == "xlsx":
                    book = load_workbook(BytesIO(content))
                    self.assertEqual(book["Table 1"].cell(2, 2).value, "000098765432109876")
                    self.assertEqual(book["Table 1"].cell(2, 2).data_type, "s")
                else:
                    self.assertIn("000098765432109876", content.decode("utf-8"))

    def test_project_cleanup_accounts_for_and_cascades_fusion_data(self):
        from ocr_workbench.maintenance import ProjectMaintenance
        result_id, _ = self.create()
        maintenance = ProjectMaintenance(self.store, TaskQueue(self.store, self.root))
        usage = maintenance.usage(self.project["id"])
        self.assertGreater(usage["fusion_bytes"], 0)
        maintenance.delete(self.project["id"], self.project["name"])
        for table in ("fusion_inputs", "fusion_dependencies", "fusion_evidence", "fusion_issues", "fusion_submissions"):
            self.assertEqual(self.store.rows("SELECT * FROM " + table), [])

    def test_cancelled_cpu_task_must_exit_before_project_deletion(self):
        from ocr_workbench.maintenance import ProjectMaintenance
        result_id, _ = self.create()
        task_id = self.store.result(result_id)['task_id']
        with self.store.transaction() as db: db.execute("UPDATE tasks SET status='cancelled' WHERE id=?",(task_id,))
        self.cpu.current = task_id
        maintenance = ProjectMaintenance(self.store, TaskQueue(self.store,self.root), self.cpu)
        with self.assertRaisesRegex(ValueError,'仍在运行'):
            maintenance.delete(self.project['id'],self.project['name'])
        self.assertTrue(self.store.result(result_id))

    def test_preprocessed_dependencies_bind_success_and_cancellation_to_one_version(self):
        from ocr_workbench.imaging import prepare_task
        ids=self.store.enqueue(self.project['id'],[self.photo['active_version']],['glm','paddlevl'],
            preprocess=[{'kind':'rotate','degrees':90}],fusion_policy=self.policy,request_id='preprocessed-fusion-request')
        task=prepare_task(self.store,self.store.claim())
        self.assertNotEqual(task['version_id'],self.photo['active_version'])
        self.store.complete(task['id'],{'engine':task['engine'],'text':'prepared','tables':[],
            'blocks':[],'project_image_version':task['version_id'],'image':{'width':400,'height':600}})
        other=self.store.claim()
        with self.store.transaction() as db: db.execute("UPDATE tasks SET status='cancelled' WHERE id=?",(other['id'],))
        self.assertTrue(self.cpu.step())
        fused=self.store.one('tasks',ids[-1]);self.assertEqual(fused['status'],'succeeded')
        sources=self.store.fusion_sources(fused['id'])
        self.assertEqual({s['version_id'] for s in sources},{task['version_id']})
        self.assertEqual({s['status'] for s in sources},{'succeeded','cancelled'})

    def test_multiwindow_revision_conflict_preserves_the_second_pending_issue(self):
        result_id,_=self.create();self.adopt(result_id)
        issue=self.store.review_issues(result_id)['issues'][0]
        window_a=self.decision(result_id,issue,'manual',request_id='window-a-request',value='A')
        window_b=self.decision(result_id,issue,'manual',request_id='window-b-request',value='B')
        self.store.decide_issue(result_id,issue['id'],window_a)
        with self.assertRaises(Conflict): self.store.decide_issue(result_id,issue['id'],window_b)
        self.assertEqual(self.store.result(result_id)['edited']['tables'][0]['cells'][-1]['text'],'A')
        self.assertEqual(len(self.store.rows('SELECT * FROM fusion_decisions')),1)

    def test_workspace_copy_preserves_frozen_sources_decisions_and_exports(self):
        import shutil
        import zipfile
        from ocr_workbench.exporting import build_export
        result_id, _ = self.create()
        self.adopt(result_id)
        issue = self.store.review_issues(result_id)['issues'][0]
        self.store.decide_issue(result_id, issue['id'], self.decision(result_id, issue))
        sources = self.store.fusion_sources(self.store.result(result_id)['task_id'])
        relocated = self.root / '搬迁 工作区'
        shutil.copytree(self.store.root, relocated)
        moved = Store(relocated)
        self.assertEqual(moved.review_issues(result_id)['counts']['resolved'], 1)
        self.assertEqual(moved.review_issues(result_id)['position'], issue['id'])
        self.assertEqual(moved.fusion_sources(moved.result(result_id)['task_id']), sources)
        version = moved.one('versions', self.photo['active_version'])
        self.assertTrue((relocated / version['path']).is_file())
        with zipfile.ZipFile(build_export(moved, [result_id], 'json')) as archive:
            metadata = json.loads(archive.read('fusion-sources.json'))['results'][0]
            evidence = json.loads(archive.read(metadata['file']))
            self.assertEqual(evidence['review_summary']['resolved'], 1)
            self.assertEqual(len(evidence['sources']), 3)


if __name__ == "__main__":
    unittest.main()
