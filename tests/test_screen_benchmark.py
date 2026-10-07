"""Regression scoring cannot hide cell shifts, numeric changes, spans or failures."""
from pathlib import Path
import sys
import copy
import random
import json
import pytest

import numpy as np
from PIL import Image

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts/screen_benchmark'))
from generate import homography,points,corners,render_screen,chinese,DEVELOPMENT_FAMILIES,TEST_FAMILIES
from score import score_sample,score_run,summarize
from run import first_pass_images,aggregate_results,select_input_samples,validate_context_binding,digest,write
from context import verify_context_record
from review_context import coverage,inspect_preparation
from validate_run import validate_run
from acceptance import compare,readback,performance

def truth():
    return {'id':'test','cohort':'chinese','level':'normal','group_id':'templateA','source_id':'sourceA','tables':[
        {'rows':2,'columns':2,'cells':[{'row':0,'column':0,'row_span':1,'column_span':2,'text':'金额（元）'},
        {'row':1,'column':0,'row_span':1,'column_span':1,'text':'000001'},
        {'row':1,'column':1,'row_span':1,'column_span':1,'text':'-1,234.56'}]}],
        'context_targets':[],'ui_regions':[{'text':'Excel SUM 100.00'}]}

def result():return {'tables':copy.deepcopy(truth()['tables'])}

def test_transform_maps_whole_ui_and_cells_with_inverse():
    h=homography([[0,0],[100,0],[100,80],[0,80]],[[10,12],[97,4],[92,91],[3,81]])
    for polygon in [corners([2,4,30,32]),corners([40,50,90,75])]:
        assert np.allclose(points(np.linalg.inv(h),points(h,polygon)),polygon,atol=.001)

def test_split_families_and_all_twelve_have_different_structure():
    assert not set(DEVELOPMENT_FAMILIES)&set(TEST_FAMILIES)
    signatures=[]
    for family in DEVELOPMENT_FAMILIES+TEST_FAMILIES:
        _,tables,context,_,_=chinese(family,0,'development' if family in DEVELOPMENT_FAMILIES else 'test')
        signatures.append(tuple((t['rows'],t['columns'],tuple((c['row'],c['column'],c['row_span'],c['column_span'],c['text']) for c in t['cells'] if c['row']==0)) for t in tables))
    assert len(set(signatures))==12

def test_exact_good_and_failure_denominator():
    good=score_sample(truth(),result());failed=score_sample(truth(),None,'failed')
    assert good['metrics']['perfect_images']==1
    report=summarize([good,failed]);assert report['totals']['images']==2
    assert report['ratios']['whole_image_correct_rate']==.5
    assert failed['metrics']['critical_exact_cells']==0
    assert report['independent_source_count']==1

def test_swapped_identifiers_and_amounts_count_as_errors_even_same_text_set():
    value=result();cells=value['tables'][0]['cells'];cells[1]['text'],cells[2]['text']=cells[2]['text'],cells[1]['text']
    scored=score_sample(truth(),value)
    assert scored['metrics']['critical_exact_cells']==0
    assert scored['metrics']['critical_error_images']==1
    assert scored['metrics']['perfect_images']==0

def test_leading_zero_negative_decimal_changes_never_normalized_away():
    value=result();value['tables'][0]['cells'][1]['text']='1';value['tables'][0]['cells'][2]['text']='1,23456'
    assert score_sample(truth(),value)['metrics']['critical_exact_cells']==0

def test_spans_omissions_and_extra_ui_are_separate():
    value=result();value['tables'][0]['cells'][0]['column_span']=1
    value['tables'][0]['cells'].pop();value['tables'][0]['cells'].append({'row':3,'column':0,'text':'Excel'})
    scored=score_sample(truth(),value)['metrics']
    assert scored['exact_span_cells']==1
    assert scored['omitted_cells']==1 and scored['extra_cells']==1 and scored['ui_intrusion_cells']==1
    assert scored['perfect_images']==0

def test_occluded_values_excluded_only_from_visible_denominator():
    sample=truth();sample['tables'][0]['cells'][1]['visibility']='occluded'
    value=result();value['tables'][0]['cells'][1]['text']=''
    metrics=score_sample(sample,value)['metrics']
    assert metrics['critical_cells']==2 and metrics['visible_critical_cells']==1
    assert metrics['critical_exact_cells']==1 and metrics['visible_critical_exact_cells']==1

def test_no_table_ui_hallucination_counts_extra_table():
    sample={**truth(),'tables':[],'expect_no_table':True}
    assert score_sample(sample,{'tables':[]})['metrics']['no_table_correct']==1
    hallucinated=score_sample(sample,result())['metrics'];assert hallucinated['extra_tables']==1 and hallucinated['perfect_images']==0

def test_ui_intrusion_outside_tables_still_fails_whole_image():
    value={**result(),'text':'Excel 文件 开始 =SUM(C2:C11) 300.00'}
    scored=score_sample(truth(),value)['metrics']
    assert scored['critical_exact_cells']==2 and scored['ui_intrusion_text_markers']>0
    assert scored['perfect_images']==0

def test_automatic_delivery_false_negative_is_visible_in_acceptance_metrics():
    value=result();value['tables'][0]['cells'][1]['text']='00001'
    value['quality']={'automatic_delivery':True,'issues':[]}
    metrics=score_sample(truth(),value)['metrics']
    assert metrics['automatic_delivered_images']==1 and metrics['automatic_error_images']==1
    assert metrics['automatic_critical_error_images']==1

def test_real_screen_challenge_mask_and_geometry():
    base,tables,context,_,_=chinese('detail',0,'development')
    im,mapped,ctx,ui,occ,visibility,params,matrix,screen_matrix=render_screen(base,tables,context,'hard',random.Random(13),'occlusion')
    assert occ.getbbox() is not None and occ.size==im.size and visibility.size==im.size
    assert any(c['visibility']=='occluded' for t in mapped for c in t['cells'])
    assert params['effect_order'][0]=='compose_ui' and params['effect_order'][-1]=='camera_blur_jpeg'
    assert not np.allclose(screen_matrix,np.eye(3))

def test_multiple_detected_regions_aggregate_all_tables():
    meta={'primary_id':'full','prepared_images':[{'id':'full','kind':'rectified','path':'F:/full.png'},
        {'id':'one','kind':'table_crop','path':'F:/one.png'},{'id':'two','kind':'table_crop','path':'F:/two.png'}]}
    chosen=first_pass_images(meta);assert [p['id'] for p in chosen]==['one','two']
    merged=aggregate_results([result(),result()],chosen);assert len(merged['tables'])==2
    assert merged['tables'][1]['prepared_image_id']=='two'

def test_xlsx_readback_detects_digit_loss_formula_and_merge_damage(tmp_path):
    from openpyxl import Workbook
    assert tmp_path.resolve().drive.upper()=='F:'
    book=Workbook();sheet=book.active;sheet.title='Table 1'
    sheet.cell(1,1,'金额（元）').data_type='s';sheet.merge_cells('A1:B1')
    sheet.cell(2,1,'000001').data_type='s';sheet.cell(2,2,'-1,234.56').data_type='s'
    target=tmp_path/'actual.xlsx';book.save(target)
    assert readback(target,result())['passes']
    sheet.cell(2,1,1);sheet.unmerge_cells('A1:B1');book.save(target)
    damaged=readback(target,result());kinds={e['kind'] for e in damaged['errors']}
    assert {'literal_value','unsafe_type','merged_spans'}<=kinds

def test_no_formula_execution_allowed_in_xlsx_readback(tmp_path):
    from openpyxl import Workbook
    assert tmp_path.resolve().drive.upper()=='F:'
    value={'tables':[{'rows':1,'columns':1,'cells':[{'row':0,'column':0,'row_span':1,'column_span':1,'text':'=SUM(1,2)'}]}]}
    book=Workbook();book.active.cell(1,1,'=SUM(1,2)');target=tmp_path/'formula.xlsx';book.save(target)
    check=readback(target,value)
    assert any(e['kind']=='unsafe_type' for e in check['errors'])

def saved_context(tmp_path):
    assert tmp_path.resolve().drive.upper()=='F:'
    sample={'id':'input','image_sha256':'sha','width':100,'height':80}
    raw={'status':'success','image_version':'sha','coordinate_space':'exif_normalized_original','image':{'width':100,'height':80}}
    evidence={'layout':[],'text':[],'provenance':{'contains_reference_input':False}}
    write(tmp_path/'context.json',raw);write(tmp_path/'evidence.json',evidence)
    record={'id':'input','status':'success','image_sha256':'sha','context_sha256':digest(tmp_path/'context.json'),'evidence_sha256':digest(tmp_path/'evidence.json')}
    write(tmp_path/'record.json',record)
    return sample,raw,evidence,record

def test_context_identity_hash_and_normalized_evidence_verified(tmp_path):
    sample,raw,evidence,record=saved_context(tmp_path)
    assert verify_context_record(tmp_path,sample,normalize=lambda value:evidence)[1]==raw
    write(tmp_path/'evidence.json',{**evidence,'text':[{'text':'fake'}]})
    with pytest.raises(ValueError,match='Saved evidence changed'):verify_context_record(tmp_path,sample)

def test_context_failed_dimensions_image_version_and_reference_input_rejected(tmp_path):
    sample,raw,evidence,record=saved_context(tmp_path)
    record['status']='failed';write(tmp_path/'record.json',record)
    with pytest.raises(ValueError,match='not successful'):verify_context_record(tmp_path,sample)
    record['status']='success';raw['image_version']='wrong';write(tmp_path/'context.json',raw)
    record['context_sha256']=digest(tmp_path/'context.json');write(tmp_path/'record.json',record)
    with pytest.raises(ValueError,match='image version differs'):verify_context_record(tmp_path,sample)
    raw['image_version']='sha';raw['image']['width']=99;write(tmp_path/'context.json',raw)
    record['context_sha256']=digest(tmp_path/'context.json');write(tmp_path/'record.json',record)
    with pytest.raises(ValueError,match='dimensions differ'):verify_context_record(tmp_path,sample)
    raw['image']['width']=100;write(tmp_path/'context.json',raw);evidence['provenance']['contains_reference_input']=True
    write(tmp_path/'evidence.json',evidence);record.update(context_sha256=digest(tmp_path/'context.json'),evidence_sha256=digest(tmp_path/'evidence.json'));write(tmp_path/'record.json',record)
    with pytest.raises(ValueError,match='isolated from references'):verify_context_record(tmp_path,sample)

def test_context_mismatched_worker_or_dataset_cannot_mix():
    binding={'inputs_sha256':'inputs','worker_source_sha256':'worker','engine_config_sha256':'config'}
    context={**binding,'stage':'screen-context','geometry_provider':'context','labels_or_ideal_regions_provided':False}
    validate_context_binding(context,binding)
    for field in binding:
        with pytest.raises(ValueError,match=field):validate_context_binding({**context,field:'wrong'},binding)
    with pytest.raises(ValueError,match='inference context'):validate_context_binding({**context,'labels_or_ideal_regions_provided':True},binding)

def test_explicit_context_reuse_retains_old_source_and_still_requires_input_config_identity():
    binding={'inputs_sha256':'inputs','worker_source_sha256':'new_worker','engine_config_sha256':'config','context_worker_source_sha256':'old_worker'}
    context={'inputs_sha256':'inputs','worker_source_sha256':'old_worker','engine_config_sha256':'config',
             'stage':'screen-context','geometry_provider':'context','labels_or_ideal_regions_provided':False}
    validate_context_binding(context,binding,source_policy='recorded')
    with pytest.raises(ValueError,match='worker_source_sha256'):validate_context_binding(context,binding)
    with pytest.raises(ValueError,match='inputs_sha256'):validate_context_binding({**context,'inputs_sha256':'wrong'},binding,source_policy='recorded')
    with pytest.raises(ValueError,match='Recorded context worker source'):validate_context_binding({**context,'worker_source_sha256':'fake'},binding,source_policy='recorded')

def test_runner_uses_shared_policy_when_available_and_frozen_legacy_fallback_otherwise():
    meta={'primary_id':'original','prepared_images':[{'id':'original','kind':'original'},{'id':'one','kind':'table_crop'},{'id':'two','kind':'table_crop'}]}
    assert [item['id'] for item in first_pass_images(meta)]==['one','two']
    assert [item['id'] for item in first_pass_images(meta,selector=lambda value:[value['prepared_images'][0]])]==['original']

def test_offline_crop_review_preserves_full_field_identity_and_ui():
    selected=[{'id':'crop','kind':'table_crop','source_polygon':[[10,10],[80,10],[80,80],[10,80]]}]
    assert coverage({'polygon':[[20,20],[30,20],[30,30],[20,30]]},selected)==1
    assert coverage({'polygon':[[70,20],[90,20],[90,30],[70,30]]},selected)==.5
    reference={'id':'sample','tables':[{'cells':[{'row':0,'column':0,'text':'001','polygon':[[70,20],[90,20],[90,30],[70,30]],'visibility':'visible'}]}],
               'context_targets':[{'kind':'unit','text':'万元','polygon':[[1,1],[9,1],[9,9],[1,9]]}],
               'ui_regions':[{'kind':'toolbar','polygon':[[10,10],[80,10],[80,20],[10,20]]}]}
    metadata={'primary_id':'crop','prepared_images':selected,'business_regions':[],'excluded_regions':[],'quality':{}}
    report=inspect_preparation(reference,metadata)
    assert report['metrics']['images_with_visible_cell_crop_loss']==1
    assert report['metrics']['images_with_context_crop_loss']==1
    assert report['metrics']['images_with_ui_geometry']==1
    assert report['lost_cells'][0]['geometric_coverage']==.5

def test_resident_load_is_counted_once_and_context_charged_once():
    records=[{'status':'success','inference_seconds':4,'model_load_seconds':20,'engine_session':{'id':'resident','load_seconds':20}} for _ in range(3)]
    summary={'loads':[{'session_id':'resident','wall_seconds':21,'reported_load_seconds':20}],'attempt_wall_seconds':40,'completed':3,'total':3,'cache_resume_separate':False}
    context={'loads':[],'attempt_wall_seconds':15,'completed':3,'total':3,'cache_resume_separate':False}
    report=performance(records,summary,context,[{'inference_seconds':2,'model_load_seconds':5},{'inference_seconds':2,'model_load_seconds':0}])
    assert report['primary_model_load_seconds']==20
    assert report['context_model_load_seconds']==5
    assert report['full_sequential_route_wall_seconds']==55
    assert performance(records)['primary_model_load_seconds']==20
    summary['cache_resume_separate']=True
    assert performance(records,summary,context)['full_sequential_route_wall_seconds'] is None

def test_independent_provenance_validator_rejects_posthoc_transcription(tmp_path):
    assert tmp_path.resolve().drive.upper()=='F:'
    input_image=tmp_path/'image.png';Image.new('RGB',(10,10),'white').save(input_image)
    manifest=tmp_path/'inputs.json';sample={'id':'sample','image':str(input_image),'image_sha256':digest(input_image)}
    write(manifest,{'truth_available_to_inference':False,'samples':[sample]})
    runroot=tmp_path/'run';folder=runroot/'sample';attempt=folder/'attempt-0001';region=attempt/'region-00'
    write(runroot/'binding.json',{'inputs_sha256':digest(manifest),'labels_or_ideal_regions_provided':False})
    value={**result(),'status':'success','text':'000001 -1,234.56'}
    write(region/'result.json',value);write(folder/'result.json',value)
    record={'id':'sample','status':'success','image_sha256':sample['image_sha256'],'output':str(attempt),
            'prepared_images':[{'path':str(input_image),'sha256':digest(input_image)}],'result_sha256':digest(folder/'result.json')}
    write(folder/'record.json',record)
    assert validate_run(manifest,runroot)['is_final']
    value['tables'][0]['cells'][1]['text']='00001';write(folder/'result.json',value)
    record['result_sha256']=digest(folder/'result.json');write(folder/'record.json',record)
    with pytest.raises(ValueError,match='transcription/structure differs'):validate_run(manifest,runroot)

def test_explicit_sample_selection_retains_manifest_order_and_rejects_ambiguous_ranges():
    samples=[{'id':key} for key in ['first','middle','last']]
    assert select_input_samples(samples,include_ids=['last','first'])==[samples[0],samples[2]]
    assert select_input_samples(samples,offset=1,limit=1)==[samples[1]]
    assert select_input_samples(samples)==samples
    for include_ids,message in [(['first','first'],'Duplicate requested'),(['absent'],'Unknown sample'),([], 'nonempty')]:
        with pytest.raises(ValueError,match=message):select_input_samples(samples,include_ids=include_ids)
    with pytest.raises(ValueError,match='Duplicate sample IDs in input'):select_input_samples([samples[0],samples[0]],include_ids=['first'])
    for arguments in [{'offset':1},{'limit':1}]:
        with pytest.raises(ValueError,match='offset/limit'):select_input_samples(samples,include_ids=['first'],**arguments)

def test_bound_subset_score_and_provenance_keep_missing_selected_items_in_denominator(tmp_path):
    image=tmp_path/'image.png';Image.new('RGB',(10,10),'white').save(image)
    samples=[{'id':key,'image':str(image),'image_sha256':digest(image)} for key in ['one','two','three']]
    inputs=tmp_path/'inputs.json';write(inputs,{'truth_available_to_inference':False,'samples':samples})
    references=[{**truth(),'id':sample['id']} for sample in samples]
    truth_path=tmp_path/'truth.json';write(truth_path,{'samples':references})
    runroot=tmp_path/'run';binding={'inputs_sha256':digest(inputs),'labels_or_ideal_regions_provided':False,'include_sample_ids':['three','one']}
    write(runroot/'binding.json',binding)
    write(runroot/'one'/'record.json',{'id':'one','status':'failed','image_sha256':digest(image),'error':'deliberate test failure'})
    report=score_run(truth_path,runroot)
    assert report['include_sample_ids']==['one','three']
    assert report['summary']['totals']['images']==2
    assert report['summary']['totals']['missing_images']==1
    assert report['summary']['totals']['failed_images']==1
    assert not report['completed_snapshot_is_final']
    provenance=validate_run(inputs,runroot)
    assert provenance['total']==2 and provenance['failed']==1 and provenance['missing']==1 and not provenance['is_final']
    assert 'diagnostic subset' in provenance['scope']
    for invalid in [['one','one'],['unknown']]:
        write(runroot/'binding.json',{**binding,'include_sample_ids':invalid})
        with pytest.raises(ValueError):score_run(truth_path,runroot)
        with pytest.raises(ValueError):validate_run(inputs,runroot)

def test_runner_sample_cli_binds_full_manifest_reuses_one_session_and_preserves_legacy_binding(tmp_path,monkeypatch):
    import types
    import run as runner
    worker=tmp_path/'src';worker.mkdir()
    bundle=tmp_path/'bundle';write(bundle/'config'/'engines.json',{})
    samples=[]
    for key in ['one','two','three']:
        image=tmp_path/(key+'.png');Image.new('RGB',(10,10),'white').save(image)
        samples.append({'id':key,'image':str(image),'image_sha256':digest(image)})
    inputs=tmp_path/'inputs.json';write(inputs,{'truth_available_to_inference':False,'samples':samples})
    calls=[];timeouts=[]
    class FakeAdapter:
        def __init__(self,*args,**kwargs):self.ipc=types.SimpleNamespace(name='one-resident-session');self.load_seconds=2;self.recognition_timeout=900
        def load(self):calls.append('load')
        def recognize(self,image,output,**kwargs):
            calls.append(image.stem);timeouts.append(self.recognition_timeout)
            if self.recognition_timeout==180 and image.stem=='two':raise TimeoutError('test request timeout')
            value={'status':'success','text':'recognized','tables':[],'engine_session':{'id':self.ipc.name}}
            write(output/'result.json',value)
            return value
        def unload(self):calls.append('unload')
    monkeypatch.setitem(sys.modules,'ocr_workbench.adapter',types.SimpleNamespace(EngineAdapter=FakeAdapter))
    monkeypatch.setattr(runner,'gpu_sample',lambda:{'test_stub':True})
    # Runtime isolation is verified separately; this fixture creates no child processes.
    monkeypatch.setattr(runner,'isolate_runtime',lambda output:None)
    runroot=tmp_path/'subset-run'
    base=['run.py','--inputs',str(inputs),'--bundle',str(bundle),'--worker-source',str(worker),'--engine','hunyuan']
    monkeypatch.setattr(sys,'argv',base+['--output',str(runroot),'--sample-id','three','--sample-id','one'])
    runner.main()
    binding=json.loads((runroot/'binding.json').read_text('utf-8'))
    assert binding['inputs_sha256']==digest(inputs) and binding['include_sample_ids']==['one','three']
    assert calls==['load','one','three','unload']
    assert not (runroot/'two').exists() and validate_run(inputs,runroot)['is_final']
    calls.clear();runner.main()
    assert not calls and json.loads((runroot/'run-summary.json').read_text('utf-8'))['resumed_records']==2
    monkeypatch.setattr(sys,'argv',base+['--output',str(runroot),'--sample-id','two'])
    with pytest.raises(ValueError,match='immutable'):runner.main()
    legacy=tmp_path/'legacy-run'
    monkeypatch.setattr(sys,'argv',base+['--output',str(legacy)])
    runner.main()
    assert json.loads((legacy/'binding.json').read_text('utf-8'))=={
        'inputs_sha256':digest(inputs),'worker_source_sha256':runner.source_hash(worker),'engine':'hunyuan',
        'engine_config_sha256':digest(bundle/'config/engines.json'),'bundle':str(bundle),
        'offset':0,'limit':None,'preprocess_module_sha256':None,'first_pass':True,'labels_or_ideal_regions_provided':False}
    assert all(timeout==900 for timeout in timeouts)
    calls.clear();timeouts.clear();bounded=tmp_path/'bounded-run'
    bounded_arguments=base+['--output',str(bounded),'--request-timeout-seconds','180',
                           '--sample-id','three','--sample-id','two','--sample-id','one']
    monkeypatch.setattr(sys,'argv',bounded_arguments)
    runner.main()
    assert timeouts==[180,180,180]
    assert calls==['load','one','two','unload','load','three','unload']
    assert json.loads((bounded/'binding.json').read_text('utf-8'))['request_timeout_seconds']==180
    provenance=validate_run(inputs,bounded)
    assert provenance['total']==3 and provenance['success']==2 and provenance['failed']==1 and provenance['is_final']
    truth_path=tmp_path/'truth.json';write(truth_path,{'samples':[{**truth(),'id':sample['id']} for sample in samples]})
    score=score_run(truth_path,bounded)
    assert score['summary']['totals']['images']==3 and score['summary']['totals']['failed_images']==1
    monkeypatch.setattr(sys,'argv',[part if part!='180' else '181' for part in bounded_arguments])
    with pytest.raises(ValueError,match='immutable'):runner.main()
    for timeout in ['9','901']:
        monkeypatch.setattr(sys,'argv',base+['--output',str(tmp_path/'invalid-run'),'--request-timeout-seconds',timeout])
        with pytest.raises(ValueError,match='10..900'):runner.main()

def test_subset_reusing_larger_context_keeps_costs_separate_from_route_wall(tmp_path):
    references=[{**truth(),'id':key} for key in ['one','two','three']]
    truth_path=tmp_path/'truth.json';write(truth_path,{'samples':references})
    inputs=tmp_path/'inputs.json';write(inputs,{'truth_available_to_inference':False,'samples':[{'id':key} for key in ['one','two','three']]})
    runroot=tmp_path/'run';contextroot=tmp_path/'context'
    write(contextroot/'binding.json',{'inputs_sha256':digest(inputs)})
    write(contextroot/'run-summary.json',{'attempt_wall_seconds':15,'completed':3,'total':3,'loads':[]})
    write(runroot/'binding.json',{'inputs_sha256':digest(inputs),'include_sample_ids':['one','three'],
          'context_run':str(contextroot),'context_binding_sha256':digest(contextroot/'binding.json')})
    for key in ['one','three']:
        write(runroot/key/'record.json',{'id':key,'status':'failed','wall_seconds':4,'error':'test only'})
    write(runroot/'run-summary.json',{'attempt_wall_seconds':8,'completed':2,'total':2,'loads':[]})
    report=compare(truth_path,{'subset':runroot})['runs']['subset']
    assert report['score']['summary']['totals']['images']==2
    assert report['performance']['primary_stage_invocation_wall_seconds']==8
    assert report['performance']['context_stage_invocation_wall_seconds']==15
    assert report['performance']['full_sequential_route_wall_seconds'] is None
