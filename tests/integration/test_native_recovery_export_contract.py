"""Independent source-index readback against persisted SQLite/file identities.

Small engineering fixtures prove the rejection contract, not OCR quality.
The workbook is constructed directly, without SOURCE_COLUMNS/export_xlsx.
"""
from contextlib import closing
import hashlib
import io
from pathlib import Path
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from scripts.evaluate_screen_native_recovery import check_source_index


@pytest.fixture
def source_index(tmp_path):
    image_path = tmp_path / 'frozen-image.png'
    image_path.write_bytes(b'persisted-image-engineering-fixture')
    image_hash = hashlib.sha256(image_path.read_bytes()).hexdigest()
    db_path = tmp_path / 'workbench.sqlite3'
    with sqlite3.connect(db_path) as db:
        db.executescript('''
            CREATE TABLE images(id TEXT, name TEXT, active_version TEXT);
            CREATE TABLE versions(id TEXT, image_id TEXT, path TEXT, sha256 TEXT);
            CREATE TABLE tasks(id TEXT, image_id TEXT, version_id TEXT);
        ''')
        db.execute('INSERT INTO images VALUES(?,?,?)', ('imageA', '原图.png', 'versionA'))
        db.execute('INSERT INTO versions VALUES(?,?,?,?)', ('versionA', 'imageA', image_path.name, image_hash))
        db.execute('INSERT INTO tasks VALUES(?,?,?)', ('taskA', 'imageA', 'versionA'))

    class StoredIdentities:
        def one(self, table, identity):
            assert table in {'images', 'versions', 'tasks'}
            with closing(sqlite3.connect(db_path)) as db:
                db.row_factory = sqlite3.Row
                return dict(db.execute('SELECT * FROM ' + table + ' WHERE id=?', (identity,)).fetchone())

        def file(self, relative):
            return tmp_path / relative

    item = {'image_id': 'imageA', 'input_version_id': 'versionA', 'source_name': '原图.png',
            'source_hash': image_hash, 'result_id': 'resultA', 'result_revision': 0}
    actual = {'id': 'resultA', 'task_id': 'taskA', 'revision': 0,
              'original': {'engine': 'hunyuan', 'project_image_version': 'versionA'}}
    labels = ['工作簿', '工作表', '原图', '图片 ID', '图像版本', '输入 SHA-256',
              '引擎', '结果 ID', '校对 revision', '表格序号', '文档 ID']
    book = Workbook()
    sheet = book.active
    sheet.title = '来源索引'
    sheet.append(labels)
    for table_index in (1, 2):
        sheet.append(['0001-原图.xlsx', f'Table {table_index}', '原图.png', 'imageA', 'versionA',
                      image_hash, 'hunyuan', 'resultA', '0', str(table_index), ''])
    for row in sheet:
        for cell in row:
            cell.data_type = 's'
            cell.number_format = '@'
    payload = io.BytesIO()
    book.save(payload)
    book.close()
    loaded = load_workbook(io.BytesIO(payload.getvalue()), data_only=False)
    try:
        yield loaded['来源索引'], StoredIdentities(), item, actual, db_path, image_path
    finally:
        loaded.close()


def test_source_index_reads_two_tables_and_zero_revision_from_actual_sqlite(source_index):
    sheet, store, item, actual, _, _ = source_index
    result = check_source_index(sheet, '0001-原图.xlsx', item, actual, store, 2)
    assert result['source_index_rows'] == 2 and result['input_file_sha256_verified']
    assert result['rows'][0]['校对 revision'] == '0'
    assert sheet.cell(2, 11).value is None  # Nonapplicable document metadata is allowed.


@pytest.mark.parametrize('label', ['工作簿', '工作表', '原图', '图片 ID', '图像版本',
    '输入 SHA-256', '引擎', '结果 ID', '校对 revision', '表格序号'])
def test_source_index_rejects_each_blank_core_identity(source_index, label):
    sheet, store, item, actual, _, _ = source_index
    labels = [cell.value for cell in sheet[1]]
    sheet.cell(2, labels.index(label) + 1).value = None
    with pytest.raises(ValueError, match='actual DB identity'):
        check_source_index(sheet, '0001-原图.xlsx', item, actual, store, 2)


def test_source_index_rejects_changed_version_bytes_despite_matching_db_and_sheet(source_index):
    sheet, store, item, actual, _, image_path = source_index
    image_path.write_bytes(b'different-current-image')
    with pytest.raises(ValueError, match='actual DB/file state'):
        check_source_index(sheet, '0001-原图.xlsx', item, actual, store, 2)


def test_source_index_rejects_different_actual_version_owner(source_index):
    sheet, store, item, actual, db_path, _ = source_index
    with sqlite3.connect(db_path) as db:
        db.execute("UPDATE versions SET image_id='foreign-image'")
    with pytest.raises(ValueError, match='actual DB/file state'):
        check_source_index(sheet, '0001-原图.xlsx', item, actual, store, 2)
