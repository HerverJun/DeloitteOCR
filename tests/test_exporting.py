import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
import zipfile
from openpyxl import load_workbook
from ocr_workbench.exporting import build_export


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.values = {}
        for key in ["a", "b"]:
            table = {
                "rows": 2,
                "columns": 2,
                "cells": [
                    {
                        "row": 0,
                        "column": 0,
                        "row_span": 1,
                        "column_span": 2,
                        "text": "校对" + key,
                    },
                    {
                        "row": 1,
                        "column": 0,
                        "row_span": 1,
                        "column_span": 1,
                        "text": "00123456789012345678",
                    },
                    {
                        "row": 1,
                        "column": 1,
                        "row_span": 1,
                        "column_span": 1,
                        "text": "=1+1",
                    },
                ],
            }
            self.values[key] = {
                "id": key,
                "original": {"text": "原始"},
                "edited": {"text": "校对" + key, "tables": [table, table]},
            }
        self.store = SimpleNamespace(root=self.root, result=self.values.__getitem__)

    def tearDown(self):
        self.temp.cleanup()

    def test_default_separate_workbooks_preserve_tables_strings_and_merges(self):
        path = build_export(self.store, ["a", "b", "a"], "xlsx")
        with zipfile.ZipFile(path) as archive:
            self.assertEqual(archive.namelist(), ["0001.xlsx", "0002.xlsx"])
            for name, key in zip(archive.namelist(), ["a", "b"]):
                book = load_workbook(io.BytesIO(archive.read(name)))
                self.assertEqual(len(book.worksheets), 2)
                self.assertEqual(book.active["A1"].value, "校对" + key)
                self.assertEqual(
                    str(next(iter(book.active.merged_cells.ranges))), "A1:B1"
                )
                self.assertEqual(book.active["A2"].value, "00123456789012345678")
                self.assertEqual(book.active["B2"].data_type, "s")

    def test_aggregate_is_explicit_and_single_remains_workbook(self):
        merged = build_export(self.store, ["a", "b"], "xlsx", True)
        self.assertEqual(len(load_workbook(merged).worksheets), 4)
        single = build_export(self.store, ["a"], "xlsx")
        self.assertEqual(len(load_workbook(single).worksheets), 2)

    def test_text_json_and_markdown_use_saved_edits(self):
        for format in ["txt", "md", "json"]:
            with zipfile.ZipFile(
                build_export(self.store, ["a", "b"], format)
            ) as archive:
                for name, key in zip(archive.namelist(), ["a", "b"]):
                    content = archive.read(name).decode("utf-8")
                    self.assertIn("校对" + key, content)
                    if format == "json":
                        self.assertEqual(
                            json.loads(content)["original"], {"text": "原始"}
                        )
                    else:
                        self.assertNotIn("原始", content)

    def test_failed_batch_leaves_no_partial_archive(self):
        self.values["b"]["edited"]["tables"] = []
        with self.assertRaisesRegex(ValueError, "没有结构化表格"):
            build_export(self.store, ["a", "b"], "xlsx")
        self.assertEqual(list((self.root / "exports").iterdir()), [])
        self.values["a"]["edited"]["tables"] = []
        with self.assertRaisesRegex(ValueError, "No recognized tables"):
            build_export(self.store, ["a", "b"], "xlsx", True)


if __name__ == "__main__":
    unittest.main()
