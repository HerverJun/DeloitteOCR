import tempfile
from pathlib import Path
import unittest
from ocr_workbench.tables import parse_tables, export_xlsx


class TableTests(unittest.TestCase):
    def test_spans_empty_cells_and_exact_excel_strings(self):
        from openpyxl import load_workbook
        tables = parse_tables('<table><tr><th colspan="3">标题</th></tr>'
                              '<tr><td rowspan="2">00123456789012345678</td><td></td><td>=1+1</td></tr>'
                              '<tr><td>A&amp;B<br>C</td><td>02</td></tr></table>')
        table = tables[0]
        self.assertEqual((table['rows'], table['columns']), (3, 3))
        self.assertEqual(table['cells'][-2]['text'], 'A&B\nC')
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / '结果 空格.xlsx'
            export_xlsx(tables, path)
            sheet = load_workbook(path).active
            self.assertEqual(sheet['A2'].value, '00123456789012345678')
            self.assertEqual(sheet['C2'].value, '=1+1')
            self.assertEqual(sheet['C2'].data_type, 's')
            self.assertEqual(sheet['C3'].value, '02')
            self.assertEqual({str(r) for r in sheet.merged_cells.ranges}, {'A1:C1', 'A2:A3'})

    def test_multiple_tables_and_markdown(self):
        self.assertEqual(len(parse_tables('<table><tr><td>1</td></tr></table><table><tr><td>2</td></tr></table>')), 2)
        t = parse_tables('| 编号 | 值 |\n| --- | --- |\n| 0001 | A\\|B |')[0]
        self.assertEqual(t['rows'], 2)
        self.assertEqual(t['cells'][-1]['text'], 'A|B')

    def test_truncated_invalid_or_overlapping_table_fails(self):
        for text in ['<table><tr><td>x', '<table><tr><td rowspan="0">x</td></tr></table>',
                     '<table><tr><td>x</td><td rowspan="2">y</td></tr><tr><td colspan="2">z</td></tr></table>']:
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_tables(text)

    def test_caption_preserved_without_inventing_a_span(self):
        from openpyxl import load_workbook
        t=parse_tables('<table><caption>=标题</caption><tr><td>0001</td></tr></table>')[0]
        self.assertEqual(t['caption'],'=标题')
        self.assertEqual(t['rows'],1)
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'caption.xlsx'
            export_xlsx([t],path)
            sheet=load_workbook(path).active
            self.assertEqual(sheet['A1'].value,'=标题')
            self.assertEqual(sheet['A1'].data_type,'s')
            self.assertEqual(sheet['A2'].value,'0001')
            self.assertEqual(len(sheet.merged_cells.ranges),0)


if __name__ == '__main__':
    unittest.main()
