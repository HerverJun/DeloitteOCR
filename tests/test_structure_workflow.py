from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from ocr_workbench.coordinates import box_polygon
from ocr_workbench.editing import tables_html
from ocr_workbench.imaging import add_image
from ocr_workbench.store import Store, Conflict
from ocr_workbench.structure_diagnostics import prepare_candidates, identify_tables, local_variants, preserve_values
from ocr_workbench.structure_store import record_candidates, refresh_proposals, decide_structure, structure_view
from ocr_workbench.tables import parse_tables


def table(rows):
    return {"rows": len(rows), "columns": len(rows[0]), "cells": [
        {"row": r,"column": c,"row_span": 1,"column_span": 1,"text": value}
        for r, values in enumerate(rows) for c,value in enumerate(values)]}


def prediction(values, *, merged=False):
    cells, blocks = [], []
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            if merged and r == 0 and c == 1:
                continue
            box = [c*100,r*40,(c+1)*100,(r+1)*40]
            cells.append({"cell_id": str(len(cells)), "row_id": r, "column_id": c,
                "rowspan_val": 1,"colspan_val": 2 if merged and r == 0 else 1,
                "bbox": [0,0,200,40] if merged and r == 0 else box})
            if value:
                blocks.append({"id": f"token-{r}-{c}", "text": value,"kind": "text","granularity": "word",
                    "polygon": box_polygon([box[0]+5,box[1]+5,box[2]-5,box[3]-5])})
    return {"tf_table_cells": cells,"source_semantics": "raw_structure","model_sha256": "fixed-test-model"}, blocks


class StructureDiagnosticsTests(unittest.TestCase):
    def test_raw_ids_numbering_and_literal_numbers_survive(self):
        pred, blocks = prediction([["金额","编号"],["-0.01","000123"]])
        result = prepare_candidates(pred,blocks,source_result="ocr",image_version="v",width=200,height=80)
        cells = result["tables"][0]["skeleton"]["cells"]
        self.assertEqual([c["text"] for c in cells], ["金额","编号","-0.01","000123"])
        self.assertFalse(result["contributes_to_votes"])
        self.assertEqual(cells[3]["structure_source"]["original_cell_id"], "3")
        self.assertFalse(cells[3]["structure_source"]["index_mapping"]["compression_applied"])

    def test_cross_cell_line_is_not_divided_or_duplicated(self):
        pred, _ = prediction([["a","b"]])
        blocks = [{"id":"line","text":"001 002","kind":"text","polygon":box_polygon([10,5,190,35])}]
        result = prepare_candidates(pred,blocks,source_result="ocr",image_version="v",width=200,height=40)
        t = result["tables"][0]
        self.assertEqual(t["unassigned_token_ids"],["line"])
        self.assertEqual([c["text"] for c in t["skeleton"]["cells"]],["",""])

    def test_incomplete_grid_is_not_compressed(self):
        pred, blocks = prediction([["a","b"],["c","d"],["e","f"]])
        pred["tf_table_cells"] = [c for c in pred["tf_table_cells"] if c["row_id"] != 1]
        value = prepare_candidates(pred,blocks,source_result="ocr",image_version="v",width=200,height=120)["tables"][0]
        self.assertEqual(value["skeleton"]["rows"],3)
        self.assertEqual(len(value["skeleton"]["cells"]),4)
        self.assertIn("incomplete_grid",value["reason_codes"])

    def test_same_values_do_not_identify_two_tables(self):
        pred, blocks = prediction([["same","heading"],["1","2"]])
        candidate = prepare_candidates(pred,blocks,source_result="ocr",image_version="v",width=200,height=80)["tables"][0]
        current = table([["same","heading"],["1","2"]])
        self.assertEqual(identify_tables([current,current],[candidate,candidate]),[None,None])

    def test_local_merge_preserves_other_rows(self):
        current = table([["header",""],["0001","-2.00"]])
        proposed = deepcopy(current)
        proposed["cells"][0]["column_span"] = 2
        proposed["cells"].pop(1)
        variants = local_variants(current,proposed)
        self.assertEqual(len(variants),1)
        self.assertEqual(variants[0]["kind"],"merge")
        self.assertEqual(variants[0]["table"]["cells"][1:],[current["cells"][2],current["cells"][3]])

    def test_native_original_regions_disambiguate_same_text_and_reject_stale_evidence(self):
        from ocr_workbench.fusion_alignment import source_tables
        current = table([["same", "heading"], ["1", "2"]])
        for cell in current['cells']:
            cell['native_content'] = {'version':'native-adopted-fragments-v1', 'image_version':'v'}
            cell['structure_source'] = {'image_version':'v'}
        second = deepcopy(current)
        current['region_polygon'] = box_polygon([0, 0, 200, 80])
        second['region_polygon'] = box_polygon([0, 100, 200, 180])
        original = {'tables':[current, second], 'text':tables_html([current, second]),
                    'blocks':[], 'image':{'width':200, 'height':200}, 'project_image_version':'v',
                    'document':{'structure_preview':True, 'structure_text_source':'pdf-native'}}
        candidates = [{'skeleton':t, 'polygon':t['region_polygon']} for t in original['tables']]
        recovered = source_tables({'original':original})
        self.assertEqual([m[1] for m in identify_tables(recovered, candidates)], [0, 1])
        original['tables'][0]['native_content'] = {}  # Table-level metadata cannot authorize cells.
        original['tables'][0]['cells'][0]['native_content']['image_version'] = 'stale'
        self.assertIsNone(source_tables({'original':original})[0]['region_polygon'])
        original['document']['structure_preview'] = False
        self.assertTrue(all(t['region_polygon'] is None for t in source_tables({'original':original})))

    def test_headers_roundtrip(self):
        t = parse_tables('<table><tr><th colspan="2">2026</th></tr><tr><td>001</td><td>-0.10</td></tr></table>')[0]
        self.assertTrue(t["cells"][0]["is_header"])
        self.assertTrue(parse_tables(tables_html([t]))[0]["cells"][0]["is_header"])


class StructureStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name)/"workspace")
        self.project = self.store.project("结构试验")
        path = Path(self.temp.name)/"image.png"
        Image.new("RGB",(200,120),"white").save(path)
        self.image = add_image(self.store,self.project["id"],"image.png",path)
        task = self.store.enqueue(self.project["id"],[self.image["active_version"]],["paddlevl"])[0]
        self.store.claim()
        original = table([["项目","金额"],["收入","001.00"]])
        text = tables_html([original])
        original["source"] = parse_tables(text)[0]["source"]
        self.store.complete(task,{"engine":"paddlevl","text":text,"tables":[original],"blocks":[],"image":{"width":200,"height":120}})
        self.result_id = self.store.one("tasks",task)["result_id"]
        self.pred, self.blocks = prediction([["项目","金额"],["收入","001.00"],["税额","-0.01"]])
        self.version = self.store.one("versions",self.image["active_version"])
        with self.store.transaction() as db:
            record_candidates(db,self.result_id,self.version,self.pred,self.blocks)

    def tearDown(self):
        self.temp.cleanup()

    def proposal(self):
        value = self.store.result(self.result_id)
        view = refresh_proposals(self.store,self.result_id,value["revision"])
        return next(p for p in view["proposals"] if p["kind"] == "replace_table" and p["state"] == "pending")

    def body(self,p,request="decision-001",action="accept"):
        return {"request_id":request,"action":action,"revision":p["revision"],"basis":p["basis"],"version_id":self.version["id"]}

    def test_apply_preserves_source_and_undo_redo(self):
        p = self.proposal()
        self.assertTrue(p["can_apply"],p["conflicts"])
        saved = decide_structure(self.store,self.result_id,p["id"],self.body(p))
        self.assertEqual(saved["edited"]["tables"][0]["rows"],3)
        self.assertEqual(saved["edited"]["tables"][0]["cells"][-1]["text"],"-0.01")
        self.assertEqual(saved["edited"]["tables"][0]["cells"][-1]["structure_source"]["token_ids"],["token-2-1"])
        back = self.store.history(self.result_id,-1,saved["revision"])
        self.assertEqual(back["edited"]["tables"][0]["rows"],2)
        self.assertFalse(any(p["state"] == "accepted" for p in structure_view(self.store,self.result_id)["proposals"]))
        redone = self.store.history(self.result_id,1,back["revision"])
        self.assertEqual(redone["edited"],saved["edited"])

    def test_request_replay_is_idempotent_and_payload_bound(self):
        p = self.proposal()
        a = decide_structure(self.store,self.result_id,p["id"],self.body(p))
        b = decide_structure(self.store,self.result_id,p["id"],self.body(p))
        self.assertEqual(a,b)
        with self.assertRaises(Conflict):
            decide_structure(self.store,self.result_id,p["id"],self.body(p,action="keep"))

    def test_old_preservation_policy_cannot_be_applied_from_cached_proposal(self):
        p=self.proposal()
        with self.store.transaction() as db:
            payload=json.loads(db.execute('SELECT payload FROM structure_proposals WHERE id=?',(p['id'],)).fetchone()[0])
            payload['version']='structure-review-v2'
            db.execute('UPDATE structure_proposals SET payload=? WHERE id=?',(json.dumps(payload),p['id']))
        cached=next(x for x in structure_view(self.store,self.result_id)['proposals'] if x['id']==p['id'])
        self.assertFalse(cached['can_apply'])
        with self.assertRaisesRegex(Conflict,'保全规则'):
            decide_structure(self.store,self.result_id,p['id'],self.body(p))

    def test_stale_revision_and_image_cannot_apply(self):
        p = self.proposal()
        edit = self.store.result(self.result_id)["edited"]
        edit["tables"][0]["cells"][3]["text"] = "001.05"
        self.store.save(self.result_id,edit,0)
        with self.assertRaises(Conflict):
            decide_structure(self.store,self.result_id,p["id"],self.body(p))
        next_p = self.proposal()
        self.assertTrue(next_p["can_apply"], next_p["conflicts"])
        self.assertEqual(next_p["proposed_tables"][0]["cells"][3]["text"],"001.05")
        with self.store.transaction() as db:
            db.execute("UPDATE versions SET sha256='changed' WHERE id=?",(self.version["id"],))
        with self.assertRaises(Conflict):
            decide_structure(self.store,self.result_id,next_p["id"],self.body(next_p))

    def test_keep_defer_and_unrelated_edit_invalidation(self):
        p = self.proposal()
        decide_structure(self.store,self.result_id,p["id"],self.body(p,action="defer"))
        self.assertEqual(structure_view(self.store,self.result_id)["proposals"][0]["state"],"deferred")
        decide_structure(self.store,self.result_id,p["id"],self.body(p,request="keep-002",action="keep"))
        edit = self.store.result(self.result_id)["edited"]
        edit["text"] += "\n核对注记"
        self.store.save(self.result_id,edit,0)
        view = refresh_proposals(self.store,self.result_id,1)
        self.assertTrue(view["proposals"])
        self.assertTrue(all(v["state"] == "kept" for v in view["proposals"]))

    def test_blank_recovered_cell_needs_explicit_acknowledgement(self):
        pred, blocks = prediction([["项目","金额"],["收入","001.00"],["税额",""]])
        with self.store.transaction() as db:
            record_candidates(db,self.result_id,self.version,pred,blocks)
        p = self.proposal()
        self.assertTrue(p["unverified_empty_cells"])
        with self.assertRaises(ValueError):
            decide_structure(self.store,self.result_id,p["id"],self.body(p))
        saved = decide_structure(self.store,self.result_id,p["id"],{**self.body(p),"acknowledge_unverified_empty":True})
        self.assertEqual(saved["edited"]["tables"][0]["cells"][-1]["structure_source"]["text_state"],"unverified_empty")

    def test_document_queue_groups_provider_alternatives(self):
        from ocr_workbench.document_review import document_review_queue
        self.proposal()
        doc = self.store.page_for_version(self.version["id"])["document_id"]
        queue = document_review_queue(self.store,doc)
        self.assertEqual(queue["total"],1)
        self.assertEqual(queue["tasks"][0]["kind"],"structure")
        self.assertFalse(queue["empty_means_correct"])

    def test_export_readback_preserves_amounts_headers_and_full_lineage(self):
        import io
        import zipfile
        from openpyxl import load_workbook
        from ocr_workbench.exporting import build_export
        p = self.proposal()
        saved = decide_structure(self.store,self.result_id,p['id'],self.body(p))
        edit = saved['edited']
        edit['tables'][0]['cells'][0]['is_header'] = True
        self.store.save(self.result_id,edit,saved['revision'])
        archive_path = build_export(self.store,[self.result_id],'xlsx')
        with zipfile.ZipFile(archive_path) as archive:
            workbook = load_workbook(io.BytesIO(archive.read('OCR-result.xlsx')))
            self.assertEqual(workbook['Table 1']['B2'].value,'001.00')
            self.assertEqual(workbook['Table 1']['B3'].value,'-0.01')
            self.assertEqual(workbook['Table 1']['B3'].data_type,'s')
            self.assertTrue(workbook['Table 1']['A1'].font.bold)
            source = json.loads(archive.read(f'sources/structure-{self.result_id}.json'))
            self.assertEqual(source['revision'],2)
            self.assertEqual(source['adopted_tables'][0]['cells'][-1]['source']['token_ids'],['token-2-1'])
            self.assertEqual(source['candidates'][0]['payload']['prediction'],self.pred)
            index = workbook['来源索引']
            columns = [c.value for c in index[1]]
            self.assertEqual(index.cell(2,columns.index('原文页码')+1).value,'1')

    def test_geometry_is_recomputed_on_new_structure(self):
        from ocr_workbench.geometry import geometry_view
        p = self.proposal()
        saved = decide_structure(self.store,self.result_id,p['id'],self.body(p))
        view = geometry_view(self.store,self.result_id)
        self.assertEqual(view['revision'],saved['revision'])
        self.assertTrue(view['evidence'])
        self.assertTrue(all(e['details'].get('structure_revision') == saved['revision'] for e in view['evidence']))

    def test_new_candidates_invalidate_check_and_previous_pending_decisions(self):
        p = self.proposal()
        pred,blocks = prediction([['项目','金额'],['收入','001.00'],['税额','-0.02']])
        with self.store.transaction() as db:
            record_candidates(db,self.result_id,self.version,pred,blocks)
        self.assertEqual(self.store.rows('SELECT * FROM structure_checks'),[])
        with self.assertRaises(Conflict):
            decide_structure(self.store,self.result_id,p['id'],self.body(p))
        self.assertEqual(self.proposal()['proposed_tables'][0]['cells'][-1]['text'],'-0.02')

    def test_check_after_second_provider_keeps_unchanged_alternatives(self):
        first=refresh_proposals(self.store,self.result_id,0)
        original_ids={p['id'] for p in first['proposals']}
        with self.store.transaction() as db:
            record_candidates(db,self.result_id,self.version,{**self.pred,'component':'tableformer-raw'},self.blocks)
        refreshed=refresh_proposals(self.store,self.result_id,0)
        self.assertEqual({p['provider'] for p in refreshed['proposals']},{'paddle','tableformer-raw'})
        self.assertTrue(original_ids <= {p['id'] for p in refreshed['proposals']})
        self.assertEqual({p['id'] for p in refreshed['proposals']},{p['id'] for p in refresh_proposals(self.store,self.result_id,0)['proposals']})

    def test_manual_cell_binding_survives_structural_adoption(self):
        from ocr_workbench.geometry import bind_manual,geometry_view
        target={'kind':'cell','table':0,'row':1,'column':1}
        polygon=box_polygon([100,40,200,80])
        bind_manual(self.store,self.result_id,{'revision':0,'version_id':self.version['id'],'target':target,'polygon':polygon})
        p=self.proposal();saved=decide_structure(self.store,self.result_id,p['id'],self.body(p))
        evidence=geometry_view(self.store,self.result_id,target)['evidence']
        self.assertTrue(any(e['source']=='manual' and e['polygon']==polygon for e in evidence))

    def test_legacy_candidate_import_requires_original_image_and_artifact_hash(self):
        import hashlib
        from ocr_workbench.store import encoded,uid,now
        from ocr_workbench.document_store import structure_fingerprint
        pred={**self.pred,'image_version':self.version['id'],'image_sha256':self.version['sha256'],'ocr_blocks':self.blocks}
        artifact=self.store.root/'legacy.json';artifact.write_text(encoded(pred),encoding='utf-8')
        details={'artifact':'legacy.json','artifact_sha256':hashlib.sha256(artifact.read_bytes()).hexdigest()}
        with self.store.transaction() as db:
            db.execute('DELETE FROM structure_candidates')
            db.execute('INSERT INTO geometry_evidence VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (uid(),self.result_id,self.version['id'],None,self.version['sha256'],structure_fingerprint(self.store.result(self.result_id)['edited']),
                 'paddle-table-local-v2','test',encoded({'kind':'table','table':0}),None,encoded(details),'valid',now()))
        original=artifact.read_bytes();artifact.write_bytes(original+b' ')
        self.assertEqual(refresh_proposals(self.store,self.result_id,0)['candidates'],[])
        artifact.write_bytes(original)
        self.assertEqual(len(refresh_proposals(self.store,self.result_id,0)['candidates']),1)

    def test_structure_routes_require_authentication_and_revision(self):
        import threading,time,urllib.request,urllib.error
        import uvicorn
        from ocr_workbench.service import create_app
        root=Path(__file__).resolve().parents[1]
        app=create_app(root,self.store.root,'structure-test-token',start_queue=False)
        server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=0,log_level='error'))
        thread=threading.Thread(target=server.run);thread.start()
        try:
            deadline=time.monotonic()+10
            while not server.started:
                if time.monotonic()>deadline:raise TimeoutError('Test service failed to start')
                time.sleep(.01)
            url=f'http://127.0.0.1:{server.servers[0].sockets[0].getsockname()[1]}/api/results/{self.result_id}/structure/check'
            def post(revision,authenticated=True):
                headers={'Content-Type':'application/json'}
                if authenticated:headers['Authorization']='Bearer structure-test-token'
                request=urllib.request.Request(url,data=json.dumps({'revision':revision}).encode(),headers=headers)
                opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
                try:
                    with opener.open(request,timeout=10) as response:return response.status,json.load(response)
                except urllib.error.HTTPError as error:return error.code,None
            self.assertEqual(post(0,False)[0],401)
            self.assertEqual(post(99)[0],409)
            status,value=post(0)
            self.assertEqual(status,200);self.assertTrue(value['proposals'])
        finally:
            server.should_exit=True;thread.join(timeout=15)
        self.assertFalse(thread.is_alive())

    def test_native_adoption_retains_exact_units_without_duplicate_pdf_text(self):
        from ocr_workbench.store import encoded,history_encoded
        from ocr_workbench.structure_export import native_structure_units
        from ocr_workbench.pdf_export import positioned_content
        pred,blocks = prediction([['Account','Amount'],['Revenue','001.00']])
        pred['tf_table_cells'][0]['label'] = 'ched'
        text = 'OCR artifact\n'
        native=[{'id':'overlap-ocr','text':'OCR artifact','source':'ppocr','source_kind':'scan',
                 'kind':'text','polygon':blocks[0]['polygon'],'text_range':[0,len(text)-1]}]
        for block in blocks:
            start=len(text);text+=block['text']+'\n'
            native.append({**block,'source':'pdf-native','source_kind':'native','text_range':[start,len(text)-1]})
        raw={'origin':'document','engine':'pdf-native','text':text,'tables':[],'blocks':native,'document':{},
             'project_image_version':self.version['id']}
        edit={'text':text,'tables':[]}
        with self.store.transaction() as db:
            db.execute('DELETE FROM structure_candidates')
            db.execute('UPDATE results SET original=?,edited=? WHERE id=?',(encoded(raw),encoded(edit),self.result_id))
            db.execute('UPDATE edits SET value=? WHERE result_id=?',(history_encoded(edit),self.result_id))
            record_candidates(db,self.result_id,self.version,pred,native)
        p=next(p for p in refresh_proposals(self.store,self.result_id,0)['proposals'] if p['kind']=='native_table')
        saved=decide_structure(self.store,self.result_id,p['id'],self.body(p))
        self.assertTrue(saved['edited']['tables'][0]['cells'][0]['is_header'])
        with self.store.transaction() as db:
            units=native_structure_units(db,saved,self.version)
        positioned,missing,changed=positioned_content(saved,[],self.version,units)
        self.assertEqual(missing,[])
        self.assertFalse(changed)
        table_units=[u for u in positioned if u['target']['kind']=='cell']
        self.assertEqual([u['text'] for u in table_units],['Account','Amount','Revenue','001.00'])
        self.assertTrue(all(u['native_preserved'] for u in table_units))
        saved['edited']['tables'][0]['cells'][3]['text']='001.05'
        with self.store.transaction() as db:
            units=native_structure_units(db,saved,self.version)
        self.assertNotIn((0,1,1),units)
        self.assertTrue(positioned_content(saved,[],self.version,units)[1])

    def test_unknown_manual_value_blocks_instead_of_disappearing(self):
        edit = self.store.result(self.result_id)['edited']
        edit['tables'][0]['cells'][0]['text'] = '人工新增合同号0008765'
        edit['tables'][0]['source'] = None
        self.store.save(self.result_id,edit,0)
        p = self.proposal()
        self.assertFalse(p['can_apply'])
        self.assertTrue(any(c.get('text') == '人工新增合同号0008765' for c in p['conflicts']))

    def test_schema9_upgrade_backup_and_failure_rollback(self):
        from unittest.mock import patch
        import sqlite3
        from ocr_workbench.structure_store import migrate_v10
        root = Path(self.temp.name)/'old'
        with patch('ocr_workbench.store.SCHEMA_VERSION',9):
            Store(root)
        def broken(db):
            migrate_v10(db)
            raise OSError('migration interrupted')
        with patch('ocr_workbench.structure_store.migrate_v10',broken), self.assertRaises(OSError):
            Store(root)
        from contextlib import closing
        with closing(sqlite3.connect(root/'workbench.sqlite3')) as db:
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0],9)
            self.assertFalse(db.execute("SELECT 1 FROM sqlite_master WHERE name='structure_candidates'").fetchone())
        upgraded = Store(root)
        from ocr_workbench.store import SCHEMA_VERSION
        self.assertEqual(upgraded.rows('PRAGMA user_version')[0]['user_version'],SCHEMA_VERSION)
        self.assertTrue(list((root/'database-backups').glob('*.sqlite3')))


class GroupTests(unittest.TestCase):
    def test_group_replacement_keeps_notes_and_removes_old_tables(self):
        from ocr_workbench.structure_groups import replace_group
        original = table([['a','b'],['c','d']])
        edit = {'text':'Before\n'+tables_html([original])+'\nAfter','tables':[original]}
        one,two = table([['a','b']]),table([['c','d']])
        split = replace_group(edit,[0],[one,two])
        self.assertEqual(len(parse_tables(split['text'])),2)
        self.assertTrue(split['text'].startswith('Before\n'))
        self.assertTrue(split['text'].endswith('\nAfter'))
        merged = replace_group(split,[0,1],[original])
        self.assertEqual(len(parse_tables(merged['text'])),1)
        split['text'] = split['text'].replace('</table>\n\n<table>','</table>\nINTERVENING NOTE\n<table>')
        for t,p in zip(split['tables'],parse_tables(split['text'])):
            t['source'] = p['source']
        with self.assertRaises(ValueError):
            replace_group(split,[0,1],[original])

    def test_local_insert_is_proposed_but_split_needs_adopted_fragments(self):
        before = table([['Account','Amount'],['Revenue','001.00']])
        pred,blocks = prediction([['Account','Amount'],['Revenue','001.00'],['Tax','-0.01']])
        after = prepare_candidates(pred,blocks,source_result='ocr',image_version='v',width=200,height=120)['tables'][0]['skeleton']
        self.assertEqual(local_variants(before,after)[0]['kind'],'insert_rows')
        merged = deepcopy(before)
        merged['cells'][0]['column_span'] = 2
        merged['cells'][0]['text'] = 'Account Amount'
        merged['cells'].pop(1)
        split,conflicts,_ = preserve_values(merged,{'rows':2,'columns':2,'cells':after['cells'][:4]},merged,
            prepare_candidates(pred,blocks,source_result='ocr',image_version='v',width=200,height=120)['tokens'])
        self.assertTrue(any(c['kind']=='current_value_unmapped' for c in conflicts))
        self.assertEqual([c['text'] for c in split['cells'][:2]],['Account','Amount'])


if __name__ == '__main__':
    unittest.main()
