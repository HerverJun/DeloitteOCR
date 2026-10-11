"""Batch identity, independent XLSX coverage, failures and revision protection."""
from copy import deepcopy
import io
import json
import shutil
from pathlib import Path
import tempfile
import threading
import unittest
import zipfile

from PIL import Image
from openpyxl import load_workbook

from ocr_workbench.imaging import add_image
from ocr_workbench.screen_batches import ScreenBatches, fingerprint, validate_policy
from ocr_workbench.store import Conflict, Store


TABLE = {'rows': 3, 'columns': 2, 'caption': '单位：元', 'cells': [
    {'row': r, 'column': c, 'row_span': 1, 'column_span': 1, 'text': text, 'is_header': r == 0}
    for r, row in enumerate([['编号', '金额'], ['00123456789012345678', '123.45'], ['00002', '=1+1']])
    for c, text in enumerate(row)]}


class Queue:
    def __init__(self, store):
        self.store, self.wake = store, threading.Event()

    def action(self, project, action, task_ids):
        status = {'pause': 'paused', 'resume': 'queued', 'cancel': 'cancelled'}[action]
        with self.store.transaction() as db:
            for key in task_ids:
                db.execute("UPDATE tasks SET status=? WHERE id=? AND project_id=? AND status NOT IN ('succeeded','failed','cancelled')", (status, key, project))


def prepare(source, folder, **kwargs):
    target = Path(folder) / 'primary.png'
    with Image.open(source) as image:
        image.save(target)
        w, h = image.size
    matrix = [1, 0, 0, 0, 1, 0, 0, 0, 1]
    return {'width': w, 'height': h, 'primary_id': 'primary',
            'quality': {'flags': [], 'unreadable': False}, 'business_regions': [],
            'prepared_images': [{'id': 'primary', 'path': str(target), 'kind': 'original',
                                 'width': w, 'height': h, 'prepared_to_source': matrix}]}


def select(candidates, *args, **kwargs):
    selected = candidates[0]
    return {'selected_id': selected['id'], 'selected_result': deepcopy(selected['raw']),
            'automatic_delivery': True, 'review_required': False,
            'assessments': [{'id': selected['id'], 'issues': []}], 'conflicts': [], 'modifications': [], 'upgrade_kind': None}


class ScreenBatchTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='native-screen-test-', dir='F:/OCR/output/ocr-screen-benchmark-20261007/temp')
        self.root = Path(self.temporary.name)
        self.store = Store(self.root / 'workspace')
        self.project = self.store.project('屏拍批次')['id']
        self.queue = Queue(self.store)
        self.manager = ScreenBatches(self.store, self.queue, self.root, prepare=prepare, select=select)
        self.measured_defaults = self.manager.defaults
        # Existing engineering scenarios intentionally exercise the legacy
        # route. Measured-default tests opt into the real frozen config below.
        self.manager.defaults = lambda: {}

    def tearDown(self):
        self.manager.stop()
        self.temporary.cleanup()

    def images(self, count):
        versions = []
        for i in range(count):
            source = self.root / f'input-{i}.png'
            Image.new('RGB', (100, 70), (i % 255, 255, 255)).save(source)
            image = add_image(self.store, self.project, '同名.png', source)
            versions.append(image['active_version'])
        return versions

    def submit(self, versions, **policy):
        return self.manager.submit(self.project, {'version_ids': versions, 'request_id': 'request-screen-1',
            'policy': {'primary_engine': 'glm', 'max_upgrades': 0, **policy}})['batch_id']

    def finish(self, *, failure=False):
        while task := self.store.claim():
            if failure:
                with self.store.transaction() as db:
                    db.execute("UPDATE tasks SET status='failed',error='model unavailable' WHERE id=?", (task['id'],))
            else:
                self.store.complete(task['id'], {'status': 'success', 'text': '业务数据', 'tables': [deepcopy(TABLE)],
                    'blocks': [], 'engine': task['engine'], 'elapsed_seconds': .1, 'project_image_version': task['version_id']})

    def test_240_inputs_yield_240_independent_excel_files_and_stable_names(self):
        versions = self.images(240)
        batch = self.submit(versions)
        for _ in range(10):
            self.manager.step()
            self.finish()
        self.manager.step()
        state = self.manager.get(self.project, batch)
        self.assertEqual(state['counts'], {'total': 240, 'completed': 240, 'needs_review': 0, 'failed': 0, 'cancelled': 0, 'pending': 0})
        self.assertEqual(len(state['items']), 30)
        path, manifest = self.manager.export(self.project, batch)
        self.assertEqual(manifest['success_files'], 240)
        with zipfile.ZipFile(path) as archive:
            xlsx = [name for name in archive.namelist() if name.endswith('.xlsx')]
            self.assertEqual(len(xlsx), 240)
            book = load_workbook(io.BytesIO(archive.read(xlsx[0])))
            sheet = book.worksheets[0]
            self.assertEqual(sheet['A3'].value, '00123456789012345678')
            self.assertEqual(sheet['B4'].value, '=1+1')
            self.assertEqual(sheet['B4'].data_type, 's')
        self.assertEqual({x['input_version_id'] for x in manifest['items']}, set(versions))
        # Synthetic OCR payloads verify engineering/export only, never OCR accuracy.

    def test_submission_idempotency_and_conflicting_scope(self):
        versions = self.images(2)
        first = self.submit(versions)
        self.assertEqual(first, self.submit(versions))
        with self.assertRaises(Conflict):
            self.submit(versions[:1])
        self.assertEqual(self.manager.lookup(self.project, 'request-screen-1')['batch_id'], first)

    def test_math_display_export_preserves_engine_ledger_and_pending_review(self):
        version = self.images(1)[0]
        batch = self.submit([version], render_math_text=True, require_human_review=True)
        self.manager.step()
        task = self.store.claim()
        table = deepcopy(TABLE)
        table['cells'][3]['text'] = r'-1.00$^{\ast}$'
        payload = {'status': 'success', 'text': r'original model: -1.00$^{\ast}$',
                   'tables': [table], 'blocks': [], 'engine': task['engine'],
                   'project_image_version': version}
        self.store.complete(task['id'], payload)
        engine_result_id = self.store.one('tasks', task['id'])['result_id']
        engine_before = deepcopy(self.store.result(engine_result_id))
        self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.assertEqual(detail['status'], 'needs_review')
        original = detail['result']['original']
        self.assertEqual(original['text'], payload['text'])
        self.assertEqual(original['tables'][0]['cells'][3]['text'], '-1.00*')
        change = original['screen']['table_text_rendering']['changes'][0]
        self.assertEqual((change['before'], change['after']), (r'-1.00$^{\ast}$', '-1.00*'))
        self.assertEqual(original['screen']['sources'][0]['result_id'], engine_result_id)
        self.assertEqual(self.store.result(engine_result_id), engine_before)
        path, manifest = self.manager.export(self.project, batch, {'include_pending_review': True})
        self.assertEqual(manifest['success_files'], 1)
        with zipfile.ZipFile(path) as archive:
            book = load_workbook(io.BytesIO(archive.read(next(n for n in archive.namelist() if n.endswith('.xlsx')))))
            self.assertEqual(book.worksheets[0]['B3'].value, '-1.00*')
            self.assertEqual(book.worksheets[0]['B4'].value, '=1+1')
            self.assertEqual(book.worksheets[0]['B4'].data_type, 's')

    def test_math_display_requires_explicit_boolean_and_legacy_opt_in(self):
        self.assertFalse(validate_policy({})['render_math_text'])
        for invalid in (0, 1, 'true'):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                validate_policy({'render_math_text': invalid})
        self.assertTrue(validate_policy({'render_math_text': True})['render_math_text'])

    def test_lost_submission_recovers_frozen_defaults_after_policy_update(self):
        versions = self.images(1)
        self.manager.defaults = lambda: {'primary_engine': 'hunyuan', 'max_upgrades': 0, 'group_size': 32,
                                        'strategy_version': 'measured-first'}
        body = {'version_ids': versions, 'request_id': 'request-default-recovery', 'policy': {'primary_engine': 'auto'}}
        first = self.manager.submit(self.project, body)
        self.manager.defaults = lambda: {'primary_engine': 'glm', 'max_upgrades': 2, 'group_size': 16,
                                        'strategy_version': 'measured-next'}
        recovered = self.manager.submit(self.project, body)
        self.assertEqual(recovered['batch_id'], first['batch_id'])
        self.assertEqual(recovered['policy'], first['policy'])
        with self.assertRaises(Conflict):
            self.manager.submit(self.project, {**body, 'policy': {'primary_engine': 'glm'}})
        with self.assertRaisesRegex(ValueError, '下载时'):
            self.manager.submit(self.project, {**body, 'request_id': 'other-export-options', 'export_options': {'numeric': True}})

    def test_failed_only_retry_preserves_prior_failure_and_completed_result(self):
        versions = self.images(2)
        batch = self.submit(versions)
        self.manager.step()
        tasks = self.store.rows('SELECT * FROM tasks ORDER BY created,id')
        success, failed = tasks
        with self.store.transaction() as db:
            db.execute("UPDATE tasks SET status='running' WHERE id=?", (success['id'],))
            db.execute("UPDATE tasks SET status='failed',error='unavailable' WHERE id=?", (failed['id'],))
        self.store.complete(success['id'], {'text': 'data', 'tables': [deepcopy(TABLE)], 'blocks': [], 'engine': 'glm'})
        self.manager.step()
        original = self.manager.detail(self.project, batch, success['version_id'])['result_id']
        self.manager.action(self.project, batch, {'action': 'retry_failed', 'request_id': 'retry-screen-1'})
        self.manager.step()
        self.finish()
        self.manager.step()
        self.assertEqual(self.manager.get(self.project, batch)['counts']['completed'], 2)
        self.assertEqual(self.store.one('tasks', failed['id'])['status'], 'failed')
        self.assertEqual(self.manager.detail(self.project, batch, success['version_id'])['result_id'], original)
        self.assertEqual(len(self.store.rows('SELECT * FROM screen_tasks')), 3)

    def test_pause_before_preparation_and_resume_has_no_duplicate_tasks(self):
        batch = self.submit(self.images(2))
        self.manager.action(self.project, batch, {'action': 'pause'})
        self.assertFalse(self.manager.step())
        self.manager.action(self.project, batch, {'action': 'resume'})
        self.manager.step()
        self.finish()
        self.manager.step()
        self.assertEqual(self.manager.get(self.project, batch)['status'], 'succeeded')
        self.assertEqual(len(self.store.rows('SELECT * FROM screen_tasks')), 2)

    def test_reopen_recovers_completed_tasks_without_new_recognition(self):
        batch = self.submit(self.images(1))
        self.manager.step()
        self.finish()
        reopened = ScreenBatches(Store(self.store.root), self.queue, self.root, prepare=prepare, select=select)
        reopened.step()
        self.assertEqual(reopened.get(self.project, batch)['status'], 'succeeded')
        self.assertEqual(len(self.store.rows('SELECT * FROM screen_tasks')), 1)
        reopened.stop()

    def test_stale_result_revision_cannot_be_exported_or_reviewed(self):
        version = self.images(1)[0]
        batch = self.submit([version])
        self.manager.step(); self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        value = deepcopy(detail['result']['edited'])
        value['tables'][0]['cells'][2]['text'] = 'changed'
        self.store.save(detail['result_id'], value, 0)
        with self.assertRaises(Conflict):
            self.manager.export(self.project, batch)
        with self.assertRaises(Conflict):
            self.manager.action(self.project, batch, {'action': 'review', 'version_ids': [version],
                'result_id': detail['result_id'], 'expected_revision': 0, 'state': 'accepted'})

    def test_cross_project_scope_and_file_traversal_rejected(self):
        versions = self.images(1)
        batch = self.submit(versions)
        other = self.store.project('其他项目')['id']
        with self.assertRaises(KeyError):
            self.manager.get(other, batch)
        with self.assertRaises(KeyError):
            self.manager.file(self.project, batch, versions[0], '../outside')

    def test_explicit_review_is_separate_from_automatic_delivery(self):
        version = self.images(1)[0]
        batch = self.submit([version])
        self.manager.step(); self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.assertFalse(detail['human_confirmed'])
        self.manager.action(self.project, batch, {'action': 'review', 'version_ids': [version],
            'result_id': detail['result_id'], 'expected_revision': 0, 'state': 'accepted'})
        self.assertTrue(self.manager.detail(self.project, batch, version)['human_confirmed'])

    def test_invalid_policy_and_no_table_never_make_blank_workbook(self):
        for policy in [{'max_upgrades': 3}, {'primary_engine': 'ppocr'}, {'budget_seconds': True}]:
            with self.assertRaises(ValueError):
                validate_policy(policy)
        batch = self.submit(self.images(1))
        with self.assertRaises(ValueError):
            self.manager.export(self.project, batch)

    def test_request_timeout_is_strict_and_frozen_in_new_batch(self):
        for value in [True, False, 9, 901, 180.0, '180']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                validate_policy({'request_timeout_seconds': value})
        self.assertEqual(validate_policy({})['request_timeout_seconds'], 180)
        for value in (10, 900):
            self.assertEqual(validate_policy({'request_timeout_seconds': value})['request_timeout_seconds'], value)
        batch = self.submit(self.images(1))
        self.assertEqual(self.manager.get(self.project, batch)['policy']['request_timeout_seconds'], 180)
        self.manager.step()
        metadata = json.loads(self.store.rows('SELECT metadata FROM screen_tasks')[0]['metadata'])
        self.assertEqual(metadata['request_timeout_seconds'], 180)

    def test_pre_cap_submission_recovers_original_frozen_policy(self):
        versions = self.images(1)
        batch = self.submit(versions)
        with self.store.transaction() as db:
            policy = json.loads(db.execute('SELECT policy FROM screen_batches WHERE id=?', (batch,)).fetchone()['policy'])
            policy.pop('request_timeout_seconds')
            policy.pop('engine_route')
            policy.pop('technical_fallback_reserve_seconds')
            signature = fingerprint({'version_ids': versions, 'policy': policy,
                                     'export_options': None, 'parent_batch_id': None})
            db.execute('UPDATE screen_batches SET policy=?,payload_hash=? WHERE id=?',
                       (json.dumps(policy), signature, batch))
        self.assertEqual(self.submit(versions), batch)
        self.assertNotIn('request_timeout_seconds', self.manager.get(self.project, batch)['policy'])
        self.assertNotIn('engine_route', self.manager.get(self.project, batch)['policy'])
        self.assertNotIn('technical_fallback_reserve_seconds', self.manager.get(self.project, batch)['policy'])
        with self.assertRaises(Conflict):
            self.submit(versions, request_timeout_seconds=180)

    def test_measured_policy_is_frozen_and_new_flags_are_strict(self):
        self.manager.defaults = self.measured_defaults
        version = self.images(1)[0]
        first = self.manager.submit(self.project, {'version_ids': [version],
            'request_id': 'measured-default-route', 'policy': {}})
        policy = first['policy']
        self.assertEqual(policy, validate_policy(self.measured_defaults()))
        self.assertEqual((policy['primary_engine'], policy['fallback_engines'], policy['max_upgrades']),
                         ('hunyuan', ['glm'], 1))
        self.assertEqual((policy['upgrade_trigger'], policy['require_human_review']), ('failure_only', True))
        self.assertEqual((policy['group_size'], policy['budget_seconds'], policy['request_timeout_seconds'],
                          policy['batch_budget_seconds']), (32, 360, 180, 36000))
        self.assertTrue(policy['layout_evidence'] and policy['perspective'])
        self.assertFalse(policy['allow_text_fixes'] or policy['visual_review'])
        for invalid in ('always', True, [], None):
            with self.subTest(trigger=invalid), self.assertRaises(ValueError):
                validate_policy({'upgrade_trigger': invalid})
        for invalid in (0, 1, 'true', None):
            with self.subTest(human=invalid), self.assertRaises(ValueError):
                validate_policy({'require_human_review': invalid})

    def test_failure_only_does_not_upgrade_a_quality_warning(self):
        def questionable(candidates, *args, **kwargs):
            value = select(candidates)
            value.update(automatic_delivery=False, review_required=True, upgrade_kind='low_quality')
            value['assessments'][0]['issues'] = [{'code': 'low_quality', 'message': '测试质量问题'}]
            return value
        self.manager.select_override = questionable
        version = self.images(1)[0]
        batch = self.submit([version], primary_engine='hunyuan', fallback_engines=['glm'], max_upgrades=1,
            upgrade_trigger='failure_only', require_human_review=True)
        self.manager.step(); self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.assertEqual((detail['status'], detail['stage']), ('needs_review', 'quality_review'))
        self.assertEqual({issue['category'] for issue in detail['issues'] if 'category' in issue},
                         {'quality', 'human_confirmation'})
        self.assertEqual([task['engine'] for task in self.store.rows(
            'SELECT t.engine FROM screen_tasks s JOIN tasks t ON t.id=s.task_id WHERE s.batch_id=?', (batch,))], ['hunyuan'])

    def test_initial_precision_checkbox_uses_glm_without_upgrading_only_for_human_review(self):
        self.manager.defaults = self.measured_defaults
        version = self.images(1)[0]
        body = {'version_ids': [version], 'request_id': 'initial-precision-checkbox',
                'policy': {'primary_engine': 'auto', 'mode': 'accuracy'}}
        first = self.manager.submit(self.project, body)
        policy = first['policy']
        self.assertEqual((policy['primary_engine'], policy['fallback_engines'], policy['max_upgrades']),
                         ('glm', ['hunyuan'], 1))
        self.assertEqual((policy['upgrade_trigger'], policy['require_human_review']), ('quality', True))
        self.assertEqual(policy['strategy_version'], 'screen-glm-quality-human-v1-20261007')
        self.assertEqual(self.manager.submit(self.project, body)['batch_id'], first['batch_id'])
        # The measured defaults first require independent layout context. This
        # deterministic payload tests routing only, not recognition quality.
        self.manager.step()
        context = self.store.claim()
        self.assertEqual(context['kind'], 'screen_context')
        raw = {'status': 'success', 'image': {'width': 100, 'height': 70},
               'regions': [{'box': [0, 0, 100, 70]}], 'ocr_blocks': []}
        with self.store.transaction() as db:
            db.execute("UPDATE tasks SET status='succeeded',finished=started WHERE id=?", (context['id'],))
            db.execute('UPDATE screen_contexts SET raw=? WHERE task_id=?', (json.dumps(raw), context['id']))
        self.manager.step(); self.finish(); self.manager.step()
        tasks = self.store.rows('''SELECT t.engine FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? ORDER BY s.pass_index''', (first['batch_id'],))
        self.assertEqual([task['engine'] for task in tasks], ['glm'])
        detail = self.manager.detail(self.project, first['batch_id'], version)
        self.assertEqual((detail['status'], detail['stage']), ('needs_review', 'human_confirmation'))

    def test_precision_resolution_preserves_explicit_and_legacy_policy_choices(self):
        defaults = self.measured_defaults()
        for value, expected in [
            ({'mode': 'accuracy'}, ('glm', 'quality')),
            ({'mode': 'accuracy', 'primary_engine': 'hunyuan'}, ('hunyuan', 'failure_only')),
            ({'mode': 'accuracy', 'primary_engine': 'auto', 'upgrade_trigger': 'failure_only'}, ('hunyuan', 'failure_only')),
        ]:
            with self.subTest(value=value):
                policy = validate_policy(value, defaults)
                self.assertEqual((policy['primary_engine'], policy['upgrade_trigger']), expected)
        legacy = {key: value for key, value in defaults.items()
                  if key not in {'upgrade_trigger', 'require_human_review'}}
        policy = validate_policy({'primary_engine': 'auto', 'mode': 'accuracy'}, legacy)
        self.assertEqual((policy['primary_engine'], policy['require_human_review']), ('hunyuan', False))
        frozen = validate_policy({'primary_engine': 'auto', 'mode': 'accuracy'}, defaults,
                                 resolve_precision=False)
        self.assertEqual((frozen['primary_engine'], frozen['upgrade_trigger']), ('hunyuan', 'failure_only'))

    def test_failure_only_uses_one_glm_fallback_after_a_technical_failure(self):
        version = self.images(1)[0]
        batch = self.submit([version], primary_engine='hunyuan', fallback_engines=['glm'], max_upgrades=1,
            upgrade_trigger='failure_only', require_human_review=True)
        self.manager.step(); self.finish(failure=True); self.manager.step()
        self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.assertEqual((detail['status'], detail['stage']), ('needs_review', 'human_confirmation'))
        self.assertIsNone(detail['error'])
        tasks = self.store.rows('''SELECT t.engine,t.status FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? ORDER BY s.pass_index''', (batch,))
        self.assertEqual([(task['engine'], task['status']) for task in tasks], [('hunyuan', 'failed'), ('glm', 'succeeded')])
        self.assertFalse(detail['human_confirmed'])

    def test_automatic_checks_still_require_explicit_human_confirmation(self):
        version = self.images(1)[0]
        batch = self.submit([version], upgrade_trigger='failure_only', require_human_review=True)
        self.manager.step(); self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.assertEqual((detail['status'], detail['stage']), ('needs_review', 'human_confirmation'))
        self.assertEqual([issue['code'] for issue in detail['issues']], ['human_confirmation_required'])
        self.assertFalse(detail['human_confirmed'])
        with self.assertRaises(ValueError):
            self.manager.export(self.project, batch)
        _, manifest = self.manager.export(self.project, batch, {'include_pending_review': True})
        self.assertEqual(manifest['items'][0]['status'], 'needs_review')
        self.assertFalse(manifest['items'][0]['human_confirmed'])
        self.manager.action(self.project, batch, {'action': 'review', 'version_ids': [version],
            'result_id': detail['result_id'], 'expected_revision': detail['result_revision'], 'state': 'accepted'})
        confirmed = self.manager.detail(self.project, batch, version)
        self.assertEqual(confirmed['status'], 'completed')
        self.assertTrue(confirmed['human_confirmed'])
        self.assertNotIn('human_confirmation_required', [issue['code'] for issue in confirmed['issues']])

    def test_old_persisted_policy_keeps_fingerprint_and_legacy_accuracy_pass(self):
        versions = self.images(1)
        body = {'version_ids': versions, 'request_id': 'old-persisted-policy', 'policy': {
            'primary_engine': 'hunyuan', 'fallback_engines': ['glm'], 'max_upgrades': 1, 'mode': 'accuracy'}}
        batch = self.manager.submit(self.project, body)['batch_id']
        with self.store.transaction() as db:
            previous = json.loads(db.execute('SELECT policy FROM screen_batches WHERE id=?', (batch,)).fetchone()['policy'])
            for key in ('upgrade_trigger', 'require_human_review'):
                previous.pop(key)
            signature = fingerprint({'version_ids': versions, 'policy': previous,
                'export_options': None, 'parent_batch_id': None})
            db.execute('UPDATE screen_batches SET policy=?,payload_hash=? WHERE id=?', (json.dumps(previous), signature, batch))
        self.manager.defaults = self.measured_defaults
        recovered = self.manager.submit(self.project, body)
        self.assertEqual(recovered['batch_id'], batch)
        self.assertEqual(recovered['policy'], previous)
        self.assertEqual(self.store.rows('SELECT payload_hash FROM screen_batches WHERE id=?', (batch,))[0]['payload_hash'], signature)
        with self.assertRaises(Conflict):
            self.manager.submit(self.project, {**body, 'policy': {**body['policy'], 'require_human_review': False}})
        self.manager.step(); self.finish(); self.manager.step()
        self.assertEqual([task['engine'] for task in self.store.rows('''SELECT t.engine FROM screen_tasks s
            JOIN tasks t ON t.id=s.task_id WHERE s.batch_id=? ORDER BY s.pass_index''', (batch,))], ['hunyuan', 'glm'])
        self.finish(); self.manager.step()
        self.assertEqual(self.manager.get(self.project, batch)['counts']['completed'], 1)

    def test_precision_subset_explicitly_uses_glm_quality_and_human_review(self):
        from unittest.mock import patch
        version = self.images(1)[0]
        parent = self.submit([version], primary_engine='hunyuan', fallback_engines=['glm'], max_upgrades=1,
            upgrade_trigger='failure_only', require_human_review=True)
        self.manager.step(); self.finish(); self.manager.step()
        # Controlled layout evidence exercises the legacy entry -> bounded
        # issue route. It is not a claim about this fixture's OCR accuracy.
        item = self.manager._items(parent)[0]
        preparation = json.loads(item['preparation'])
        preparation['business_regions'] = [{'id': 'controlled-grid', 'table_box': [0, 0, 100, 70],
            'sources': ['pixel_grid'], 'grid_reliable': True,
            'observed_grid_rows': TABLE['rows'], 'observed_grid_columns': TABLE['columns'] + 1}]
        self.manager._update(parent, version, preparation=json.dumps(preparation))
        before = self.store.result(item['result_id'])
        command = {'action': 'upgrade', 'version_ids': [version], 'request_id': 'manual-glm-precision'}
        with patch('ocr_workbench.screen_resolution.precision_capability',
                   return_value={'available': True, 'engine': 'glm', 'basis': 'controlled-test-only'}):
            child = self.manager.action(self.project, parent, command)
            self.assertEqual(self.manager.action(self.project, parent, command)['batch_id'], child['batch_id'])
        self.assertEqual(child['status'], 'submitted')
        self.assertEqual((child['policy']['primary_engine'], child['policy']['max_upgrades'],
                          child['policy']['require_human_review']), ('glm', 0, True))
        self.assertEqual(child['policy']['fallback_engines'], [])
        self.assertLessEqual(child['budget']['reserved_seconds'], 180)
        self.assertFalse(child['selection_applied'])
        self.assertFalse(child['values_modified'])
        self.assertEqual(child['human_review'], 'not_performed')
        self.assertEqual(child['parent_batch_id'], parent)
        self.assertEqual(self.store.result(before['id']), before)
        self.manager.step(); self.finish(); self.manager.step()
        tasks = self.store.rows('''SELECT t.engine FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? ORDER BY s.pass_index''', (child['batch_id'],))
        self.assertEqual([task['engine'] for task in tasks], ['glm'])
        self.assertEqual(self.manager.get(self.project, child['batch_id'])['counts']['needs_review'], 1)
        self.assertEqual(self.manager.get(self.project, parent)['counts']['needs_review'], 1)

    def test_selected_precision_attempt_has_lineage_and_is_idempotent(self):
        from unittest.mock import patch
        versions = self.images(2)
        batch = self.submit(versions, primary_engine='hunyuan')
        self.manager.step(); self.finish(); self.manager.step()
        item = next(item for item in self.manager._items(batch) if item['input_version_id'] == versions[1])
        preparation = json.loads(item['preparation'])
        preparation['business_regions'] = [{'id': 'controlled-grid', 'table_box': [0, 0, 100, 70],
            'sources': ['pixel_grid'], 'grid_reliable': True,
            'observed_grid_rows': TABLE['rows'], 'observed_grid_columns': TABLE['columns'] + 1}]
        self.manager._update(batch, versions[1], preparation=json.dumps(preparation))
        before = self.store.result(item['result_id'])
        command = {'action': 'upgrade', 'version_ids': [versions[1]], 'request_id': 'precision-only-one'}
        with patch('ocr_workbench.screen_resolution.precision_capability',
                   return_value={'available': True, 'engine': 'glm', 'basis': 'controlled-test-only'}):
            improved = self.manager.action(self.project, batch, command)
            replay = self.manager.action(self.project, batch, command, allow_new=False)
        self.assertEqual(improved['status'], 'submitted')
        self.assertEqual(improved['parent_batch_id'], batch)
        self.assertEqual(improved['attempt']['counts']['total'], 1)
        self.assertEqual([entry['input_version_id'] for entry in improved['bindings']], [versions[1]])
        self.assertEqual(improved['policy']['mode'], 'accuracy')
        self.assertNotEqual(improved['batch_id'], batch)
        self.assertEqual(replay['batch_id'], improved['batch_id'])
        self.assertEqual(self.store.result(before['id']), before)
        self.assertEqual(len(self.store.rows('SELECT id FROM screen_batches WHERE parent_batch_id=?', (batch,))), 1)
        self.assertEqual(self.manager.get(self.project, batch)['counts']['completed'], 2)

    def test_pending_filter_and_collection_preserve_all_inputs(self):
        versions = self.images(3)
        batch = self.submit(versions)
        self.manager.step()
        self.assertEqual(self.manager.get(self.project, batch, status='pending')['total_items'], 3)
        self.assertEqual(self.manager.list(self.project)['batches'][0]['batch_id'], batch)
        with self.assertRaises(Conflict):
            self.manager.action(self.project, batch, {'action': 'review', 'version_ids': [versions[0]], 'state': 'accepted'})

    def test_exact_content_cache_reuses_inference_but_delivers_each_input(self):
        version = self.images(1)[0]
        batch = self.submit([version])
        self.manager.step(); self.finish(); self.manager.step()
        duplicate = self.root / 'duplicate.png'
        shutil.copyfile(self.store.file(self.store.one('versions', version)['path']), duplicate)
        image = add_image(self.store, self.project, '重复.png', duplicate)
        second = self.manager.submit(self.project, {'version_ids': [image['active_version']], 'request_id': 'duplicate-content',
            'policy': {'primary_engine': 'glm', 'max_upgrades': 0}})['batch_id']
        self.manager.step()
        self.assertEqual(self.manager.get(self.project, second)['counts']['completed'], 1)
        task = self.store.rows("SELECT * FROM tasks WHERE kind='screen_cached'")[0]
        self.assertEqual(task['input_version_id'], image['active_version'])
        self.assertNotEqual(self.manager.detail(self.project, second, image['active_version'])['result_id'],
                            self.manager.detail(self.project, batch, version)['result_id'])

    def test_numeric_export_is_per_image_and_preserves_identifiers(self):
        version = self.images(1)[0]
        batch = self.submit([version])
        self.manager.step(); self.finish(); self.manager.step()
        detail = self.manager.detail(self.project, batch, version)
        self.manager.action(self.project, batch, {'action': 'review', 'version_ids': [version], 'result_id': detail['result_id'],
            'expected_revision': detail['result_revision'], 'state': 'accepted'})
        path, _ = self.manager.export(self.project, batch, {'numeric_columns': {version: {'0': [0, 1]}}})
        with zipfile.ZipFile(path) as archive:
            name = next(n for n in archive.namelist() if n.endswith('.xlsx'))
            book = load_workbook(io.BytesIO(archive.read(name)))
            self.assertEqual(book.worksheets[0]['A3'].value, '00123456789012345678')
            self.assertEqual(book.worksheets[0]['B3'].value, 123.45)
            self.assertEqual(book.worksheets[0]['B4'].data_type, 's')
        with self.assertRaises(ValueError):
            self.manager.export(self.project, batch, {'numeric_columns': {'0': [1]}})

    def test_workbook_source_index_matches_actual_database_identity_for_every_table(self):
        versions = self.images(2)
        batch = self.submit(versions)
        self.manager.step()
        while task := self.store.claim():
            tables = [deepcopy(TABLE)]
            if task['version_id'] == versions[0]:
                tables.append(deepcopy(TABLE))
            self.store.complete(task['id'], {'status':'success','engine':'glm','text':'业务数据',
                'tables':tables,'blocks':[],'project_image_version':task['version_id']})
        self.manager.step()
        archive_path, manifest = self.manager.export(self.project, batch)
        identities = {row['version_id']: row for row in self.store.rows('''
            SELECT i.id image_id,i.name,v.id version_id,v.sha256,s.result_id,r.revision,
                   json_extract(r.original,'$.engine') engine
            FROM screen_items s JOIN images i ON i.id=s.image_id
            JOIN versions v ON v.id=s.input_version_id JOIN results r ON r.id=s.result_id
            WHERE s.batch_id=?''', (batch,))}
        with zipfile.ZipFile(archive_path) as archive:
            for item in manifest['items']:
                identity = identities[item['input_version_id']]
                book = load_workbook(io.BytesIO(archive.read(item['file_name'])))
                try:
                    index = book['来源索引']
                    headers = [cell.value for cell in index[1]]
                    data = [{headers[n]:cell.value for n,cell in enumerate(row)}
                            for row in index.iter_rows(min_row=2)]
                    result = self.store.result(identity['result_id'])
                    self.assertEqual(len(data),len(result['edited']['tables']))
                    for number, source in enumerate(data,1):
                        expected = {'工作簿':item['file_name'], '工作表':f'Table {number}',
                            '原图':identity['name'], '图片 ID':identity['image_id'],
                            '图像版本':identity['version_id'], '输入 SHA-256':identity['sha256'],
                            '引擎':identity['engine'], '结果 ID':identity['result_id'],
                            '校对 revision':str(identity['revision']), '表格序号':str(number)}
                        for label,value in expected.items():
                            self.assertEqual(source[label],value,label)
                            self.assertIsNotNone(source[label],label)
                        self.assertIn(source['工作表'],book.sheetnames)
                        self.assertEqual(book[source['工作表']]['A3'].value,'00123456789012345678')
                finally:
                    book.close()

    def test_context_stage_finishes_before_any_primary_task_and_recovers(self):
        versions = self.images(2)
        batch = self.submit(versions, layout_evidence=True)
        self.manager.step()
        tasks = self.store.rows('SELECT * FROM tasks')
        self.assertEqual({t['kind'] for t in tasks}, {'screen_context'})
        self.assertEqual(len(tasks), 2)
        raw = {'status': 'success', 'component': 'shared-geometry-context', 'image': {'width': 100, 'height': 70},
               'layout_blocks': [], 'regions': [], 'ocr_blocks': [], 'ocr_source': 'PP-OCRv6-auxiliary'}
        with self.store.transaction() as db:
            for task in tasks:
                db.execute("UPDATE tasks SET status='succeeded' WHERE id=?", (task['id'],))
                db.execute('UPDATE screen_contexts SET raw=? WHERE task_id=?', (json.dumps(raw), task['id']))
        self.manager.step(); self.finish(); self.manager.step()
        self.assertEqual(self.manager.get(self.project, batch)['counts']['completed'], 2)
        self.assertEqual(len(self.store.rows('SELECT * FROM screen_contexts')), 2)

    def test_expired_batch_budget_preserves_completed_and_stops_pending(self):
        versions = self.images(2)
        batch = self.submit(versions)
        self.manager.step()
        task = self.store.claim()
        self.store.complete(task['id'], {'text': 'data', 'tables': [deepcopy(TABLE)], 'blocks': [], 'engine': 'glm'})
        self.manager.step()
        with self.store.transaction() as db:
            db.execute('UPDATE screen_batches SET deadline=0.1 WHERE id=?', (batch,))
        self.manager.step()
        state = self.manager.get(self.project, batch)
        self.assertEqual(state['counts']['completed'], 1)
        self.assertEqual(state['counts']['failed'], 1)
        self.assertEqual(state['counts']['pending'], 0)

    def test_v18_cannot_be_opened_by_v17_without_modifying_database(self):
        import hashlib
        from unittest.mock import patch
        self.manager.stop()
        database = self.store.root / 'workbench.sqlite3'
        before = hashlib.sha256(database.read_bytes()).hexdigest()
        with patch('ocr_workbench.store.SCHEMA_VERSION', 17), self.assertRaises(ValueError):
            Store(self.store.root)
        self.assertEqual(hashlib.sha256(database.read_bytes()).hexdigest(), before)

    def test_project_cleanup_covers_batch_exports_and_cascades_state(self):
        from ocr_workbench.maintenance import ProjectMaintenance
        version = self.images(1)[0]
        batch = self.submit([version])
        self.manager.step(); self.finish(); self.manager.step()
        self.manager.export(self.project, batch)
        maintenance = ProjectMaintenance(self.store, self.queue)
        self.assertIn(self.manager.root / batch, maintenance.paths(self.project))
        self.assertEqual(maintenance.safe(str((self.manager.root / batch).relative_to(self.store.root)), self.project), self.manager.root / batch)
        self.manager.stop()
        self.queue.status = lambda: {'task_id': None}
        maintenance.delete(self.project, '屏拍批次')
        self.assertFalse((self.manager.root / batch).exists())
        self.assertFalse(self.store.rows('SELECT * FROM screen_batches'))


if __name__ == '__main__':
    unittest.main()
