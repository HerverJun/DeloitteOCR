"""Lossless table cell text, spans and explicit Excel string values."""
from html.parser import HTMLParser
import re


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self.table = None
        self.cell = None
        self.row = -1
        self.col = 0
        self.occupied = set()
        self.in_caption = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'table':
            if self.table is not None:
                raise ValueError('Nested tables are not supported')
            self.table = {'cells': [], 'rows': 0, 'columns': 0, 'caption': ''}
            self.row, self.col, self.occupied = -1, 0, set()
        elif tag == 'caption' and self.table is not None:
            self.in_caption = True
        elif tag == 'tr' and self.table is not None:
            self.row += 1
            self.col = 0
        elif tag in {'td', 'th'} and self.table is not None:
            if self.cell is not None or self.row < 0:
                raise ValueError('Malformed table cell')
            while (self.row, self.col) in self.occupied:
                self.col += 1
            rs, cs = int(attrs.get('rowspan', '1')), int(attrs.get('colspan', '1'))
            if not (1 <= rs <= 1000 and 1 <= cs <= 1000):
                raise ValueError('Invalid table span')
            for r in range(self.row, self.row + rs):
                for c in range(self.col, self.col + cs):
                    if (r, c) in self.occupied:
                        raise ValueError('Overlapping table cells')
                    self.occupied.add((r, c))
            self.cell = {'row': self.row, 'column': self.col, 'row_span': rs,
                         'column_span': cs, 'text': '', 'confidence': None, 'polygon': None}
            self.table['rows'] = max(self.table['rows'], self.row + rs)
            self.table['columns'] = max(self.table['columns'], self.col + cs)
            self.col += cs
        elif tag == 'br' and self.cell is not None:
            self.cell['text'] += '\n'

    def handle_data(self, data):
        if self.cell is not None:
            self.cell['text'] += data
        elif self.in_caption:
            self.table['caption'] += data

    def handle_endtag(self, tag):
        if tag == 'caption':
            self.in_caption = False
        elif tag in {'td', 'th'} and self.cell is not None:
            self.table['cells'].append(self.cell)
            self.cell = None
        elif tag == 'table' and self.table is not None:
            if self.cell is not None:
                raise ValueError('Unclosed table cell')
            self.tables.append(self.table)
            self.table = None


def parse_tables(text):
    parser = TableParser()
    parser.feed(text)
    parser.close()
    if parser.table is not None:
        raise ValueError('Truncated HTML table')
    if parser.tables:
        return parser.tables
    # Markdown has no span information; never infer merges.
    tables, rows = [], []
    for line in text.splitlines() + ['']:
        if line.strip().startswith('|') and line.strip().endswith('|'):
            fields = re.split(r'(?<!\\)\|', line.strip()[1:-1])
            if all(re.fullmatch(r'\s*:?-{3,}:?\s*', f) for f in fields):
                continue
            rows.append([f.strip().replace('\\|', '|') for f in fields])
        elif rows:
            tables.append({'rows': len(rows), 'columns': max(map(len, rows)),
                           'cells': [{'row': r, 'column': c, 'row_span': 1, 'column_span': 1,
                                      'text': value, 'confidence': None, 'polygon': None}
                                     for r, row in enumerate(rows) for c, value in enumerate(row)]})
            rows = []
    return tables


def export_xlsx(tables, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for index, table in enumerate(tables, 1):
        sheet = book.create_sheet(f'Table {index}')
        offset = 1 if table.get('caption') else 0
        if offset:
            sheet.cell(1,1).value = table['caption']
            sheet.cell(1,1).data_type = 's'
        for cell in table['cells']:
            row, col = cell['row'] + 1 + offset, cell['column'] + 1
            if len(cell['text']) > 32767:
                raise ValueError('Cell exceeds Excel text limit; raw JSON retains the full text')
            out = sheet.cell(row, col)
            out.value = cell['text']
            out.data_type = 's'
            out.number_format = '@'
            out.alignment = Alignment(wrap_text=True, vertical='top')
            if cell['row_span'] > 1 or cell['column_span'] > 1:
                sheet.merge_cells(start_row=row, start_column=col,
                                  end_row=row + cell['row_span'] - 1,
                                  end_column=col + cell['column_span'] - 1)
        for col in range(1, table['columns'] + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24
    if not tables:
        raise ValueError('No recognized tables; refusing to export an empty workbook')
    book.save(path)
