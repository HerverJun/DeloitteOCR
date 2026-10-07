"""Durable screen-table batches, with bounded GPU work and explicit delivery state."""
from __future__ import annotations

from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import threading
import time
import zipfile

from ocr_workbench.store import Conflict, encoded, now, uid

VERSION = 'screen-batch-v1'
TERMINAL = {'completed', 'needs_review', 'failed', 'cancelled'}
ENGINES = {'glm', 'hunyuan', 'paddlevl'}


def fingerprint(value):
    return hashlib.sha256(encoded(value).encode('utf-8')).hexdigest()


def screen_seconds_spent(store, batch_id, version_id, attempt=0):
    timestamp = now()
    rows = store.rows('''SELECT COALESCE(SUM(s.spent_seconds + MAX(0,(julianday(COALESCE(t.finished,?))-julianday(t.started))*86400)),0) seconds
        FROM tasks t JOIN screen_tasks s ON s.task_id=t.id
        WHERE s.batch_id=? AND s.input_version_id=? AND s.attempt=? AND t.started IS NOT NULL''',
        (timestamp, batch_id, version_id, attempt))
    spent = rows[0]['seconds']
    if attempt == 0:
        context = store.rows('''SELECT COALESCE(SUM(s.spent_seconds + MAX(0,(julianday(COALESCE(t.finished,?))-julianday(t.started))*86400)),0) seconds
            FROM tasks t JOIN screen_contexts s ON s.task_id=t.id
            WHERE s.batch_id=? AND s.input_version_id=? AND t.started IS NOT NULL''', (timestamp, batch_id, version_id))
        spent += context[0]['seconds']
    return spent


def validate_policy(value, defaults=None):
    value = value or {}
    if not isinstance(value, dict):
        raise ValueError('批次策略必须是对象')
    permitted = {'primary_engine', 'fallback_engines', 'max_upgrades', 'mode',
                 'group_size', 'budget_seconds', 'request_timeout_seconds', 'perspective', 'allow_text_fixes',
                 'visual_review', 'strategy_version', 'layout_evidence', 'batch_budget_seconds'}
    if set(value) - permitted:
        raise ValueError('未知的屏拍批次策略参数')
    result = {**(defaults or {}), **value}
    result.setdefault('primary_engine', 'glm')
    if result['primary_engine'] == 'auto':
        result['primary_engine'] = (defaults or {}).get('primary_engine', 'glm')
    result.setdefault('fallback_engines', ['paddlevl', 'hunyuan'])
    result.setdefault('max_upgrades', 2)
    result.setdefault('mode', 'auto')
    result.setdefault('group_size', 32)
    result.setdefault('budget_seconds', 300)
    result.setdefault('request_timeout_seconds', 180)
    result.setdefault('batch_budget_seconds', 36000)
    result.setdefault('perspective', True)
    result.setdefault('allow_text_fixes', False)
    result.setdefault('visual_review', False)
    result.setdefault('layout_evidence', False)
    result.setdefault('strategy_version', VERSION + '-provisional')
    if result['primary_engine'] not in ENGINES:
        raise ValueError('首轮模型必须具备表格结构输出能力')
    choices = result['fallback_engines']
    if (not isinstance(choices, list) or len(choices) > 3
            or len(set(choices)) != len(choices) or not set(choices) <= ENGINES):
        raise ValueError('升级模型清单无效')
    result['fallback_engines'] = [x for x in choices if x != result['primary_engine']]
    for key, lower, upper in [('max_upgrades', 0, 2), ('group_size', 1, 64),
                              ('budget_seconds', 10, 1800), ('request_timeout_seconds', 10, 900),
                              ('batch_budget_seconds', 10, 864000)]:
        if type(result[key]) is not int or not lower <= result[key] <= upper:
            raise ValueError('批次预算或分组大小无效')
    if result['mode'] not in {'auto', 'accuracy'}:
        raise ValueError('未知处理模式')
    for key in ('perspective', 'allow_text_fixes', 'visual_review', 'layout_evidence'):
        if type(result[key]) is not bool:
            raise ValueError('处理开关必须为布尔值')
    if result['visual_review']:
        raise ValueError('局部视觉审校仍在评测中，请通过复核工作流调用；批次自动审校尚未启用')
    if not isinstance(result['strategy_version'], str) or not 1 <= len(result['strategy_version']) <= 120:
        raise ValueError('策略版本无效')
    return result


class ScreenBatches:
    def __init__(self, store, queue, bundle, *, prepare=None, assess=None, select=None, hold_wal=True):
        self.store, self.queue, self.bundle = store, queue, Path(bundle)
        self.prepare_override, self.assess_override, self.select_override = prepare, assess, select
        self.guard = threading.RLock()
        self.stopping, self.wake = threading.Event(), threading.Event()
        self.thread = None
        self.worker_error = None
        self.wal_keeper = None
        config = self.bundle / 'config/engines.json'
        source = Path(__file__).resolve().parent
        self.model_identity = fingerprint({'configuration': hashlib.sha256(config.read_bytes()).hexdigest() if config.is_file() else None,
            'worker_source': [(p.name, hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(source.glob('*.py'))]})
        self.root = store.root / 'screen-batches'
        if self.root.absolute().drive.upper() != 'F:' or self.root.resolve().drive.upper() != 'F:':
            raise ValueError('屏拍批次数据必须保存于 F 盘')
        self.root.mkdir(exist_ok=True)
        from ocr_workbench.screen_schema import migrate_screen_schema
        with store.transaction() as db:
            migrate_screen_schema(db)
        if hold_wal:
            self._hold_wal()

    def _hold_wal(self):
        # Keep the WAL open without holding a transaction. FULL commits remain
        # durable; repeated short connections no longer checkpoint every write.
        if self.wal_keeper is None:
            self.wal_keeper = sqlite3.connect(self.store.root / 'workbench.sqlite3', check_same_thread=False)
            self.wal_keeper.execute('PRAGMA journal_mode=WAL').fetchall()
            self.wal_keeper.execute('SELECT count(*) FROM projects').fetchall()

    def defaults(self):
        path = Path(__file__).resolve().parents[2] / 'config/screen-policy.json'
        return json.loads(path.read_text('utf-8')) if path.is_file() else {}

    def start(self):
        if self.thread and self.thread.is_alive():
            return
        self.stopping.clear()
        self._hold_wal()
        self.thread = threading.Thread(target=self._run, name='screen-batch-coordinator', daemon=True)
        self.thread.start()

    def stop(self):
        self.stopping.set()
        self.wake.set()
        if self.thread:
            self.thread.join(timeout=15)
        with self.store.lock:
            if self.wal_keeper is not None:
                self.wal_keeper.close()
                self.wal_keeper = None

    def _run(self):
        while not self.stopping.is_set():
            try:
                worked = self.step()
                self.worker_error = None
            except Exception as error:
                import traceback
                self.worker_error = str(error)
                with (self.root / 'coordinator-errors.log').open('a', encoding='utf-8') as out:
                    out.write(traceback.format_exc() + '\n')
                worked = False
            if not worked:
                self.wake.wait(.35)
                self.wake.clear()

    def _batch(self, project_id, batch_id):
        records = self.store.rows('SELECT * FROM screen_batches WHERE id=?', (batch_id,))
        if not records:
            raise KeyError('批次不存在')
        row = records[0]
        if row['project_id'] != project_id:
            raise KeyError('批次不属于当前项目')
        row['policy'] = json.loads(row['policy'])
        return row

    def list(self, project_id):
        self.store.one('projects', project_id)
        summaries = []
        for row in self.store.rows('SELECT id FROM screen_batches WHERE project_id=? ORDER BY created DESC LIMIT 100', (project_id,)):
            summary = self.get(project_id, row['id'])
            summaries.append({key: value for key, value in summary.items() if key not in {'items', 'task_ids'}})
        return {'batches': summaries}

    def submit(self, project_id, body, *, parent_batch_id=None):
        if not isinstance(body, dict) or set(body) - {'version_ids', 'request_id', 'policy', 'export_options'}:
            raise ValueError('无效的屏拍批次请求')
        versions, request_id = body.get('version_ids'), body.get('request_id')
        if (not isinstance(versions, list) or not 1 <= len(versions) <= 1000
                or any(not isinstance(x, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,81}', x) for x in versions)
                or len(set(versions)) != len(versions)):
            raise ValueError('请选择 1–1000 个不同的图片版本')
        if not isinstance(request_id, str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,120}', request_id):
            raise ValueError('批次缺少有效请求编号')
        if body.get('export_options') is not None:
            raise ValueError('请在下载时明确选择导出选项；创建批次不接受导出参数')
        self.store.one('projects', project_id)
        with self.guard, self.store.file_lock, self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            previous = db.execute('SELECT * FROM screen_batches WHERE project_id=? AND request_id=?', (project_id, request_id)).fetchone()
            # Recover the original policy after defaults change for future batches.
            saved_policy = json.loads(previous['policy']) if previous else None
            policy = validate_policy(body.get('policy'), saved_policy if previous else self.defaults())
            if (previous and 'request_timeout_seconds' not in saved_policy
                    and 'request_timeout_seconds' not in (body.get('policy') or {})):
                policy.pop('request_timeout_seconds')
            signature = fingerprint({'version_ids': versions, 'policy': policy, 'export_options': None, 'parent_batch_id': parent_batch_id})
            if previous:
                if previous['payload_hash'] != signature:
                    raise Conflict('请求编号已绑定不同的批次范围或策略')
                batch_id = previous['id']
            else:
                batch_id = uid()
                db.execute('INSERT INTO screen_batches(id,project_id,request_id,payload_hash,policy,status,created,updated,parent_batch_id,deadline) VALUES(?,?,?,?,?,?,?,?,?,?)',
                           (batch_id, project_id, request_id, signature, encoded(policy), 'running', now(), now(), parent_batch_id, time.time() + policy['batch_budget_seconds']))
                for position, key in enumerate(versions):
                    version = db.execute('SELECT v.*,i.project_id,i.name,i.active_version FROM versions v JOIN images i ON i.id=v.image_id WHERE v.id=?', (key,)).fetchone()
                    if not version or version['project_id'] != project_id or version['active_version'] != key:
                        raise Conflict('图片范围或当前版本已变化，请重新选择')
                    db.execute('''INSERT INTO screen_items(batch_id,input_version_id,image_id,position,source_hash,source_name,status,stage,updated)
                                  VALUES(?,?,?,?,?,?,?,?,?)''',
                               (batch_id, key, version['image_id'], position, version['sha256'], version['name'], 'queued', 'preparing', now()))
        self.wake.set()
        return self.get(project_id, batch_id)

    def lookup(self, project_id, request_id):
        self.store.one('projects', project_id)
        rows = self.store.rows('SELECT id FROM screen_batches WHERE project_id=? AND request_id=?', (project_id, request_id))
        return {'submitted': bool(rows), 'batch_id': rows[0]['id'] if rows else None}

    def _items(self, batch_id):
        return self.store.rows('SELECT * FROM screen_items WHERE batch_id=? ORDER BY position', (batch_id,))

    def _update(self, batch_id, version_id, **changes):
        permitted = {'status', 'stage', 'preparation', 'candidates', 'issues', 'pass_index',
                     'result_id', 'result_revision', 'review_state', 'error', 'started', 'attempt', 'evidence'}
        if set(changes) - permitted:
            raise ValueError('无效内部批次状态')
        changes['updated'] = now()
        with self.store.transaction() as db:
            db.execute('UPDATE screen_items SET ' + ','.join(key + '=?' for key in changes)
                       + ' WHERE batch_id=? AND input_version_id=?', (*changes.values(), batch_id, version_id))

    def _prepare_item(self, batch, item):
        from ocr_workbench.screen_imaging import prepare_screen_image
        prepare = self.prepare_override or prepare_screen_image
        source = self.store.one('versions', item['input_version_id'])
        actual = hashlib.sha256(self.store.file(source['path']).read_bytes()).hexdigest()
        if actual != item['source_hash']:
            raise Conflict('输入图像内容已变化')
        folder = self.root / batch['id'] / item['input_version_id'] / 'prepared'
        folder.mkdir(parents=True, exist_ok=True)
        evidence = json.loads(item['evidence']) if item.get('evidence') else None
        tick = time.perf_counter()
        info = prepare(self.store.file(source['path']), folder, evidence=evidence, perspective=batch['policy']['perspective'])
        info['preparation_seconds'] = time.perf_counter() - tick
        info.setdefault('warnings', []).extend((evidence or {}).get('warnings', []))
        if any(w.get('code') == 'context_failed' for w in info['warnings']):
            info['quality']['flags'].append({'code': 'context_failed', 'severity': 'high', 'message': '布局或文字证据不可用，请核对业务范围及界面排除'})
        for prepared in info['prepared_images']:
            path = Path(prepared['path']).resolve(strict=True)
            if not path.is_relative_to(folder.resolve()):
                raise ValueError('预处理图像路径越界')
            prepared['path'] = str(path)
        self._update(batch['id'], item['input_version_id'], preparation=encoded(info), status='prepared', started=now())
        return self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?', (batch['id'], item['input_version_id']))[0]

    def _contexts(self, batch, items):
        """Complete the resident layout/text stage before switching to OCR."""
        from ocr_workbench.screen_evidence import normalize_screen_context
        worked = False
        with self.store.transaction() as db:
            for item in items:
                if item['status'] != 'queued' or item.get('evidence'):
                    continue
                old = db.execute('SELECT * FROM screen_contexts WHERE batch_id=? AND input_version_id=?', (batch['id'], item['input_version_id'])).fetchone()
                if old:
                    continue
                key = uid()
                source = db.execute('''SELECT c.task_id FROM screen_contexts c
                    JOIN tasks t ON t.id=c.task_id JOIN screen_items i
                    ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                    WHERE c.batch_id=? AND i.source_hash=? AND t.kind='screen_context'
                    AND t.status IN ('queued','running','succeeded') ORDER BY i.position LIMIT 1''',
                    (batch['id'], item['source_hash'])).fetchone()
                db.execute('''INSERT INTO tasks(id,project_id,image_id,version_id,input_version_id,engine,batch,status,phase,kind,created)
                              VALUES(?,?,?,?,?,?,?,?,?,?,?)''', (key, batch['project_id'], item['image_id'], item['input_version_id'],
                    item['input_version_id'], 'geometry', batch['id'] + '-stage-0-context', 'queued',
                    '等待相同图片定位结果' if source else '定位业务表区',
                    'screen_context_cached' if source else 'screen_context', now()))
                pending_raw = encoded({'cache_pending_source_task_id': source['task_id'],
                    'cache_key': fingerprint([item['source_hash'], self.model_identity, 'context'])}) if source else None
                db.execute('INSERT INTO screen_contexts(task_id,batch_id,input_version_id,raw) VALUES(?,?,?,?)',
                           (key, batch['id'], item['input_version_id'], pending_raw))
                worked = True
        worked = self._resolve_context_reuse(batch, items) or worked
        rows = self.store.rows('''SELECT c.*,t.status,t.error FROM screen_contexts c JOIN tasks t ON t.id=c.task_id WHERE c.batch_id=?''', (batch['id'],))
        pending = False
        for row in rows:
            item = next(i for i in items if i['input_version_id'] == row['input_version_id'])
            if item['status'] != 'queued' or item.get('evidence'):
                continue
            if row['status'] == 'succeeded' and row['raw']:
                context = json.loads(row['raw'])
                evidence = normalize_screen_context(context)
                if context.get('cache_hit'):
                    evidence['provenance']['cache_hit'] = context['cache_hit']
                self._update(batch['id'], row['input_version_id'], evidence=encoded(evidence))
                worked = True
            elif row['status'] in {'failed', 'cancelled'}:
                self._update(batch['id'], row['input_version_id'], evidence=encoded({'layout': [], 'text': [],
                    'warnings': [{'code': 'context_failed', 'message': row['error'] or row['status']}]}))
                worked = True
            elif row['status'] in {'interrupted', 'paused'}:
                self._update(batch['id'], row['input_version_id'], status='paused', stage='interrupted')
                worked = True
            else:
                pending = True
        self.queue.wake.set()
        return worked, pending

    def _resolve_context_reuse(self, batch, items):
        """Reuse only exact original bytes; every target keeps its own binding."""
        worked = False
        by_version = {item['input_version_id']: item for item in items}
        waiters = self.store.rows('''SELECT c.*,t.status FROM screen_contexts c JOIN tasks t ON t.id=c.task_id
            WHERE c.batch_id=? AND t.kind='screen_context_cached' AND t.status='queued' ORDER BY t.created,t.id''', (batch['id'],))
        for row in waiters:
            item = by_version[row['input_version_id']]
            pending = json.loads(row['raw'] or '{}')
            source_id = pending.get('cache_pending_source_task_id')
            sources = self.store.rows('''SELECT c.raw,c.artifact,c.input_version_id,t.status,t.id,i.source_hash
                FROM screen_contexts c JOIN tasks t ON t.id=c.task_id JOIN screen_items i
                ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                WHERE c.task_id=? AND c.batch_id=?''', (source_id, batch['id']))
            source = sources[0] if sources else None
            context = json.loads(source['raw']) if source and source['status']=='succeeded' and source['raw'] else None
            if source and source['status'] in {'queued','running'}:
                continue
            if (source and source['source_hash']==item['source_hash'] and context
                    and context.get('status')=='success'):
                version = self.store.one('versions', item['input_version_id'])
                if hashlib.sha256(self.store.file(version['path']).read_bytes()).hexdigest()!=item['source_hash']:
                    raise Conflict('重复图片的输入内容已变化')
                raw = deepcopy(context)
                raw.update(image_version=item['input_version_id'], inference_seconds=0, load_seconds=0,
                    screen_context_input={'image_version':item['input_version_id'],
                        'input_path':str(self.store.file(version['path'])), 'source_sha256':item['source_hash'],
                        'coordinate_space':'exif_normalized_original'},
                    cache_hit={'source_task_id':source['id'], 'source_image_version':source['input_version_id'],
                        'source_artifact':source['artifact'], 'source_raw_sha256':fingerprint(context),
                        'key':pending['cache_key'], 'reuse_scope':'same_batch_exact_input'})
                from ocr_workbench.atomic_files import write_json
                artifact = self.root/batch['id']/item['input_version_id']/'context-reuse.json'
                artifact.parent.mkdir(parents=True,exist_ok=True)
                write_json(artifact,raw,durable=True)
                with self.store.transaction() as db:
                    if db.execute("UPDATE tasks SET status='succeeded',phase='复用相同图片定位结果',finished=? WHERE id=? AND status='queued'",
                                  (now(),row['task_id'])).rowcount:
                        db.execute('UPDATE screen_contexts SET raw=?,artifact=? WHERE task_id=?',
                            (encoded(raw),str(artifact.relative_to(self.store.root)),row['task_id']))
                        worked = True
            else:
                # A failed/cancelled/individually paused leader must not strand
                # an unrelated active input or manufacture successful evidence.
                with self.store.transaction() as db:
                    replacement = db.execute('''SELECT c.task_id FROM screen_contexts c JOIN tasks t ON t.id=c.task_id
                        JOIN screen_items i ON i.batch_id=c.batch_id AND i.input_version_id=c.input_version_id
                        WHERE c.batch_id=? AND i.source_hash=? AND t.kind='screen_context'
                        AND t.status IN ('queued','running') AND t.id!=? ORDER BY i.position LIMIT 1''',
                        (batch['id'],item['source_hash'],row['task_id'])).fetchone()
                    if replacement:
                        pending['cache_pending_source_task_id']=replacement['task_id']
                        db.execute('UPDATE screen_contexts SET raw=? WHERE task_id=?',(encoded(pending),row['task_id']))
                    else:
                        db.execute("UPDATE tasks SET kind='screen_context',phase='定位业务表区' WHERE id=? AND status='queued'",(row['task_id'],))
                        db.execute('UPDATE screen_contexts SET raw=NULL,artifact=NULL WHERE task_id=?',(row['task_id'],))
                    worked = True
        return worked

    def _resolve_ocr_reuse(self, batch):
        """Resolve exact prepared-pixel dependencies without sharing source maps."""
        worked = False
        waiters = self.store.rows('''SELECT s.*,t.status,t.engine FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? AND t.kind='screen_cached' AND t.status IN ('queued','running') ORDER BY t.created,t.id''',(batch['id'],))
        for row in waiters:
            metadata = json.loads(row['metadata'])
            source_id = metadata.get('cache_pending_source_task_id')
            sources = self.store.rows('''SELECT t.id,t.status,t.result_id,t.input_version_id,s.metadata FROM tasks t
                JOIN screen_tasks s ON s.task_id=t.id WHERE t.id=?''',(source_id,))
            source = sources[0] if sources else None
            if source and source['status'] in {'queued','running'}:
                continue
            if (source and source['status']=='succeeded' and source['result_id']
                    and json.loads(source['metadata']).get('cache_key')==metadata.get('cache_key')):
                if (metadata.get('prepared_pixel_sha256') and hashlib.sha256(
                        self.store.file(row['prepared_path']).read_bytes()).hexdigest()!=metadata['prepared_pixel_sha256']):
                    with self.store.transaction() as db:
                        db.execute("UPDATE tasks SET status='failed',phase='输入已变化',error='重复复用的预处理像素已变化',finished=? WHERE id=? AND status IN ('queued','running')",
                                   (now(),row['task_id']))
                    worked = True
                    continue
                raw = deepcopy(self.store.result(source['result_id'])['original'])
                version = self.store.one('versions',row['input_version_id'])
                raw.update(project_image_version=row['input_version_id'], screen_input=metadata,
                    version_operations=version['operations'], elapsed_seconds=0,
                    cache_hit={'source_task_id':source['id'],'source_result_id':source['result_id'],
                        'source_image_version':source['input_version_id'],'key':metadata['cache_key']})
                with self.store.transaction() as db:
                    changed = db.execute("UPDATE tasks SET status='running',phase='复用相同图片结果' WHERE id=? AND status IN ('queued','running')",
                                         (row['task_id'],)).rowcount
                if changed:
                    worked = self.store.complete(row['task_id'],raw) or worked
            else:
                with self.store.transaction() as db:
                    replacement = db.execute('''SELECT s.task_id FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
                        WHERE s.batch_id=? AND json_extract(s.metadata,'$.cache_key')=?
                        AND t.kind!='screen_cached' AND t.status IN ('queued','running') AND t.id!=?
                        ORDER BY t.created,t.id LIMIT 1''',(batch['id'],metadata.get('cache_key'),row['task_id'])).fetchone()
                    if replacement:
                        metadata['cache_pending_source_task_id']=replacement['task_id']
                    else:
                        metadata.pop('cache_pending_source_task_id',None)
                        db.execute("UPDATE tasks SET kind='ocr',status='queued',phase='屏拍表格识别' WHERE id=? AND status IN ('queued','running')",(row['task_id'],))
                    db.execute('UPDATE screen_tasks SET metadata=? WHERE task_id=?',(encoded(metadata),row['task_id']))
                    worked = True
        if worked:
            self.queue.wake.set()
        return worked

    def _enqueue(self, batch, items, engine, pass_index):
        from ocr_workbench.screen_imaging import select_screen_first_pass
        with self.guard, self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            for item in items:
                info = json.loads(item['preparation'])
                chosen = select_screen_first_pass(info)
                if not chosen:
                    raise ValueError('预处理没有提供有效识别图像')
                for image in chosen:
                    prior = db.execute('SELECT task_id FROM screen_tasks WHERE batch_id=? AND input_version_id=? AND pass_index=? AND prepared_id=? AND attempt=?',
                                       (batch['id'], item['input_version_id'], pass_index, image['id'], item['attempt'])).fetchone()
                    if prior:
                        continue
                    task_id = self.store._enqueue(db, batch['project_id'], [item['input_version_id']], [engine],
                                                  request_id='screen-' + fingerprint([batch['id'], item['input_version_id'], pass_index, image['id'], item['attempt']])[:64])[0]
                    db.execute('UPDATE tasks SET batch=?,phase=? WHERE id=?',
                               (batch['id'] + '-stage-' + str(pass_index + 1) + '-' + engine, '屏拍表格识别', task_id))
                    pixel_hash = hashlib.sha256(Path(image['path']).read_bytes()).hexdigest()
                    metadata = {**image, 'input_version_id': item['input_version_id'], 'batch_id': batch['id'],
                                'prepared_pixel_sha256':pixel_hash,
                                'strategy_version': batch['policy']['strategy_version'],
                                'timeout_seconds': batch['policy']['budget_seconds'],
                                'request_timeout_seconds': batch['policy'].get('request_timeout_seconds', 900),
                                'attempt': item['attempt'],
                                'cache_key': fingerprint([pixel_hash,
                                                          engine, self.model_identity, batch['policy']['strategy_version']])}
                    path = str(Path(image['path']).relative_to(self.store.root))
                    db.execute('INSERT INTO screen_tasks(task_id,batch_id,input_version_id,pass_index,prepared_id,prepared_path,metadata,attempt) VALUES(?,?,?,?,?,?,?,?)',
                               (task_id, batch['id'], item['input_version_id'], pass_index, image['id'], path, encoded(metadata), item['attempt']))
                    hit = db.execute('SELECT r.task_id FROM screen_cache c JOIN results r ON r.id=c.result_id WHERE c.cache_key=?', (metadata['cache_key'],)).fetchone()
                    if not hit:
                        hit = db.execute('''SELECT s.task_id FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
                            WHERE s.batch_id=? AND json_extract(s.metadata,'$.cache_key')=?
                            AND t.id!=? AND t.kind!='screen_cached' AND t.status IN ('queued','running','succeeded')
                            ORDER BY t.created,t.id LIMIT 1''',(batch['id'],metadata['cache_key'],task_id)).fetchone()
                    if hit:
                        metadata['cache_pending_source_task_id']=hit['task_id']
                        db.execute('UPDATE screen_tasks SET metadata=? WHERE task_id=?',(encoded(metadata),task_id))
                        db.execute("UPDATE tasks SET kind='screen_cached',phase='等待相同图片结果' WHERE id=?", (task_id,))
                db.execute('UPDATE screen_items SET status=?,stage=?,pass_index=?,updated=? WHERE batch_id=? AND input_version_id=?',
                           ('running', 'first_pass' if pass_index == 0 else 'upgrading', pass_index, now(), batch['id'], item['input_version_id']))
        self._resolve_ocr_reuse(batch)
        self.queue.wake.set()

    def _candidate(self, batch, item, rows):
        from ocr_workbench.screen_quality import map_screen_result_to_source
        tables, texts, blocks, sources = [], [], [], []
        warnings, finishes = [], []
        truncated = False
        preparation = json.loads(item['preparation'])
        evidence = json.loads(item['evidence']) if item.get('evidence') else {}
        ancillary = []
        seen_ancillary = set()
        for region in preparation.get('business_regions', []):
            for part in region.get('ancillary', []):
                area = part['box']
                for token in evidence.get('text', []):
                    box = token.get('box')
                    text = token.get('text', '').strip()
                    if not box or not text:
                        continue
                    intersection = max(0, min(area[2], box[2])-max(area[0], box[0])) * max(0,min(area[3],box[3])-max(area[1],box[1]))
                    if intersection < .6 * (box[2]-box[0])*(box[3]-box[1]):
                        continue
                    key = (text, tuple(box))
                    if key in seen_ancillary:
                        continue
                    seen_ancillary.add(key)
                    ancillary.append({'kind': part['kind'], 'text': text, 'box': box,
                        'source': token.get('source'), 'confidence': token.get('confidence'), 'human_confirmed': False})
        elapsed = 0.0
        for row in rows:
            result = self.store.result(row['result_id'])
            raw = deepcopy(result['original'])
            metadata = json.loads(row['metadata'])
            if metadata.get('cache_key') and not raw.get('cache_hit'):
                with self.store.transaction() as db:
                    db.execute('INSERT OR IGNORE INTO screen_cache VALUES(?,?,?)', (metadata['cache_key'], result['id'], now()))
            raw['tables'] = deepcopy(result['edited'].get('tables', []))
            raw = map_screen_result_to_source(raw, metadata)
            tables.extend(raw.get('tables', []))
            texts.append(result['edited'].get('text', ''))
            blocks.extend(raw.get('blocks', []))
            elapsed += float(raw.get('elapsed_seconds', 0) or 0)
            warnings.extend(raw.get('warnings', []))
            finishes.append(raw.get('finish_reason'))
            truncated = truncated or bool(raw.get('truncated')) or raw.get('finish_reason') in {'length', 'max_tokens', 'token_limit'}
            sources.append({'result_id': result['id'], 'revision': result['revision'], 'engine': row['engine'],
                            'task_id': row['task_id'], 'prepared_id': row['prepared_id'],
                            'transform': metadata.get('prepared_to_source'), 'prepared_metadata': metadata,
                            'warnings': raw.get('warnings', []), 'finish_reason': raw.get('finish_reason'),
                            'cache_hit': deepcopy(raw.get('cache_hit'))})
        return {'id': fingerprint([x['result_id'] for x in sources]), 'engine': rows[0]['engine'],
                'raw': {'status': 'success', 'engine': rows[0]['engine'], 'text': '\n\n'.join(texts), 'tables': tables,
                        'blocks': blocks, 'warnings': warnings, 'truncated': truncated,
                        'finish_reason': 'length' if truncated else next((v for v in finishes if v), None),
                        'elapsed_seconds': elapsed, 'project_image_version': item['input_version_id'],
                        'origin': 'screen_batch', 'screen': {'version': VERSION, 'batch_id': batch['id'],
                        'input_sha256': item['source_hash'], 'strategy_version': batch['policy']['strategy_version'],
                        'sources': sources, 'prepared': preparation, 'business_metadata': ancillary,
                        'context_provenance': deepcopy(evidence.get('provenance'))}}}

    def _final_result(self, batch, item, raw):
        content = fingerprint(raw)
        with self.store.transaction() as db:
            db.execute('BEGIN IMMEDIATE')
            previous = db.execute('SELECT task_id FROM screen_final_tasks WHERE batch_id=? AND input_version_id=? AND content_hash=?',
                                  (batch['id'], item['input_version_id'], content)).fetchone()
            if previous:
                task_id = previous['task_id']
                task = db.execute('SELECT * FROM tasks WHERE id=?', (task_id,)).fetchone()
                if task['result_id']:
                    return task['result_id']
                db.execute("UPDATE tasks SET status='running',error=NULL WHERE id=?", (task_id,))
            else:
                task_id = uid()
                db.execute('''INSERT INTO tasks(id,project_id,image_id,version_id,input_version_id,engine,batch,status,phase,kind,created)
                              VALUES(?,?,?,?,?,?,?,?,?,?,?)''',
                           (task_id, batch['project_id'], item['image_id'], item['input_version_id'], item['input_version_id'],
                            raw.get('engine', batch['policy']['primary_engine']), batch['id'], 'running', '保存最终结果', 'screen_final', now()))
                db.execute('INSERT INTO screen_final_tasks VALUES(?,?,?,?)', (batch['id'], item['input_version_id'], content, task_id))
        if not self.store.complete(task_id, raw):
            raise Conflict('最终结果保存被中断，请恢复批次')
        result_id = self.store.one('tasks', task_id)['result_id']
        # Preserve manual selections/edits. Replace only untouched intermediate results of this batch.
        with self.store.transaction() as db:
            selected = db.execute('SELECT r.id,r.revision,r.task_id FROM selections s JOIN results r ON r.id=s.result_id WHERE s.image_id=?', (item['image_id'],)).fetchone()
            owned = selected and (db.execute('SELECT 1 FROM screen_tasks WHERE task_id=? AND batch_id=?', (selected['task_id'], batch['id'])).fetchone()
                                  or db.execute('SELECT 1 FROM screen_final_tasks WHERE task_id=? AND batch_id=?', (selected['task_id'], batch['id'])).fetchone())
            if owned and selected['revision'] == 0 and item['review_state'] != 'accepted':
                db.execute('UPDATE selections SET result_id=? WHERE image_id=?', (result_id, item['image_id']))
        return result_id

    def _collect(self, batch, item):
        from ocr_workbench.screen_quality import select_screen_candidates
        rows = self.store.rows('''SELECT s.*,t.status,t.engine,t.result_id,t.error FROM screen_tasks s JOIN tasks t ON t.id=s.task_id
            WHERE s.batch_id=? AND s.input_version_id=? AND s.pass_index=? AND s.attempt=? ORDER BY s.prepared_id''',
                               (batch['id'], item['input_version_id'], item['pass_index'], item['attempt']))
        if not rows or any(row['status'] not in {'succeeded', 'failed', 'cancelled', 'interrupted', 'paused'} for row in rows):
            return False
        if any(row['status'] in {'paused', 'interrupted'} for row in rows):
            self._update(batch['id'], item['input_version_id'], status='paused', stage='interrupted')
            return True
        candidates = json.loads(item['candidates'])
        spent = screen_seconds_spent(self.store, batch['id'], item['input_version_id'], item['attempt'])
        spent += json.loads(item['preparation']).get('preparation_seconds', 0) if item['preparation'] and item['attempt'] == 0 else 0
        successful = [row for row in rows if row['status'] == 'succeeded' and row['result_id']]
        if len(successful) == len(rows):
            candidate = self._candidate(batch, item, successful)
            if not any(c['id'] == candidate['id'] for c in candidates):
                candidates.append(candidate)
        if not candidates:
            if (item['pass_index'] < len(batch['policy']['fallback_engines']) and item['pass_index'] < batch['policy']['max_upgrades']
                    and spent < batch['policy']['budget_seconds']):
                self._update(batch['id'], item['input_version_id'], status='awaiting_upgrade', stage='technical_retry',
                             error='；'.join(str(r.get('error') or r['status']) for r in rows)[:500])
            else:
                self._update(batch['id'], item['input_version_id'], status='failed', stage='failed',
                             error='；'.join(str(r.get('error') or r['status']) for r in rows)[:500])
            return True
        select = self.select_override or select_screen_candidates
        evidence = json.loads(item['evidence']) if item.get('evidence') else None
        selected = select(candidates, json.loads(item['preparation']), text_evidence=(evidence or {}).get('text'), allow_text_fixes=batch['policy']['allow_text_fixes'])
        raw = selected['selected_result']
        raw.setdefault('screen', {})['selection'] = {key: selected.get(key) for key in
            ('selected_id', 'conflicts', 'modifications', 'automatic_delivery', 'review_required', 'upgrade_kind')}
        chosen_assessment = selected.get('selected_assessment') or next((a for a in selected.get('assessments', []) if a['id'] == selected.get('selected_id')), {})
        issues = list(chosen_assessment.get('issues', []))
        if selected.get('conflicts'):
            issues.extend(selected['conflicts'])
        result_id = self._final_result(batch, item, raw)
        needs_upgrade = (not selected.get('automatic_delivery', False) or batch['policy']['mode'] == 'accuracy' and item['pass_index'] == 0)
        can_upgrade = (item['pass_index'] < batch['policy']['max_upgrades']
                       and item['pass_index'] < len(batch['policy']['fallback_engines'])
                       and spent < batch['policy']['budget_seconds'])
        if needs_upgrade and can_upgrade:
            state, stage = 'awaiting_upgrade', selected.get('upgrade_kind') or 'accuracy_check'
        else:
            state = 'completed' if selected.get('automatic_delivery', False) else 'needs_review'
            stage = 'completed' if state == 'completed' else 'review'
            if needs_upgrade and not can_upgrade:
                issues.append({'code': 'upgrade_budget_reached', 'message': '已达到本图升级预算，保留最佳结果待核对'})
        self._update(batch['id'], item['input_version_id'], candidates=encoded(candidates), issues=encoded(issues),
                     result_id=result_id, result_revision=self.store.result(result_id)['revision'], status=state, stage=stage)
        return True

    def step(self):
        worked = False
        with self.guard, self.store.file_lock:
            batches = self.store.rows("SELECT id,project_id FROM screen_batches WHERE status='running' ORDER BY created")
            for record in batches:
                batch = self._batch(record['project_id'], record['id'])
                items = self._items(batch['id'])
                if batch.get('deadline') and time.time() >= batch['deadline']:
                    self.action(batch['project_id'], batch['id'], {'action': 'cancel'})
                    for item in items:
                        if item['status'] not in TERMINAL:
                            self._update(batch['id'], item['input_version_id'], status='needs_review' if item['result_id'] else 'failed',
                                         stage='batch_budget_reached', error='已达到整批时间预算，成果已保留')
                    self._reconcile(batch)
                    worked = True
                    continue
                worked = self._resolve_ocr_reuse(batch) or worked
                if batch['policy']['layout_evidence']:
                    changed, pending = self._contexts(batch, items)
                    worked = worked or changed
                    if pending:
                        continue
                    items = self._items(batch['id'])
                todo = [item for item in items if item['status'] == 'queued'][:batch['policy']['group_size']]
                if todo:
                    prepared = []
                    with ThreadPoolExecutor(max_workers=2) as pool:
                        futures = [(item, pool.submit(self._prepare_item, batch, item)) for item in todo]
                        for item, future in futures:
                            try:
                                prepared.append(future.result())
                            except Exception as error:
                                self._update(batch['id'], item['input_version_id'], status='failed', stage='preparation_failed', error=str(error)[:500])
                    if prepared:
                        self._enqueue(batch, prepared, batch['policy']['primary_engine'], 0)
                    worked = True
                # Prepared items can survive a crash between preparation and enqueue.
                prepared = [item for item in self._items(batch['id']) if item['status'] == 'prepared']
                if prepared:
                    self._enqueue(batch, prepared, batch['policy']['primary_engine'], 0)
                    worked = True
                for item in self._items(batch['id']):
                    if item['status'] == 'running':
                        worked = self._collect(batch, item) or worked
                current = self._items(batch['id'])
                # Finish the resident engine group before scheduling another model.
                active = any(i['status'] in {'queued', 'prepared', 'running'} for i in current)
                if not active:
                    upgrades = [i for i in current if i['status'] == 'awaiting_upgrade']
                    grouped = {}
                    for item in upgrades:
                        next_pass = item['pass_index'] + 1
                        choices = batch['policy']['fallback_engines']
                        if next_pass > len(choices):
                            self._update(batch['id'], item['input_version_id'], status='needs_review' if item['result_id'] else 'failed', stage='review')
                            continue
                        grouped.setdefault((choices[next_pass - 1], next_pass), []).append(item)
                    for (engine, index), group in grouped.items():
                        self._enqueue(batch, group, engine, index)
                        worked = True
                self._reconcile(batch)
        return worked

    def _reconcile(self, batch):
        states = Counter(item['status'] for item in self._items(batch['id']))
        if states and all(state in TERMINAL for state in states):
            status = ('succeeded' if states['completed'] == sum(states.values()) else
                      'cancelled' if states['cancelled'] == sum(states.values()) else
                      'failed' if states['failed'] == sum(states.values()) else 'partial')
            with self.store.transaction() as db:
                db.execute('UPDATE screen_batches SET status=?,updated=? WHERE id=?', (status, now(), batch['id']))

    def get(self, project_id, batch_id, *, offset=0, limit=30, status=None):
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 100:
            raise ValueError('批次分页范围无效')
        batch = self._batch(project_id, batch_id)
        items = self._items(batch_id)
        stale = False
        active_versions = {row['id']: row['active_version'] for row in self.store.rows('SELECT id,active_version FROM images WHERE project_id=?', (project_id,))}
        for item in items:
            if active_versions.get(item['image_id']) != item['input_version_id'] and item['status'] in {'completed', 'needs_review'}:
                stale = True
                item.update(status='needs_review', stage='input_changed',
                    issues=encoded([*json.loads(item['issues']), {'code': 'input_version_changed', 'severity': 'high', 'message': '原图版本已变化，旧结果不能交付'}]))
        counts = Counter(item['status'] for item in items)
        coverage = {key: counts[key] for key in ('completed', 'needs_review', 'failed', 'cancelled')}
        coverage.update(total=len(items), pending=len(items) - sum(coverage.values()))
        if status is not None and status not in TERMINAL | {'pending', 'queued', 'running', 'paused', 'awaiting_upgrade', 'prepared'}:
            raise ValueError('未知图片状态筛选')
        filtered = [item for item in items if not status or item['status'] == status or status == 'pending' and item['status'] not in TERMINAL]
        page = [self._display(item) for item in filtered[offset:offset + limit]]
        from ocr_workbench.screen_eta import estimate_screen_eta
        timed_tasks = self.store.rows('''SELECT t.id,t.kind,t.engine,t.status,t.started,t.finished,t.created,s.pass_index,s.spent_seconds
            FROM screen_tasks s JOIN tasks t ON t.id=s.task_id WHERE s.batch_id=?
            UNION ALL SELECT t.id,t.kind,t.engine,t.status,t.started,t.finished,t.created,-1,c.spent_seconds
            FROM screen_contexts c JOIN tasks t ON t.id=c.task_id WHERE c.batch_id=?''',(batch_id,batch_id))
        eta = estimate_screen_eta(timed_tasks,batch_status=batch['status'],pending_images=coverage['pending'])
        return {'batch_id': batch_id, 'project_id': project_id, 'status': 'partial' if stale and batch['status'] in {'succeeded', 'partial'} else batch['status'],
                'parent_batch_id': batch.get('parent_batch_id'),
                'counts': coverage, 'policy': batch['policy'], 'items': page, 'total_items': len(filtered), 'eta':eta,
                'next_offset': offset + limit if offset + limit < len(filtered) else None,
                'created': batch['created'], 'updated': batch['updated'], 'worker_error': self.worker_error,
                'task_ids': [r['task_id'] for r in self.store.rows('SELECT task_id FROM screen_tasks WHERE batch_id=?', (batch_id,))][:100]}

    @staticmethod
    def _display(item, *, detail=False):
        value = {key: item[key] for key in ('input_version_id', 'image_id', 'source_name', 'status', 'stage',
                 'pass_index', 'attempt', 'result_id', 'result_revision', 'review_state', 'error', 'updated')}
        value['issues'] = json.loads(item['issues'])
        value['candidates'] = [{'id': c['id'], 'engine': c['engine'], 'table_count': len(c['raw'].get('tables', []))}
                               for c in json.loads(item['candidates'])]
        value['human_confirmed'] = item['review_state'] == 'accepted'
        if detail:
            value['preparation'] = json.loads(item['preparation']) if item['preparation'] else None
        return value

    def detail(self, project_id, batch_id, version_id):
        self._batch(project_id, batch_id)
        rows = self.store.rows('SELECT * FROM screen_items WHERE batch_id=? AND input_version_id=?', (batch_id, version_id))
        if not rows:
            raise KeyError('批次没有该图片')
        value = self._display(rows[0], detail=True)
        value['result'] = self.store.result(rows[0]['result_id']) if rows[0]['result_id'] else None
        value['candidate_results'] = json.loads(rows[0]['candidates'])
        value['image_url'] = '/api/versions/' + version_id + '/image'
        return value

    def action(self, project_id, batch_id, body):
        if not isinstance(body, dict):
            raise ValueError('批次操作无效')
        action = body.get('action')
        if action not in {'pause', 'resume', 'cancel', 'retry_failed', 'review', 'upgrade'}:
            raise ValueError('未知批次操作')
        with self.guard, self.store.file_lock:
            batch = self._batch(project_id, batch_id)
            request_id = body.get('request_id')
            if request_id is not None:
                if not isinstance(request_id, str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,120}', request_id):
                    raise ValueError('操作请求编号无效')
                old = self.store.rows('SELECT * FROM screen_actions WHERE batch_id=? AND request_id=?', (batch_id, request_id))
                if old:
                    if old[0]['payload_hash'] != fingerprint(body):
                        raise Conflict('操作编号已用于不同的参数')
                    return json.loads(old[0]['receipt'])
            items = self._items(batch_id)
            selected = body.get('version_ids')
            if selected is not None:
                if not isinstance(selected, list) or not selected or len(set(selected)) != len(selected) or not set(selected) <= {i['input_version_id'] for i in items}:
                    raise ValueError('操作范围不属于该批次')
                items = [i for i in items if i['input_version_id'] in selected]
            if action == 'upgrade':
                if selected is None or not request_id:
                    raise ValueError('提高精度必须指定图片和操作请求编号')
                if any(i['status'] not in TERMINAL for i in items):
                    raise Conflict('请等待指定图片的当前尝试结束')
                policy = deepcopy(batch['policy'])
                engines = [policy['primary_engine'], *policy['fallback_engines']]
                if len(engines) > 1:
                    policy['primary_engine'], policy['fallback_engines'] = engines[1], [engines[0], *engines[2:]]
                policy.update(mode='accuracy', max_upgrades=min(2, len(policy['fallback_engines'])))
                receipt = self.submit(project_id, {'version_ids': [i['input_version_id'] for i in items],
                    'request_id': 'upgrade-' + fingerprint([batch_id, request_id])[:64], 'policy': policy}, parent_batch_id=batch_id)
                with self.store.transaction() as db:
                    db.execute('INSERT INTO screen_actions VALUES(?,?,?,?)', (batch_id, request_id, fingerprint(body), encoded(receipt)))
                return receipt
            if action == 'review':
                if len(items) != 1:
                    raise ValueError('单次复核必须指定一张图片')
                self._review(batch, items[0], body)
            elif action == 'retry_failed':
                for item in items:
                    if item['status'] != 'failed':
                        continue
                    # A technical retry gets a new pass identity, preserving every failed task.
                    self._update(batch_id, item['input_version_id'], status='prepared' if item['preparation'] else 'queued',
                                 stage='retry', pass_index=0, attempt=item['attempt'] + 1, error=None)
                with self.store.transaction() as db:
                    db.execute("UPDATE screen_batches SET status='running',updated=? WHERE id=?", (now(), batch_id))
            else:
                active = [i for i in items if i['status'] not in TERMINAL]
                task_ids = [r['task_id'] for i in active for r in self.store.rows('SELECT task_id FROM screen_tasks WHERE batch_id=? AND input_version_id=?', (batch_id, i['input_version_id']))]
                task_ids.extend(r['task_id'] for i in active for r in self.store.rows('SELECT task_id FROM screen_contexts WHERE batch_id=? AND input_version_id=?', (batch_id, i['input_version_id'])))
                if task_ids:
                    self.queue.action(project_id, action, task_ids)
                for item in active:
                    existing = self.store.rows('SELECT task_id FROM screen_tasks WHERE batch_id=? AND input_version_id=? AND attempt=?',
                                               (batch_id, item['input_version_id'], item['attempt']))
                    self._update(batch_id, item['input_version_id'], status='cancelled' if action == 'cancel' else
                                 'paused' if action == 'pause' else 'running' if existing else 'prepared' if item['preparation'] else 'queued',
                                 stage='cancelled' if action == 'cancel' else 'paused' if action == 'pause' else 'resuming')
                with self.store.transaction() as db:
                    if selected is None and action == 'pause':
                        db.execute('UPDATE screen_batches SET paused_at=COALESCE(paused_at,?) WHERE id=?', (time.time(), batch_id))
                    if selected is None and action == 'resume' and batch.get('paused_at'):
                        db.execute('UPDATE screen_batches SET deadline=deadline+?,paused_at=NULL WHERE id=?', (time.time()-batch['paused_at'], batch_id))
                    db.execute('UPDATE screen_batches SET status=?,updated=? WHERE id=?',
                               ('paused' if action == 'pause' and selected is None else 'running', now(), batch_id))
            self._reconcile(batch)
            receipt = self.get(project_id, batch_id)
            if request_id:
                with self.store.transaction() as db:
                    db.execute('INSERT INTO screen_actions VALUES(?,?,?,?)', (batch_id, request_id, fingerprint(body), encoded(receipt)))
            self.wake.set()
            return receipt

    def _review(self, batch, item, body):
        if item['status'] not in {'completed', 'needs_review'}:
            raise Conflict('图片仍在处理或未产生结果，请等待当前尝试结束')
        state = body.get('state')
        if state not in {'accepted', 'retain', 'deferred'}:
            raise ValueError('复核状态无效')
        result_id = body.get('result_id', item['result_id'])
        if not result_id or result_id != item['result_id']:
            raise Conflict('最终候选已变化，请重新加载；局部中间结果不能替代整图结果')
        result = self.store.result(result_id)
        if type(body.get('expected_revision')) is not int or result['revision'] != body['expected_revision']:
            raise Conflict('复核内容已变化，请重新加载')
        image = self.store.one('images', item['image_id'])
        if image['active_version'] != item['input_version_id']:
            raise Conflict('原图版本已变化，不能复核旧批次')
        candidate_id = body.get('candidate_id')
        if candidate_id:
            candidate = next((c for c in json.loads(item['candidates']) if c['id'] == candidate_id), None)
            if not candidate:
                raise ValueError('整图候选不存在')
            result_id = self._final_result(batch, item, deepcopy(candidate['raw']))
            result = self.store.result(result_id)
        edits = body.get('edits') or []
        if not isinstance(edits, list) or len(edits) > 100:
            raise ValueError('复核编辑范围无效')
        if edits:
            value = deepcopy(result['edited'])
            for edit in edits:
                if (not isinstance(edit, dict) or any(type(edit.get(k)) is not int or edit[k] < 0 for k in ('table_index', 'row', 'column'))
                        or edit['table_index'] >= len(value.get('tables', []))):
                    raise ValueError('单元格编辑范围无效')
                table = value['tables'][edit['table_index']]
                cells = [c for c in table['cells'] if (c['row'], c['column']) == (edit['row'], edit['column'])]
                if len(cells) != 1 or cells[0]['text'] != edit['before'] or not isinstance(edit['after'], str):
                    raise Conflict('待修改单元格已变化')
                cells[0]['text'] = edit['after']
            result = self.store.save(result_id, value, result['revision'])
        with self.store.lock:
            with self.store.transaction() as db:
                current = db.execute('SELECT revision FROM results WHERE id=?', (result_id,)).fetchone()
                version = db.execute('SELECT active_version FROM images WHERE id=?', (item['image_id'],)).fetchone()
                if current['revision'] != result['revision'] or version['active_version'] != item['input_version_id']:
                    raise Conflict('复核内容或图片已变化')
                db.execute('INSERT INTO selections VALUES(?,?) ON CONFLICT(image_id) DO UPDATE SET result_id=excluded.result_id', (item['image_id'], result_id))
            self.store.set_review(item['image_id'], 'confirmed' if state == 'accepted' else 'question',
                                  result_id, result['revision'], item['input_version_id'])
        self._update(batch['id'], item['input_version_id'], result_id=result_id, result_revision=result['revision'],
                     review_state=state, status='completed' if state == 'accepted' else 'needs_review', stage='completed' if state == 'accepted' else 'review')

    def file(self, project_id, batch_id, version_id, file_id):
        detail = self.detail(project_id, batch_id, version_id)
        for image in (detail['preparation'] or {}).get('prepared_images', []):
            if image['id'] == file_id:
                path = Path(image['path']).resolve(strict=True)
                if not path.is_relative_to((self.root / batch_id / version_id).resolve()):
                    raise ValueError('图像路径越界')
                return path
        raise KeyError('预处理图像不存在')

    def export(self, project_id, batch_id, body=None):
        from ocr_workbench.tables import export_xlsx
        body = body or {}
        if set(body) - {'include_pending_review', 'numeric_columns'} or type(body.get('include_pending_review', False)) is not bool:
            raise ValueError('导出参数无效')
        batch = self._batch(project_id, batch_id)
        numeric = body.get('numeric_columns')
        if numeric is not None and (not isinstance(numeric, dict)
                or not set(numeric) <= {i['input_version_id'] for i in self._items(batch_id)}
                or any(not isinstance(v, dict) for v in numeric.values())):
            raise ValueError('数值列须按本批次图片版本逐项指定')
        include_review = body.get('include_pending_review', False)
        folder = self.root / batch_id / 'exports' / uid()
        folder.mkdir(parents=True)
        archive_path = folder / 'screen-tables.zip'
        manifest = {'format': VERSION + '-export', 'batch_id': batch_id, 'project_id': project_id,
                    'policy': batch['policy'], 'created': now(), 'items': [], 'success_files': 0}
        with self.guard, self.store.file_lock, zipfile.ZipFile(archive_path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            snapshots = []
            for item in self._items(batch_id):
                entry = self._display(item)
                entry['file_name'] = None
                if item['result_id'] and (item['status'] == 'completed' or include_review and item['status'] == 'needs_review'):
                    result = self.store.result(item['result_id'])
                    if result['revision'] != item['result_revision']:
                        raise Conflict('最终结果已编辑，请刷新后重新导出')
                    if self.store.one('images', item['image_id'])['active_version'] != item['input_version_id']:
                        raise Conflict('图片版本已变化，请重新处理或复核')
                    if result['edited'].get('tables'):
                        stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', Path(item['source_name']).stem).strip(' .') or 'image'
                        name = f'{item["position"] + 1:04d}-{stem[:100]}.xlsx'
                        target = folder / name
                        export_xlsx(result['edited']['tables'], target, source_rows=[
                            {'image': item['source_name'], 'result_id': result['id'], 'engine': result['original'].get('engine', ''),
                             'version_id': item['input_version_id'], 'revision': result['revision']}])
                        metadata = result['original'].get('screen', {}).get('business_metadata', [])
                        if metadata:
                            self._export_metadata(target, metadata)
                        numeric = (body.get('numeric_columns') or {}).get(item['input_version_id'])
                        if numeric:
                            if item['review_state'] != 'accepted':
                                raise Conflict('数值列转换须先确认当前图片结果')
                            self._numeric_columns(target, result['edited']['tables'], numeric)
                        archive.write(target, name)
                        snapshots.append((item['image_id'], item['input_version_id'], result['id'], result['revision']))
                        entry['file_name'] = name
                        manifest['success_files'] += 1
                    else:
                        entry['export_error'] = '没有结构化表格，不生成空白 Excel'
                manifest['items'].append(entry)
            if not manifest['success_files']:
                raise ValueError('尚无可交付表格，可先复核或明确导出待复核结果')
            for image_id, version_id, result_id, revision in snapshots:
                if (self.store.one('images', image_id)['active_version'] != version_id
                        or self.store.result(result_id)['revision'] != revision):
                    raise Conflict('导出期间结果或图片已变化，请重新导出')
            archive.writestr('manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2))
        (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        return archive_path, manifest

    @staticmethod
    def _export_metadata(path, metadata):
        from openpyxl import load_workbook
        book = load_workbook(path)
        sheet = book.create_sheet('业务说明')
        sheet.append(['类别', '识别原文', '来源', '原图坐标', '核对状态'])
        for part in metadata:
            sheet.append([part['kind'], part['text'], part.get('source', ''), encoded(part.get('box')), '独立文字识别，未逐项人工确认'])
            sheet.cell(sheet.max_row, 2).data_type = 's'
        sheet.column_dimensions['B'].width = 70
        book.save(path)

    @staticmethod
    def _numeric_columns(path, tables, specification):
        from openpyxl import load_workbook
        from ocr_workbench.financial_checks import parse_amount, IDENTIFIERS
        from ocr_workbench.screen_quality import _field_kind
        if not isinstance(specification, dict) or len(specification) > 100:
            raise ValueError('请明确指定已确认的数值列')
        book = load_workbook(path)
        for table_index, columns in specification.items():
            if not str(table_index).isdecimal() or not 0 <= int(table_index) < len(tables):
                raise ValueError('数值列的表格范围无效')
            table = tables[int(table_index)]
            if not isinstance(columns, list) or any(type(c) is not int or not 0 <= c < table['columns'] for c in columns):
                raise ValueError('数值列无效')
            sheet = book.worksheets[int(table_index)]
            offset = 1 if table.get('caption') else 0
            for cell in table['cells']:
                header = ' '.join(c.get('text', '') for c in table['cells'] if c['column'] == cell['column'] and (c.get('is_header') or c['row'] == 0))
                if cell['column'] not in columns or cell.get('is_header') or _field_kind(table, cell) == 'id' or IDENTIFIERS.search(header):
                    continue
                parsed = parse_amount(cell['text'])
                if parsed['kind'] not in {'amount', 'number', 'percent'}:
                    continue
                amount = parsed.get('value')
                if amount is None or sum(ch.isdigit() for ch in cell['text']) > 15:
                    continue
                target = sheet.cell(cell['row'] + 1 + offset, cell['column'] + 1)
                target.value = float(amount)
                target.number_format = '0' + ('.' + '0' * parsed.get('places', 0) if parsed.get('places', 0) else '') + ('%' if parsed.get('percent') else '')
        book.save(path)
