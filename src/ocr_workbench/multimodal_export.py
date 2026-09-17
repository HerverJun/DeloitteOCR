"""Deterministic reports over a pinned review/result snapshot, never model code."""
import json
from pathlib import Path
import shutil
import tempfile

from ocr_workbench.geometry_contract import fingerprint

DECISIONS = {"keep": "建议保留", "replace": "建议修改", "uncertain": "模型存疑"}
STATES = {"pending": "待处理", "accepted": "已采用", "rejected": "已拒绝", "question": "人工存疑", "stale": "已过期"}
NOTICE = "本清单记录模型建议与人工决定。模型建议不是原文事实；未采用的建议不改变正式结果，清单为空不代表原文无误。过期建议保留最近关联的内容与修订，不能据此判断当前原文。"


def capture_report(store, result_id):
    from ocr_workbench.multimodal_store import review_sources

    with store.transaction() as db:
        db.execute("BEGIN")
        row = db.execute("""SELECT r.*,t.image_id,t.version_id,i.name,p.page_number,p.document_id
            FROM results r JOIN tasks t ON t.id=r.task_id JOIN images i ON i.id=t.image_id
            LEFT JOIN pages p ON p.image_id=i.id WHERE r.id=?""", (result_id,)).fetchone()
        if row is None:
            raise KeyError("识别结果不存在")
        result = dict(row)
        result["original"] = json.loads(result["original"])
        result["edited"] = json.loads(result["edited"])
        evidence = review_sources(db, result) or {
            "schema_version": 1, "result_id": result_id, "revision": result["revision"],
            "automatic_adoption": False, "contributes_to_votes": False,
            "requests": [], "proposals": [], "decisions": [],
        }
        evidence["document"] = {"image_id": row["image_id"], "name": row["name"],
            "image_version": result["original"].get("project_image_version") or row["version_id"],
            "page_number": row["page_number"], "document_id": row["document_id"]}
        evidence["adopted_content_sha256"] = fingerprint(result["edited"])
        evidence["notice"] = NOTICE
        # Task lifecycle belongs to this same read snapshot, including failures before proposals.
        tasks = {r["id"]: dict(r) for r in db.execute("""SELECT t.id,t.status,t.phase,t.error,t.created,t.started,t.finished
            FROM tasks t JOIN multimodal_requests m ON m.task_id=t.id WHERE m.result_id=?""", (result_id,))}
        for request in evidence["requests"]:
            request["task"] = tasks.get(request["task_id"], {})
        return evidence


def target_label(target):
    if target.get("kind") == "cell":
        return f"表 {target['table']+1} / 行 {target['row']+1} / 列 {target['column']+1}"
    if target.get("kind") == "text":
        return f"文字位置 {target['start']}–{target['end']}（Unicode 字符，右端不含）"
    return "未知目标"


def _rows(report):
    requests = {r["task_id"]: r for r in report["requests"]}
    yield ["建议编号", "目标", "模型原建议", "人工状态", "审校时原文", "模型建议内容", "最近关联的目标内容",
           "模型理由", "证据范围", "图像版本", "模型配置", "任务编号", "起始修订", "当前关联修订"]
    for item in report["proposals"]:
        request = requests.get(item["task_id"], {})
        yield [item["id"], target_label(item["target"]), DECISIONS.get(item["decision"], item["decision"]),
            STATES.get(item["state"], item["state"]), item["before"], item["after"], item["current_value"],
            item["reason"], item["evidence"].get("reason", ""), item["evidence"].get("version_id", ""),
            request.get("snapshot", {}).get("model_id", ""), item["task_id"], str(item["basis_revision"]), str(item["revision"])]


def _markdown(report):
    def cell(value):
        return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")

    lines = ["# 视觉审校清单", "", NOTICE, "", f"结果：{report['result_id']} · 修订：{report['revision']}",
             f"已采用内容 SHA-256：{report['adopted_content_sha256']}", "", "## 建议与决定", ""]
    for index, row in enumerate(_rows(report)):
        lines.append("| " + " | ".join(cell(value) for value in row) + " |")
        if index == 0:
            lines.append("| " + " | ".join("---" for _ in row) + " |")
    lines.extend(["", "## 任务记录", ""])
    for request in report["requests"]:
        task = request["task"]
        lines.append(f"- {cell(request['task_id'])}：{cell(task.get('status', ''))} · {cell(task.get('phase', ''))}")
        if task.get("error"):
            lines.append(f"  错误：{cell(task['error'])}")
    return "\n".join(lines) + "\n"


def _workbook(report, path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "校验清单"
    expected = []

    def append_literal(ws, values):
        values = ["" if value is None else str(value) for value in values]
        if any(len(value) > 32767 for value in values):
            raise ValueError("清单内容超过 Excel 单格限制，请使用 JSON 格式")
        ws.append(values)
        for cell in ws[ws.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        return values

    for row in _rows(report):
        expected.append(append_literal(sheet, row))
    tasks = book.create_sheet("任务与来源")
    append_literal(tasks, ["任务编号", "模型配置", "状态", "阶段", "错误", "配置指纹", "图像 SHA-256"])
    for request in report["requests"]:
        snap, task = request.get("snapshot", {}), request["task"]
        append_literal(tasks, [request["task_id"], snap.get("model_id"), task.get("status"), task.get("phase"), task.get("error"),
                              snap.get("config", {}).get("config_sha256"), snap.get("image_sha256")])
    notes = book.create_sheet("说明")
    append_literal(notes, ["说明", NOTICE])
    for key in ("result_id", "revision", "adopted_content_sha256"):
        append_literal(notes, [key, report[key]])
    append_literal(notes, ["页面身份", json.dumps(report["document"], ensure_ascii=False)])
    append_literal(notes, ["完整追溯", "JSON 清单另含锁定配置、原始模型响应、原始/当前位置及逐条决定；Excel 是供核对的文字清单。"])
    for ws in book:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="286E55")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 28
    for col in ("E", "F", "G", "H"):
        sheet.column_dimensions[col].width = 42
    notes.column_dimensions["B"].width = 100
    book.save(path)
    book.close()
    # File generation is verified against the captured strings, including leading zeros and formula-like input.
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        got = [["" if c.value is None else str(c.value) for c in row] for row in check["校验清单"].iter_rows()]
        if got != expected or any(c.data_type == "f" for ws in check for row in ws.iter_rows() for c in row):
            raise RuntimeError("审校清单 Excel 读回与来源快照不一致")
    finally:
        check.close()


def build_review_report(store, result_id, format="xlsx"):
    if format not in {"json", "md", "xlsx"}:
        raise ValueError("校验清单支持 JSON、Markdown 或 XLSX")
    report = capture_report(store, result_id)
    parent = store.root / "exports"
    parent.mkdir(exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix="export-", dir=parent))
    try:
        path = folder / ("OCR-review-report." + format)
        if format == "json":
            path.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), "utf-8")
        elif format == "md":
            path.write_text(_markdown(report), "utf-8")
        else:
            _workbook(report, path)
        return path
    except BaseException:
        shutil.rmtree(folder)
        raise
