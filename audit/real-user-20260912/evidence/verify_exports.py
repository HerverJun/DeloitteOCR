"""Read retained UI responses; never rewrite the exported artifacts."""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import sys
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parent
GLM = "6993a7e01e8140439975951ec0b480ed"
PP = "d94ee08e0f834775ad38c32ff8951b1b"
ROAD = "dd6dbfe75d6b4f2fb048c03c5e8d8229"
checks = []
observations = []


def check(key, passed, detail):
    checks.append({"id": key, "passed": bool(passed), "detail": detail})


def workbook_data(source):
    book = openpyxl.load_workbook(source, data_only=False)
    sheets = []
    for sheet in book.worksheets:
        sheets.append({
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
            "values": [list(row) for row in sheet.values],
            "cells": {
                cell.coordinate: {"value": cell.value, "data_type": cell.data_type}
                for row in sheet for cell in row
                if not isinstance(cell, openpyxl.cell.cell.MergedCell)
            },
            "formula_cells": [cell.coordinate for row in sheet for cell in row if cell.data_type == "f"],
        })
    book.close()
    return sheets


class HtmlTables(HTMLParser):
    """Independent structural extraction of HTML tables embedded in Markdown."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self.current = None
        self.cell = None
        self.occupied = set()
        self.row = -1
        self.column = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "table":
            self.current = {"cells": [], "rows": 0, "columns": 0}
            self.occupied, self.row = set(), -1
        elif tag == "tr" and self.current is not None:
            self.row += 1
            self.column = 0
            self.current["rows"] = max(self.current["rows"], self.row + 1)
        elif tag in ("td", "th") and self.current is not None:
            while (self.row, self.column) in self.occupied:
                self.column += 1
            self.cell = {
                "row": self.row, "column": self.column,
                "row_span": int(attrs.get("rowspan", 1)),
                "column_span": int(attrs.get("colspan", 1)), "text": "",
            }
        elif tag == "br" and self.cell is not None:
            self.cell["text"] += "\n"

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None:
            cell = self.cell
            self.current["cells"].append(cell)
            for r in range(cell["row"], cell["row"] + cell["row_span"]):
                for c in range(cell["column"], cell["column"] + cell["column_span"]):
                    self.occupied.add((r, c))
            self.column += cell["column_span"]
            self.current["columns"] = max(self.current["columns"], self.column)
            self.cell = None
        elif tag == "table" and self.current is not None:
            self.tables.append(self.current)
            self.current = None


def normalized_table(table):
    return {
        "rows": table["rows"], "columns": table["columns"],
        "cells": sorted([
            {k: c[k] for k in ("row", "column", "row_span", "column_span", "text")}
            for c in table["cells"]
        ], key=lambda c: (c["row"], c["column"])),
    }


inventory = []
for receipt_file in sorted((ROOT / "exports").glob("*/receipt.json")):
    receipt = json.loads(receipt_file.read_text(encoding="utf-8"))
    file = receipt_file.parent / Path(receipt["preserved_file"]).name
    content = file.read_bytes()
    digest = sha256(content).hexdigest()
    record = {
        "file": str(file.relative_to(ROOT)), "receipt": str(receipt_file.relative_to(ROOT)),
        "captured_at": receipt["captured_at"], "parameters": receipt["parameters"],
        "bytes": len(content), "sha256": digest,
        "receipt_hashes_match": digest == receipt["preserved_sha256"] == receipt["source_sha256"],
        "receipt_length_matches": len(content) == receipt["bytes"],
        "receipt_byte_identical": receipt["byte_identical"],
        "capture_method": receipt["capture_method"],
    }
    if file.suffix == ".xlsx":
        record["workbook"] = workbook_data(BytesIO(content))
    elif file.suffix == ".zip":
        with ZipFile(BytesIO(content)) as archive:
            record["zip_members"] = archive.namelist()
            record["zip_bad_member"] = archive.testzip()
            record["zip_workbooks"] = {
                name: workbook_data(BytesIO(archive.read(name)))
                for name in archive.namelist() if name.endswith(".xlsx")
            }
            if "sources.json" in archive.namelist():
                record["sources"] = json.loads(archive.read("sources.json"))
    elif file.suffix == ".json":
        record["json"] = json.loads(content)
    else:
        record["text"] = content.decode("utf-8-sig")
    inventory.append(record)

check("retained_bytes_match_all_receipts", all(x["receipt_hashes_match"] and x["receipt_length_matches"] and x["receipt_byte_identical"] for x in inventory), {"files": len(inventory)})
road = [x for x in inventory if x["parameters"]["result_ids"] == [ROAD] and "text" in x]
check("earliest_road_txt_contains_edit_marker", bool(road) and "QA 校对标记 00123" in road[0]["text"], road[0]["file"] if road else "missing")

glm_xlsx = [x for x in inventory if x["parameters"]["result_ids"] == [GLM] and "workbook" in x]
pp_txt = [x for x in inventory if x["parameters"]["result_ids"] == [PP] and x["parameters"]["format"] == "txt"]
originals = json.loads((ROOT / "spotcheck-original-results.json").read_text(encoding="utf-8"))
original_by_id = {x["result_id"]: x["original"] for x in originals}
if glm_xlsx and pp_txt:
    first = glm_xlsx[0]
    index = next(x for x in first["workbook"] if x["name"] == "来源索引")
    mapping = dict(zip(index["values"][0], index["values"][1]))
    check("preview_and_selected_export_use_distinct_results", mapping["结果 ID"] == GLM and mapping["引擎"] == "glm" and pp_txt[0]["parameters"]["result_ids"] == [PP], {"preview": mapping, "selected_result": PP})
    check("selected_ppocr_txt_matches_its_original", pp_txt[0]["text"].replace("\r\n", "\n") == original_by_id[PP]["text"], {"file": pp_txt[0]["file"], "characters": len(pp_txt[0]["text"]), "normalization": "Windows CRLF export compared to LF model text; no other normalization"})

jsons = [x for x in inventory if x["parameters"]["result_ids"] == [GLM] and isinstance(x.get("json"), dict)]
mds = [x for x in inventory if x["parameters"]["result_ids"] == [GLM] and x["parameters"]["format"] == "md"]
if len(glm_xlsx) >= 2 and jsons and mds:
    # The first edited trio is retained even when later multi-table files arrive.
    edited_xlsx, exported_json, exported_md = glm_xlsx[1], jsons[0], mds[0]
    sheet = next(x for x in edited_xlsx["workbook"] if x["name"] != "来源索引")
    value = exported_json["json"]
    expected = {"B3": "00123", "C3": "90071992547409931234", "D3": "=1+1"}
    observations.append({"id": "glm_intermediate_editing_state", "classification": "test_operation_observation_not_product_bug", "expected_before_operator_clarification": expected, "actual_B3_through_E3": {a: sheet["cells"][a] for a in ("B3", "C3", "D3", "E3")}, "file": edited_xlsx["file"], "operator_clarification": "The test inserted at B3 and later deleted the last row/column rather than the inserted row/column. This artifact is an intentional retained intermediate state, not evidence of an export defect."})
    actual_identifier_cells = {a: sheet["cells"][a] for a in ("B3", "D3", "E3")}
    check("actual_identifier_values_remain_strings", all(c["data_type"] == "s" and isinstance(c["value"], str) for c in actual_identifier_cells.values()), actual_identifier_cells)
    check("no_formulas_in_edited_workbook", all(not x["formula_cells"] for x in edited_xlsx["workbook"]), edited_xlsx["file"])
    check("edited_table_has_six_rows_eight_columns", (sheet["rows"], sheet["columns"]) == (6, 8), {"rows": sheet["rows"], "columns": sheet["columns"]})
    first_table = next(x for x in glm_xlsx[0]["workbook"] if x["name"] != "来源索引")
    observations[-1]["header_merges"] = {"before": first_table["merged"], "edited": sheet["merged"]}
    parser = HtmlTables()
    parser.feed(exported_md["text"])
    check("markdown_tables_match_json_edited_tables", [normalized_table(t) for t in parser.tables] == [normalized_table(t) for t in value["edited"]["tables"]], {"md": exported_md["file"], "json": exported_json["file"], "revision": value["revision"]})
    table = value["edited"]["tables"][0]
    expected_merges = []
    mismatches = []
    for cell in table["cells"]:
        address = f"{openpyxl.utils.get_column_letter(cell['column'] + 1)}{cell['row'] + 1}"
        actual = sheet["cells"][address]["value"]
        if (actual if actual is not None else "") != cell["text"]:
            mismatches.append({"address": address, "xlsx": actual, "json": cell["text"]})
        if cell["row_span"] > 1 or cell["column_span"] > 1:
            end = f"{openpyxl.utils.get_column_letter(cell['column'] + cell['column_span'])}{cell['row'] + cell['row_span']}"
            expected_merges.append(address + ":" + end)
    check("xlsx_matches_json_edited_cells_and_merges", not mismatches and sheet["merged"] == sorted(expected_merges), {"cell_mismatches": mismatches, "merges": sheet["merged"], "revision": value["revision"]})
    check("json_original_matches_pre_edit_snapshot", value["original"] == original_by_id.get(GLM), {"result_id": GLM, "snapshot": "spotcheck-original-results.json"})
    check("json_original_retains_unedited_values", [c["text"] for c in value["original"]["tables"][0]["cells"] if c["row"] == 2 and c["column"] in (1, 2, 3)] == ["4", "516.1", "± 56.5"], {"edited_text_equals_original_text": value["edited"]["text"] == value["original"]["text"], "note": "Structured table edits are independent of the original text field."})

def data_sheets(book):
    return [s for s in book if s["name"] != "来源索引"]


def source_rows(book):
    index = next(s for s in book if s["name"] == "来源索引")
    return [dict(zip(index["values"][0], row)) for row in index["values"][1:]]


all_books = []
for item in inventory:
    if "workbook" in item:
        all_books.append((item, Path(item["file"]).name, item["workbook"]))
    for name, book in item.get("zip_workbooks", {}).items():
        all_books.append((item, name, book))
source_errors = []
for item, filename, book in all_books:
    mapping = source_rows(book)
    if len(mapping) != len(data_sheets(book)):
        source_errors.append({"file": item["file"], "reason": "source row count differs from table count"})
    for row in mapping:
        source = original_by_id.get(row["结果 ID"])
        valid = (
            row["工作簿"] == filename
            and row["工作表"] in {s["name"] for s in data_sheets(book)}
            and row["结果 ID"] in item["parameters"]["result_ids"]
            and source is not None
            and row["引擎"] == source["engine"]
            and row["图像版本"] == source["project_image_version"]
            and row["输入 SHA-256"] == source["image"]["sha256"]
        )
        if not valid:
            source_errors.append({"file": item["file"], "row": row})
check("all_workbook_source_indexes_reference_correct_inputs", not source_errors, {"workbooks_including_zip_members": len(all_books), "errors": source_errors})
check("all_workbooks_have_no_formula_cells", all(not s["formula_cells"] for _, _, book in all_books for s in book), {"workbooks_including_zip_members": len(all_books)})

road_books = [x for x in inventory if x["parameters"]["result_ids"] == [ROAD] and "workbook" in x and not x["parameters"]["confirmed_only"]]
if road_books:
    latest_road = road_books[-1]
    tables = data_sheets(latest_road["workbook"])
    check("road_single_image_exports_two_tables", len(tables) == 2 and len(source_rows(latest_road["workbook"])) == 2, {"file": latest_road["file"], "sheets": [s["name"] for s in tables]})
    first = tables[0]
    offset = 1 if first["values"][0][0] == "路牌人工校对" else 0
    expected_road = {f"A{offset+1}": "00123", f"B{offset+1}": "90071992547409931234", f"C{offset+1}": "=1+1"}
    road_values_ok = all(first["cells"].get(a, {}).get("value") == v and first["cells"].get(a, {}).get("data_type") == "s" for a, v in expected_road.items())
    multiline = first["cells"].get(f"A{offset+2}", {}).get("value")
    check("latest_road_tsv_identifiers_and_multiline", road_values_ok and isinstance(multiline, str) and "\n" in multiline, {"file": latest_road["file"], "caption_rows": offset, "expected": expected_road, "actual_cells": first["cells"], "classification_if_failed": "requires_browser_or_saved_revision_check_before_attributing_to_product"})
    for prior in road_books:
        first_prior = data_sheets(prior["workbook"])[0]
        if not any(c["value"] == "00123" for c in first_prior["cells"].values()):
            observations.append({"id": "road_export_with_empty_first_table", "file": prior["file"], "source": source_rows(prior["workbook"]), "values": first_prior["values"], "classification": "test_operation_observation_not_product_bug", "operator_clarification": "Fast successive edits were grouped into the same autosave history entry; the test undo also reverted TSV insertion. The operator re-pasted, waited for saved state, switched between both tables to verify the values, and exported revision 11."})

for item in inventory:
    if "zip_workbooks" in item:
        sources = item.get("sources", {}).get("tables", [])
        source_tuples = {(s["workbook"], s["sheet"], s["result_id"], str(s["revision"]), str(s["table_index"])) for s in sources}
        expected_tuples = {(name, s["工作表"], s["结果 ID"], s["校对 revision"], s["表格序号"]) for name, book in item["zip_workbooks"].items() for s in source_rows(book)}
        check("zip_sources_match_workbook_indexes:" + Path(item["file"]).parent.name, item["zip_bad_member"] is None and source_tuples == expected_tuples and len(sources) == 3 and len(item["zip_workbooks"]) == 2, {"file": item["file"], "members": item["zip_members"], "source_table_count": len(sources)})
    elif "workbook" in item and len(item["parameters"]["result_ids"]) == 2 and item["parameters"]["aggregate"]:
        mapping = source_rows(item["workbook"])
        check("aggregate_contains_all_three_tables:" + Path(item["file"]).parent.name, len(data_sheets(item["workbook"])) == 3 and [s["结果 ID"] for s in mapping] == [GLM, ROAD, ROAD], {"file": item["file"], "sources": mapping})
    elif "workbook" in item and item["parameters"]["confirmed_only"]:
        mapping = source_rows(item["workbook"])
        check("confirmed_only_contains_road_two_tables:" + Path(item["file"]).parent.name, item["parameters"]["result_ids"] == [ROAD] and len(mapping) == 2 and all(s["结果 ID"] == ROAD for s in mapping), {"file": item["file"], "parameters": item["parameters"], "sources": mapping})

standalone_tables = {}
for item in inventory:
    if "workbook" not in item or len(item["parameters"]["result_ids"]) != 1 or item["parameters"]["confirmed_only"]:
        continue
    book = {s["name"]: s for s in data_sheets(item["workbook"])}
    for row in source_rows(item["workbook"]):
        standalone_tables[(row["结果 ID"], row["校对 revision"], row["表格序号"])] = (item["file"], book[row["工作表"]])
batch_fidelity = []
for item, filename, book in all_books:
    if len(item["parameters"]["result_ids"]) == 1 and not item["parameters"]["confirmed_only"]:
        continue
    by_name = {s["name"]: s for s in data_sheets(book)}
    for row in source_rows(book):
        key = (row["结果 ID"], row["校对 revision"], row["表格序号"])
        prior = standalone_tables.get(key)
        current = by_name[row["工作表"]]
        equal = prior is not None and all(current[field] == prior[1][field] for field in ("rows", "columns", "merged", "values", "cells", "formula_cells"))
        batch_fidelity.append({"file": item["file"], "workbook": filename, "sheet": row["工作表"], "key": key, "matches_standalone": equal, "standalone": prior[0] if prior else None})
check("batch_and_confirmed_tables_match_same_revision_standalone", bool(batch_fidelity) and all(x["matches_standalone"] for x in batch_fidelity), batch_fidelity)

extra_selected_id = "66d7c221ad944168b20fdeec0c6e01ae"
extra_selected = [x for x in inventory if x["parameters"]["result_ids"] == [extra_selected_id] and x["parameters"]["format"] == "txt"]
if extra_selected:
    snapshot_file = ROOT / "selected-export-source-snapshot.json"
    if not snapshot_file.exists():
        database = Path(json.loads((ROOT / "tasks-latest.json").read_text(encoding="utf-8"))["database"])
        with sqlite3.connect(database.as_uri() + "?mode=ro", uri=True) as connection:
            connection.execute("PRAGMA query_only=ON")
            row = connection.execute("SELECT original,edited,revision FROM results WHERE id=?", (extra_selected_id,)).fetchone()
        snapshot = {"id": extra_selected_id, "original": json.loads(row[0]), "edited": json.loads(row[1]), "revision": row[2], "database_access": "SQLite mode=ro and PRAGMA query_only=ON", "captured_at": datetime.now(timezone.utc).isoformat()}
        snapshot_file.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    snapshot = json.loads(snapshot_file.read_text(encoding="utf-8"))
    check("selected_txt_works_while_current_image_unrecognized", extra_selected[0]["text"].replace("\r\n", "\n") == snapshot["edited"]["text"] and not snapshot["edited"]["tables"], {"file": extra_selected[0]["file"], "result_id": extra_selected_id, "source_snapshot": snapshot_file.name, "browser_observation": "Current image 科研表格.png had no result; selected 路牌 原图.jpg result exported through enabled 导出所选."})
    review_exports = [x for x in extra_selected if "QA REVIEW_ONLY 保存导出" in x["text"]]
    if review_exports:
        review_export = review_exports[-1]
        check("review_only_txt_contains_saved_edit", review_export["text"].replace("\r\n", "\n").rstrip("\n") == snapshot["edited"]["text"].rstrip("\n") + "\nQA REVIEW_ONLY 保存导出", {"file": review_export["file"], "result_id": extra_selected_id, "source_snapshot": snapshot_file.name, "marker": "QA REVIEW_ONLY 保存导出", "browser_observation": "Review-only mode: appended marker, waited for saved status, exported TXT, then restored original text and saved."})

report = {
    "verified_at": datetime.now(timezone.utc).isoformat(), "python": sys.executable,
    "method": "Read-only SHA-256, receipt, openpyxl data_type/merge inspection, independent HTML table parse, and pre-edit original snapshot comparison.",
    "release_integrity_acceptance": False,
    "checks": checks,
    "observations": observations,
    "summary": {"passed": sum(x["passed"] for x in checks), "failed": sum(not x["passed"] for x in checks), "artifact_count": len(inventory)},
    "inventory": inventory,
    "interpretation": "The first edited XLSX/MD/JSON agree with saved GLM revision 7; changed placement/merges came from test insertion/deletion at different positions. Road revision 10's empty first table came from test undo of a grouped autosave history entry. Both are retained intermediate observations, not product bugs. Latest road revision 11 passes identifiers, literal formula text and newline checks. Original OCR JSON remains identical to the pre-edit snapshot. Earlier batch files correctly contain their captured revisions, not the subsequently edited road revision 11.",
}
(ROOT / "export-verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
lines = ["# UI 原始导出文件核验", "", f"核验时间：{report['verified_at']}。已读取 {len(inventory)} 份 UI 实际导出文件及 receipt。原始导出文件未改动。", "", f"已通过 {report['summary']['passed']} 项检查，{report['summary']['failed']} 项尚不满足预期。GLM 三种导出与已保存校对一致、原始 JSON 不变；行列与合并变化经操作人确认来自测试操作，归为中间编辑状态，不作为产品缺陷。", "", "| 检查 | 结果 |", "| --- | --- |"]
for item in checks:
    lines.append(f"| {item['id']} | {'通过' if item['passed'] else '不符合预期'} |")
lines += ["", "实测说明：", "", "- 最新路牌导出为 `20260912T074050.234526Z-42a60311/OCR-result.xlsx`，revision 11。第一表 A1 是标题，所以数据从第 2 行开始：A2=`00123`，B2=`90071992547409931234`，C2=`=1+1`，均是字符串；A3 为含真实换行的 `多行\\n文本`，B3 为 `Yuyuan Rd.`，C3 为空。第二表独立保留。", "- GLM revision 7 中间状态为 6×8，B3=`00123`、D3=`90071992547409931234`、E3=`=1+1`，均是字符串；MD/JSON/XLSX 单元格和合并结构一致。", "- GLM 中间状态的空列/行与合并变化由测试插入、删除位置不同造成；路牌 revision 10 的空第一表由测试撤销时一并撤销了合并保存的 TSV 操作造成。操作人已确认，两者均保留为测试观察，不记为产品 bug。", "- GLM 预览 XLSX 使用 `6993a7e01e8140439975951ec0b480ed` / GLM / revision 0；随后所选 TXT 使用 `d94ee08e0f834775ad38c32ff8951b1b` / PP-OCR，除 Windows CRLF 外内容等于其原文。", "- 路牌最早 TXT 含完整 `QA 校对标记 00123`。GLM JSON original 与编辑前快照完全相等。", "- ZIP 包含 2 个逐图工作簿与 sources.json，共 3 张数据表；汇总工作簿同样有 3 表及来源索引；confirmed_only 实际请求只含路牌结果，输出其 2 张表。批量与单独导出的相同 revision 数据、类型与合并完全一致。", "- 较早批量文件固定在 GLM revision 7 / 路牌 revision 10，晚于它们的路牌 revision 11 不应反向改变这些历史导出。", "", "文件清单：", "", "| 文件 | 字节 | receipt SHA-256 |", "| --- | ---: | --- |"]
for item in inventory:
    lines.append(f"| [{item['file']}]({item['file']}) | {item['bytes']} | {'一致' if item['receipt_hashes_match'] else '不一致'} |")
lines += ["", "数据详情、所有单元格类型、完整来源索引及检查证据见 `export-verification.json`。这是本机源码/UI 导出核验，不代表便携发布包完整性验收。", ""]
(ROOT / "export-verification.md").write_text("\n".join(lines), encoding="utf-8")
print(json.dumps({"summary": report["summary"], "failed": [x for x in checks if not x["passed"]]}, ensure_ascii=False, indent=2))
