"""Lossless table cell text, spans and explicit Excel string values."""
from html.parser import HTMLParser
import hashlib
import re


def code_ranges(text):
    """Protect fenced, indented and inline code from structural interpretation."""
    ranges, offset, fence, opened = [], 0, None, 0
    for line in text.splitlines(keepends=True):
        marker = re.match(r' {0,3}(`{3,}|~{3,})(.*)', line)
        if fence:
            if marker and marker[1][0] == fence[0] and len(marker[1]) >= len(fence) and not marker[2].strip():
                ranges.append((opened, offset + len(line)))
                fence = None
        elif marker:
            fence, opened = marker[1], offset
        elif line.startswith(('    ', '\t')):
            ranges.append((offset, offset + len(line)))
        else:
            # A code span closes only with a run of the same number of backticks.
            runs = list(re.finditer(r'`+', line))
            i = 0
            while i < len(runs):
                closing = next((j for j in range(i + 1, len(runs))
                                if len(runs[j][0]) == len(runs[i][0])), None)
                if closing is None:
                    i += 1
                else:
                    ranges.append((offset + runs[i].start(), offset + runs[closing].end()))
                    i = closing + 1
        offset += len(line)
    if fence:
        ranges.append((opened, len(text)))
    return ranges


def table_source(text, start, end, format):
    return {'start': start, 'end': end, 'format': format,
            'sha256': hashlib.sha256(text[start:end].encode('utf-8')).hexdigest()}


class TableParser(HTMLParser):
    def __init__(self, text=''):
        super().__init__(convert_charrefs=True)
        self.text = text
        self.line_offsets = [0] + [m.end() for m in re.finditer('\n', text)]
        self.start = 0
        self.tables = []
        self.table = None
        self.cell = None
        self.row = -1
        self.col = 0
        self.occupied = set()
        self.in_caption = False
        self.code = code_ranges(text)

    def tag_is_code(self):
        offset = self.source_offset()
        return any(start <= offset < end and not (
            self.table is not None and self.text.startswith(('    ', '\t'), start))
            for start, end in self.code)

    def source_offset(self):
        line, column = self.getpos()
        return self.line_offsets[line - 1] + column

    def handle_starttag(self, tag, attrs):
        if self.tag_is_code():
            if self.cell is not None:
                self.cell['text'] += self.get_starttag_text()
            return
        attrs = dict(attrs)
        if tag == 'table':
            if self.table is not None:
                raise ValueError('Nested tables are not supported')
            self.table = {'cells': [], 'rows': 0, 'columns': 0, 'caption': ''}
            self.start = self.source_offset()
            self.row, self.col, self.occupied = -1, 0, set()
        elif tag == 'caption' and self.table is not None:
            self.in_caption = True
        elif tag == 'tr' and self.table is not None:
            self.row += 1
            self.col = 0
        elif tag in {'td', 'th'} and self.table is not None:
            if self.cell is not None or self.row < 0:
                raise ValueError('Malformed table cell')
            while (self.row, self.col) in self.occupied:
                self.col += 1
            rs, cs = int(attrs.get('rowspan', '1')), int(attrs.get('colspan', '1'))
            if not (1 <= rs <= 1000 and 1 <= cs <= 1000):
                raise ValueError('Invalid table span')
            for r in range(self.row, self.row + rs):
                for c in range(self.col, self.col + cs):
                    if (r, c) in self.occupied:
                        raise ValueError('Overlapping table cells')
                    self.occupied.add((r, c))
            self.cell = {'row': self.row, 'column': self.col, 'row_span': rs,
                         'column_span': cs, 'text': '', 'confidence': None, 'polygon': None}
            self.table['rows'] = max(self.table['rows'], self.row + rs)
            self.table['columns'] = max(self.table['columns'], self.col + cs)
            self.col += cs
        elif tag == 'br' and self.cell is not None:
            self.cell['text'] += '\n'

    def handle_data(self, data):
        if self.cell is not None:
            self.cell['text'] += data
        elif self.in_caption:
            self.table['caption'] += data

    def handle_endtag(self, tag):
        if self.tag_is_code():
            if self.cell is not None:
                self.cell['text'] += f'</{tag}>'
            return
        if tag == 'caption':
            self.in_caption = False
        elif tag in {'td', 'th'} and self.cell is not None:
            self.table['cells'].append(self.cell)
            self.cell = None
        elif tag == 'table' and self.table is not None:
            if self.cell is not None:
                raise ValueError('Unclosed table cell')
            end = self.text.find('>', self.source_offset()) + 1
            self.table['source'] = table_source(self.text, self.start, end, 'html')
            self.tables.append(self.table)
            self.table = None
            self.in_caption = False


def parse_tables(text, *, warnings=None):
    parser = TableParser(text)
    invalid_start = None
    try:
        parser.feed(text)
        parser.close()
        if parser.table is not None:
            raise ValueError('Truncated HTML table')
    except ValueError as error:
        if warnings is None:
            raise
        invalid_start = parser.start
        warnings.append({'code': 'table_parse_failed', 'stage': 'table_parse',
                         'message': '部分表格结构无法解析，原文与原始输出已保留，可继续校对或导出 TXT / JSON。',
                         'detail': str(error),
                         'source': table_source(text, invalid_start, len(text), 'html')})
    # Parse Markdown only outside HTML tables, then merge by source position.
    tables = []
    offset = 0
    for table in parser.tables:
        tables.extend(parse_markdown_tables(text, offset, table['source']['start']))
        tables.append(table)
        offset = table['source']['end']
    tables.extend(parse_markdown_tables(text, offset, invalid_start if invalid_start is not None else len(text)))
    return tables


def parse_markdown_tables(text, start, end):
    # A table requires a header and matching delimiter row. Inconsistent rows
    # remain literal text; Markdown never supplies merge information.
    protected = code_ranges(text)
    lines, offset = [], start
    for line in text[start:end].splitlines(keepends=True):
        stripped = line.strip()
        fields = None
        if stripped:
            separators = []
            for match in re.finditer(r'\|', line):
                position = offset + match.start()
                slashes = len(line[:match.start()]) - len(line[:match.start()].rstrip('\\'))
                if slashes % 2 == 0 and not any(a <= position < b for a, b in protected):
                    separators.append(match.start())
            if separators:
                bounds = [-1] + separators + [len(line.rstrip('\r\n'))]
                fields = [line[a + 1:b].strip() for a, b in zip(bounds, bounds[1:])]
                if not fields[0]:
                    fields.pop(0)
                if fields and not fields[-1]:
                    fields.pop()
                if not fields:
                    fields = None
        lines.append((offset, offset + len(line.rstrip('\r\n')), fields))
        offset += len(line)
    tables, i = [], 0
    while i + 1 < len(lines):
        header, delimiter = lines[i][2], lines[i + 1][2]
        if not header or not delimiter or len(header) != len(delimiter) or not all(
                re.fullmatch(r':?-{3,}:?', field) for field in delimiter):
            i += 1
            continue
        rows, first, last = [header], lines[i][0], lines[i + 1][1]
        i += 2
        while i < len(lines) and lines[i][2] is not None and len(lines[i][2]) == len(header):
            rows.append(lines[i][2])
            last = lines[i][1]
            i += 1
        tables.append({'rows': len(rows), 'columns': len(header),
                       'source': table_source(text, first, last, 'markdown'),
                       'cells': [{'row': r, 'column': c, 'row_span': 1, 'column_span': 1,
                                  'text': value.replace('\\|', '|'), 'confidence': None, 'polygon': None}
                                 for r, row in enumerate(rows) for c, value in enumerate(row)]})
    return tables


SOURCE_COLUMNS = (
    ('workbook', '工作簿'), ('sheet', '工作表'), ('image_name', '原图'),
    ('image_id', '图片 ID'), ('image_version', '图像版本'),
    ('input_sha256', '输入 SHA-256'), ('engine', '引擎'),
    ('engine_name', '引擎名称'), ('engine_package', '引擎包'),
    ('model_revisions', '模型版本'), ('result_id', '结果 ID'),
    ('revision', '校对 revision'), ('table_index', '表格序号'),
    ('origin', '结果来源类型'), ('policy_version', '融合策略版本'),
    ('policy_sha256', '融合策略指纹'), ('review_summary', '融合校对状态'),
)


def export_xlsx(tables, path, *, source_rows=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for index, table in enumerate(tables, 1):
        sheet = book.create_sheet(f'Table {index}')
        offset = 1 if table.get('caption') else 0
        if offset:
            if len(table['caption']) > 32767:
                raise ValueError('表格标题超过 Excel 的 32,767 字符限制；可导出 TXT / JSON 或缩短标题后重试')
            sheet.cell(1,1).value = table['caption']
            sheet.cell(1,1).data_type = 's'
        for cell in table['cells']:
            row, col = cell['row'] + 1 + offset, cell['column'] + 1
            if len(cell['text']) > 32767:
                raise ValueError('单元格超过 Excel 的 32,767 字符限制；可导出 TXT / JSON 或缩短单元格后重试')
            out = sheet.cell(row, col)
            out.value = cell['text']
            out.data_type = 's'
            out.number_format = '@'
            out.alignment = Alignment(wrap_text=True, vertical='top')
            if cell['row_span'] > 1 or cell['column_span'] > 1:
                sheet.merge_cells(start_row=row, start_column=col,
                                  end_row=row + cell['row_span'] - 1,
                                  end_column=col + cell['column_span'] - 1)
        for col in range(1, table['columns'] + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24
    if not book.worksheets:
        raise ValueError('No recognized tables; refusing to export an empty workbook')
    if source_rows is not None:
        index = book.create_sheet('来源索引')
        index.freeze_panes = 'A2'
        rows = [[label for _, label in SOURCE_COLUMNS]] + [
            [str(source.get(key, '')) for key, _ in SOURCE_COLUMNS] for source in source_rows]
        for r, values in enumerate(rows, 1):
            for c, value in enumerate(values, 1):
                out = index.cell(r, c, value)
                out.data_type = 's'
                out.number_format = '@'
                out.alignment = Alignment(wrap_text=True, vertical='top')
        index.auto_filter.ref = index.dimensions
        for c in range(1, len(SOURCE_COLUMNS) + 1):
            index.column_dimensions[get_column_letter(c)].width = 28
    book.save(path)
