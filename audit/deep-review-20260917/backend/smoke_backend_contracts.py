"""A single end-to-end persisted edit/export/cleanup chain in an isolated workspace."""
from pathlib import Path
from copy import deepcopy
from types import SimpleNamespace
import hashlib
import json
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'src'))

from PIL import Image
from openpyxl import load_workbook
from ocr_workbench.store import Store, Conflict
from ocr_workbench.imaging import add_image, prepare_task
from ocr_workbench.tables import parse_tables
from ocr_workbench.exporting import build_export
from ocr_workbench.maintenance import ProjectMaintenance


def main():
    checks = []
    with tempfile.TemporaryDirectory(prefix='ocr-backend-contracts-') as temporary:
        folder = Path(temporary)
        store = Store(folder / '独立项目 空格')
        project = store.project('持久化与导出审计')
        other = store.project('保留项目')
        source = folder / '原图.png'
        Image.new('RGB', (64, 40), 'white').save(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        image = add_image(store, project['id'], source.name, source)
        assert hashlib.sha256(store.file(image['original_path']).read_bytes()).hexdigest() == source_hash
        checks.append('Unicode workspace import preserves original bytes')

        jobs = store.enqueue(project['id'], [image['active_version']], ['ppocr', 'glm'], [{'kind': 'rotate', 'degrees': 90}])
        text = '原始正文\n<table><tr><td colspan="2">原始表头</td></tr><tr><td>00123456789012345678</td><td>=1+1</td></tr></table>'
        raw = {'text': text, 'tables': parse_tables(text), 'blocks': []}
        prepared = []
        for _ in jobs:
            task = prepare_task(store, store.claim())
            prepared.append(task['version_id'])
            assert store.complete(task['id'], {**raw, 'engine': task['engine'], 'project_image_version': task['version_id']})
        assert len(set(prepared)) == 1
        assert len(store.rows('SELECT id FROM versions')) == 2
        checks.append('Batch preprocessing persists one shared transformed version across engines')

        result_id = store.one('tasks', jobs[0])['result_id']
        original = store.result(result_id)['original']
        edit = deepcopy(store.result(result_id)['edited'])
        edit['text'] = edit['text'].replace('原始正文', '人工正文')
        edit['tables'][0]['cells'][0]['text'] = '人工表头'
        saved = store.save(result_id, edit, 0)
        try:
            store.save(result_id, edit, 0)
            raise AssertionError('stale save was accepted')
        except Conflict:
            pass
        store = Store(store.root)
        assert store.result(result_id)['edited'] == edit
        back = store.history(result_id, -1, saved['revision'])
        assert back['edited']['text'] == text
        restored = store.history(result_id, 1, back['revision'])
        assert restored['edited'] == edit
        assert restored['original'] == original
        checks.append('Save/reopen/conflict/undo/redo preserves immutable original and edited content')

        for fmt in ('txt', 'md', 'json', 'xlsx'):
            output = build_export(store, [result_id], fmt)
            if fmt == 'xlsx':
                book = load_workbook(output)
                sheet = book['Table 1']
                assert sheet['A1'].value == '人工表头'
                assert sheet['A2'].value == '00123456789012345678'
                assert sheet['B2'].value == '=1+1' and sheet['B2'].data_type == 's'
                assert str(next(iter(sheet.merged_cells.ranges))) == 'A1:B1'
                assert book['来源索引']['K2'].value == result_id
                book.close()
            elif fmt == 'json':
                value = json.loads(output.read_text('utf-8'))
                assert value['edited'] == edit and value['original'] == original
            else:
                value = output.read_text('utf-8')
                assert '人工正文' in value and '人工表头' in value and '原始表头' not in value
        checks.append('TXT/Markdown/JSON/XLSX exports preserve saved revisions, merges, identifiers and literal formulas')

        # Explicitly adopt the selected test result before confirming.
        with store.transaction() as db:
            db.execute('UPDATE selections SET result_id=? WHERE image_id=?', (result_id, image['id']))
        store.set_review(image['id'], 'confirmed', result_id, restored['revision'], prepared[0])
        build_export(store, [result_id], 'json', confirmed_only=True)
        updated = deepcopy(edit)
        updated['text'] += '\n新增正文'
        updated = store.save(result_id, updated, restored['revision'])
        before = set((store.root / 'exports').iterdir())
        try:
            build_export(store, [result_id], 'json', confirmed_only=True)
            raise AssertionError('stale confirmation accepted')
        except Conflict:
            pass
        assert set((store.root / 'exports').iterdir()) == before
        checks.append('Confirmed-only export rejects stale review and removes the incomplete export')

        reverted = store.history(result_id, -1, updated['revision'])
        branch = deepcopy(reverted['edited'])
        branch['text'] += '\n另一个分支'
        branch = store.save(result_id, branch, reverted['revision'])
        assert not branch['can_redo'] and branch['original'] == original
        checks.append('Editing after undo truncates redo without changing original')

        try:
            store.file('../outside.txt')
            raise AssertionError('workspace traversal accepted')
        except ValueError:
            pass
        maintenance = ProjectMaintenance(store, SimpleNamespace(status=lambda: {'task_id': None}))
        assert maintenance.orphans()['count'] == 0
        deletion = maintenance.delete(project['id'], project['name'])
        assert deletion['deleted'] and store.one('projects', other['id'])['name'] == '保留项目'
        assert not store.rows('SELECT id FROM images') and not store.rows('SELECT result_id FROM edits')
        checks.append('Path containment, orphan inventory and project cleanup preserve unrelated project')
    evidence = {'passed': len(checks), 'checks': checks, 'scope': 'synthetic data, isolated workspace, no GPU or browser'}
    Path(__file__).with_name('smoke-results.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
