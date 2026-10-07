"""Bounded workbook/schema discovery, in the existing data worker runtime."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook


def text(value):
    return None if value is None else str(value)[:512]


def inspect_input(request):
    path = Path(request['path']).resolve(strict=True)
    if path.drive.upper() != 'F:' or path.stat().st_size > 64 * 1024 * 1024:
        raise ValueError('Input storage or size is unsupported')
    options = request.get('options', {})
    header = options.get('header_row', 1)
    if type(header) is not int or not 1 <= header <= 100:
        raise ValueError('Invalid header row')
    format = request['format']
    sheets = []
    if format == 'xlsx':
        with zipfile.ZipFile(path) as archive:
            if sum(file.file_size for file in archive.infolist()) > 512 * 1024 * 1024:
                raise ValueError('Workbook expands beyond limit')
        # Verified asset blobs deliberately have hash-only names. Read bytes
        # rather than asking openpyxl to infer a format from that filename.
        book = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.sheetnames) > 50:
                raise ValueError('Too many sheets')
            for sheet in book:
                if sheet.max_column > 256:
                    raise ValueError('Too many columns')
                rows = list(sheet.iter_rows(min_row=header, max_row=header + 5, values_only=True))
                sheets.append({'name': sheet.title, 'rows': max(0, sheet.max_row - header),
                               'columns': [{'column_id': 'c' + str(i + 1), 'header': text(value)}
                                           for i, value in enumerate(rows[0] if rows else [])],
                               'sample': [[text(cell) for cell in row] for row in rows[1:]]})
        finally:
            book.close()
    elif format == 'csv':
        encoding, delimiter = options.get('encoding', 'utf-8-sig'), options.get('delimiter', ',')
        if encoding not in {'utf-8', 'utf-8-sig', 'gb18030', 'utf-16'} or delimiter not in {',', '\t', ';', '|'}:
            raise ValueError('Invalid CSV options')
        rows = []
        with path.open('r', newline='', encoding=encoding) as stream:
            reader = csv.reader(stream, delimiter=delimiter)
            for index, row in enumerate(reader, start=1):
                if index < header:
                    continue
                if len(row) > 256:
                    raise ValueError('Too many columns')
                rows.append([text(cell) for cell in row])
                if len(rows) == 6:
                    break
        sheets.append({'name': None, 'rows': None,
                       'columns': [{'column_id': 'c' + str(i + 1), 'header': value}
                                   for i, value in enumerate(rows[0] if rows else [])], 'sample': rows[1:]})
    elif format == 'parquet':
        import duckdb
        with duckdb.connect(':memory:') as db:
            description = db.execute('DESCRIBE SELECT * FROM read_parquet(?)', [str(path)]).fetchall()
            if len(description) > 256:
                raise ValueError('Too many columns')
            sample = db.execute('SELECT * FROM read_parquet(?) LIMIT 5', [str(path)]).fetchall()
        sheets.append({'name': None, 'rows': None,
                       'columns': [{'column_id': 'c' + str(i + 1), 'header': row[0], 'type': row[1]}
                                   for i, row in enumerate(description)],
                       'sample': [[text(cell) for cell in row] for row in sample]})
    else:
        raise ValueError('Unsupported data format')
    return {'format': format, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'sheets': sheets, 'requires_sheet_selection': len(sheets) > 1, 'header_row': header}


if __name__ == '__main__':
    request, output = map(Path, sys.argv[1:3])
    try:
        result = {'ok': True, 'data': inspect_input(json.loads(request.read_text(encoding='utf-8')))}
        serialized = json.dumps(result, ensure_ascii=False, allow_nan=False)
        if len(serialized.encode('utf-8')) > 256 * 1024:
            raise ValueError('Schema and sample exceed bounded response limit')
        output.write_text(serialized, encoding='utf-8')
    except Exception:
        output.write_text(json.dumps({'ok': False, 'error': '无法探查资料；请检查格式、编码和工作表大小'}), encoding='utf-8')
        sys.exit(1)
