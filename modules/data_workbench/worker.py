"""Private subprocess entry. Only fixed table import and summary queries run."""

from __future__ import annotations

import csv
import hashlib
import json
import sys
import zipfile
from importlib.metadata import version
from pathlib import Path

import duckdb
from openpyxl import load_workbook


def _validate_runtime() -> None:
    lock = Path(__file__).resolve().parents[2] / "config/runtime-locks/data.json"
    config = json.loads(lock.read_text("utf-8"))
    if sys.version_info[:2] != (3, 12) or config.get("schema_version") != "1.0":
        raise ValueError("Data runtime version differs")
    for package, expected in config["dependencies"].items():
        if version(package) != expected:
            raise ValueError("Data dependency version differs")


def _excel_to_csv(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source) as archive:
        parts = archive.infolist()
        if len(parts) > 10_000 or sum(part.file_size for part in parts) > 256 * 1024 * 1024:
            raise ValueError("Workbook exceeds limits")
    book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = book.active
        if sheet is None:
            raise ValueError("Workbook has no sheet")
        with destination.open("x", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index > 100_000 or len(row) > 1024:
                    raise ValueError("Workbook exceeds limits")
                writer.writerow(["" if value is None else str(value) for value in row])
                if stream.tell() > 128 * 1024 * 1024:
                    raise ValueError("Workbook exceeds limits")
    finally:
        book.close()


def run(request: dict) -> dict:
    _validate_runtime()
    source = Path(request["source"]).resolve(strict=True)
    database = Path(request["database"])
    temp = Path(request["temp"])
    report = Path(request["report"])
    if (not source.is_file() or source.parent != database.parent or temp.parent != database.parent
            or report.parent != database.parent or report.name != "report.tmp"
            or source.name not in {"input.csv", "input.parquet", "input.xlsx"}
            or database.name != "analysis.duckdb" or temp.name != "temp"):
        raise ValueError("Invalid private job layout")
    if hashlib.sha256(source.read_bytes()).hexdigest() != request["source_sha256"]:
        raise ValueError("Staged source changed")
    temp.mkdir(mode=0o700)
    readable = source
    if source.suffix == ".xlsx":
        readable = source.parent / "workbook.csv"
        _excel_to_csv(source, readable)
    # The only path exception is the exact staged asset. External access and
    # automatic extension installation/loading remain disabled for all queries.
    config = {"threads": "2", "memory_limit": "256MB", "max_temp_directory_size": "256MB",
              "temp_directory": str(temp),
              "autoinstall_known_extensions": "false",
              "autoload_known_extensions": "false", "allow_community_extensions": "false"}
    connection = duckdb.connect(str(database), config=config)
    try:
        connection.execute("SET allowed_paths = ?", [[str(readable)]])
        connection.execute("SET enable_external_access = false")
        connection.execute("SET lock_configuration = true")
        if readable.suffix == ".parquet":
            connection.execute("CREATE TABLE source AS SELECT * FROM read_parquet(?)", [str(readable)])
        else:
            connection.execute("CREATE TABLE source AS SELECT * FROM read_csv(?)", [str(readable)])
        rows = connection.execute("SELECT count(*) FROM source").fetchone()[0]
        columns = [row[1] for row in connection.execute("PRAGMA table_info('source')").fetchall()]
        if len(columns) > 256 or any(len(name) > 256 for name in columns):
            raise ValueError("Table schema exceeds limits")
        security = connection.execute("SELECT current_setting('enable_external_access'), "
                                      "current_setting('autoinstall_known_extensions'), "
                                      "current_setting('autoload_known_extensions')").fetchone()
        return {"job_id": request["job_id"], "source_sha256": request["source_sha256"],
                "rows": rows, "columns": columns,
                "external_access": security[0], "extension_autoinstall": security[1],
                "extension_autoload": security[2]}
    finally:
        connection.close()


if __name__ == "__main__":
    try:
        payload = sys.stdin.buffer.read(8193)
        if len(payload) > 8192:
            raise ValueError("Invalid private request")
        request = json.loads(payload)
        result = run(request)
        Path(request["report"]).write_bytes(json.dumps(result, ensure_ascii=True,
                                                      separators=(",", ":")).encode("utf-8"))
    except Exception:
        # Raw file, parser and extension errors can reveal sensitive paths.
        sys.exit(1)
