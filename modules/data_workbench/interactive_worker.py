"""One isolated process performs one data-workbench operation.

The platform supplies a private session directory and an immutable request.
Only successful replies advance the session's visible revision pointers.
"""
from __future__ import annotations

from contextlib import closing
from dataclasses import asdict
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import sys
import traceback
import uuid

from openpyxl import load_workbook

from modules.data_workbench.importing import import_csv, import_xlsx
from modules.data_workbench.importing.parquet import import_parquet
from modules.data_workbench.recipes import RecipeStore, Step
from modules.data_workbench.comparison import ComparisonSpec, compare_exact
from modules.data_workbench.exporting.results import (
    SourceDescriptor, _headers, _rows, _csv, _xlsx, _parquet, _readback,
)


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def write_json(path, value):
    temporary = path.with_name(path.name + '.tmp')
    temporary.write_text(encoded(value), encoding='utf-8')
    os.replace(temporary, path)


def views(root):
    db = sqlite3.connect(root / 'views.sqlite3', timeout=10)
    db.executescript('''
        CREATE TABLE IF NOT EXISTS table_views(revision TEXT PRIMARY KEY, metadata TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS table_rows(revision TEXT, ordinal INTEGER, payload TEXT,
            PRIMARY KEY(revision,ordinal));
        CREATE TABLE IF NOT EXISTS result_views(id TEXT PRIMARY KEY, metadata TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS result_rows(id TEXT, ordinal INTEGER, category TEXT, payload TEXT,
            PRIMARY KEY(id,ordinal));
        CREATE INDEX IF NOT EXISTS result_category ON result_rows(id,category,ordinal);
    ''')
    return db


def history(store, revision):
    result = []
    with closing(store._connect()) as db:
        current = revision
        while current:
            row = db.execute('SELECT parent_revision,step_json,event_json FROM recipe_revisions '
                             'WHERE revision_id=?', (current,)).fetchone()
            if row is None or len(result) > 250:
                raise ValueError('操作历史不可用或过长')
            event = json.loads(row['event_json'])
            result.append({'revision': current, 'step': json.loads(row['step_json']) if row['step_json'] else None,
                           'input_rows': event['input_rows'], 'output_rows': event['output_rows'],
                           'errors_total': event['errors_total']})
            current = row['parent_revision']
    return list(reversed(result))


def cache_table(root, store, revision):
    value = store.get(revision)
    snapshot = value['snapshot']
    metadata = {'revision': revision, 'columns': snapshot['columns'], 'rows': len(snapshot['rows']),
                'event': {key: value['event'][key] for key in
                          ('operator', 'input_rows', 'output_rows', 'errors_total', 'error_sample')},
                'duplicate_groups': len(value['event']['duplicate_groups'])}
    with closing(views(root)) as db, db:
        db.execute('INSERT OR REPLACE INTO table_views VALUES (?,?)', (revision, encoded(metadata)))
        db.execute('DELETE FROM table_rows WHERE revision=?', (revision,))
        db.executemany('INSERT INTO table_rows VALUES (?,?,?)',
                       ((revision, i, encoded(row)) for i, row in enumerate(snapshot['rows'])))
    return metadata


def spec_from(value):
    allowed = {'key_pairs', 'compare_pairs', 'amount_pairs', 'type_policy', 'null_policy',
               'amount_tolerance', 'trim_text', 'casefold_text'}
    if not isinstance(value, dict) or set(value) - allowed:
        raise ValueError('比对参数无效')
    value = dict(value)
    for key in ('key_pairs', 'compare_pairs', 'amount_pairs'):
        value[key] = tuple(tuple(pair) for pair in value.get(key, []))
    return ComparisonSpec(**value)


def descriptor(table):
    source = table['source']
    return SourceDescriptor(source['asset_id'], 1, source['sha256'], source['format'], source.get('sheet_name'))


def do_operation(request):
    root = Path(request['session_dir']).resolve(strict=True)
    state = request['state']
    payload = request['payload']
    action = request['action']
    store = RecipeStore(root / 'recipes.sqlite3')
    result = {}
    if action == 'import':
        side = payload['side']
        if side not in ('left', 'right'):
            raise ValueError('请选择左右表')
        upload = request['upload']
        path = (root / upload['relative']).resolve(strict=True)
        if not path.is_relative_to(root / 'uploads'):
            raise ValueError('文件不可用')
        options = dict(payload.get('options', {}))
        sheets = []
        if upload['format'] == 'xlsx':
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheets = workbook.sheetnames
            finally:
                workbook.close()
            name = options.get('sheet_name') or sheets[0]
            table = import_xlsx(path, source_sha256=upload['sha256'], sheet_name=name,
                                header_row=options.get('header_row', 1))
            options['sheet_name'] = name
        elif upload['format'] == 'csv':
            table = import_csv(path, source_sha256=upload['sha256'], encoding=options.get('encoding', 'utf-8-sig'),
                               delimiter=options.get('delimiter', ','), header_row=options.get('header_row', 1))
        else:
            table = import_parquet(path, source_sha256=upload['sha256'])
        revision = store.import_table(table)
        metadata = cache_table(root, store, revision)
        state['tables'][side] = {'revision': revision, 'upload_id': upload['id'], 'name': upload['name'],
                                 'format': upload['format'], 'options': options, 'sheets': sheets,
                                 'warnings': list(table.warnings), 'history': history(store, revision),
                                 'versions': history(store, revision),
                                 'columns': metadata['columns'], 'rows': metadata['rows'],
                                 'source': {'asset_id': upload['id'], 'sha256': upload['sha256'],
                                            'name': upload['name'], 'format': upload['format'],
                                            'sheet_name': options.get('sheet_name') if upload['format'] == 'xlsx' else None}}
        state['result_id'] = None
    elif action in {'step', 'undo', 'select_revision', 'save_recipe', 'replay'}:
        side = payload['side']
        if side not in ('left', 'right') or not state['tables'].get(side):
            raise ValueError('请先导入表格')
        table = state['tables'][side]
        revision = table['revision']
        if action == 'step':
            revision = store.apply(revision, Step(payload['operator'], payload['parameters']))
        elif action == 'undo':
            revision = store.undo(revision)
        elif action == 'select_revision':
            revision = payload['revision']
            if revision not in {item['revision'] for item in table.get('versions', table['history'])}:
                raise ValueError('请选择当前表的历史版本')
        elif action == 'save_recipe':
            steps = [item['step'] for item in history(store, revision) if item['step']]
            if not steps:
                raise ValueError('请先添加清洗步骤')
            name = str(payload.get('name', '')).strip()
            if not name or len(name) > 100 or len(state['recipes']) >= 100:
                raise ValueError('配方名称或数量超过限制')
            original = store.get(table['history'][0]['revision'])['snapshot']
            recipe = {'id': 'p' + uuid.uuid4().hex, 'name': name, 'steps': steps,
                      'columns': original['columns']}
            state['recipes'].append(recipe)
            return {'state': state, 'result': {'recipe_id': recipe['id']}}
        else:
            recipe = next((item for item in state['recipes'] if item['id'] == payload['recipe_id']), None)
            if recipe is None:
                raise ValueError('配方不存在')
            mapping = payload['mapping']
            known = {column['column_id'] for column in table['columns']}
            if not isinstance(mapping, dict) or any(value not in known for value in mapping.values()):
                raise ValueError('请明确选择目标列映射')
            for step in recipe['steps']:
                parameters = dict(step['parameters'])
                if 'column' in parameters:
                    parameters['column'] = mapping[parameters['column']]
                if 'columns' in parameters:
                    parameters['columns'] = [mapping[column] for column in parameters['columns']]
                revision = store.apply(revision, Step(step['operator'], parameters))
        metadata = cache_table(root, store, revision)
        lineage = history(store, revision)
        versions = {item['revision']: item for item in table.get('versions', table['history'])}
        versions.update({item['revision']: item for item in lineage})
        table.update({'revision': revision, 'history': history(store, revision),
                      'versions': list(versions.values()),
                      'columns': metadata['columns'], 'rows': metadata['rows']})
        state['result_id'] = None
    elif action == 'compare':
        left, right = state['tables']['left'], state['tables']['right']
        left_snapshot = store.get(left['revision'])['snapshot']
        right_snapshot = store.get(right['revision'])['snapshot']
        compared = compare_exact(left_snapshot, right_snapshot, left_revision=left['revision'],
                                 right_revision=right['revision'], spec=spec_from(payload['spec']))
        headers = _headers(left_snapshot, right_snapshot)
        rows, _ = _rows(compared, left_snapshot, right_snapshot, descriptor(left), descriptor(right))
        summary = {'id': compared.comparison_id, 'coverage': compared.coverage, 'policy': compared.policy,
                   'left': left, 'right': right, 'headers': headers, 'row_count': len(rows)}
        with closing(views(root)) as db, db:
            db.execute('INSERT OR REPLACE INTO result_views VALUES (?,?)', (compared.comparison_id, encoded(summary)))
            db.execute('DELETE FROM result_rows WHERE id=?', (compared.comparison_id,))
            db.executemany('INSERT INTO result_rows VALUES (?,?,?,?)',
                           ((compared.comparison_id, i, row[0], encoded(dict(zip(headers, row))))
                            for i, row in enumerate(rows)))
        state['result_id'] = compared.comparison_id
        state['comparison'] = {'left_revision': left['revision'], 'right_revision': right['revision'],
                               'spec': compared.policy}
        state['comparisons'] = [item for item in state['comparisons'] if item['id'] != compared.comparison_id]
        state['comparisons'].insert(0, {'id': compared.comparison_id, 'left_name': left['name'],
                                      'right_name': right['name'], 'coverage': compared.coverage})
        state['comparisons'] = state['comparisons'][:50]
    elif action == 'export_table':
        table = state['tables'][payload['side']]
        snapshot = store.get(table['revision'])['snapshot']
        folder = root / 'exports' / request['job_id']
        folder.mkdir(parents=True, exist_ok=False)
        headers = ['source_row', *[column['column_id'] + ' | ' + (column['header'] or column['column_id'])
                                   for column in snapshot['columns']]]
        rows = [[str(row['source_row']), *['' if cell['value'] is None else str(cell['value'])
                                         for cell in row['cells']]] for row in snapshot['rows']]
        files = {name: folder / name for name in
                 ('cleaned.csv', 'cleaned.xlsx', 'cleaned.parquet', 'summary.json', 'provenance.json')}
        _csv(files['cleaned.csv'], headers, rows)
        _xlsx(files['cleaned.xlsx'], headers, rows)
        _parquet(files['cleaned.parquet'], headers, rows)
        write_json(files['summary.json'], {'revision': table['revision'], 'source': table['source'],
                   'columns': snapshot['columns'], 'rows': len(rows), 'recipe': table['history']})
        write_json(files['provenance.json'], {'source': table['source'], 'revision': table['revision'],
                   'rows': snapshot['rows']})
        _readback({'results.' + suffix: files['cleaned.' + suffix] for suffix in ('csv', 'xlsx', 'parquet')}, headers, rows)
        exported = {'id': request['job_id'], 'result_id': None, 'files': []}
        for name, path in files.items():
            with path.open('rb') as stream:
                digest = hashlib.file_digest(stream, 'sha256').hexdigest()
            exported['files'].append({'name': name, 'relative': path.relative_to(root).as_posix(),
                                      'sha256': digest, 'bytes': path.stat().st_size})
        state['exports'].insert(0, exported)
        result = {'export_id': exported['id']}
    elif action == 'export':
        identity = payload['result_id']
        if identity not in {item['id'] for item in state['comparisons']}:
            raise ValueError('比对结果不可用')
        with closing(views(root)) as db:
            summary = json.loads(db.execute('SELECT metadata FROM result_views WHERE id=?', (identity,)).fetchone()[0])
        left, right = summary['left'], summary['right']
        left_snapshot = store.get(left['revision'])['snapshot']
        right_snapshot = store.get(right['revision'])['snapshot']
        compared = compare_exact(left_snapshot, right_snapshot, left_revision=left['revision'],
                                 right_revision=right['revision'], spec=spec_from(summary['policy']))
        if compared.comparison_id != identity:
            raise ValueError('结果版本不一致')
        folder = root / 'exports' / request['job_id']
        folder.mkdir(parents=True, exist_ok=False)
        headers = _headers(left_snapshot, right_snapshot)
        rows, provenance = _rows(compared, left_snapshot, right_snapshot, descriptor(left), descriptor(right))
        files = {name: folder / name for name in
                 ('results.csv', 'results.xlsx', 'results.parquet', 'summary.json', 'provenance.json')}
        _csv(files['results.csv'], headers, rows)
        _xlsx(files['results.xlsx'], headers, rows)
        _parquet(files['results.parquet'], headers, rows)
        write_json(files['summary.json'], {'comparison_id': identity, 'coverage': compared.coverage,
                   'policy': compared.policy, 'left_source': left['source'], 'right_source': right['source'],
                   'left_revision': left['revision'], 'right_revision': right['revision'],
                   'left_recipe': left['history'], 'right_recipe': right['history'],
                   'csv_text_policy': 'formula-like text is prefixed with apostrophe'})
        write_json(files['provenance.json'], {'comparison_id': identity, 'rows': provenance})
        _readback(files, headers, rows)
        exported = {'id': request['job_id'], 'result_id': identity, 'files': []}
        for name, path in files.items():
            with path.open('rb') as stream:
                digest = hashlib.file_digest(stream, 'sha256').hexdigest()
            exported['files'].append({'name': name, 'relative': path.relative_to(root).as_posix(),
                                      'sha256': digest, 'bytes': path.stat().st_size})
        state['exports'].insert(0, exported)
        result = {'export_id': exported['id']}
    else:
        raise ValueError('不支持的操作')
    return {'state': state, 'result': result}


if __name__ == '__main__':
    request_path, response_path = map(Path, sys.argv[1:3])
    try:
        value = do_operation(json.loads(request_path.read_text(encoding='utf-8')))
        write_json(response_path, {'ok': True, **value})
    except Exception as error:
        response_path.with_suffix('.log').write_text(traceback.format_exc(), encoding='utf-8')
        message = (str(error) if isinstance(error, ValueError) else '处理失败，请检查文件格式、工作表和列映射。')
        write_json(response_path, {'ok': False, 'error': message[:300]})
        sys.exit(1)
